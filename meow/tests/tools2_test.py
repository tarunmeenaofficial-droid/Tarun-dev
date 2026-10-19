import os, time
os.environ["MEOW_CAPTIONS"]="0"
exec(open("/tmp/meow_tests/harness.py").read().split("from fastapi.testclient")[0])
import pymupdf as fitz
from PIL import Image
import numpy as np
T="/tmp/meow_tests/"
# form
d=fitz.open(); pg=d.new_page()
pg.insert_text((50,80),"Name:",fontsize=12); w=fitz.Widget(); w.field_name="Full Name"; w.field_type=fitz.PDF_WIDGET_TYPE_TEXT; w.rect=fitz.Rect(100,65,300,85); pg.add_widget(w)
w=fitz.Widget(); w.field_name="agree"; w.field_type=fitz.PDF_WIDGET_TYPE_CHECKBOX; w.rect=fitz.Rect(100,100,115,115); pg.add_widget(w)
d.save(T+"form2.pdf")
# compare versions
for name,txt in (("v1.pdf","The rent is 8000 rupees per month and is due on the 5th.\nTenant: Rahul Sharma."),("v2.pdf","The rent is 9000 rupees per month and is due on the 1st.\nTenant: Rahul Sharma. Deposit two months.")):
    d=fitz.open(); d.new_page().insert_text((50,80),txt,fontsize=12); d.save(T+name)
ids={}
for n in ["form2.pdf","v1.pdf","v2.pdf","scan.pdf","doc.pdf","photo.jpg"]:
    rec=g["ingest"](n, open(T+n,"rb").read()); ids[n]=rec["id"]
print("form fields in extract:", "Form fields" in g["FILES"][ids["form2.pdf"]]["text"], g["FILES"][ids["form2.pdf"]]["text"][-160:])
def run(op, files, params={}):
    t=time.time()
    try:
        rec,msg=g["run_tool"](op,[ids[f] for f in files],params); print(f"{op}: {msg} → {rec['name']} ({time.time()-t:.1f}s)"); return rec
    except Exception as e: print(op,"ERROR",e)
r=run("pdf_fill_form",["form2.pdf"],{"values":'{"full name":"Tarun Meena","Agree":true}'})
with fitz.open(r["path"]) as dd: print("   values:", [(w.field_name,w.field_value) for w in dd[0].widgets()])
run("pdf_fill_form",["form2.pdf"],{"values":'{"phone":"1"}'})
r=run("pdf_compare",["v1.pdf","v2.pdf"])
r=run("pdf_make_searchable",["scan.pdf"])
with fitz.open(r["path"]) as dd: print("   searchable text:", dd[0].get_text()[:120].replace("\n"," | "))
r=run("pdf_to_word",["scan.pdf"])
import docx; print("   docx paras:", [p.text for p in docx.Document(r["path"]).paragraphs][:4])
r=run("pdf_to_word",["doc.pdf"])
print("   docx paras:", [p.text for p in docx.Document(r["path"]).paragraphs if p.text.strip()][:4])
r=run("image_restore",["photo.jpg"],{"upscale":"no"})
