# ╔════════════════════════════════════════════════════════════════════╗
# ║  MEOW 🐈 — your own AI assistant, running on a Kaggle GPU          ║
# ║  Fast / Think modes · PDFs, images, docs, sheets · file tools      ║
# ╚════════════════════════════════════════════════════════════════════╝
#
# Run it with the one-cell launcher (kaggle_launcher.py): it installs what's missing, downloads the newest
# meow.py from GitHub and starts it. Keys come from the launcher or Kaggle Secrets — never from this file.

import os, re, gc, io, json, time, uuid, html, zipfile, datetime, threading, mimetypes
from typing import List

from fastapi import FastAPI, File, UploadFile, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, StreamingResponse, FileResponse
from starlette.concurrency import run_in_threadpool
from pydantic import BaseModel
from PIL import Image, ImageOps, ImageEnhance, ImageFilter, ExifTags
try:
    import pymupdf as fitz  # PyMuPDF
except ImportError:
    import fitz
import uvicorn

try:
    from pillow_heif import register_heif_opener  # iPhone HEIC photos
    register_heif_opener()
except Exception:
    pass

# ==========================================================
# CONFIG
# ==========================================================
USER_NAME = "Tarun"
REPO_ID = "HauhauCS/Qwen3.8-27B-Uncensored-HauhauCS-Aggressive-MTP-GGUF"
GGUF_FILENAME = "Qwen3.8-27B-Uncensored-HauhauCS-Aggressive-IQ4_XS.gguf"
MODEL_LABEL = "Qwen3.8 · 27B"
N_CTX = 16384              # falls back to 8192 / 4096 automatically if VRAM is short
MAX_NEW_FAST = 2048
MAX_NEW_THINK = 6144
ENABLE_CAPTIONS = os.environ.get('MEOW_CAPTIONS','1')=='1'    # BLIP image captions on CPU (one-time ~1.9 GB download)
PORT = 8000
MEOW_VERSION = "2026.10.19.2"
MEOW_CHANGELOG = [   # newest first — shown in the app under "What's new"
    {"v": "2026.10.19.2", "items": ["Start Meow from anywhere with one click (Kaggle API) and get a phone notification with the link",
                                     "Password protection", "Auto power-off when idle to save GPU hours", "What's new panel"]},
    {"v": "2026.10.19", "items": ["One-cell Kaggle launcher with automatic updates from GitHub", "Sharper OCR spacing for IDs"]},
    {"v": "2026.10.18", "items": ["Reddit/Deep research keeps searching until your question is answered (up to 3 rounds)",
                                   "Comments ranked by relevance; promo/scam comments filtered", "Python code interpreter with charts",
                                   "YouTube, Wikipedia, papers, GitHub, Stack Overflow, Hacker News sources; reads pasted links",
                                   "Mermaid diagrams, HTML preview, follow-up chips, edit & resend, voice mode",
                                   "PDF→Word, searchable scans, form filling, PDF compare, photo restore"]},
    {"v": "2026.10.17", "items": ["Web + Reddit search with citations", "Photo-real text editing and magic eraser"]},
]

# Keys are NOT stored in this file (it lives on GitHub). Meow reads them from environment variables (set by the
# launcher cell) or from Kaggle → Add-ons → Secrets, using exactly these names.
NGROK_AUTH_TOKEN = os.environ.get("NGROK_AUTH_TOKEN", "")
SERPER_API_KEY = os.environ.get("SERPER_API_KEY", "")
TAVILY_API_KEY = os.environ.get("TAVILY_API_KEY", "")
REDDIT_CLIENT_ID = os.environ.get("REDDIT_CLIENT_ID", "")          # optional — only if Reddit approves your API request
REDDIT_CLIENT_SECRET = os.environ.get("REDDIT_CLIENT_SECRET", "")
GITHUB_TOKEN = os.environ.get("GITHUB_TOKEN", "")                  # optional — higher GitHub search limits
BRAVE_API_KEY = os.environ.get("BRAVE_API_KEY", "")                # optional
WORK_DIR = "/kaggle/working/meow_files" if os.path.isdir("/kaggle/working") else "/tmp/meow_files"
MAX_FILE_CHARS = 250_000
os.makedirs(WORK_DIR, exist_ok=True)

# ==========================================================
# MODEL (reuses the one already in GPU memory when possible)
# ==========================================================
def _load_llm(ctx):
    from huggingface_hub import hf_hub_download
    from llama_cpp import Llama
    import glob
    local = glob.glob(f"/kaggle/input/**/{GGUF_FILENAME}", recursive=True)   # model attached as a Kaggle input → no download
    if local:
        path = local[0]; print(f"📦 Using model from {os.path.dirname(path)}")
    else:
        try:
            import hf_transfer  # noqa: F401 — 5-10× faster downloads when installed
            os.environ.setdefault("HF_HUB_ENABLE_HF_TRANSFER", "1")
        except ImportError:
            pass
        path = hf_hub_download(repo_id=REPO_ID, filename=GGUF_FILENAME)
    kw = dict(model_path=path, n_gpu_layers=-1, n_ctx=ctx, n_batch=512, verbose=False)
    try:
        return Llama(flash_attn=True, **kw)
    except TypeError:
        return Llama(**kw)

_G = globals()
if _G.get("llm") is not None and _G["llm"].n_ctx() >= N_CTX:
    print(f"✅ Reusing model already on GPU (ctx {_G['llm'].n_ctx()})")
else:
    if _G.get("llm") is not None:
        print("♻️  Reloading model with a bigger context window…")
        _G["llm"] = None
        gc.collect()
    for _ctx in (N_CTX, 8192, 4096):
        try:
            print(f"🔄 Loading {GGUF_FILENAME} onto GPU (ctx {_ctx})…")
            llm = _load_llm(_ctx)
            print("✅ Model loaded!")
            break
        except Exception as _e:
            print(f"⚠️  ctx={_ctx} failed: {_e}")
            gc.collect()
    else:
        raise RuntimeError("Model failed to load — restart the session and try again.")

CTX = llm.n_ctx()
_TEMPLATE = (getattr(llm, "metadata", None) or {}).get("tokenizer.chat_template", "") or ""
USE_CHATML = ("<|im_start|>" in _TEMPLATE) or not _TEMPLATE   # Qwen family = ChatML
LLM_LOCK = threading.Lock()
STOP = set()
THINK_RE = re.compile(r"<think>.*?(?:</think>|$)", re.S)


def ntok(s: str) -> int:
    if not s:
        return 0
    return len(llm.tokenize(s.encode("utf-8", "ignore"), add_bos=False, special=True))


def trim_tokens(s: str, n: int) -> str:
    if n <= 0:
        return ""
    toks = llm.tokenize(s.encode("utf-8", "ignore"), add_bos=False, special=False)
    if len(toks) <= n:
        return s
    return llm.detokenize(toks[:n]).decode("utf-8", errors="ignore") + "\n…[truncated to fit context]"

# ==========================================================
# FILE STORE + EXTRACTION
# ==========================================================
FILES = {}
IMG_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tiff", ".tif", ".heic", ".heif"}
SHEET_EXT = {".xlsx", ".xlsm", ".xls", ".csv", ".tsv"}


def safe_name(n):
    n = os.path.basename(n or "file")
    return re.sub(r"[^\w.\- ()]+", "_", n).strip()[:120] or "file"


def kind_of(name):
    ext = os.path.splitext(name)[1].lower()
    if ext == ".pdf": return "pdf"
    if ext in IMG_EXT: return "image"
    if ext == ".docx": return "docx"
    if ext == ".pptx": return "pptx"
    if ext in SHEET_EXT: return "sheet"
    if ext == ".zip": return "zip"
    return "text"


def new_path(name):
    fid = uuid.uuid4().hex[:8]
    return fid, os.path.join(WORK_DIR, f"{fid}_{safe_name(name)}")


def register(fid, path, name, origin="upload"):
    name = safe_name(name)
    rec = dict(id=fid, name=name, path=path, kind=kind_of(name), size=os.path.getsize(path),
               origin=origin, created=time.time(), text=None, meta={})
    try:
        if rec["kind"] == "pdf":
            with fitz.open(path) as d:
                rec["meta"]["pages"] = d.page_count
                if d.needs_pass: rec["meta"]["locked"] = True
        elif rec["kind"] == "image":
            with Image.open(path) as im:
                rec["meta"]["dims"] = f"{im.width}×{im.height}"
    except Exception:
        pass
    FILES[fid] = rec
    return rec


def public(rec):
    return {k: rec[k] for k in ("id", "name", "kind", "size", "origin", "meta")} | {
        "url": f"/api/files/{rec['id']}",
        "preview": (rec["text"] or "")[:240] if rec.get("text") else "",
    }


# ---------- image understanding (OCR + caption + colors + EXIF) ----------
_OCR, _CAP = None, None
FEATURES = {"ocr": False, "captions": False}


def ocr_image(img):
    global _OCR
    try:
        import numpy as np
        if _OCR is None:
            from rapidocr_onnxruntime import RapidOCR
            _OCR = RapidOCR()
            FEATURES["ocr"] = True
        im = img.convert("RGB")
        if max(im.size) > 2000:
            im.thumbnail((2000, 2000))
        res, _ = _OCR(np.array(im))
        return "\n".join(r[1] for r in (res or []) if float(r[2]) > 0.5)
    except Exception:
        return ""


def caption_image(img):
    global _CAP
    if not ENABLE_CAPTIONS:
        return ""
    try:
        from transformers import BlipProcessor, BlipForConditionalGeneration
        if _CAP is None:
            m = "Salesforce/blip-image-captioning-large"
            _CAP = (BlipProcessor.from_pretrained(m), BlipForConditionalGeneration.from_pretrained(m).eval())
            FEATURES["captions"] = True
        proc, model = _CAP
        im = img.convert("RGB")
        im.thumbnail((768, 768))
        import torch
        with torch.no_grad():
            out = model.generate(**proc(im, return_tensors="pt"), max_new_tokens=60, num_beams=3)
        return proc.decode(out[0], skip_special_tokens=True).strip()
    except Exception:
        return ""


def dominant_colors(img, k=5):
    try:
        q = img.convert("RGB").resize((80, 80)).quantize(colors=k, method=Image.Quantize.MEDIANCUT)
        pal, counts = q.getpalette(), sorted(q.getcolors(), reverse=True)
        total = sum(c for c, _ in counts)
        return ", ".join(f"#{pal[i*3]:02x}{pal[i*3+1]:02x}{pal[i*3+2]:02x} ({c*100//total}%)" for c, i in counts[:k])
    except Exception:
        return ""


def analyze_image(path):
    with Image.open(path) as raw:
        fmt = raw.format
        exif = {}
        try:
            for k, v in raw.getexif().items():
                tag = ExifTags.TAGS.get(k, k)
                if tag in ("Make", "Model", "DateTime", "Software", "LensModel"):
                    exif[tag] = str(v)
        except Exception:
            pass
        im = ImageOps.exif_transpose(raw)
        im.load()
    lines = [f"Dimensions: {im.width}×{im.height}px · {fmt or 'unknown'} · mode {im.mode}"]
    cap = caption_image(im)
    if cap: lines.append(f"Scene description (auto-caption): {cap}")
    txt = ocr_image(im)
    lines.append("Text found in image (OCR):\n" + txt if txt else "Text found in image (OCR): none")
    col = dominant_colors(im)
    if col: lines.append(f"Dominant colors: {col}")
    if exif: lines.append("EXIF: " + ", ".join(f"{k}={v}" for k, v in exif.items()))
    return "\n".join(lines)


# ---------- documents ----------
def _extract_pdf(rec):
    out, ocr_pages = [], 0
    with fitz.open(rec["path"]) as doc:
        if doc.needs_pass:
            return "(This PDF is password-protected. Ask the user for the password and use pdf_unlock.)"
        md = doc.metadata or {}
        if md.get("title"): out.append(f"Title: {md['title']}")
        for i, page in enumerate(doc):
            t = page.get_text("text").strip()
            if len(t) < 25 and ocr_pages < 15:          # scanned page → OCR it
                pix = page.get_pixmap(dpi=150)
                t = ocr_image(Image.frombytes("RGB", (pix.width, pix.height), pix.samples)).strip()
                ocr_pages += 1
                if t: t = "(OCR) " + t
            out.append(f"[Page {i+1}]\n{t or '(no text)'}")
            try:
                ims = page_layout(page)["images"]
                if ims:
                    out.append(f"[Page {i+1} images: " + "; ".join(
                        f"#{m['index']} at ({m['bbox'][0]:.0f},{m['bbox'][1]:.0f})-({m['bbox'][2]:.0f},{m['bbox'][3]:.0f})" for m in ims[:12]) + "]")
            except Exception:
                pass
        try:                                             # fillable form → list its fields for pdf_fill_form
            fields = [w for pg in doc for w in (pg.widgets() or [])]
            if fields:
                out.append("[Form fields (use pdf_fill_form): " + "; ".join(
                    f"{w.field_name!r} ({w.field_type_string}{', options: ' + '/'.join(map(str, (w.choice_values or [])[:8])) if w.choice_values else ''})"
                    f" = {w.field_value!r}" for w in fields[:60]) + "]")
        except Exception:
            pass
    return "\n\n".join(out)


def _extract_docx(p):
    import docx
    d = docx.Document(p)
    out = [para.text for para in d.paragraphs if para.text.strip()]
    for ti, t in enumerate(d.tables, 1):
        out.append(f"\n[Table {ti}]")
        for row in t.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)


def _extract_pptx(p):
    from pptx import Presentation
    out = []
    for i, s in enumerate(Presentation(p).slides, 1):
        texts = [sh.text_frame.text for sh in s.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
        notes = ""
        if s.has_notes_slide:
            notes = s.notes_slide.notes_text_frame.text.strip()
        out.append(f"[Slide {i}]\n" + "\n".join(texts) + (f"\nSpeaker notes: {notes}" if notes else ""))
    return "\n\n".join(out)


def _extract_sheet(p):
    import pandas as pd
    ext = os.path.splitext(p)[1].lower()
    sheets = {"data": pd.read_csv(p, sep=None, engine="python")} if ext in (".csv", ".tsv") else pd.read_excel(p, sheet_name=None)
    out = []
    for name, df in sheets.items():
        out.append(f"[Sheet: {name}] {df.shape[0]} rows × {df.shape[1]} columns\nColumns: {', '.join(map(str, df.columns))}")
        out.append(df.head(300).to_csv(index=False))
        num = df.select_dtypes("number")
        if not num.empty:
            out.append("Summary statistics:\n" + num.describe().round(3).to_csv())
    return "\n".join(out)


def _extract_zip(p):
    with zipfile.ZipFile(p) as z:
        names = z.namelist()
    return f"ZIP archive with {len(names)} entries:\n" + "\n".join(names[:300])


def _extract_text(p):
    data = open(p, "rb").read()
    if b"\x00" in data[:4096] and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return f"(Binary file, {len(data)} bytes — not readable as text.)"
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except Exception:
            continue
    return ""


def extract(rec):
    k, p = rec["kind"], rec["path"]
    try:
        if k == "pdf": t = _extract_pdf(rec)
        elif k == "image": t = analyze_image(p)
        elif k == "docx": t = _extract_docx(p)
        elif k == "pptx": t = _extract_pptx(p)
        elif k == "sheet": t = _extract_sheet(p)
        elif k == "zip": t = _extract_zip(p)
        else: t = _extract_text(p)
    except Exception as e:
        t = f"(Could not read this file: {e})"
    rec["text"] = (t or "").strip()[:MAX_FILE_CHARS]
    return rec["text"]


def ingest(name, data):
    fid, path = new_path(name)
    with open(path, "wb") as f:
        f.write(data)
    rec = register(fid, path, name, origin="upload")
    extract(rec)
    return rec

