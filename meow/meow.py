# ╔════════════════════════════════════════════════════════════════════╗
# ║  MEOW 🐈 — your own AI assistant, running on a Kaggle GPU          ║
# ║  Fast / Think modes · PDFs, images, docs, sheets · file tools      ║
# ╚════════════════════════════════════════════════════════════════════╝
#
# Run it with the one-cell launcher (kaggle_launcher.py): it installs what's missing, downloads the newest
# meow.py from GitHub and starts it. Keys come from the launcher or Kaggle Secrets — never from this file.

import os, re, gc, io, json, time, uuid, html, zipfile, datetime, threading, mimetypes
from typing import List

from fastapi import FastAPI, File, UploadFile, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, StreamingResponse, FileResponse
from starlette.concurrency import run_in_threadpool
from pydantic import BaseModel
from PIL import Image, ImageOps, ImageEnhance, ImageFilter, ExifTags
try:
    import pymupdf as fitz  # PyMuPDF
except ImportError:
    import fitz
import uvicorn

try:
    from pillow_heif import register_heif_opener  # iPhone HEIC photos
    register_heif_opener()
except Exception:
    pass

# ==========================================================
# CONFIG
# ==========================================================
USER_NAME = "Tarun"
REPO_ID = "HauhauCS/Qwen3.8-27B-Uncensored-HauhauCS-Aggressive-MTP-GGUF"
GGUF_FILENAME = "Qwen3.8-27B-Uncensored-HauhauCS-Aggressive-IQ4_XS.gguf"
MODEL_LABEL = "Qwen3.8 · 27B"
N_CTX = 16384              # falls back to 8192 / 4096 automatically if VRAM is short
MAX_NEW_FAST = 2048
MAX_NEW_THINK = 6144
ENABLE_CAPTIONS = os.environ.get('MEOW_CAPTIONS','1')=='1'    # BLIP image captions on CPU (one-time ~1.9 GB download)
PORT = 8000
MEOW_VERSION = "2026.10.19.3"
MEOW_CHANGELOG = [   # newest first — shown in the app under "What's new"
    {"v": "2026.10.19.3", "items": ["Your own private copy of the model (Kaggle dataset + private GitHub repo): Meow keeps working "
                                     "even if the Hugging Face repo disappears, and starts faster from the Kaggle copy"]},
    {"v": "2026.10.19.2", "items": ["Start Meow from anywhere with one click (Kaggle API) and get a phone notification with the link",
                                     "Password protection", "Auto power-off when idle to save GPU hours", "What's new panel"]},
    {"v": "2026.10.19", "items": ["One-cell Kaggle launcher with automatic updates from GitHub", "Sharper OCR spacing for IDs"]},
    {"v": "2026.10.18", "items": ["Reddit/Deep research keeps searching until your question is answered (up to 3 rounds)",
                                   "Comments ranked by relevance; promo/scam comments filtered", "Python code interpreter with charts",
                                   "YouTube, Wikipedia, papers, GitHub, Stack Overflow, Hacker News sources; reads pasted links",
                                   "Mermaid diagrams, HTML preview, follow-up chips, edit & resend, voice mode",
                                   "PDF→Word, searchable scans, form filling, PDF compare, photo restore"]},
    {"v": "2026.10.17", "items": ["Web + Reddit search with citations", "Photo-real text editing and magic eraser"]},
]

# Keys are NOT stored in this file (it lives on GitHub). Meow reads them from environment variables (set by the
# launcher cell) or from Kaggle → Add-ons → Secrets, using exactly these names.
NGROK_AUTH_TOKEN = os.environ.get("NGROK_AUTH_TOKEN", "")
SERPER_API_KEY = os.environ.get("SERPER_API_KEY", "")
TAVILY_API_KEY = os.environ.get("TAVILY_API_KEY", "")
REDDIT_CLIENT_ID = os.environ.get("REDDIT_CLIENT_ID", "")          # optional — only if Reddit approves your API request
REDDIT_CLIENT_SECRET = os.environ.get("REDDIT_CLIENT_SECRET", "")
GITHUB_TOKEN = os.environ.get("GITHUB_TOKEN", "")                  # optional — higher GitHub search limits
BRAVE_API_KEY = os.environ.get("BRAVE_API_KEY", "")                # optional
WORK_DIR = "/kaggle/working/meow_files" if os.path.isdir("/kaggle/working") else "/tmp/meow_files"
MAX_FILE_CHARS = 250_000
os.makedirs(WORK_DIR, exist_ok=True)

# ==========================================================
# MODEL (reuses the one already in GPU memory when possible)
# ==========================================================
def _model_from_github(dest_dir="/tmp/meow_model", workers=4):
    """Download YOUR private backup (release `model` in MEOW_MODEL_GH, made by `start_meow.py backup`):
    manifest.json + <2 GiB parts, fetched in parallel straight into place and checked against their sha256."""
    import hashlib, requests
    from concurrent.futures import ThreadPoolExecutor
    repo, tok = os.environ["MEOW_MODEL_GH"], os.environ.get("MEOW_MODEL_TOKEN") or os.environ.get("GITHUB_TOKEN", "")
    api = os.environ.get("MEOW_GH_API", "https://api.github.com")
    H = {"Authorization": f"Bearer {tok}"} if tok else {}
    r = requests.get(f"{api}/repos/{repo}/releases/tags/model", headers=H, timeout=30); r.raise_for_status()
    assets = {a["name"]: a["url"] for a in r.json()["assets"]}
    get = lambda name, **kw: requests.get(assets[name], headers={**H, "Accept": "application/octet-stream"}, timeout=(30, 600), **kw)
    man = get("manifest.json").json()
    os.makedirs(dest_dir, exist_ok=True)
    dst = os.path.join(dest_dir, man["file"])
    if os.path.exists(dst) and os.path.getsize(dst) == man["size"]: return dst
    tmp = dst + ".download"
    with open(tmp, "wb") as f: f.truncate(man["size"])
    offs, o = [], 0
    for p in man["parts"]: offs.append(o); o += p["size"]
    def one(i):
        p = man["parts"][i]
        for attempt in range(4):
            try:
                h, n = hashlib.sha256(), 0
                with get(p["name"], stream=True) as rr, open(tmp, "r+b") as f:
                    rr.raise_for_status(); f.seek(offs[i])
                    for c in rr.iter_content(8 << 20): f.write(c); h.update(c); n += len(c)
                if n == p["size"] and h.hexdigest() == p["sha256"]: return
                raise IOError(f"{p['name']} came back damaged")
            except Exception as e:
                if attempt == 3: raise
                print(f"   retrying {p['name']}: {str(e)[:100]}"); time.sleep(5 * (attempt + 1))
    t = time.time()
    print(f"⬇️  Downloading your backup from github.com/{repo} ({len(man['parts'])} parts, {man['size'] / 2**30:.1f} GiB)…")
    with ThreadPoolExecutor(workers) as ex: list(ex.map(one, range(len(man["parts"]))))
    os.replace(tmp, dst)
    print(f"✅ Backup downloaded in {time.time() - t:.0f}s")
    return dst

def _load_llm(ctx):
    from huggingface_hub import hf_hub_download
    from llama_cpp import Llama
    import glob
    local = glob.glob(f"/kaggle/input/**/{GGUF_FILENAME}", recursive=True)   # model attached as a Kaggle input → no download
    if local:
        path = local[0]; print(f"📦 Using model from {os.path.dirname(path)}")
    else:
        os.environ.setdefault("HF_XET_HIGH_PERFORMANCE", "1")   # fastest download path in current huggingface_hub
        try:
            path = hf_hub_download(repo_id=REPO_ID, filename=GGUF_FILENAME)
        except Exception as e:                                   # repo gone / HF down → your own private GitHub copy
            if not os.environ.get("MEOW_MODEL_GH"): raise
            print(f"⚠️  Hugging Face download failed ({str(e)[:120]}) — using your GitHub backup")
            path = _model_from_github()
    kw = dict(model_path=path, n_gpu_layers=-1, n_ctx=ctx, n_batch=512, verbose=False)
    try:
        return Llama(flash_attn=True, **kw)
    except TypeError:
        return Llama(**kw)

_G = globals()
if _G.get("llm") is not None and _G["llm"].n_ctx() >= N_CTX:
    print(f"✅ Reusing model already on GPU (ctx {_G['llm'].n_ctx()})")
else:
    if _G.get("llm") is not None:
        print("♻️  Reloading model with a bigger context window…")
        _G["llm"] = None
        gc.collect()
    for _ctx in (N_CTX, 8192, 4096):
        try:
            print(f"🔄 Loading {GGUF_FILENAME} onto GPU (ctx {_ctx})…")
            llm = _load_llm(_ctx)
            print("✅ Model loaded!")
            break
        except Exception as _e:
            print(f"⚠️  ctx={_ctx} failed: {_e}")
            gc.collect()
    else:
        raise RuntimeError("Model failed to load — restart the session and try again.")

CTX = llm.n_ctx()
_TEMPLATE = (getattr(llm, "metadata", None) or {}).get("tokenizer.chat_template", "") or ""
USE_CHATML = ("<|im_start|>" in _TEMPLATE) or not _TEMPLATE   # Qwen family = ChatML
LLM_LOCK = threading.Lock()
STOP = set()
THINK_RE = re.compile(r"<think>.*?(?:</think>|$)", re.S)


def ntok(s: str) -> int:
    if not s:
        return 0
    return len(llm.tokenize(s.encode("utf-8", "ignore"), add_bos=False, special=True))


def trim_tokens(s: str, n: int) -> str:
    if n <= 0:
        return ""
    toks = llm.tokenize(s.encode("utf-8", "ignore"), add_bos=False, special=False)
    if len(toks) <= n:
        return s
    return llm.detokenize(toks[:n]).decode("utf-8", errors="ignore") + "\n…[truncated to fit context]"

# ==========================================================
# FILE STORE + EXTRACTION
# ==========================================================
FILES = {}
IMG_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tiff", ".tif", ".heic", ".heif"}
SHEET_EXT = {".xlsx", ".xlsm", ".xls", ".csv", ".tsv"}


def safe_name(n):
    n = os.path.basename(n or "file")
    return re.sub(r"[^\w.\- ()]+", "_", n).strip()[:120] or "file"


def kind_of(name):
    ext = os.path.splitext(name)[1].lower()
    if ext == ".pdf": return "pdf"
    if ext in IMG_EXT: return "image"
    if ext == ".docx": return "docx"
    if ext == ".pptx": return "pptx"
    if ext in SHEET_EXT: return "sheet"
    if ext == ".zip": return "zip"
    return "text"


def new_path(name):
    fid = uuid.uuid4().hex[:8]
    return fid, os.path.join(WORK_DIR, f"{fid}_{safe_name(name)}")


def register(fid, path, name, origin="upload"):
    name = safe_name(name)
    rec = dict(id=fid, name=name, path=path, kind=kind_of(name), size=os.path.getsize(path),
               origin=origin, created=time.time(), text=None, meta={})
    try:
        if rec["kind"] == "pdf":
            with fitz.open(path) as d:
                rec["meta"]["pages"] = d.page_count
                if d.needs_pass: rec["meta"]["locked"] = True
        elif rec["kind"] == "image":
            with Image.open(path) as im:
                rec["meta"]["dims"] = f"{im.width}×{im.height}"
    except Exception:
        pass
    FILES[fid] = rec
    return rec


def public(rec):
    return {k: rec[k] for k in ("id", "name", "kind", "size", "origin", "meta")} | {
        "url": f"/api/files/{rec['id']}",
        "preview": (rec["text"] or "")[:240] if rec.get("text") else "",
    }


# ---------- image understanding (OCR + caption + colors + EXIF) ----------
_OCR, _CAP = None, None
FEATURES = {"ocr": False, "captions": False}


def ocr_image(img):
    global _OCR
    try:
        import numpy as np
        if _OCR is None:
            from rapidocr_onnxruntime import RapidOCR
            _OCR = RapidOCR()
            FEATURES["ocr"] = True
        im = img.convert("RGB")
        if max(im.size) > 2000:
            im.thumbnail((2000, 2000))
        res, _ = _OCR(np.array(im))
        return "\n".join(r[1] for r in (res or []) if float(r[2]) > 0.5)
    except Exception:
        return ""


def caption_image(img):
    global _CAP
    if not ENABLE_CAPTIONS:
        return ""
    try:
        from transformers import BlipProcessor, BlipForConditionalGeneration
        if _CAP is None:
            m = "Salesforce/blip-image-captioning-large"
            _CAP = (BlipProcessor.from_pretrained(m), BlipForConditionalGeneration.from_pretrained(m).eval())
            FEATURES["captions"] = True
        proc, model = _CAP
        im = img.convert("RGB")
        im.thumbnail((768, 768))
        import torch
        with torch.no_grad():
            out = model.generate(**proc(im, return_tensors="pt"), max_new_tokens=60, num_beams=3)
        return proc.decode(out[0], skip_special_tokens=True).strip()
    except Exception:
        return ""


def dominant_colors(img, k=5):
    try:
        q = img.convert("RGB").resize((80, 80)).quantize(colors=k, method=Image.Quantize.MEDIANCUT)
        pal, counts = q.getpalette(), sorted(q.getcolors(), reverse=True)
        total = sum(c for c, _ in counts)
        return ", ".join(f"#{pal[i*3]:02x}{pal[i*3+1]:02x}{pal[i*3+2]:02x} ({c*100//total}%)" for c, i in counts[:k])
    except Exception:
        return ""


def analyze_image(path):
    with Image.open(path) as raw:
        fmt = raw.format
        exif = {}
        try:
            for k, v in raw.getexif().items():
                tag = ExifTags.TAGS.get(k, k)
                if tag in ("Make", "Model", "DateTime", "Software", "LensModel"):
                    exif[tag] = str(v)
        except Exception:
            pass
        im = ImageOps.exif_transpose(raw)
        im.load()
    lines = [f"Dimensions: {im.width}×{im.height}px · {fmt or 'unknown'} · mode {im.mode}"]
    cap = caption_image(im)
    if cap: lines.append(f"Scene description (auto-caption): {cap}")
    txt = ocr_image(im)
    lines.append("Text found in image (OCR):\n" + txt if txt else "Text found in image (OCR): none")
    col = dominant_colors(im)
    if col: lines.append(f"Dominant colors: {col}")
    if exif: lines.append("EXIF: " + ", ".join(f"{k}={v}" for k, v in exif.items()))
    return "\n".join(lines)


# ---------- documents ----------
def _extract_pdf(rec):
    out, ocr_pages = [], 0
    with fitz.open(rec["path"]) as doc:
        if doc.needs_pass:
            return "(This PDF is password-protected. Ask the user for the password and use pdf_unlock.)"
        md = doc.metadata or {}
        if md.get("title"): out.append(f"Title: {md['title']}")
        for i, page in enumerate(doc):
            t = page.get_text("text").strip()
            if len(t) < 25 and ocr_pages < 15:          # scanned page → OCR it
                pix = page.get_pixmap(dpi=150)
                t = ocr_image(Image.frombytes("RGB", (pix.width, pix.height), pix.samples)).strip()
                ocr_pages += 1
                if t: t = "(OCR) " + t
            out.append(f"[Page {i+1}]\n{t or '(no text)'}")
            try:
                ims = page_layout(page)["images"]
                if ims:
                    out.append(f"[Page {i+1} images: " + "; ".join(
                        f"#{m['index']} at ({m['bbox'][0]:.0f},{m['bbox'][1]:.0f})-({m['bbox'][2]:.0f},{m['bbox'][3]:.0f})" for m in ims[:12]) + "]")
            except Exception:
                pass
        try:                                             # fillable form → list its fields for pdf_fill_form
            fields = [w for pg in doc for w in (pg.widgets() or [])]
            if fields:
                out.append("[Form fields (use pdf_fill_form): " + "; ".join(
                    f"{w.field_name!r} ({w.field_type_string}{', options: ' + '/'.join(map(str, (w.choice_values or [])[:8])) if w.choice_values else ''})"
                    f" = {w.field_value!r}" for w in fields[:60]) + "]")
        except Exception:
            pass
    return "\n\n".join(out)


def _extract_docx(p):
    import docx
    d = docx.Document(p)
    out = [para.text for para in d.paragraphs if para.text.strip()]
    for ti, t in enumerate(d.tables, 1):
        out.append(f"\n[Table {ti}]")
        for row in t.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)


def _extract_pptx(p):
    from pptx import Presentation
    out = []
    for i, s in enumerate(Presentation(p).slides, 1):
        texts = [sh.text_frame.text for sh in s.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
        notes = ""
        if s.has_notes_slide:
            notes = s.notes_slide.notes_text_frame.text.strip()
        out.append(f"[Slide {i}]\n" + "\n".join(texts) + (f"\nSpeaker notes: {notes}" if notes else ""))
    return "\n\n".join(out)


def _extract_sheet(p):
    import pandas as pd
    ext = os.path.splitext(p)[1].lower()
    sheets = {"data": pd.read_csv(p, sep=None, engine="python")} if ext in (".csv", ".tsv") else pd.read_excel(p, sheet_name=None)
    out = []
    for name, df in sheets.items():
        out.append(f"[Sheet: {name}] {df.shape[0]} rows × {df.shape[1]} columns\nColumns: {', '.join(map(str, df.columns))}")
        out.append(df.head(300).to_csv(index=False))
        num = df.select_dtypes("number")
        if not num.empty:
            out.append("Summary statistics:\n" + num.describe().round(3).to_csv())
    return "\n".join(out)


def _extract_zip(p):
    with zipfile.ZipFile(p) as z:
        names = z.namelist()
    return f"ZIP archive with {len(names)} entries:\n" + "\n".join(names[:300])


def _extract_text(p):
    data = open(p, "rb").read()
    if b"\x00" in data[:4096] and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return f"(Binary file, {len(data)} bytes — not readable as text.)"
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except Exception:
            continue
    return ""


def extract(rec):
    k, p = rec["kind"], rec["path"]
    try:
        if k == "pdf": t = _extract_pdf(rec)
        elif k == "image": t = analyze_image(p)
        elif k == "docx": t = _extract_docx(p)
        elif k == "pptx": t = _extract_pptx(p)
        elif k == "sheet": t = _extract_sheet(p)
        elif k == "zip": t = _extract_zip(p)
        else: t = _extract_text(p)
    except Exception as e:
        t = f"(Could not read this file: {e})"
    rec["text"] = (t or "").strip()[:MAX_FILE_CHARS]
    return rec["text"]


def ingest(name, data):
    fid, path = new_path(name)
    with open(path, "wb") as f:
        f.write(data)
    rec = register(fid, path, name, origin="upload")
    extract(rec)
    return rec


# ==========================================================
# FILE TOOLS (used by the Tools panel AND by the model)
# ==========================================================
TOOLS = {}


def P(name, label, type="text", default="", options=None, hint=""):
    d = dict(name=name, label=label, type=type, default=default, hint=hint)
    if options: d["options"] = options
    return d


def tool(op, group, label, desc, accepts=(), multi=False, params=()):
    def deco(fn):
        TOOLS[op] = dict(op=op, group=group, label=label, desc=desc, accepts=list(accepts),
                         multi=multi, params=list(params), fn=fn)
        return fn
    return deco


def _num(p, k, d):
    try: return float(p.get(k, d))
    except Exception: return float(d)


def _s(p, k, d=""):
    v = p.get(k, d)
    return str(d if v is None else v)


def _rgb(h, d=(0.5, 0.5, 0.5)):
    h = str(h or "").strip().lstrip("#")
    if len(h) == 3: h = "".join(c * 2 for c in h)
    try: return tuple(int(h[i:i+2], 16) / 255 for i in (0, 2, 4))
    except Exception: return d


def _stem(rec): return os.path.splitext(rec["name"])[0]


def parse_pages(spec, n):
    spec = str(spec or "").replace(" ", "").lower()
    if spec in ("", "all"): return list(range(n))
    res = []
    for part in spec.split(","):
        if not part: continue
        if "-" in part:
            a, b = part.split("-", 1)
            a = int(a) if a else 1
            b = int(b) if b else n
            res += list(range(a - 1, b))
        else:
            res.append(int(part) - 1)
    res = [x for x in res if 0 <= x < n]
    if not res: raise ValueError(f"No valid pages in '{spec}' (document has {n} pages).")
    return res


def _open_pdf(rec, pw=""):
    d = fitz.open(rec["path"])
    if d.needs_pass and not d.authenticate(pw):
        raise ValueError(f"{rec['name']} is password-protected — use 'Unlock PDF' first.")
    return d


def _save_pdf(doc, name, **kw):
    fid, path = new_path(name)
    opts = dict(garbage=3, deflate=True); opts.update(kw)
    doc.save(path, **opts); doc.close()
    return register(fid, path, name, origin="tool")


def _img_as_pdf(path):
    im = ImageOps.exif_transpose(Image.open(path))
    im = _flatten(im)
    buf = io.BytesIO(); im.save(buf, "JPEG", quality=92)
    with fitz.open(stream=buf.getvalue(), filetype="jpg") as idoc:
        return fitz.open("pdf", idoc.convert_to_pdf())


def _flatten(im):
    if im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA")
        bg = Image.new("RGB", im.size, (255, 255, 255)); bg.paste(im, mask=im.getchannel("A"))
        return bg
    return im.convert("RGB") if im.mode != "RGB" else im


PDF, IMG = ["pdf"], ["image"]

# ---------------- PDF ----------------
@tool("pdf_merge", "PDF", "Merge PDFs", "Combine several PDFs (and images) into one PDF, in the order given.", ["pdf", "image"], multi=True)
def t_merge(recs, p):
    out = fitz.open()
    for r in recs:
        src = _img_as_pdf(r["path"]) if r["kind"] == "image" else _open_pdf(r)
        out.insert_pdf(src); src.close()
    n = out.page_count
    return _save_pdf(out, "merged.pdf"), f"Merged {len(recs)} files · {n} pages"


@tool("pdf_extract_pages", "PDF", "Extract pages", "Keep only the chosen pages, e.g. 1-3,7.", PDF,
      params=[P("pages", "Pages", default="1", hint="e.g. 1-3,7,10-")])
def t_extract(recs, p):
    d = _open_pdf(recs[0]); keep = parse_pages(p.get("pages"), d.page_count); d.select(keep)
    return _save_pdf(d, f"{_stem(recs[0])}_pages.pdf"), f"Extracted {len(keep)} pages"


@tool("pdf_delete_pages", "PDF", "Delete pages", "Remove the chosen pages.", PDF,
      params=[P("pages", "Pages to delete", default="1", hint="e.g. 2,5-6")])
def t_delete(recs, p):
    d = _open_pdf(recs[0]); n = d.page_count; rm = sorted(set(parse_pages(p.get("pages"), n)))
    if len(rm) >= n: raise ValueError("That would delete every page.")
    d.delete_pages(rm)
    return _save_pdf(d, f"{_stem(recs[0])}_edited.pdf"), f"Deleted {len(rm)} pages · {n-len(rm)} left"


@tool("pdf_reorder", "PDF", "Reorder pages", "Rearrange pages into a new order (pages may repeat).", PDF,
      params=[P("order", "New order", default="", hint="e.g. 3,1,2,4-")])
def t_reorder(recs, p):
    d = _open_pdf(recs[0]); order = parse_pages(p.get("order"), d.page_count); d.select(order)
    return _save_pdf(d, f"{_stem(recs[0])}_reordered.pdf"), f"Reordered → {len(order)} pages"


@tool("pdf_rotate", "PDF", "Rotate pages", "Rotate pages clockwise by 90/180/270°.", PDF,
      params=[P("pages", "Pages", default="all"), P("angle", "Angle", "select", "90", ["90", "180", "270"])])
def t_rotate(recs, p):
    d = _open_pdf(recs[0]); ang = int(_num(p, "angle", 90)) % 360
    pages = parse_pages(p.get("pages"), d.page_count)
    for i in pages: d[i].set_rotation((d[i].rotation + ang) % 360)
    return _save_pdf(d, f"{_stem(recs[0])}_rotated.pdf"), f"Rotated {len(pages)} pages by {ang}°"


@tool("pdf_watermark", "PDF", "Watermark", "Stamp diagonal text across every page.", PDF,
      params=[P("text", "Text", default="CONFIDENTIAL"), P("opacity", "Opacity (0-1)", "number", "0.18"),
              P("size", "Font size", "number", "60"), P("color", "Color", default="#7c3aed")])
def t_watermark(recs, p):
    d = _open_pdf(recs[0]); text = _s(p, "text", "CONFIDENTIAL")
    fs, op, col = _num(p, "size", 60), max(0.02, min(1, _num(p, "opacity", 0.18))), _rgb(p.get("color"), (0.5, 0.5, 0.5))
    tw = fitz.get_text_length(text, fontname="helv", fontsize=fs)
    for page in d:
        r = page.rect; c = fitz.Point(r.x0 + r.width / 2, r.y0 + r.height / 2)
        page.insert_text(fitz.Point(c.x - tw / 2, c.y + fs / 3), text, fontsize=fs, fontname="helv",
                         color=col, fill_opacity=op, morph=(c, fitz.Matrix(45)), overlay=True)
    n = d.page_count
    return _save_pdf(d, f"{_stem(recs[0])}_watermarked.pdf"), f"Watermarked {n} pages"


@tool("pdf_add_text", "PDF", "Add text", "Write text onto a page at x,y (points from top-left; A4 ≈ 595×842).", PDF,
      params=[P("page", "Page", "number", "1"), P("text", "Text", "textarea", "", hint="Text to add"),
              P("x", "X", "number", "72"), P("y", "Y", "number", "72"), P("size", "Font size", "number", "12"),
              P("color", "Color", default="#000000")])
def t_add_text(recs, p):
    d = _open_pdf(recs[0]); i = int(_num(p, "page", 1)) - 1
    if not 0 <= i < d.page_count: raise ValueError(f"Page must be 1–{d.page_count}")
    d[i].insert_text(fitz.Point(_num(p, "x", 72), _num(p, "y", 72)), _s(p, "text"), fontsize=_num(p, "size", 12),
                     fontname="helv", color=_rgb(p.get("color"), (0, 0, 0)))
    return _save_pdf(d, f"{_stem(recs[0])}_edited.pdf"), f"Added text on page {i+1}"


@tool("pdf_page_numbers", "PDF", "Page numbers", "Add 'Page X of N' to every page.", PDF,
      params=[P("position", "Position", "select", "bottom-center", ["bottom-center", "bottom-right", "top-right"]),
              P("size", "Font size", "number", "10")])
def t_pagenum(recs, p):
    d = _open_pdf(recs[0]); n = d.page_count; fs = _num(p, "size", 10); pos = _s(p, "position", "bottom-center")
    for i, page in enumerate(d):
        r = page.rect; label = f"Page {i+1} of {n}"; w = fitz.get_text_length(label, "helv", fs)
        x = r.width - w - 36 if "right" in pos else (r.width - w) / 2
        y = 30 if pos.startswith("top") else r.height - 24
        page.insert_text(fitz.Point(x, y), label, fontsize=fs, fontname="helv", color=(0.35, 0.35, 0.35))
    return _save_pdf(d, f"{_stem(recs[0])}_numbered.pdf"), f"Numbered {n} pages"


def _find_hits(page, find):
    return page.search_for(find) if find else []


def _span_style(page, rect):
    """Font, size, colour and baseline of the original text under `rect`."""
    best, area = None, 0
    for b in page.get_text("dict", clip=rect + (-2, -2, 2, 2))["blocks"]:
        for l in b.get("lines", []):
            for sp in l["spans"]:
                if not sp["text"].strip(): continue
                ov = fitz.Rect(sp["bbox"]) & rect
                a = ov.get_area() if not ov.is_empty else 0
                if a > area: best, area = sp, a
    return best


def _base14(font, flags):
    f = (font or "").lower()
    bold = bool(flags & 16) or any(k in f for k in ("bold", "black", "heavy", "semibold"))
    ital = bool(flags & 2) or "italic" in f or "oblique" in f
    if (flags & 8) or any(k in f for k in ("courier", "mono", "consol")):
        return {(0, 0): "cour", (1, 0): "cobo", (0, 1): "coit", (1, 1): "cobi"}[(bold, ital)]
    if (flags & 4) or any(k in f for k in ("times", "serif", "roman", "georgia", "garamond", "cambria", "book")) and "sans" not in f:
        return {(0, 0): "tiro", (1, 0): "tibo", (0, 1): "tiit", (1, 1): "tibi"}[(bold, ital)]
    return {(0, 0): "helv", (1, 0): "hebo", (0, 1): "heit", (1, 1): "hebi"}[(bold, ital)]


def _embedded_font(doc, page, span_font, text, cache):
    """Reuse the PDF's own embedded font when it contains every needed glyph."""
    norm = lambda n: re.sub(r"[\s_\-,]", "", (n or "").split("+")[-1]).lower()
    want = norm(span_font)
    for xref, ext, ftype, basefont, *_ in page.get_fonts():
        if norm(basefont) != want or ext in ("n/a", ""): continue
        key = (xref, text)
        if key in cache: return cache[key]
        try:
            buf = doc.extract_font(xref)[-1]
            f = fitz.Font(fontbuffer=buf)
            ok = buf and all(f.has_glyph(ord(c)) for c in text if not c.isspace())
            cache[key] = (f"MW{xref}", buf) if ok else None
        except Exception:
            cache[key] = None
        return cache[key]
    return None


@tool("pdf_replace_text", "PDF", "Find & replace text", "Replace a phrase everywhere, keeping the original font, size and colour (text PDFs, not scans).", PDF,
      params=[P("find", "Find", default=""), P("replace", "Replace with", default="")])
def t_replace(recs, p):
    d = _open_pdf(recs[0]); find, rep = _s(p, "find"), _s(p, "replace")
    if not find: raise ValueError("Tell me what text to find.")
    total, matched_fonts, cache = 0, set(), {}
    for page in d:
        hits = _find_hits(page, find)
        if not hits: continue
        jobs = []
        for r in hits:
            sp = _span_style(page, r) or {}
            st = _style_of(sp) if sp else dict(font=None, size=r.height * 0.8, color="#000000", family="sans", bold=False, italic=False)
            origin = fitz.Point(r.x0, sp["origin"][1]) if sp.get("origin") else fitz.Point(r.x0, r.y1 - st["size"] * 0.22)
            jobs.append((origin, st))
            matched_fonts.add((sp.get("font") or "?").split("+")[-1])
            page.add_redact_annot(r, fill=False)
        _apply_redactions(page, **_redact_kw())
        for origin, st in jobs:
            write_run(d, page, origin, rep, st, cache, keep_font=True)
        total += len(hits)
    if not total:                                   # scanned / image-only PDF → edit the text inside the images
        n, fonts = scan_replace_doc(d, find, rep)
        if not n: d.close(); raise ValueError(f"'{find}' wasn't found in the PDF (text or images).")
        return _save_pdf(d, f"{_stem(recs[0])}_replaced.pdf"), f"Replaced {n} line(s) inside page images · font ≈ {fonts or 'matched'}"
    return _save_pdf(d, f"{_stem(recs[0])}_replaced.pdf"), f"Replaced {total} occurrence(s) · matched font: {', '.join(sorted(matched_fonts))}"


@tool("pdf_redact", "PDF", "Redact text", "Black out every occurrence of a phrase (removes it permanently).", PDF,
      params=[P("find", "Text to redact", default="")])
def t_redact(recs, p):
    d = _open_pdf(recs[0]); find = _s(p, "find"); total = 0
    if not find: raise ValueError("Tell me what text to redact.")
    for page in d:
        hits = _find_hits(page, find)
        for r in hits: page.add_redact_annot(r, fill=(0, 0, 0))
        if hits: page.apply_redactions(); total += len(hits)
    if not total: d.close(); raise ValueError(f"'{find}' wasn't found.")
    return _save_pdf(d, f"{_stem(recs[0])}_redacted.pdf"), f"Redacted {total} occurrence(s)"


@tool("pdf_highlight", "PDF", "Highlight text", "Highlight every occurrence of a phrase.", PDF,
      params=[P("find", "Text to highlight", default="")])
def t_highlight(recs, p):
    d = _open_pdf(recs[0]); find = _s(p, "find"); total = 0
    for page in d:
        hits = _find_hits(page, find)
        if hits: page.add_highlight_annot(hits); total += len(hits)
    if not total: d.close(); raise ValueError(f"'{find}' wasn't found.")
    return _save_pdf(d, f"{_stem(recs[0])}_highlighted.pdf"), f"Highlighted {total} occurrence(s)"


@tool("pdf_compress", "PDF", "Compress PDF", "Shrink file size (downsamples big images).", PDF,
      params=[P("quality", "Image quality", "select", "medium", ["high", "medium", "low"])])
def t_compress(recs, p):
    d = _open_pdf(recs[0]); q = _s(p, "quality", "medium")
    dpi, jq = {"high": (150, 80), "medium": (110, 65), "low": (80, 45)}.get(q, (110, 65))
    try: d.rewrite_images(dpi_threshold=dpi + 30, dpi_target=dpi, quality=jq)
    except Exception: pass
    rec = _save_pdf(d, f"{_stem(recs[0])}_compressed.pdf", garbage=4, clean=True, deflate_images=True, deflate_fonts=True)
    before, after = recs[0]["size"], rec["size"]
    return rec, f"{before/1e6:.2f} MB → {after/1e6:.2f} MB ({100 - after*100//max(before,1)}% smaller)"


@tool("pdf_to_images", "PDF", "PDF → images", "Render pages as PNG images (zip for multiple pages).", PDF,
      params=[P("pages", "Pages", default="all"), P("dpi", "DPI", "number", "150")])
def t_pdf2img(recs, p):
    d = _open_pdf(recs[0]); pages = parse_pages(p.get("pages"), d.page_count); dpi = int(max(50, min(400, _num(p, "dpi", 150))))
    stem = _stem(recs[0])
    if len(pages) == 1:
        fid, path = new_path(f"{stem}_p{pages[0]+1}.png"); d[pages[0]].get_pixmap(dpi=dpi).save(path); d.close()
        return register(fid, path, f"{stem}_p{pages[0]+1}.png", "tool"), "Rendered 1 page"
    fid, path = new_path(f"{stem}_images.zip")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for i in pages: z.writestr(f"{stem}_p{i+1}.png", d[i].get_pixmap(dpi=dpi).tobytes("png"))
    d.close()
    return register(fid, path, f"{stem}_images.zip", "tool"), f"Rendered {len(pages)} pages"


@tool("pdf_to_text", "PDF", "PDF → text", "Export all text as a .txt file.", PDF)
def t_pdf2txt(recs, p):
    r = recs[0]; txt = r["text"] if r.get("text") else extract(r)
    fid, path = new_path(f"{_stem(r)}.txt"); open(path, "w", encoding="utf-8").write(txt)
    return register(fid, path, f"{_stem(r)}.txt", "tool"), f"{len(txt):,} characters"


@tool("pdf_protect", "PDF", "Password-protect", "Encrypt the PDF with a password (AES-256).", PDF,
      params=[P("password", "Password", "password", "")])
def t_protect(recs, p):
    pw = _s(p, "password")
    if not pw: raise ValueError("Enter a password.")
    d = _open_pdf(recs[0])
    return _save_pdf(d, f"{_stem(recs[0])}_protected.pdf", encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw=pw, user_pw=pw), "Encrypted with AES-256"


@tool("pdf_unlock", "PDF", "Unlock PDF", "Remove the password from a PDF you know the password to.", PDF,
      params=[P("password", "Current password", "password", "")])
def t_unlock(recs, p):
    d = fitz.open(recs[0]["path"])
    if d.needs_pass and not d.authenticate(_s(p, "password")): d.close(); raise ValueError("Wrong password.")
    rec = _save_pdf(d, f"{_stem(recs[0])}_unlocked.pdf", encryption=fitz.PDF_ENCRYPT_NONE)
    extract(rec)
    return rec, "Password removed"


@tool("images_to_pdf", "PDF", "Images → PDF", "Turn images into a PDF, one per page, in order.", IMG, multi=True)
def t_img2pdf(recs, p):
    out = fitz.open()
    for r in recs:
        src = _img_as_pdf(r["path"]); out.insert_pdf(src); src.close()
    n = out.page_count
    return _save_pdf(out, "images.pdf"), f"{n} pages"


# ---------------- IMAGES ----------------
EXT_FMT = {"jpg": "JPEG", "jpeg": "JPEG", "png": "PNG", "webp": "WEBP", "bmp": "BMP", "gif": "GIF", "tiff": "TIFF"}


def _load_img(rec):
    im = Image.open(rec["path"]); im = ImageOps.exif_transpose(im); im.load(); return im


def _src_ext(rec):
    e = os.path.splitext(rec["name"])[1].lstrip(".").lower()
    return e if e in EXT_FMT else ("jpg" if e in ("heic", "heif") else "png")


def _save_img(im, stem, ext, quality=92):
    ext = ext.lower().replace("jpeg", "jpg")
    if ext not in EXT_FMT: ext = "png"
    fmt = EXT_FMT[ext]; name = f"{stem}.{ext}"
    if fmt in ("JPEG", "BMP"): im = _flatten(im)
    elif im.mode not in ("RGB", "RGBA", "L", "LA", "P"): im = im.convert("RGBA")
    fid, path = new_path(name)
    kw = {"quality": int(quality), "optimize": True} if fmt in ("JPEG", "WEBP") else ({"optimize": True} if fmt == "PNG" else {})
    im.save(path, fmt, **kw)
    return register(fid, path, name, origin="tool")


@tool("img_resize", "Image", "Resize", "Resize to a width/height (leave one blank to keep proportions).", IMG,
      params=[P("width", "Width px", "number", "1080"), P("height", "Height px", "number", "")])
def t_resize(recs, p):
    im = _load_img(recs[0]); w = p.get("width") or ""; h = p.get("height") or ""
    w = int(float(w)) if str(w).strip() else 0; h = int(float(h)) if str(h).strip() else 0
    if not w and not h: raise ValueError("Give a width or a height.")
    if not h: h = round(im.height * w / im.width)
    if not w: w = round(im.width * h / im.height)
    out = im.resize((w, h), Image.LANCZOS)
    return _save_img(out, f"{_stem(recs[0])}_{w}x{h}", _src_ext(recs[0])), f"{im.width}×{im.height} → {w}×{h}"


@tool("img_crop", "Image", "Crop", "Crop to a box in pixels (left, top, right, bottom).", IMG,
      params=[P("left", "Left", "number", "0"), P("top", "Top", "number", "0"), P("right", "Right", "number", ""), P("bottom", "Bottom", "number", "")])
def t_crop(recs, p):
    im = _load_img(recs[0])
    box = (int(_num(p, "left", 0)), int(_num(p, "top", 0)),
           int(_num(p, "right", im.width) or im.width), int(_num(p, "bottom", im.height) or im.height))
    out = im.crop(box)
    return _save_img(out, f"{_stem(recs[0])}_crop", _src_ext(recs[0])), f"Cropped to {out.width}×{out.height}"


@tool("img_rotate", "Image", "Rotate", "Rotate clockwise by any angle.", IMG, params=[P("angle", "Degrees", "number", "90")])
def t_irotate(recs, p):
    a = _num(p, "angle", 90); im = _load_img(recs[0]).rotate(-a, expand=True, resample=Image.BICUBIC)
    return _save_img(im, f"{_stem(recs[0])}_rot", _src_ext(recs[0])), f"Rotated {a:g}°"


@tool("img_flip", "Image", "Flip / mirror", "Mirror horizontally or vertically.", IMG,
      params=[P("direction", "Direction", "select", "horizontal", ["horizontal", "vertical"])])
def t_flip(recs, p):
    im = _load_img(recs[0])
    im = ImageOps.mirror(im) if _s(p, "direction", "horizontal") == "horizontal" else ImageOps.flip(im)
    return _save_img(im, f"{_stem(recs[0])}_flip", _src_ext(recs[0])), "Flipped"


@tool("img_convert", "Image", "Convert format", "Convert to PNG, JPG, WEBP or PDF (great for iPhone HEIC).", IMG,
      params=[P("format", "Format", "select", "png", ["png", "jpg", "webp", "pdf"])])
def t_convert(recs, p):
    fmt = _s(p, "format", "png").lower()
    if fmt == "pdf":
        src = _img_as_pdf(recs[0]["path"]); return _save_pdf(src, f"{_stem(recs[0])}.pdf"), "Converted to PDF"
    return _save_img(_load_img(recs[0]), _stem(recs[0]), fmt), f"Converted to {fmt.upper()}"


@tool("img_compress", "Image", "Compress image", "Reduce file size with adjustable quality.", IMG,
      params=[P("quality", "Quality (10-95)", "number", "70"), P("format", "Format", "select", "jpg", ["jpg", "webp"]),
              P("max_side", "Max side px (optional)", "number", "")])
def t_icompress(recs, p):
    im = _load_img(recs[0]); ms = p.get("max_side")
    if str(ms or "").strip(): im.thumbnail((int(float(ms)), int(float(ms))), Image.LANCZOS)
    rec = _save_img(im, f"{_stem(recs[0])}_small", _s(p, "format", "jpg"), int(max(10, min(95, _num(p, "quality", 70)))))
    return rec, f"{recs[0]['size']/1e6:.2f} MB → {rec['size']/1e6:.2f} MB"


@tool("img_grayscale", "Image", "Black & white", "Convert to grayscale.", IMG)
def t_gray(recs, p):
    return _save_img(ImageOps.grayscale(_load_img(recs[0])), f"{_stem(recs[0])}_bw", _src_ext(recs[0])), "Converted to grayscale"


@tool("img_enhance", "Image", "Enhance", "Adjust brightness, contrast, sharpness, saturation (1 = unchanged).", IMG,
      params=[P("brightness", "Brightness", "number", "1.05"), P("contrast", "Contrast", "number", "1.1"),
              P("sharpness", "Sharpness", "number", "1.3"), P("saturation", "Saturation", "number", "1.1")])
def t_enhance(recs, p):
    im = _load_img(recs[0])
    base = im.convert("RGBA") if "A" in im.getbands() else im.convert("RGB")
    alpha = base.getchannel("A") if base.mode == "RGBA" else None
    rgb = base.convert("RGB")
    for k, E in (("brightness", ImageEnhance.Brightness), ("contrast", ImageEnhance.Contrast),
                 ("sharpness", ImageEnhance.Sharpness), ("saturation", ImageEnhance.Color)):
        rgb = E(rgb).enhance(max(0, _num(p, k, 1)))
    if alpha: rgb.putalpha(alpha)
    return _save_img(rgb, f"{_stem(recs[0])}_enhanced", _src_ext(recs[0])), "Enhanced"


@tool("img_upscale", "Image", "Upscale", "Enlarge 2× / 3× / 4× with sharpening.", IMG,
      params=[P("scale", "Scale", "select", "2", ["2", "3", "4"])])
def t_upscale(recs, p):
    im = _load_img(recs[0]); s = int(_num(p, "scale", 2))
    out = im.resize((im.width * s, im.height * s), Image.LANCZOS).filter(ImageFilter.UnsharpMask(radius=2, percent=80, threshold=2))
    return _save_img(out, f"{_stem(recs[0])}_x{s}", _src_ext(recs[0])), f"{im.width}×{im.height} → {out.width}×{out.height}"


@tool("img_remove_bg", "Image", "Remove background", "Cut out the subject to a transparent PNG (needs: pip install rembg).", IMG)
def t_rembg(recs, p):
    try:
        from rembg import remove
    except Exception:
        raise ValueError("Background removal needs one extra package — run: !pip install -q rembg onnxruntime")
    out = remove(_load_img(recs[0]))
    return _save_img(out, f"{_stem(recs[0])}_nobg", "png"), "Background removed"


# ---------------- CREATE ----------------
PDF_CSS = """
body { font-family: sans-serif; font-size: 11pt; line-height: 1.5; color: #1f2937; }
h1 { font-size: 22pt; color: #4c1d95; margin-bottom: 4pt; }
h2 { font-size: 15pt; color: #5b21b6; margin-top: 14pt; }
h3 { font-size: 12.5pt; color: #374151; margin-top: 10pt; }
p { margin: 5pt 0; }
li { margin: 2pt 0; }
code { font-family: monospace; background-color: #f3f4f6; }
pre { font-family: monospace; font-size: 9pt; background-color: #f3f4f6; padding: 6pt; }
table { border-collapse: collapse; }
th, td { border: 1px solid #d1d5db; padding: 4pt 6pt; }
th { background-color: #ede9fe; }
blockquote { color: #4b5563; margin-left: 12pt; }
"""


def _md_to_html(md_text):
    try:
        import markdown
        return markdown.markdown(md_text, extensions=["tables", "fenced_code", "sane_lists"])
    except Exception:
        return "<pre>" + html.escape(md_text) + "</pre>"


@tool("make_pdf", "Create", "Create PDF", "Write a nicely formatted PDF from Markdown content (reports, notes, letters).",
      params=[P("title", "Title", default="Meow Report"), P("content", "Content (Markdown)", "textarea", "", hint="# Heading\nYour text…")])
def t_make_pdf(recs, p):
    title, body = _s(p, "title", "Document"), _s(p, "content")
    if not body.strip(): raise ValueError("No content given.")
    doc_html = f"<h1>{html.escape(title)}</h1>" + _md_to_html(body) if title.strip() else _md_to_html(body)
    fname = safe_name((title.strip() or "document")[:60]) + ".pdf"
    fid, path = new_path(fname)
    story = fitz.Story(html=doc_html, user_css=PDF_CSS)
    writer = fitz.DocumentWriter(path)
    mediabox = fitz.paper_rect("a4"); where = mediabox + (56, 56, -56, -56)
    more = True
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    rec = register(fid, path, fname, "tool")
    return rec, f"{rec['meta'].get('pages', '?')} page PDF created"


@tool("make_file", "Create", "Create text file", "Save any text file — .md, .py, .csv, .html, .json, .txt …",
      params=[P("filename", "Filename", default="notes.md"), P("content", "Content", "textarea", "")])
def t_make_file(recs, p):
    name = safe_name(_s(p, "filename", "file.txt"))
    if "." not in name: name += ".txt"
    fid, path = new_path(name)
    open(path, "w", encoding="utf-8").write(_s(p, "content"))
    rec = register(fid, path, name, "tool"); rec["text"] = _s(p, "content")[:MAX_FILE_CHARS]
    return rec, f"{rec['size']:,} bytes"


def run_tool(op, file_ids, params):
    spec = TOOLS.get(op)
    if not spec: raise ValueError(f"Unknown tool '{op}'.")
    recs = [FILES[i] for i in (file_ids or []) if i in FILES]
    if spec["accepts"]:
        if not recs: raise ValueError("This tool needs a file — the file id wasn't found (was the server restarted?).")
        bad = [r["name"] for r in recs if r["kind"] not in spec["accepts"]]
        if bad: raise ValueError(f"{', '.join(bad)} can't be used with {spec['label']} (needs {'/'.join(spec['accepts'])}).")
        if not spec["multi"]: recs = recs[:1]
    rec, msg = spec["fn"](recs, params or {})
    return rec, msg


# ==========================================================
# PRO PDF EDITING ENGINE (visual editor + AI tools)
# ==========================================================
FAM_CSS = {"serif": "serif", "mono": "monospace", "sans": "sans-serif"}


def _latin1(t):
    try:
        t.encode("latin-1"); return True
    except UnicodeEncodeError:
        return False


def _hex(c):
    if isinstance(c, int):
        return f"#{c >> 16 & 255:02x}{c >> 8 & 255:02x}{c & 255:02x}"
    return "#%02x%02x%02x" % tuple(int(max(0, min(1, v)) * 255) for v in c[:3])


def _family_of(font, flags):
    f = (font or "").lower()
    if (flags & 8) or any(k in f for k in ("courier", "mono", "consol")): return "mono"
    if "sans" in f or "arial" in f or "helvet" in f or "calibri" in f or "verdana" in f: return "sans"
    if (flags & 4) or any(k in f for k in ("times", "serif", "roman", "georgia", "garamond", "cambria", "book", "minion")): return "serif"
    return "sans"


def _style_of(sp):
    font, flags = sp.get("font", ""), sp.get("flags", 0)
    lf = font.lower()
    return dict(font=font, size=sp.get("size", 11), color=_hex(sp.get("color", 0)), family=_family_of(font, flags),
                bold=bool(flags & 16) or any(k in lf for k in ("bold", "black", "heavy", "semibold")),
                italic=bool(flags & 2) or "italic" in lf or "oblique" in lf)


def _b14(family, bold, italic):
    t = {"mono": ("cour", "cobo", "coit", "cobi"), "serif": ("tiro", "tibo", "tiit", "tibi"),
         "sans": ("helv", "hebo", "heit", "hebi")}[family if family in ("mono", "serif") else "sans"]
    return t[(1 if bold else 0) + (2 if italic else 0)]


def write_run(doc, page, origin, text, style, cache, keep_font=True):
    """Write text at a baseline origin in the given style. Reuses the PDF's own
    embedded font when possible; falls back to standard fonts, and to HTML
    rendering (full Unicode: Hindi, ₹, emoji-free scripts) when needed."""
    size = float(style.get("size") or 11)
    color = _rgb(style.get("color"), (0, 0, 0))
    lines = str(text).split("\n")
    lh = size * 1.2
    if all(_latin1(l) for l in lines):
        emb = _embedded_font(doc, page, style.get("font"), text, cache) if keep_font and style.get("font") else None
        if emb:
            page.insert_font(fontname=emb[0], fontbuffer=emb[1]); fname = emb[0]
        else:
            fname = _b14(style.get("family"), style.get("bold"), style.get("italic"))
        for i, l in enumerate(lines):
            if l:
                page.insert_text(fitz.Point(origin.x, origin.y + i * lh), l, fontname=fname, fontsize=size, color=color)
        return
    css = (f"* {{font-family: {FAM_CSS.get(style.get('family'), 'sans-serif')}; font-size: {size}px; line-height: 1.2;"
           f" color: {style.get('color') or '#000'}; font-weight: {'bold' if style.get('bold') else 'normal'};"
           f" font-style: {'italic' if style.get('italic') else 'normal'}; margin: 0; padding: 0; }}")
    top = origin.y - size * 0.95
    rect = fitz.Rect(origin.x, top, max(origin.x + 40, page.rect.x1 - 8), top + lh * len(lines) + size)
    page.insert_htmlbox(rect, html.escape(str(text)).replace("\n", "<br>"), css=css)


def _redact_kw(images="none", graphics="none", text=True):
    kw = {"images": {"none": getattr(fitz, "PDF_REDACT_IMAGE_NONE", 0), "pixels": getattr(fitz, "PDF_REDACT_IMAGE_PIXELS", 2),
                     "remove": getattr(fitz, "PDF_REDACT_IMAGE_REMOVE", 1)}[images]}
    if hasattr(fitz, "PDF_REDACT_LINE_ART_NONE"):
        kw["graphics"] = {"none": fitz.PDF_REDACT_LINE_ART_NONE,
                          "covered": getattr(fitz, "PDF_REDACT_LINE_ART_REMOVE_IF_COVERED", 1)}[graphics]
    if not text and hasattr(fitz, "PDF_REDACT_TEXT_NONE"):
        kw["text"] = fitz.PDF_REDACT_TEXT_NONE
    return kw


def _apply_redactions(page, **kw):
    for drop in (None, "text", "graphics"):
        if drop: kw.pop(drop, None)
        try:
            return page.apply_redactions(**kw)
        except TypeError:
            continue
    page.apply_redactions()


def page_layout(page):
    """Editable text runs (merged spans of identical style) + placed images."""
    runs = []
    try:
        flags = fitz.TEXTFLAGS_DICT & ~fitz.TEXT_PRESERVE_IMAGES
    except Exception:
        flags = 0
    data = page.get_text("dict", flags=flags) if flags else page.get_text("dict")
    for b in data["blocks"]:
        if b.get("type") != 0: continue
        for l in b["lines"]:
            d = l.get("dir", (1, 0))
            if abs(d[0] - 1) > 0.01 or abs(d[1]) > 0.01: continue        # skip rotated text
            cur = None
            for sp in l["spans"]:
                t = sp["text"]
                st = _style_of(sp)
                key = (sp["font"], round(sp["size"], 1), sp["color"], st["bold"], st["italic"])
                if not t.strip():
                    if cur: cur["text"] += t
                    continue
                if cur and cur["key"] == key and abs(sp["origin"][1] - cur["origin"][1]) < 1 and sp["bbox"][0] - cur["bbox"][2] < sp["size"] * 0.9:
                    if sp["bbox"][0] - cur["bbox"][2] > sp["size"] * 0.15 and not cur["text"].endswith(" "): cur["text"] += " "
                    cur["text"] += t
                    cur["bbox"] = [min(cur["bbox"][0], sp["bbox"][0]), min(cur["bbox"][1], sp["bbox"][1]),
                                   max(cur["bbox"][2], sp["bbox"][2]), max(cur["bbox"][3], sp["bbox"][3])]
                else:
                    if cur: runs.append(cur)
                    cur = dict(key=key, text=t, bbox=list(sp["bbox"]), origin=list(sp["origin"]), **st)
            if cur: runs.append(cur)
    out = []
    for i, r in enumerate(runs):
        r.pop("key"); r["text"] = r["text"].rstrip()
        if r["text"]:
            r["id"] = i; r["bbox"] = [round(v, 2) for v in r["bbox"]]; out.append(r)
    imgs = []
    pr = page.rect
    for info in page.get_image_info(xrefs=True):
        bb = fitz.Rect(info["bbox"]) & pr
        if bb.is_empty or bb.width < 4 or bb.height < 4 or (info.get("width") or 9) <= 2: continue
        imgs.append(dict(xref=info.get("xref", 0), bbox=[round(v, 2) for v in bb], w=info.get("width"), h=info.get("height")))
    imgs.sort(key=lambda m: (round(m["bbox"][1] / 10), m["bbox"][0]))
    for i, m in enumerate(imgs, 1): m["index"] = i
    return {"runs": out, "images": imgs, "w": pr.width, "h": pr.height}


def _xref_uses(doc, xref):
    return sum(1 for pg in doc for i in pg.get_image_info(xrefs=True) if i.get("xref") == xref)


def _image_pil(doc, xref):
    base = fitz.Pixmap(doc, xref)
    if base.n - base.alpha >= 4:
        base = fitz.Pixmap(fitz.csRGB, base)
    try:
        smask = doc.extract_image(xref).get("smask", 0)
        if smask:
            base = fitz.Pixmap(base, fitz.Pixmap(doc, smask))
    except Exception:
        pass
    im = Image.open(io.BytesIO(base.tobytes("png"))); im.load()
    return im


def _remove_image(doc, page, xref, bbox):
    if xref and _xref_uses(doc, xref) == 1 and hasattr(page, "delete_image"):
        page.delete_image(xref)
    else:
        page.add_redact_annot(fitz.Rect(bbox), fill=False)
        _apply_redactions(page, **_redact_kw(images="pixels", text=False))


def _insert_pil(page, rect, im, keep=True):
    buf = io.BytesIO()
    if "A" in im.getbands() or im.mode == "P":
        im.convert("RGBA").save(buf, "PNG", optimize=True)
    else:
        im.convert("RGB").save(buf, "JPEG", quality=92)
    page.insert_image(fitz.Rect(rect), stream=buf.getvalue(), keep_proportion=keep, overlay=True)


IMG_ACTIONS = ["grayscale", "brighter", "darker", "contrast", "sharpen", "blur", "rotate", "flip_h", "flip_v",
               "invert", "auto_enhance", "sepia", "remove_bg"]


def filter_image(im, action, amount=None):
    a = action
    rgba = im.convert("RGBA"); alpha = rgba.getchannel("A"); rgb = rgba.convert("RGB")
    amt = float(amount) if amount not in (None, "") else None
    if a == "grayscale": rgb = ImageOps.grayscale(rgb).convert("RGB")
    elif a == "brighter": rgb = ImageEnhance.Brightness(rgb).enhance(amt or 1.2)
    elif a == "darker": rgb = ImageEnhance.Brightness(rgb).enhance(amt or 0.8)
    elif a == "contrast": rgb = ImageEnhance.Contrast(rgb).enhance(amt or 1.25)
    elif a == "sharpen": rgb = rgb.filter(ImageFilter.UnsharpMask(radius=2, percent=int((amt or 1.5) * 100), threshold=2))
    elif a == "blur": rgb = rgb.filter(ImageFilter.GaussianBlur(amt or 2))
    elif a == "invert": rgb = ImageOps.invert(rgb)
    elif a == "auto_enhance":
        rgb = ImageOps.autocontrast(rgb, cutoff=1); rgb = ImageEnhance.Color(rgb).enhance(1.1)
        rgb = rgb.filter(ImageFilter.UnsharpMask(radius=1.5, percent=60, threshold=2))
    elif a == "sepia":
        g = ImageOps.grayscale(rgb); rgb = ImageOps.colorize(g, "#2e1f0f", "#f3e3c3")
    elif a == "remove_bg":
        try:
            from rembg import remove
        except Exception:
            raise ValueError("Remove background needs: !pip install -q rembg onnxruntime")
        return remove(im.convert("RGBA"))
    elif a in ("rotate", "flip_h", "flip_v"):
        rgba = Image.merge("RGBA", (*rgb.split(), alpha))
        if a == "rotate": rgba = rgba.rotate(-(amt or 90), expand=True)
        elif a == "flip_h": rgba = ImageOps.mirror(rgba)
        else: rgba = ImageOps.flip(rgba)
        return rgba
    else:
        raise ValueError(f"Unknown image action '{a}'. Try: {', '.join(IMG_ACTIONS)}")
    rgb.putalpha(alpha)
    return rgb


def _rect_of(v):
    r = fitz.Rect(*[float(x) for x in v])
    r.normalize()
    return r


def _resolve_image(page, xref, bbox):
    """Find the image actually placed at bbox (xrefs can change between saves)."""
    best, score = None, 0
    for info in page.get_image_info(xrefs=True):
        r = fitz.Rect(info["bbox"]); ov = r & bbox
        if ov.is_empty: continue
        sc = ov.get_area() / max(1e-6, (r | bbox).get_area()) + (0.5 if info.get("xref") == xref else 0)
        if sc > score: best, score = info, sc
    if not best: raise ValueError("Couldn't find that image on the page any more — reload and try again.")
    return best.get("xref", 0), fitz.Rect(best["bbox"]) & page.rect


def apply_ops(doc, ops):
    """Run editor operations on an open document. Each op is a dict with 'op'."""
    cache = {}
    notes = []
    for o in ops:
        kind = o.get("op")
        pno = int(o.get("page", 0))
        if kind not in ("page_insert",) and not 0 <= pno < doc.page_count:
            raise ValueError(f"Page {pno + 1} doesn't exist.")
        page = doc[pno] if 0 <= pno < doc.page_count else None

        if kind == "edit_text":
            bbox = _rect_of(o["bbox"])
            sp = _span_style(page, bbox) or {}
            base = _style_of(sp) if sp else dict(font=None, size=bbox.height * 0.8, color="#000000", family="sans", bold=False, italic=False)
            style = {**base, **{k: o[k] for k in ("size", "color", "family", "bold", "italic") if o.get(k) not in (None, "")}}
            keep = all(style.get(k) == base.get(k) for k in ("family", "bold", "italic"))
            origin = fitz.Point(sp["origin"][0], sp["origin"][1]) if sp.get("origin") else fitz.Point(bbox.x0, bbox.y1 - style["size"] * 0.22)
            if sp.get("bbox"): origin.x = min(origin.x, bbox.x0 + 0.5)
            page.add_redact_annot(bbox + (0.3, 0.3, -0.3, -0.3), fill=False)
            _apply_redactions(page, **_redact_kw())
            if str(o.get("text", "")).strip():
                write_run(doc, page, origin, o["text"], style, cache, keep_font=keep)

        elif kind == "scan_text":
            notes.append(scan_text_op(doc, page, o))

        elif kind == "magic_erase":
            notes.append(magic_erase_op(doc, page, o))

        elif kind == "delete_text":
            page.add_redact_annot(_rect_of(o["bbox"]) + (0.3, 0.3, -0.3, -0.3), fill=False)
            _apply_redactions(page, **_redact_kw())

        elif kind == "add_text":
            size = float(o.get("size") or 12)
            style = dict(size=size, color=o.get("color") or "#000000", family=o.get("family") or "sans",
                         bold=bool(o.get("bold")), italic=bool(o.get("italic")))
            write_run(doc, page, fitz.Point(float(o["x"]), float(o["y"])), o.get("text", ""), style, cache, keep_font=False)

        elif kind in ("move_image", "replace_image", "image_filter", "delete_image"):
            xref, bbox = _resolve_image(page, int(o.get("xref") or 0), _rect_of(o["bbox"]))
            if kind == "replace_image":
                rec = FILES.get(o.get("file_id"))
                if not rec or rec["kind"] != "image": raise ValueError("Pick an image file to replace with.")
                new = ImageOps.exif_transpose(Image.open(rec["path"])); new.load()
                _remove_image(doc, page, xref, bbox)
                _insert_pil(page, bbox, new, keep=o.get("fit", "contain") != "stretch")
                continue
            if kind == "delete_image":
                _remove_image(doc, page, xref, bbox); continue
            if not xref:
                pix = page.get_pixmap(clip=bbox, dpi=200)
                im = Image.open(io.BytesIO(pix.tobytes("png"))); im.load()
            else:
                im = _image_pil(doc, xref)
            if kind == "image_filter":
                im = filter_image(im, o.get("action"), o.get("amount"))
                target = bbox
                keep = o.get("action") in ("rotate",)
            else:
                target, keep = _rect_of(o["rect"]), False
            if xref:
                _remove_image(doc, page, xref, bbox)
            else:
                page.add_redact_annot(bbox, fill=(1, 1, 1)); _apply_redactions(page, **_redact_kw(images="pixels", text=False))
            _insert_pil(page, target, im, keep=keep)

        elif kind == "add_image":
            rec = FILES.get(o.get("file_id"))
            if not rec or rec["kind"] != "image": raise ValueError("Image file not found.")
            im = ImageOps.exif_transpose(Image.open(rec["path"])); im.load()
            if o.get("rect"):
                r = _rect_of(o["rect"])
            else:
                pr = page.rect; w = float(o.get("width") or pr.width * 0.4); h = w * im.height / max(1, im.width)
                r = fitz.Rect((pr.width - w) / 2, (pr.height - h) / 2, (pr.width + w) / 2, (pr.height + h) / 2)
            _insert_pil(page, r, im, keep=True)

        elif kind == "ink":
            pts = [fitz.Point(float(x), float(y)) for x, y in o.get("points", [])]
            if len(pts) < 2: continue
            sh = page.new_shape()
            sh.draw_polyline(pts)
            sh.finish(color=_rgb(o.get("color"), (0, 0, 0)), width=float(o.get("width") or 2), closePath=False,
                      lineCap=1, lineJoin=1, stroke_opacity=float(o.get("opacity") or 1))
            sh.commit()

        elif kind == "shape":
            r = _rect_of(o["rect"]); sh = page.new_shape(); k = o.get("kind", "rect")
            if k == "ellipse": sh.draw_oval(r)
            elif k == "line": sh.draw_line(fitz.Point(*o["rect"][:2]), fitz.Point(*o["rect"][2:]))
            else: sh.draw_rect(r)
            fill = _rgb(o.get("fill")) if o.get("fill") else None
            sh.finish(color=_rgb(o.get("color"), (0, 0, 0)), fill=fill, width=float(o.get("width") or 1.5),
                      fill_opacity=float(o.get("fill_opacity") or 1), stroke_opacity=float(o.get("opacity") or 1))
            sh.commit()

        elif kind == "highlight":
            r = _rect_of(o["rect"])
            words = []
            for w in page.get_text("words"):
                wr = fitz.Rect(w[:4]); c = fitz.Point((wr.x0 + wr.x1) / 2, (wr.y0 + wr.y1) / 2)
                if r.contains(c): words.append(wr)
            if words:
                a = page.add_highlight_annot(words)
                a.set_colors(stroke=_rgb(o.get("color"), (1, 0.85, 0.1))); a.update()
            else:
                sh = page.new_shape(); sh.draw_rect(r)
                sh.finish(color=None, fill=_rgb(o.get("color"), (1, 0.85, 0.1)), fill_opacity=0.35); sh.commit()

        elif kind == "erase":
            r = _rect_of(o["rect"])
            page.add_redact_annot(r, fill=_rgb(o.get("fill"), (1, 1, 1)) if o.get("fill") else (1, 1, 1))
            _apply_redactions(page, **_redact_kw(images="pixels", graphics="covered"))

        elif kind == "page_rotate":
            page.set_rotation((page.rotation + int(o.get("angle") or 90)) % 360)
            try: page.remove_rotation()
            except Exception: pass
        elif kind == "page_delete":
            if doc.page_count <= 1: raise ValueError("A PDF needs at least one page.")
            doc.delete_page(pno)
        elif kind == "page_insert":
            after = int(o.get("after", doc.page_count - 1))
            ref = doc[max(0, min(after, doc.page_count - 1))].rect
            doc.new_page(after + 1, width=ref.width, height=ref.height)
        elif kind == "page_duplicate":
            doc.fullcopy_page(pno, pno + 1) if hasattr(doc, "fullcopy_page") else doc.copy_page(pno, pno + 1)
        elif kind == "page_move":
            to = int(o.get("to", pno))
            if 0 <= to < doc.page_count and to != pno:
                doc.move_page(pno, to if to < pno else (to + 1 if to + 1 < doc.page_count else -1))
        else:
            raise ValueError(f"Unknown edit '{kind}'.")
    return notes


def open_draft(rec):
    """Copy a PDF (or wrap an image as a 1-page PDF) into an editable draft."""
    if rec["kind"] == "image":
        im = _load_img(rec); k = 0.75 if max(im.size) <= 4000 else 3000 / max(im.size)
        d = fitz.open(); pg = d.new_page(width=im.width * k, height=im.height * k)
        _insert_pil(pg, pg.rect, im, keep=False)
        name = f"{_stem(rec)}_edited.pdf"; fid, path = new_path(name)
        d.save(path, garbage=3, deflate=True); d.close()
        r = register(fid, path, name, origin="draft")
        r["meta"].update(source=rec["id"], image_px=[im.width, im.height])
        return r
    d = _open_pdf(rec)
    for pg in d:
        if pg.rotation:
            try: pg.remove_rotation()
            except Exception: pass
    name = f"{_stem(rec)}_edited.pdf" if not rec["name"].endswith("_edited.pdf") else rec["name"]
    fid, path = new_path(name)
    d.save(path, garbage=3, deflate=True); d.close()
    r = register(fid, path, name, origin="draft")
    r["meta"]["source"] = rec["id"]
    return r


def edit_draft(rec, ops):
    d = fitz.open(rec["path"])
    apply_ops(d, ops)
    fid, path = new_path(rec["name"])
    d.save(path, garbage=3, deflate=True); d.close()
    r = register(fid, path, rec["name"], origin="draft")
    r["meta"]["source"] = rec["meta"].get("source")
    if rec["meta"].get("image_px"): r["meta"]["image_px"] = rec["meta"]["image_px"]
    return r


def pages_info(rec):
    with fitz.open(rec["path"]) as d:
        return [{"w": round(p.rect.width, 2), "h": round(p.rect.height, 2)} for p in d]


# ---------- AI chat tools that edit images inside PDFs ----------
def _pdf_and_img(recs):
    pdf = next((r for r in recs if r["kind"] == "pdf"), None)
    img = next((r for r in recs if r["kind"] == "image"), None)
    if not pdf: raise ValueError("Include the PDF's file id.")
    return pdf, img


def _pick_image(doc, p):
    pno = int(_num(p, "page", 1)) - 1
    if not 0 <= pno < doc.page_count: raise ValueError(f"Page must be 1–{doc.page_count}.")
    imgs = page_layout(doc[pno])["images"]
    if not imgs: raise ValueError(f"No images found on page {pno + 1}.")
    idx = int(_num(p, "image", 1))
    if not 1 <= idx <= len(imgs): raise ValueError(f"Page {pno + 1} has {len(imgs)} image(s); pick 1–{len(imgs)}.")
    return pno, imgs[idx - 1]


def _run_pdf_ops(pdf, ops, suffix):
    d = _open_pdf(pdf)
    apply_ops(d, ops)
    return _save_pdf(d, f"{_stem(pdf)}_{suffix}.pdf")


@tool("pdf_replace_image", "PDF", "Replace image in PDF", "Swap an image inside a PDF for another image, in the same spot (images are numbered per page, top to bottom).",
      ["pdf", "image"], multi=True,
      params=[P("page", "Page", "number", "1"), P("image", "Image # on page", "number", "1"), P("fit", "Fit", "select", "contain", ["contain", "stretch"])])
def t_pdf_replace_image(recs, p):
    pdf, img = _pdf_and_img(recs)
    if not img: raise ValueError("Also include the new image's file id.")
    with fitz.open(pdf["path"]) as d: pno, m = _pick_image(d, p)
    rec = _run_pdf_ops(pdf, [dict(op="replace_image", page=pno, xref=m["xref"], bbox=m["bbox"], file_id=img["id"], fit=_s(p, "fit", "contain"))], "image_replaced")
    return rec, f"Replaced image {m['index']} on page {pno + 1}"


@tool("pdf_edit_image", "PDF", "Edit image in PDF", "Apply an effect to an image inside a PDF: " + ", ".join(IMG_ACTIONS) + ".", PDF,
      params=[P("page", "Page", "number", "1"), P("image", "Image # on page", "number", "1"),
              P("action", "Effect", "select", "auto_enhance", IMG_ACTIONS), P("amount", "Amount (optional)", "number", "")])
def t_pdf_edit_image(recs, p):
    pdf = recs[0]
    with fitz.open(pdf["path"]) as d: pno, m = _pick_image(d, p)
    act = _s(p, "action", "auto_enhance")
    rec = _run_pdf_ops(pdf, [dict(op="image_filter", page=pno, xref=m["xref"], bbox=m["bbox"], action=act, amount=p.get("amount"))], "image_edited")
    return rec, f"{act.replace('_', ' ').title()} → image {m['index']}, page {pno + 1}"


@tool("pdf_delete_image", "PDF", "Delete image in PDF", "Remove an image from a PDF page.", PDF,
      params=[P("page", "Page", "number", "1"), P("image", "Image # on page", "number", "1")])
def t_pdf_delete_image(recs, p):
    pdf = recs[0]
    with fitz.open(pdf["path"]) as d: pno, m = _pick_image(d, p)
    rec = _run_pdf_ops(pdf, [dict(op="delete_image", page=pno, xref=m["xref"], bbox=m["bbox"])], "image_removed")
    return rec, f"Removed image {m['index']} from page {pno + 1}"


@tool("pdf_add_image", "PDF", "Add image to PDF", "Place an image (logo, photo, signature) on a PDF page.", ["pdf", "image"], multi=True,
      params=[P("page", "Page", "number", "1"),
              P("position", "Position", "select", "bottom-right", ["top-left", "top-center", "top-right", "center", "bottom-left", "bottom-center", "bottom-right"]),
              P("width", "Width (pt)", "number", "120"), P("margin", "Margin (pt)", "number", "36")])
def t_pdf_add_image(recs, p):
    pdf, img = _pdf_and_img(recs)
    if not img: raise ValueError("Also include the image's file id.")
    with Image.open(img["path"]) as im: ar = im.height / max(1, im.width)
    with fitz.open(pdf["path"]) as d:
        pno = int(_num(p, "page", 1)) - 1
        if not 0 <= pno < d.page_count: raise ValueError(f"Page must be 1–{d.page_count}.")
        pr = d[pno].rect
    w = _num(p, "width", 120); h = w * ar; m = _num(p, "margin", 36); pos = _s(p, "position", "bottom-right")
    x = m if "left" in pos else (pr.width - w - m if "right" in pos else (pr.width - w) / 2)
    y = m if pos.startswith("top") else (pr.height - h - m if pos.startswith("bottom") else (pr.height - h) / 2)
    rec = _run_pdf_ops(pdf, [dict(op="add_image", page=pno, file_id=img["id"], rect=[x, y, x + w, y + h])], "image_added")
    return rec, f"Added {img['name']} on page {pno + 1} ({pos})"


@tool("pdf_erase_area", "PDF", "Erase area", "White-out and permanently remove everything inside a box on a page (points from top-left).", PDF,
      params=[P("page", "Page", "number", "1"), P("x0", "Left", "number", "0"), P("y0", "Top", "number", "0"),
              P("x1", "Right", "number", "100"), P("y1", "Bottom", "number", "100")])
def t_pdf_erase(recs, p):
    pno = int(_num(p, "page", 1)) - 1
    rect = [_num(p, "x0", 0), _num(p, "y0", 0), _num(p, "x1", 100), _num(p, "y1", 100)]
    rec = _run_pdf_ops(recs[0], [dict(op="erase", page=pno, rect=rect)], "erased")
    return rec, f"Erased area on page {pno + 1}"

# ==========================================================
# PHOTO-REAL IMAGE EDITING (GPU when free)
#   text in images: OCR → LaMa background rebuild → font match (40+ families) → pixel-matched redraw
#   magic eraser (LaMa) · AI upscale (Real-ESRGAN)
# ==========================================================
from PIL import ImageDraw, ImageFont
from functools import lru_cache
import glob as _glob, tarfile as _tarfile, urllib.request as _urlreq

MODEL_DIR = os.path.join(WORK_DIR, "models")
FONT_DIR = os.path.join(WORK_DIR, "fonts")
os.makedirs(MODEL_DIR, exist_ok=True); os.makedirs(FONT_DIR, exist_ok=True)
LAMA_URL = "https://github.com/enesmsahin/simple-lama-inpainting/releases/download/v0.1.0/big-lama.pt"
ESRGAN_URL = "https://github.com/xinntao/Real-ESRGAN/releases/download/v0.1.0/RealESRGAN_x4plus.pth"
FONT_FAMILIES = ["inter", "roboto", "open-sans", "noto-sans", "lato", "poppins", "montserrat", "source-sans-3", "nunito",
                 "ubuntu", "ibm-plex-sans", "dm-sans", "manrope", "rubik", "work-sans", "fira-sans", "pt-sans", "arimo",
                 "carlito", "raleway", "oswald", "tinos", "caladea", "noto-serif", "merriweather", "pt-serif",
                 "playfair-display", "source-serif-4", "lora", "eb-garamond", "roboto-mono", "cousine", "source-code-pro",
                 "jetbrains-mono", "noto-sans-devanagari", "mukta", "hind"]
_KEEP_STYLES = ("300Light", "400Regular", "500Medium", "600SemiBold", "700Bold", "800ExtraBold", "400Regular_Italic", "700Bold_Italic")
_FONTS, _FONTS_N = None, -1
_PACK_DONE = threading.Event()


def _download(url, path):
    if os.path.exists(path) and os.path.getsize(path) > 1000: return path
    tmp = path + ".part"
    _urlreq.urlretrieve(url, tmp); os.replace(tmp, path)
    return path


def _safe(fn):
    try: return fn()
    except Exception: return None


def _fetch_font_pack():
    """Static TTFs for 37 Google font families (via the npm registry), cached in WORK_DIR/fonts."""
    try:
        for fam in FONT_FAMILIES:
            key = fam.replace("-", "")
            if _glob.glob(os.path.join(FONT_DIR, key + "__*.ttf")): continue
            try:
                meta = json.loads(_urlreq.urlopen(f"https://registry.npmjs.org/@expo-google-fonts/{fam}/latest", timeout=20).read())
                tgz = os.path.join(FONT_DIR, fam + ".tgz"); _download(meta["dist"]["tarball"], tgz)
                with _tarfile.open(tgz) as tf:
                    for m in tf.getmembers():
                        if not m.name.endswith(".ttf"): continue
                        style = m.name.split("/")[-2]
                        if style not in _KEEP_STYLES: continue
                        open(os.path.join(FONT_DIR, f"{key}__{style}.ttf"), "wb").write(tf.extractfile(m).read())
                os.remove(tgz)
            except Exception:
                continue
    finally:
        _PACK_DONE.set()


def warm_up_media():
    """Background downloads (font pack + LaMa) so the first edit is fast."""
    threading.Thread(target=_fetch_font_pack, daemon=True).start()
    threading.Thread(target=lambda: _safe(lambda: _download(LAMA_URL, os.path.join(MODEL_DIR, "big-lama.pt"))), daemon=True).start()


def _font_list():
    """Candidate fonts: list of dicts {path, family, weight, italic}."""
    global _FONTS, _FONTS_N
    paths = set(_glob.glob(os.path.join(FONT_DIR, "*__*.ttf")))
    if _FONTS is not None and _FONTS_N == len(paths): return _FONTS
    sysp = set()
    for root in ("/usr/share/fonts", "/usr/local/share/fonts"):
        sysp |= set(_glob.glob(os.path.join(root, "**", "*.[ot]tf"), recursive=True))
    try:
        import matplotlib
        sysp |= set(_glob.glob(os.path.join(os.path.dirname(matplotlib.__file__), "mpl-data", "fonts", "ttf", "DejaVu*.ttf")))
    except Exception:
        pass
    out, seen = [], set()
    for p in sorted(paths):
        try:
            fam, style = os.path.basename(p)[:-4].split("__")
            out.append(dict(path=p, family=fam, weight=int(style[:3]), italic="Italic" in style)); seen.add(fam)
        except Exception:
            continue
    good_sys = ("liberation", "dejavu", "carlito", "caladea", "freesans", "freeserif", "ubuntu")
    for p in sorted(sysp):
        n = os.path.basename(p).lower()
        if not any(g in n for g in good_sys) or any(k in n for k in ("cjk", "emoji", "symbol", "math", "condensed", "display", "extralight")): continue
        try:
            fam, sty = ImageFont.truetype(p, 20).getname()
        except Exception:
            continue
        key = re.sub(r"[^a-z]", "", fam.lower())
        if key in seen: continue
        s = (sty or "regular").lower()
        if s not in ("regular", "book", "bold", "italic", "bold italic", "medium", "semibold", "light", "oblique", "bold oblique"): continue
        w = 600 if "semibold" in s else 700 if "bold" in s else 500 if "medium" in s else 300 if "light" in s else 400
        out.append(dict(path=p, family=fam, weight=w, italic="italic" in s or "oblique" in s))
    _FONTS, _FONTS_N = out, len(paths)
    return out


@lru_cache(maxsize=8192)
def _tt(path, size):
    return ImageFont.truetype(path, max(4, round(size * 4) / 4))


@lru_cache(maxsize=16384)
def _ink_box(path, size, text):
    """True ink extents of `text` relative to its baseline origin (PIL's getbbox includes side bearings)."""
    import numpy as np
    f = _tt(path, size); l, t, r, b = f.getbbox(text, anchor="ls")
    pad = int(size) + 4
    im = Image.new("L", (int(r - l) + 2 * pad, int(b - t) + 2 * pad), 0)
    ImageDraw.Draw(im).text((pad - l, pad - t), text, font=f, fill=255, anchor="ls")
    a = np.asarray(im) > 60; ys, xs = np.where(a)
    if not len(xs): return None
    return (xs.min() - (pad - l), ys.min() - (pad - t), xs.max() + 1 - (pad - l), ys.max() + 1 - (pad - t))


def _render_mask(path, text, w, h):
    """Render text stretched exactly onto a w×h ink box → (0..1 map, font size, width ratio)."""
    import numpy as np
    ib = _ink_box(path, 48, text)
    if not ib or ib[2] - ib[0] < 1 or ib[3] - ib[1] < 1: return None
    size = 48 * h / (ib[3] - ib[1]); ib = _ink_box(path, size, text)
    if not ib: return None
    f = _tt(path, size); W, H = max(1, ib[2] - ib[0]), max(1, ib[3] - ib[1])
    im = Image.new("L", (W + 4, H + 4), 0)
    ImageDraw.Draw(im).text((2 - ib[0], 2 - ib[1]), text, font=f, fill=255, anchor="ls")
    im = im.crop((2, 2, 2 + W, 2 + H)).resize((max(1, int(w)), max(1, int(h))), Image.BILINEAR)
    return np.asarray(im, float) / 255.0, size, w / W


def match_font(strength, text, only=None):
    """Two-stage search over every installed font (family, then weight/italic), blur-aware correlation.
    `only` = list of font dicts to restrict the search to."""
    import numpy as np, cv2
    h, w = strength.shape
    a = strength - strength.mean(); na = np.sqrt((a * a).sum()) or 1.0
    blurs = (0, 0.6, 1.1) if h < 60 else (0, 1.0, 2.0)
    target_cov = (strength > 0.5).mean()

    def score(fd):
        try:
            r = _render_mask(fd["path"], text, w, h)
        except Exception:
            return None
        if not r: return None
        m0, size, sx = r; top = None
        for sig in blurs:
            m = cv2.GaussianBlur(m0, (0, 0), sig) if sig else m0
            bm = m - m.mean(); corr = float((a * bm).sum() / (na * (np.sqrt((bm * bm).sum()) or 1.0)))
            weight = target_cov / max(1e-3, (m > 0.5).mean())
            sc = corr - 0.15 * abs(np.log(max(1e-3, sx))) - 0.25 * abs(np.log(max(1e-3, weight)))
            if top is None or sc > top[0]: top = (sc, weight, sig)
        return (top[0], fd, size, sx, top[1], top[2])

    if only:
        res = [r for r in (score(fd) for fd in only) if r]
        return max(res, key=lambda r: r[0]) if res else None
    fams = {}
    for fd in _font_list(): fams.setdefault(fd["family"], []).append(fd)
    stage1 = []
    for fam, lst in fams.items():
        probe = [min(lst, key=lambda f: (f["italic"], abs(f["weight"] - 400)))]
        bold = [f for f in lst if not f["italic"] and f["weight"] >= 600]
        if bold: probe.append(min(bold, key=lambda f: abs(f["weight"] - 700)))
        for fd in probe:
            r = score(fd)
            if r: stage1.append(r)
    if not stage1: return None
    stage1.sort(key=lambda r: -r[0])
    best = stage1[0]
    for fam in {r[1]["family"] for r in stage1[:5]}:
        for fd in fams[fam]:
            r = score(fd)
            if r and r[0] > best[0]: best = r
    return best


# ---------- GPU models ----------
_TORCH = {}


def _device(need_gb=1.5):
    try:
        import torch
        if torch.cuda.is_available():
            f, i = max((torch.cuda.mem_get_info(i)[0], i) for i in range(torch.cuda.device_count()))
            if f / 2**30 >= need_gb: return f"cuda:{i}"
    except Exception:
        pass
    return "cpu"


def _lama():
    import torch
    if "lama" not in _TORCH:
        path = _download(LAMA_URL, os.path.join(MODEL_DIR, "big-lama.pt"))
        _TORCH["lama"] = (torch.jit.load(path, map_location="cpu").eval(), "cpu")
    model, dev = _TORCH["lama"]
    want = _device(2.0)
    if want != dev:
        try:
            model = model.to(want); dev = want
        except Exception:
            model = model.to("cpu"); dev = "cpu"
        _TORCH["lama"] = (model, dev)
    FEATURES["lama"] = dev
    return model, dev


def inpaint(rgb, mask, max_side=1024):
    """Rebuild masked pixels with LaMa (OpenCV fallback). rgb HxWx3 uint8, mask HxW."""
    import numpy as np, cv2
    m = (mask > 0).astype(np.uint8)
    if not m.any(): return rgb
    ys, xs = np.where(m); H, W = m.shape
    big = max(ys.max() - ys.min() + 1, xs.max() - xs.min() + 1)
    mg = int(max(48, 1.6 * big if big < 200 else 0.6 * big))
    X0, Y0 = max(0, xs.min() - mg), max(0, ys.min() - mg); X1, Y1 = min(W, xs.max() + 1 + mg), min(H, ys.max() + 1 + mg)
    crop, cm = rgb[Y0:Y1, X0:X1], m[Y0:Y1, X0:X1]
    try:
        import torch
        model, dev = _lama()
        k = min(1.0, max_side / max(crop.shape[:2]))
        ci = cv2.resize(crop, None, fx=k, fy=k, interpolation=cv2.INTER_AREA) if k < 1 else crop
        mi = cv2.resize(cm, (ci.shape[1], ci.shape[0]), interpolation=cv2.INTER_NEAREST) if k < 1 else cm
        mi = cv2.dilate(mi, np.ones((3, 3), np.uint8), iterations=1)
        ph, pw = (-ci.shape[0]) % 8, (-ci.shape[1]) % 8
        ci_p = np.pad(ci, ((0, ph), (0, pw), (0, 0)), mode="reflect"); mi_p = np.pad(mi, ((0, ph), (0, pw)))
        t = torch.from_numpy(np.ascontiguousarray(ci_p)).permute(2, 0, 1)[None].float().div(255).to(dev)
        tm = torch.from_numpy(np.ascontiguousarray(mi_p))[None, None].float().to(dev)
        with torch.no_grad():
            out = model(t, tm)[0].permute(1, 2, 0).clamp(0, 1).mul(255).byte().cpu().numpy()
        out = out[:ci.shape[0], :ci.shape[1]]
        if k < 1: out = cv2.resize(out, (crop.shape[1], crop.shape[0]), interpolation=cv2.INTER_CUBIC)
    except Exception:
        FEATURES["lama"] = "unavailable (OpenCV fallback)"
        out = cv2.inpaint(crop, cm * 255, 5, cv2.INPAINT_TELEA)
    soft = cv2.GaussianBlur(cv2.dilate(cm, np.ones((3, 3), np.uint8)).astype(np.float32), (5, 5), 0)[..., None]
    res = rgb.copy()
    res[Y0:Y1, X0:X1] = np.clip(out * soft + crop * (1 - soft), 0, 255).astype(np.uint8)
    return res


def ai_upscale(img, scale=2):
    """Real-ESRGAN ×4 (tiled; GPU when free), resized to the requested scale."""
    import numpy as np, torch
    if "esrgan" not in _TORCH:
        import spandrel
        path = _download(ESRGAN_URL, os.path.join(MODEL_DIR, "RealESRGAN_x4plus.pth"))
        _TORCH["esrgan"] = spandrel.ModelLoader().load_from_file(path).model.eval()
    dev = _device(2.5); model = _TORCH["esrgan"].to(dev); half = dev.startswith("cuda")
    if half: model = model.half()
    alpha = img.getchannel("A") if "A" in img.getbands() else None
    x = np.asarray(img.convert("RGB")).astype(np.float32) / 255
    H, W = x.shape[:2]; tile, ov = 256, 16
    out = np.zeros((H * 4, W * 4, 3), np.float32)
    for y in range(0, H, tile):
        for xx in range(0, W, tile):
            y0, x0 = max(0, y - ov), max(0, xx - ov); y1, x1 = min(H, y + tile + ov), min(W, xx + tile + ov)
            t = torch.from_numpy(np.ascontiguousarray(x[y0:y1, x0:x1])).permute(2, 0, 1)[None].to(dev)
            if half: t = t.half()
            with torch.no_grad():
                r = model(t)[0].float().clamp(0, 1).permute(1, 2, 0).cpu().numpy()
            oy, ox = (y - y0) * 4, (xx - x0) * 4; hh, ww = min(tile, H - y) * 4, min(tile, W - xx) * 4
            out[y * 4:y * 4 + hh, xx * 4:xx * 4 + ww] = r[oy:oy + hh, ox:ox + ww]
    if half: _TORCH["esrgan"] = model.float()
    res = Image.fromarray((out * 255).round().astype(np.uint8))
    if scale != 4: res = res.resize((int(W * scale), int(H * scale)), Image.LANCZOS)
    if alpha: res.putalpha(alpha.resize(res.size, Image.LANCZOS))
    FEATURES["esrgan"] = dev
    return res


# ---------- OCR & analysis ----------
def ocr_boxes(img, max_side=2600):
    """Line-level OCR: [{bbox:[x0,y0,x1,y1] (image px), text, score}]."""
    global _OCR
    import numpy as np
    if _OCR is None:
        from rapidocr_onnxruntime import RapidOCR
        _OCR = RapidOCR(); FEATURES["ocr"] = True
    im = img.convert("RGB"); k = 1.0
    if max(im.size) > max_side:
        k = max_side / max(im.size); im = im.resize((int(im.width * k), int(im.height * k)), Image.LANCZOS)
    elif max(im.size) < 900:                                  # small screenshots read better enlarged
        k = 2.0; im = im.resize((im.width * 2, im.height * 2), Image.LANCZOS)
    res, _ = _OCR(np.array(im))
    out = []
    for quad, text, score in (res or []):
        score = float(score)
        if score < 0.5 or not text.strip(): continue
        xs = [q[0] / k for q in quad]; ys = [q[1] / k for q in quad]
        out.append({"bbox": [min(xs), min(ys), max(xs), max(ys)], "text": text.strip(), "score": round(score, 3)})
    return out


def _tight_box(gray, box):
    """Shrink an OCR box to the real ink (text = the minority colour). On gradients/photos the ink is
    measured against a local background estimate so shaded background never counts as text."""
    import numpy as np, cv2
    x0, y0, x1, y1 = [int(round(v)) for v in box]
    x0, y0 = max(0, x0), max(0, y0)
    crop = gray[y0:y1, x0:x1]
    if crop.size == 0: return box, None
    _, th = cv2.threshold(crop, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
    ink = th > 0
    if ink.mean() > 0.5: ink = ~ink
    if (~ink).sum() > 20 and float(crop[~ink].std()) > 6 and min(crop.shape) >= 6:
        h = crop.shape[0]; k = min(255, (int(h * 2.2) // 2) * 2 + 1); p = k // 2 + 1
        big = cv2.copyMakeBorder(crop, p, p, p, p, cv2.BORDER_REFLECT)
        bg = cv2.medianBlur(big, k)[p:-p, p:-p].astype(np.float32)
        d = crop.astype(np.float32) - bg
        sign = 1.0 if float(np.mean(crop[ink])) > float(np.mean(crop[~ink])) else -1.0
        dd = np.clip(d * sign, 0, None)
        core = dd[ink]
        if core.size > 10:
            thr = max(10.0, 0.35 * float(np.percentile(core, 90)))
            loc = dd > thr
            if loc.sum() > 10: ink = loc
    ys, xs = np.where(ink)
    if len(xs) < 4: return box, ink
    return [x0 + xs.min(), y0 + ys.min(), x0 + xs.max() + 1, y0 + ys.max() + 1], ink[ys.min():ys.max() + 1, xs.min():xs.max() + 1]


def _style_from_pixels(rgb, ink):
    import numpy as np
    txt = rgb[ink].astype(float); bg = rgb[~ink].astype(float)
    bc = np.median(bg, axis=0) if len(bg) else np.array([255., 255., 255.])
    if len(txt):
        dist = np.linalg.norm(txt - bc, axis=1)
        core = txt[dist >= np.quantile(dist, 0.7)]
        tc = np.median(core if len(core) else txt, axis=0)
    else:
        tc = np.array([0., 0., 0.])
    bstd = float(bg.std(axis=0).mean()) if len(bg) > 8 else 0.0
    return tuple(int(v) for v in tc), tuple(int(v) for v in bc), bstd


def _ink_strength(rgb, tc, bc):
    import numpy as np
    tc, bc = np.array(tc, float), np.array(bc, float); d = tc - bc; dd = float((d * d).sum()) or 1.0
    return np.clip(((rgb.astype(float) - bc) @ d) / dd, 0, 1)


_WIDTH_FONT = None


def _width_font():
    """A neutral sans font used only to estimate relative character widths."""
    global _WIDTH_FONT
    if _WIDTH_FONT is None:
        cands = [f for f in _font_list() if not f["italic"] and f["weight"] == 400 and any(k in f["family"].lower() for k in ("arimo", "liberation sans", "roboto", "dejavu sans", "notosans"))]
        _WIDTH_FONT = _tt(cands[0]["path"], 100) if cands else ImageFont.load_default()
    return _WIDTH_FONT


def _ink_gaps(ink):
    """Empty-column runs inside a text line, split into letter gaps vs word gaps (natural break)."""
    h, w = ink.shape; cols = ink.any(axis=0); runs, x = [], 0
    while x < w:
        if not cols[x]:
            s0 = x
            while x < w and not cols[x]: x += 1
            if s0 > 0 and x < w: runs.append((s0, x))
        else:
            x += 1
    if not runs: return []
    ws = sorted(e - s for s, e in runs)
    thr = 0.3 * h
    best = 0
    for a, b in zip(ws, ws[1:]):
        if b >= 0.16 * h and a > 0 and b / a > best: best, cut = b / a, b
    if best >= 1.45: thr = min(thr, cut) if cut >= 0.16 * h else thr
    return [(s, e) for s, e in runs if e - s >= thr]


def _local_strength(region, tc, k=None):
    """Ink coverage measured against a LOCAL background estimate (works on gradients, photos, paper)."""
    import numpy as np, cv2
    h = region.shape[0]
    k = k or max(3, (int(h * 2.2) // 2) * 2 + 1)
    k = min(k, 255)
    lb = cv2.medianBlur(region, k).astype(np.float32)
    d = np.array(tc, np.float32)[None, None, :] - lb
    num = ((region.astype(np.float32) - lb) * d).sum(-1); den = (d * d).sum(-1) + 1e-3
    return np.clip(num / den, 0, 1)


def _word_gap_filter(gaps, text_nospace, path, size, scale, h):
    """Keep only gaps that are wider than the font's own letter spacing would explain."""
    f = _tt(path, size); keep = []
    pos = [f.getlength(text_nospace[:i]) * scale for i in range(len(text_nospace) + 1)]
    inks = [_ink_box(path, size, c) for c in text_nospace]
    for s0, e0 in gaps:
        j = min(range(1, len(text_nospace)), key=lambda q: abs(pos[q] - s0), default=None)
        if j is None: continue
        a, b = inks[j - 1], inks[j]
        expected = ((f.getlength(text_nospace[j - 1]) - (a[2] if a else 0)) + (b[0] if b else 0)) * scale
        if (e0 - s0) - max(0, expected) > 0.16 * size * scale: keep.append((s0, e0))
    return keep


def _respace(text, ink, font=None, gaps=None):
    """OCR often drops spaces ('Flat50%offthisweekend'); put them back where the ink has word gaps.
    Each ink segment between gaps is matched to a run of characters whose rendered ink width fits it
    best (dynamic programming over all split points, matched font's widths when known)."""
    if not text or ink is None or getattr(ink, "size", 0) == 0: return text
    gaps = _ink_gaps(ink) if gaps is None else gaps
    keep_sp = set()
    if " " in text:                                            # OCR kept some spaces: only fill in the missing ones
        words = text.split()
        if not gaps or len(gaps) <= len(words) - 1: return text
        pos = 0
        for w_ in words[:-1]:
            pos += len(w_); keep_sp.add(pos)
        orig_text, text = text, "".join(words)
    n = len(text)
    if not gaps or n < 2: return text
    gaps = sorted(gaps)[:n - 1]
    h, w = ink.shape
    f = font or _width_font()
    P = [0.0] + [f.getlength(text[:i]) for i in range(1, n + 1)]
    path = getattr(f, "path", None); fs = getattr(f, "size", 100)
    lsb, rsb = [], []
    for i, c in enumerate(text):
        ib = _ink_box(path, fs, c) if path else None
        adv = P[i + 1] - P[i]
        lsb.append(ib[0] if ib else 0.0); rsb.append(adv - ib[2] if ib else 0.0)
    edges = [0] + [v for g in gaps for v in g] + [w]
    segw = [max(0.0, edges[2 * k + 1] - edges[2 * k]) for k in range(len(gaps) + 1)]
    inkw = lambda a, b: max(1e-3, P[b] - P[a] - lsb[a] - rsb[b - 1])
    k = len(segw); INF = float("inf")

    def prior(a):                                              # how plausible a word break between text[a-1] and text[a] is
        p, q = text[a - 1], text[a]
        if p in ":;,.)]!?" and q not in ".,:;": return -0.45
        if q in "([" : return -0.1
        if p.isdigit() and q.isdigit(): return 0.9
        if p.isdigit() != q.isdigit() and p.isalnum() and q.isalnum(): return -0.05
        if p.isupper() and q.islower(): return 0.12
        if q in ".,:;%)!?": return 0.3
        return 0.0
    pri = [0.0] + [prior(a) - (0.6 if a in keep_sp else 0.0) for a in range(1, n)]
    scale = sum(segw) / inkw(0, n); cuts = None
    for _ in range(2):                                         # solve, refit the scale on the chosen split, solve again
        dp = [[INF] * (n + 1) for _ in range(k + 1)]; bk = [[0] * (n + 1) for _ in range(k + 1)]
        dp[0][0] = 0.0
        for j in range(1, k + 1):
            for i in range(j, n - (k - j) + 1):
                best, arg = INF, 0
                for a in range(j - 1, i):
                    if dp[j - 1][a] == INF: continue
                    c = dp[j - 1][a] + ((segw[j - 1] - inkw(a, i) * scale) / max(4.0, h)) ** 2 + (pri[a] if a > 0 else 0.0)
                    if c < best: best, arg = c, a
                dp[j][i], bk[j][i] = best, arg
        if dp[k][n] == INF: return text
        cuts, i = [], n
        for j in range(k, 0, -1):
            cuts.append(i); i = bk[j][i]
        cuts = sorted(cuts)                                    # segment ends
        starts = [0] + cuts[:-1]
        scale = sum(segw) / max(1e-3, sum(inkw(a, b) for a, b in zip(starts, cuts)))
    cuts = starts[1:]
    cuts = set(c for c in cuts if 0 < c < n)
    return "".join((" " if i in cuts else "") + c for i, c in enumerate(text))


def _membrane(img, hole, iters=300):
    """Smooth (harmonic) fill of `hole` pixels from their surroundings — multigrid Jacobi, no seams."""
    import numpy as np, cv2
    h, w = hole.shape
    if not hole.any(): return img
    out = img.copy()
    if min(h, w) >= 32 and iters > 40:                         # coarse solve first → fast convergence
        sh, sw = max(8, h // 4), max(8, w // 4)
        kw = (~hole).astype(np.float32)                        # mask-aware downscale (holes must not darken it)
        wsm = cv2.resize(kw, (sw, sh), interpolation=cv2.INTER_AREA)
        small = cv2.resize(img * kw[..., None], (sw, sh), interpolation=cv2.INTER_AREA) / np.maximum(wsm, 1e-6)[..., None]
        sm_hole = wsm < 0.5
        sm = _membrane(small, sm_hole, iters)
        out[hole] = cv2.resize(sm, (w, h), interpolation=cv2.INTER_LINEAR)[hole]
        iters = max(40, iters // 4)
    else:
        known = out[~hole].mean(axis=0) if (~hole).any() else 0
        out[hole] = known
    hm = hole[..., None]
    for _ in range(iters):
        out = np.where(hm, cv2.blur(out, (3, 3)), img)
    return out


def _texture_fill(orig, mask, valid, hb):
    """Background rebuild for gradients, grain and soft textures: a seamless smooth colour field (harmonic
    fill from the clean pixels around the hole) plus real grain copied from nearby clean background.
    Looks like the surrounding paper/photo instead of a smear. Returns (filled rgb, texture std)."""
    import numpy as np, cv2
    H, W = mask.shape; m = mask > 0
    ys, xs = np.where(m)
    if not len(xs): return orig.copy(), 0.0
    y0, y1, x0, x1 = ys.min(), ys.max() + 1, xs.min(), xs.max() + 1
    mh = y1 - y0; P = max(mh, hb) * 2 + 8
    Y0, Y1, X0, X1 = max(0, y0 - P), min(H, y1 + P), max(0, x0 - P), min(W, x1 + P)
    o = orig[Y0:Y1, X0:X1].astype(np.float32); mm = m[Y0:Y1, X0:X1]
    good = valid[Y0:Y1, X0:X1] & ~mm; v = good.astype(np.float32)
    s0 = max(2.0, hb * 0.08)                                   # light, ink-free smoothing of the background
    wv = cv2.GaussianBlur(v, (0, 0), s0)
    B = cv2.GaussianBlur(o * v[..., None], (0, 0), s0) / np.maximum(wv, 1e-6)[..., None]
    known = good & (wv > 0.25)
    low = _membrane(np.where(known[..., None], B, 0).astype(np.float32), ~known)
    res = o - low
    rstd = float(res[good].std(axis=0).mean()) if good.sum() > 30 else 0.0
    # grain: copy the residual from the nearest clean rows above/below (same columns → same texture scale)
    out = o.copy(); my0, my1 = y0 - Y0, y1 - Y0; cols = slice(x0 - X0, x1 - X0)
    fill_res = np.zeros((my1 - my0, x1 - x0, 3), np.float32); have = np.zeros((my1 - my0, x1 - x0), bool)
    for d in (mh + 3, -(mh + 3), mh // 2 + 2, -(mh // 2 + 2), 2 * mh + 5, -(2 * mh + 5)):
        a0, a1 = my0 + d, my1 + d
        if a0 < 0 or a1 > out.shape[0]: continue
        src_ok = good[a0:a1, cols] & ~have
        fill_res[src_ok] = res[a0:a1, cols][src_ok]; have |= src_ok
        if have.all(): break
    if not have.all():                                         # leftovers: matching random grain
        sd = res[good].std(axis=0) if good.sum() > 30 else np.zeros(3)
        nz = cv2.GaussianBlur(np.random.normal(0, 1, fill_res.shape).astype(np.float32), (0, 0), 0.7) * 1.6 * sd
        fill_res[~have] = nz[~have]
    blk = low[my0:my1, cols] + fill_res
    sub = out[my0:my1, cols]; hm = mm[my0:my1, cols]
    sub[hm] = blk[hm]
    res_img = orig.copy(); res_img[Y0:Y1, X0:X1] = np.clip(out, 0, 255).astype(np.uint8)
    return res_img, rstd


def smart_fill(rgb, mask):
    """Pick the most natural background rebuild: exact colour for flat UI, harmonic field + real grain for
    gradients/paper/soft photos, LaMa (GPU) for real structure (objects, people, detailed scenes)."""
    import numpy as np, cv2
    m = (mask > 0).astype(np.uint8)
    if not m.any(): return rgb
    ys, xs = np.where(m); hb = max(8, int(min(ys.max() - ys.min(), xs.max() - xs.min()) + 1))
    ring = (cv2.dilate(m, np.ones((15, 15), np.uint8)) > 0) & (cv2.dilate(m, np.ones((5, 5), np.uint8)) == 0)
    rp = rgb[ring].astype(np.float32)
    if len(rp) > 30 and float(rp.std(axis=0).mean()) < 3.0:
        out = rgb.copy(); out[m > 0] = np.median(rp, axis=0).astype(np.uint8); return out
    filled, rstd = _texture_fill(rgb, m, np.ones(m.shape, bool), hb)
    # structure test: strong edges around the hole mean objects/lines that a smooth fill would smear
    g = cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY)
    edges = cv2.Canny(g, 60, 150) > 0
    edge_frac = float(edges[ring].mean()) if ring.any() else 0.0
    if rstd < 9 and edge_frac < 0.04: return filled
    return inpaint(rgb, m)


def _subpixel_aa(rgb, strength):
    """Detect ClearType-style coloured fringes (Windows screenshots)."""
    import numpy as np
    edge = (strength > 0.15) & (strength < 0.85)
    if edge.sum() < 20: return False
    px = rgb[edge].astype(float)
    chroma = (np.abs(px[:, 0] - px[:, 1]) + np.abs(px[:, 1] - px[:, 2])).mean()
    core = rgb[strength > 0.9].astype(float)
    base = (np.abs(core[:, 0] - core[:, 1]) + np.abs(core[:, 1] - core[:, 2])).mean() if len(core) else 0
    return chroma > 18 and chroma > base + 12


def _noise_sigma(rgb, mask):
    import numpy as np, cv2
    g = cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY).astype(np.float32)
    hp = g - cv2.GaussianBlur(g, (0, 0), 1.5)
    v = hp[mask]
    return float(np.median(np.abs(v)) * 1.4826) if v.size > 50 else 0.0


def _jpeg_quality(patch):
    """Estimate the JPEG quality that produced `patch` (lowest re-compression error)."""
    import numpy as np
    best, err0 = None, None
    arr = np.asarray(patch, np.float32)
    for q in (98, 95, 92, 90, 87, 85, 82, 80, 75, 70, 65, 60, 50, 40):
        buf = io.BytesIO(); patch.save(buf, "JPEG", quality=q); rec = np.asarray(Image.open(buf), np.float32)
        e = float(np.abs(rec - arr).mean())
        if err0 is None or e < err0 - 0.02: best, err0 = q, e
    return best, err0


def _draw_line(draw, x, y, text, font, tracking=0.0, fill=255):
    if abs(tracking) < 0.05:
        draw.text((x, y), text, font=font, fill=fill, anchor="ls"); return
    for i, ch in enumerate(text):
        draw.text((x + font.getlength(text[:i]) + i * tracking, y), ch, font=font, fill=fill, anchor="ls")


def retext_image(img, box, new_text, orig_text=None, color=None, align="left", jpeg_hint=None):
    """Replace the text in `box` (image px) so it looks untouched: same font, weight, size, colour,
    spacing, anti-aliasing (incl. ClearType), blur, grain and JPEG texture."""
    import numpy as np, cv2
    src = img.convert("RGBA"); alpha = src.getchannel("A")
    orig = np.array(src.convert("RGB")); rgb = orig.copy()
    gray = cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY); H, W = gray.shape
    b = [max(0, box[0]), max(0, box[1]), min(W, box[2]), min(H, box[3])]
    tb, ink = _tight_box(gray, b)
    x0, y0, x1, y1 = [int(v) for v in tb]
    if ink is None or x1 - x0 < 3 or y1 - y0 < 3: raise ValueError("Couldn't find the text pixels in that area.")
    if not orig_text:
        found = ocr_boxes(Image.fromarray(orig[max(0, y0 - 6):y1 + 6, max(0, x0 - 6):x1 + 6]))
        orig_text = " ".join(f["text"] for f in found) or str(new_text)
    first = _respace(orig_text.split("\n")[0], ink)
    crop = orig[y0:y1, x0:x1]
    tc, bc, bstd = _style_from_pixels(crop, ink)
    if color: tc = tuple(int(c * 255) for c in _rgb(color))
    strength = _ink_strength(crop, tc, bc)
    th = y1 - y0
    # ---- 1. erase the old text and rebuild what was behind it
    pad = max(3, int(th * 0.4))
    X0, Y0, X1, Y1 = max(0, x0 - pad), max(0, y0 - pad), min(W, x1 + pad), min(H, y1 + pad)
    mask = np.zeros((H, W), np.uint8)
    mask[y0:y1, x0:x1] = (strength > 0.06).astype(np.uint8)
    mask = cv2.dilate(mask, np.ones((3, 3), np.uint8), iterations=max(1, th // 12))
    ring = np.zeros((H, W), bool); ring[Y0:Y1, X0:X1] = True; ring &= mask == 0
    if bstd < 3.5:                                              # flat UI / document background → exact colour
        rgb[mask > 0] = np.array(bc, np.uint8)
    else:
        rgb = inpaint(rgb, mask)
    if not str(new_text).strip():
        out = Image.fromarray(rgb); out.putalpha(alpha); return out, {"font": None, "removed": True, "lama": FEATURES.get("lama")}
    # ---- 2. identify the font
    m = match_font(strength, first)
    if not m: raise ValueError("No usable fonts found on the server.")
    score, fd, size, sx, weight, blur_sig = m
    path = fd["path"]; f = _tt(path, size)
    oib = _ink_box(path, size, first) or f.getbbox(first, anchor="ls")
    nat_w = max(1, oib[2] - oib[0]); ratio = (x1 - x0) / nat_w
    stretch, tracking = 1.0, 0.0
    if abs(ratio - 1) > 0.03:
        if abs(ratio - 1) <= 0.14 and len(first) > 2: tracking = (x1 - x0 - nat_w) / (len(first) - 1)
        else: stretch = ratio
    # ---- 3. render the new text the same way the original was rendered
    lines = str(new_text).split("\n"); lh = size * 1.2
    sub = _subpixel_aa(crop, strength) and bstd < 6
    ss = 3 if sub else 1
    fbig = _tt(path, size * ss) if ss > 1 else f
    widths = [fbig.getlength(li) + len(li) * abs(tracking) * ss for li in lines]
    m0 = int(size * ss * 0.6)
    lw = int(max(widths) + 2 * m0); lhh = int(lh * ss * len(lines) + 2 * m0)
    layer = Image.new("L", (lw, lhh), 0); d = ImageDraw.Draw(layer)
    for i, li in enumerate(lines):
        if not li.strip(): continue
        ib_l = _ink_box(path, size * ss, li)
        _draw_line(d, m0 - (ib_l[0] if ib_l else 0), m0 - oib[1] * ss + i * lh * ss, li, fbig, tracking * ss)
    A = np.asarray(layer, np.float32) / 255.0
    if ss > 1:                                                 # ClearType: 3× horizontal coverage → RGB subpixels
        A = cv2.resize(A, (A.shape[1], A.shape[0] // 3), interpolation=cv2.INTER_AREA)
        A = cv2.filter2D(A, -1, np.array([[1, 2, 3, 2, 1]], np.float32) / 9)
        w3 = A.shape[1] // 3 * 3; A = A[:, :w3]
        A = np.stack([A[:, 0::3], A[:, 1::3], A[:, 2::3]], axis=-1)
        m0x, m0y = m0 / 3, m0 / 3
    else:
        if stretch != 1.0:
            A = cv2.resize(A, (max(1, int(A.shape[1] * stretch)), A.shape[0]), interpolation=cv2.INTER_AREA if stretch < 1 else cv2.INTER_CUBIC)
        A = np.repeat(A[..., None], 3, axis=-1)
        m0x, m0y = m0 * stretch, m0
    if blur_sig: A = cv2.GaussianBlur(A, (0, 0), blur_sig)
    if weight > 1.3 and size > 18: A = cv2.dilate(A, np.ones((2, 2), np.uint8))
    A = np.clip(A, 0, 1)
    ox, oy = int(round(x0 - m0x)), int(round(y0 - m0y))
    hh, ww = A.shape[:2]
    ax0, ay0 = max(0, -ox), max(0, -oy); bx0, by0 = max(0, ox), max(0, oy)
    bx1, by1 = min(W, ox + ww), min(H, oy + hh)
    if bx1 <= bx0 or by1 <= by0: raise ValueError("Text landed outside the image.")
    Aw = A[ay0:ay0 + (by1 - by0), ax0:ax0 + (bx1 - bx0)]
    region = rgb[by0:by1, bx0:bx1].astype(np.float32)
    col = np.broadcast_to(np.array(tc, np.float32), region.shape).copy()
    core = strength > 0.85
    if core.sum() > 30:                                        # keep a vertical colour gradient if the original has one
        yy_c = np.where(core)[0].astype(np.float32); vals = crop[core].astype(np.float32)
        if np.ptp(yy_c) > 3:
            k = np.polyfit(yy_c, vals, 1)
            if np.abs(k[0]).max() * th > 25:
                yy = (np.arange(by0, by1) - y0).astype(np.float32)[:, None]
                col = np.repeat((k[0][None, :] * yy + k[1][None, :])[:, None, :], bx1 - bx0, axis=1)
    comp = region * (1 - Aw) + np.clip(col, 0, 255) * Aw
    sig = _noise_sigma(orig, ring)
    if sig > 1.2:                                               # same sensor / scan grain as the surroundings
        comp += np.random.normal(0, sig, comp.shape).astype(np.float32) * np.clip(Aw.max(axis=-1, keepdims=True) * 1.5, 0, 1)
    rgb[by0:by1, bx0:bx1] = np.clip(comp, 0, 255).astype(np.uint8)
    out = Image.fromarray(rgb)
    q = None
    if jpeg_hint or (bstd > 2 and sig > 0.8):                  # match the photo's JPEG block texture
        gx0, gy0 = max(0, (min(bx0, X0) // 8) * 8), max(0, (min(by0, Y0) // 8) * 8)
        gx1, gy1 = min(W, -(-max(bx1, X1) // 8) * 8), min(H, -(-max(by1, Y1) // 8) * 8)
        q, err = _jpeg_quality(Image.fromarray(orig[gy0:gy1, gx0:gx1]))
        if q and q < 97 and (jpeg_hint or err < 1.2):
            buf = io.BytesIO(); out.crop((gx0, gy0, gx1, gy1)).save(buf, "JPEG", quality=q)
            out.paste(Image.open(buf), (gx0, gy0))
    out.putalpha(alpha)
    return out, {"font": f"{fd['family']} {fd['weight']}{' italic' if fd['italic'] else ''}", "size_px": round(size, 1),
                 "match": round(float(score), 3), "color": "#%02x%02x%02x" % tc, "cleartype": sub, "jpeg_q": q,
                 "lama": FEATURES.get("lama")}


# ---------- line editing: change only the characters that differ ----------
_OCR_ALIKE = {"0": "0Oo", "o": "0Oo", "1": "1lI|i", "l": "1lI|i", "i": "1lI|i", "5": "5Ss", "s": "5Ss", "8": "8B", "b": "8B", "2": "2Zz", "z": "2Zz"}


def _find_pattern(find):
    parts = []
    for ch in re.sub(r"\s+", "", find):
        alike = _OCR_ALIKE.get(ch.lower())
        parts.append("[" + re.escape(alike) + "]" if alike else re.escape(ch))
    return re.compile(r"\s*".join(parts), re.I)


def _edge_softness(a):
    return float(((a > 0.15) & (a < 0.85)).mean() / max(1e-3, (a >= 0.5).mean()))


def _analyze_line(orig, box, text, only=None, respace=True):
    """Measure one text line: ink box, colours, font, size, spacing, pen origin, blur."""
    import numpy as np, cv2
    gray = cv2.cvtColor(orig, cv2.COLOR_RGB2GRAY); H, W = gray.shape
    b = [max(0, box[0]), max(0, box[1]), min(W, box[2]), min(H, box[3])]
    tb, ink = _tight_box(gray, b)
    x0, y0, x1, y1 = [int(v) for v in tb]
    if ink is None or x1 - x0 < 3 or y1 - y0 < 3: raise ValueError("Couldn't find the text pixels there.")
    raw = text
    text = _respace(raw, ink) if respace else raw
    crop = orig[y0:y1, x0:x1]
    tc, bc, bstd = _style_from_pixels(crop, ink)
    strength = _ink_strength(crop, tc, bc)
    m = match_font(strength, text, only=only)
    if not m: raise ValueError("No usable fonts found on the server.")
    if respace and text != raw:                                # second pass: real word gaps, matched font's widths
        fpath, fsize = m[1]["path"], m[2]; raw_ns = raw.replace(" ", "")
        ibr = _ink_box(fpath, fsize, raw_ns); natw = max(1.0, (ibr[2] - ibr[0]) if ibr else 1.0)
        gaps = _ink_gaps(ink); gapw = sum(e - s for s, e in gaps)
        good = _word_gap_filter(gaps, raw_ns, fpath, fsize, max(0.3, (ink.shape[1] - gapw) / natw), ink.shape[0])
        t2 = _respace(raw, ink, _tt(fpath, 100), gaps=good) if good else raw
        if t2 != text:
            m2 = match_font(strength, t2, only=only)
            if m2 and m2[0] >= m[0] - 0.01: text, m = t2, m2
    score, fd, size, sx, weight, blur_sig = m
    path = fd["path"]; f = _tt(path, size)
    ib = _ink_box(path, size, text) or f.getbbox(text, anchor="ls")
    nat = max(1.0, ib[2] - ib[0]); ratio = (x1 - x0) / nat
    stretch, tracking = 1.0, 0.0
    if abs(ratio - 1) > 0.025:
        if abs(ratio - 1) <= 0.14 and len(text) > 2: tracking = (x1 - x0 - nat) / max(1, len(text) - 1)
        else: stretch = ratio
    crisp = _render_mask(path, text, x1 - x0, y1 - y0)
    soft = _edge_softness(strength) / max(1e-3, _edge_softness(crisp[0])) if crisp else 1
    if soft < 1.25: blur_sig = 0.0
    pen0 = x0 - ib[0] * stretch                                    # x of the line's pen origin
    base = y0 - ib[1]                                              # baseline y
    segs = None                                                # measured ink extent of every word
    nsp = len([w_ for w_ in text.split(" ") if w_]) - 1
    if nsp >= 1:
        allg = _ink_gaps(ink)
        if len(allg) >= nsp:
            gg = sorted(sorted(allg, key=lambda g_: g_[0] - g_[1])[:nsp])
            if len(allg) == nsp or min(e - s_ for s_, e in gg) > 1.4 * max(1, max([e - s_ for s_, e in allg if (s_, e) not in gg] or [0])):
                ed = [0] + [v for g_ in gg for v in g_] + [ink.shape[1]]
                segs = [(x0 + ed[2 * i], x0 + ed[2 * i + 1]) for i in range(nsp + 1)]
    elif nsp == 0:
        segs = [(x0, x1)]
    return dict(segs=segs, x0=x0, y0=y0, x1=x1, y1=y1, ink=ink, strength=strength, tc=tc, bc=bc, bstd=bstd, text=text, fd=fd,
                path=path, size=size, font=f, stretch=stretch, tracking=tracking, pen0=pen0, base=base,
                weight=weight, blur=blur_sig, score=score, sub=_subpixel_aa(crop, strength) and bstd < 6)


def _pen_x(L, i):
    f, t = L["font"], L["text"]
    return L["pen0"] + (f.getlength(t[:i]) + i * L["tracking"]) * L["stretch"]


def _snap(L, x, lo, hi):
    """Move a cut position to the emptiest column nearby so no glyph gets sliced."""
    import numpy as np
    col = L["strength"].sum(axis=0); rad = max(2, int(L["size"] * 0.22))
    a, b = max(lo, int(x) - rad), min(hi, int(x) + rad)
    if b <= a: return int(round(x))
    rel = np.arange(a, b) - L["x0"]
    rel = rel[(rel >= 0) & (rel < len(col))]
    if not len(rel): return int(round(x))
    cost = col[rel] + 0.02 * np.abs(rel + L["x0"] - x)
    return int(rel[int(np.argmin(cost))] + L["x0"])


def _render_run(L, text, ss=1):
    """Coverage map (H×W×3 floats) for `text` drawn like the analysed line, plus its pen offset."""
    import numpy as np, cv2
    size, path = L["size"], L["path"]
    f = _tt(path, size * ss) if ss > 1 else L["font"]
    m0 = int(size * ss * 0.6)
    w = int(f.getlength(text) + len(text) * abs(L["tracking"]) * ss + 2 * m0)
    h = int(size * ss * 1.6 + 2 * m0)
    layer = Image.new("L", (max(1, w), max(1, h)), 0)
    _draw_line(ImageDraw.Draw(layer), m0, m0 + size * ss, text, f, L["tracking"] * ss)
    A = np.asarray(layer, np.float32) / 255.0
    if ss > 1:
        A = cv2.resize(A, (A.shape[1], A.shape[0] // 3), interpolation=cv2.INTER_AREA)
        A = cv2.filter2D(A, -1, np.array([[1, 2, 3, 2, 1]], np.float32) / 9)
        A = A[:, :A.shape[1] // 3 * 3]; A = np.stack([A[:, 0::3], A[:, 1::3], A[:, 2::3]], axis=-1)
        ox, oy = m0 / 3, m0 / 3 + size
    else:
        if L["stretch"] != 1.0:
            A = cv2.resize(A, (max(1, int(A.shape[1] * L["stretch"])), A.shape[0]), interpolation=cv2.INTER_AREA if L["stretch"] < 1 else cv2.INTER_CUBIC)
        A = np.repeat(A[..., None], 3, axis=-1)
        ox, oy = m0 * L["stretch"], m0 + size
    if L["blur"]: A = cv2.GaussianBlur(A, (0, 0), L["blur"])
    if L["weight"] > 1.3 and size > 18: A = cv2.dilate(A, np.ones((2, 2), np.uint8))
    adv = (f.getlength(text) / ss + len(text) * L["tracking"]) * L["stretch"]
    return np.clip(A, 0, 1), ox, oy, adv


def _blend(rgb, A, ox, oy, color):
    import numpy as np
    H, W = rgb.shape[:2]; hh, ww = A.shape[:2]
    ox, oy = int(round(ox)), int(round(oy))
    ax0, ay0 = max(0, -ox), max(0, -oy); bx0, by0 = max(0, ox), max(0, oy)
    bx1, by1 = min(W, ox + ww), min(H, oy + hh)
    if bx1 <= bx0 or by1 <= by0: return (0, 0, 0, 0)
    Aw = A[ay0:ay0 + (by1 - by0), ax0:ax0 + (bx1 - bx0)]
    reg = rgb[by0:by1, bx0:bx1].astype(np.float32)
    rgb[by0:by1, bx0:bx1] = np.clip(reg * (1 - Aw) + np.array(color, np.float32) * Aw, 0, 255).astype(np.uint8)
    return (bx0, by0, bx1, by1)


def _adv(L, s):
    return (L["font"].getlength(s) + len(s) * L["tracking"]) * L["stretch"]


def edit_line(img, box, line_text, start, end, new_sub, color=None, jpeg_hint=None):
    """Replace characters [start:end] of an OCR'd line with new_sub. Everything else keeps its original
    pixels; text after the edit slides so spacing stays natural. The font is measured on the words being
    edited (so 'Label: value' lines keep both styles)."""
    import numpy as np, cv2
    src = img.convert("RGBA"); alpha = src.getchannel("A")
    orig = np.array(src.convert("RGB")); rgb = orig.copy(); H, W = rgb.shape[:2]
    L = _analyze_line(orig, box, line_text)
    t = L["text"]
    if line_text != t:
        start, end = _map_idx(line_text, t, start), _map_idx(line_text, t, end, is_start=False)
    x0, y0, x1, y1 = L["x0"], L["y0"], L["x1"], L["y1"]
    # context = the words around the edit (≥4 chars) → measure the font there
    ws = t.rfind(" ", 0, start) + 1; we = t.find(" ", end); we = len(t) if we < 0 else we
    if we - ws < 4:
        nxt = t.find(" ", we + 1); prv = t.rfind(" ", 0, max(0, ws - 1)) + 1 if ws > 0 else 0
        if we < len(t): we = len(t) if nxt < 0 else nxt
        elif ws > 0: ws = prv
    C = L
    if (ws, we) != (0, len(t)):
        mid = lambda i: (_pen_x(L, i) + _pen_x(L, i + 1)) / 2     # centre of a space character
        segs = L.get("segs"); wi0, wi1 = t[:ws].count(" "), t[:we].count(" ")
        if segs and wi1 < len(segs) and t.count(" ") == len(segs) - 1:   # real word gaps measured on the ink
            cx0 = x0 if ws == 0 else int((segs[wi0 - 1][1] + segs[wi0][0]) / 2)
            cx1 = x1 if we >= len(t) else int((segs[wi1][1] + segs[wi1 + 1][0]) / 2) + 1
        else:
            cx0 = x0 if ws == 0 else _snap(L, mid(ws - 1), x0, x1)
            cx1 = x1 if we >= len(t) else _snap(L, mid(we), cx0, x1)
        try:
            cbox = [cx0 - 1, y0 - 1, cx1 + 1, y1 + 1]
            C = _analyze_line(orig, cbox, t[ws:we], respace=False)
            same = [f for f in _font_list() if f["family"] == L["fd"]["family"]]
            CL = _analyze_line(orig, cbox, t[ws:we], only=same, respace=False)   # line's family, best weight, on these words
            if CL["score"] >= C["score"] - 0.06: C = CL                # consistent with the line unless clearly different
        except Exception:
            C = L; ws, we = 0, len(t)
    tc = tuple(int(c * 255) for c in _rgb(color)) if color else C["tc"]
    cpen = lambda i: C["pen0"] + _adv(C, t[ws:i])               # pen x of char i (inside context)
    ct = t[ws:we]; cseg = C.get("segs") if C.get("segs") and len(C["segs"]) == ct.count(" ") + 1 else None
    s_word = start == ws or t[start - 1] == " "; e_word = end == we or (end < len(t) and t[end] == " ")
    xs = C["x0"] if start <= ws else _snap(C, cpen(start), C["x0"], C["x1"])
    xe = C["x1"] if end >= we else _snap(C, cpen(end), xs, C["x1"])
    if end >= we and we < len(t): xe = C["x1"]
    if cseg and s_word: xs = cseg[ct[:start - ws].count(" ")][0]    # measured ink edges beat font arithmetic
    if cseg and e_word: xe = max(xs + 1, cseg[ct[:max(0, end - ws - 1)].count(" ")][1])
    pad = max(2, int((y1 - y0) * 0.35))
    by0, by1 = max(0, y0 - pad), min(H, y1 + pad)
    tail_text = t[end:]
    ss = 3 if C["sub"] else 1
    if new_sub: A, ax, ay, adv_new = _render_run(C, new_sub, ss)
    else: A, ax, ay, adv_new = None, 0, 0, 0
    adv_old = _adv(C, t[start:end])
    pen_s = cpen(start)
    if new_sub and (start == ws or (cseg and s_word)) and new_sub[0] != " ":   # word start: align ink edges
        ib_new = _ink_box(C["path"], C["size"], new_sub)
        pen_s = xs - (ib_new[0] if ib_new else 0) * C["stretch"]
    delta = int(round(adv_new - adv_old)) if tail_text.strip() else 0
    tail, st = None, None
    if delta:
        tail = orig[by0:by1, xe:min(W, x1 + pad)].copy()
        hb = by1 - by0; wx0, wx1 = max(0, xe - 2 * hb), min(W, x1 + pad + 2 * hb)
        st = _local_strength(orig[by0:by1, wx0:wx1], L["tc"], k=min(255, (int(hb * 2.2) // 2) * 2 + 1))[:, xe - wx0:xe - wx0 + tail.shape[1]]
        st[y0 - by0:y1 - by0, :min(tail.shape[1], x1 - xe)] = np.maximum(st[y0 - by0:y1 - by0, :min(tail.shape[1], x1 - xe)], L["strength"][:, xe - x0:x1 - x0][:, :tail.shape[1]])
    ex1 = x1 if tail is not None else xe
    lpad = 0 if start > 0 else pad                             # cut sits in an empty gap → no need to bleed left
    rpad = pad if (end >= len(t) or tail is not None) else 0
    mx0, mx1 = max(0, xs - lpad), min(W, ex1 + rpad)
    hb = by1 - by0; wx0, wx1 = max(0, mx0 - 2 * hb), min(W, mx1 + 2 * hb)
    ls = _local_strength(orig[by0:by1, wx0:wx1], L["tc"], k=min(255, (int(hb * 2.2) // 2) * 2 + 1))[:, mx0 - wx0:mx1 - wx0]
    mask = np.zeros((H, W), np.uint8)
    mask[by0:by1, mx0:mx1] = (ls > 0.1).astype(np.uint8)
    mask[y0:y1, xs:ex1] |= (L["strength"][:, xs - x0:ex1 - x0] > 0.06).astype(np.uint8)
    mask = cv2.dilate(mask, np.ones((3, 3), np.uint8), iterations=max(1, (y1 - y0) // 10))
    if lpad == 0:                                              # never touch the characters before the edit,
        n_, lab, stats, _c = cv2.connectedComponentsWithStats((L["strength"] > 0.4).astype(np.uint8))
        rx = xs - x0; prev_end = -1                            # but clean the whole gap up to them (JPEG ringing)
        for ci in range(1, n_):
            l_, w_, area = stats[ci, 0], stats[ci, 2], stats[ci, 4]
            if area >= 3 and l_ + w_ / 2 < rx - 1: prev_end = max(prev_end, min(rx, l_ + w_))
        cut = max(0, x0 + prev_end + 2) if prev_end >= 0 else xs
        cut = min(cut, xs)
        if cut < xs: mask[by0:by1, cut:xs] = 1
        mask[:, :cut] = 0
    rx0, rx1 = max(0, xs - 2 * pad), min(W, ex1 + 2 * pad)
    ring = np.zeros((H, W), bool); ring[by0:by1, rx0:rx1] = True
    ring &= cv2.dilate(mask, np.ones((5, 5), np.uint8)) == 0
    anyink = _local_strength(orig[by0:by1, rx0:rx1], L["tc"], k=min(255, (int((by1 - by0) * 2.2) // 2) * 2 + 1)) > 0.06
    anyink = cv2.dilate(anyink.astype(np.uint8), np.ones((5, 5), np.uint8)) > 0
    ring[by0:by1, rx0:rx1] &= ~anyink                           # background only — no letters from neighbours
    rp = orig[ring].astype(np.float32)
    bg_std = float(rp.std(axis=0).mean()) if len(rp) > 30 else L["bstd"]
    bg_med = np.median(rp, axis=0) if len(rp) > 30 else np.array(L["bc"], np.float32)
    flat = bg_std < 4.0
    if flat:                                                   # plain UI / paper colour → exact fill
        rgb[mask > 0] = bg_med.astype(np.uint8)
    else:
        valid = np.ones((H, W), bool)
        vy0, vy1 = max(0, by0 - 3 * hb), min(H, by1 + 3 * hb); vx0, vx1 = max(0, rx0 - 3 * hb), min(W, rx1 + 3 * hb)
        reg = orig[vy0:vy1, vx0:vx1]
        ink_all = _local_strength(reg, L["tc"], k=min(255, (int(hb * 2.2) // 2) * 2 + 1)) > 0.3   # glyph cores only —
        dk = 2 * max(2, int(round(L["blur"] * 2 + (y1 - y0) * 0.06))) + 1                          # light grain stays "background"
        valid[vy0:vy1, vx0:vx1] = cv2.dilate(ink_all.astype(np.uint8), np.ones((dk, dk), np.uint8)) == 0
        filled, rstd = _texture_fill(orig, mask, valid, hb)
        if rstd < 14 or FEATURES.get("lama") is False:         # gradient / grain / soft texture → smooth field + real grain
            rgb = filled; FEATURES["fill"] = "texture"
        else:                                                  # real structure (photo) → LaMa on GPU
            rgb = inpaint(rgb, mask); FEATURES["fill"] = "lama"
    clean = rgb.copy(); boxes = []
    if new_sub:
        boxes.append(_blend(rgb, A, pen_s - ax, C["base"] - ay, tc))
    if tail is not None:                                       # put the rest of the line back, shifted
        tx = xe + delta; ww = tail.shape[1]; tx0, tx1 = max(0, tx), min(W, tx + ww)
        if tx1 > tx0:
            a_ = st[:, tx0 - tx:tx1 - tx][..., None]
            reg = rgb[by0:by1, tx0:tx1].astype(np.float32); px = tail[:, tx0 - tx:tx1 - tx].astype(np.float32)
            if flat:
                rgb[by0:by1, tx0:tx1] = np.where(a_ > 0.02, px, reg).astype(np.uint8)
            else:                                              # unmix: recover ink colour, recomposite on the new background
                bg_old = np.array(L["bc"], np.float32)
                ink_col = np.where(a_ > 0.05, (px - (1 - a_) * bg_old) / np.maximum(a_, 0.05), np.array(L["tc"], np.float32))
                rgb[by0:by1, tx0:tx1] = np.clip(reg * (1 - a_) + np.clip(ink_col, 0, 255) * a_, 0, 255).astype(np.uint8)
            boxes.append((tx0, by0, tx1, by1))
    sig = _noise_sigma(orig, ring)
    if sig > 1.2:
        for (a0, b0, a1, b1) in boxes:
            if a1 > a0:
                diff = np.abs(rgb[b0:b1, a0:a1].astype(np.float32) - clean[b0:b1, a0:a1]).max(axis=-1, keepdims=True)
                k = np.clip(diff / 40, 0, 1)
                rgb[b0:b1, a0:a1] = np.clip(rgb[b0:b1, a0:a1] + np.random.normal(0, sig, rgb[b0:b1, a0:a1].shape) * k, 0, 255).astype(np.uint8)
    out = Image.fromarray(rgb); q = None
    if jpeg_hint or (not flat and sig > 0.8):
        right = max(ex1, int(pen_s + adv_new), x1 + max(0, delta)) + pad
        gx0, gy0 = max(0, (max(0, xs - pad) // 8) * 8), max(0, (by0 // 8) * 8)
        gx1, gy1 = min(W, -(-min(W, right) // 8) * 8), min(H, -(-by1 // 8) * 8)
        if gx1 > gx0 and gy1 > gy0:
            q, err = _jpeg_quality(Image.fromarray(orig[gy0:gy1, gx0:gx1]))
            if q and q < 97 and (jpeg_hint or err < 1.2):
                buf = io.BytesIO(); out.crop((gx0, gy0, gx1, gy1)).save(buf, "JPEG", quality=q)
                out.paste(Image.open(buf), (gx0, gy0))
    out.putalpha(alpha)
    fd = C["fd"]
    return out, {"font": f"{fd['family']} {fd['weight']}{' italic' if fd['italic'] else ''}", "size_px": round(C["size"], 1),
                 "match": round(float(C["score"]), 3), "color": "#%02x%02x%02x" % tc, "cleartype": C["sub"], "jpeg_q": q,
                 "shift_px": delta, "fill": "flat" if flat else FEATURES.get("fill"), "lama": FEATURES.get("lama")}


def _map_idx(a, b, i, is_start=True):
    """Map an index in string a to the same character position in b (b = a with extra spaces)."""
    j = 0; k = 0
    while k < i and j < len(b):
        if a[k] == b[j]: k += 1; j += 1
        elif b[j] == " ": j += 1
        else: k += 1; j += 1
    if is_start:
        while j < len(b) and b[j] == " " and (i >= len(a) or a[i] != " "): j += 1
    return j


def _char_diff(old, new):
    p = 0
    while p < min(len(old), len(new)) and old[p] == new[p]: p += 1
    s = 0
    while s < min(len(old), len(new)) - p and old[-1 - s] == new[-1 - s]: s += 1
    return p, len(old) - s, new[p:len(new) - s]


def smart_retext(img, box, old_line, new_line, color=None, jpeg_hint=None):
    """Editor entry: the user typed a whole new line; only the changed characters are redrawn."""
    if "\n" in str(new_line):
        return retext_image(img, box, new_line, orig_text=old_line, color=color, jpeg_hint=jpeg_hint)
    a, b_, rep = _char_diff(old_line, str(new_line))
    if a == b_ and not rep and not color: return img, {"font": None, "unchanged": True}
    if color and a == b_ and not rep: a, b_, rep = 0, len(old_line), old_line
    return edit_line(img, box, old_line, a, b_, rep, color=color, jpeg_hint=jpeg_hint)


def replace_in_image(img, find, rep, jpeg_hint=None):
    pat = _find_pattern(find); n = 0; fonts = set()
    for bx in ocr_boxes(img):
        if not pat.search(bx["text"]): continue
        matches = list(pat.finditer(bx["text"]))
        for mt in reversed(matches):                           # right-to-left keeps earlier positions valid
            img, info = edit_line(img, bx["bbox"], bx["text"], mt.start(), mt.end(), rep, jpeg_hint=jpeg_hint)
            n += 1; fonts.add(info.get("font") or "")
    return img, n, ", ".join(sorted(f for f in fonts if f))


# ---------- PDFs: images inside pages ----------
def _scan_images(page, min_frac=0.0):
    pr = page.rect; out = []
    for info in page.get_image_info(xrefs=True):
        bb = fitz.Rect(info["bbox"]) & pr
        if bb.is_empty or not info.get("xref") or (info.get("width") or 0) <= 2 or (info.get("height") or 0) <= 2: continue
        if bb.get_area() / pr.get_area() >= min_frac: out.append((info["xref"], bb))
    return sorted(out, key=lambda t: -t[1].get_area())


def _is_jpeg_xref(doc, xref):
    try: return doc.extract_image(xref).get("ext") in ("jpeg", "jpg", "jpx")
    except Exception: return False


def scan_replace_doc(doc, find, rep):
    total, fonts = 0, set()
    for page in doc:
        for xref, bb in _scan_images(page, 0.02):
            try:
                im = _image_pil(doc, xref)
            except Exception:
                continue
            out, n, f = replace_in_image(im, find, rep, jpeg_hint=_is_jpeg_xref(doc, xref))
            if n:
                _remove_image(doc, page, xref, bb); _insert_pil(page, bb, out, keep=False)
                total += n; fonts.add(f)
    return total, ", ".join(sorted(x for x in fonts if x))


def pdf_ocr_layout(doc, pno, img_bbox=None):
    page = doc[pno]; imgs = _scan_images(page)
    if not imgs: return {"boxes": [], "image": None}
    if img_bbox:
        want = _rect_of(img_bbox)
        xref, bb = max(imgs, key=lambda t: (t[1] & want).get_area())
    else:
        xref, bb = imgs[0]
    im = _image_pil(doc, xref)
    kx, ky = bb.width / im.width, bb.height / im.height
    import numpy as np, cv2
    rgb = np.array(im.convert("RGB")); gray = cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY)
    out = []
    for i, bx in enumerate(ocr_boxes(im)):
        tb, ink = _tight_box(gray, bx["bbox"])
        x0, y0, x1, y1 = [int(v) for v in tb]
        if ink is not None and ink.size:
            tc, bc, _ = _style_from_pixels(rgb[y0:y1, x0:x1], ink)
        else:
            tc, bc = (0, 0, 0), (255, 255, 255)
        pb = [bb.x0 + bx["bbox"][0] * kx, bb.y0 + bx["bbox"][1] * ky, bb.x0 + bx["bbox"][2] * kx, bb.y0 + bx["bbox"][3] * ky]
        out.append({"id": f"s{i}", "scan": True, "text": _respace(bx["text"], ink) if ink is not None else bx["text"],
                    "bbox": [round(v, 2) for v in pb], "size": round((y1 - y0) * ky / 0.72, 1),
                    "color": "#%02x%02x%02x" % tc, "bg": "#%02x%02x%02x" % bc,
                    "family": "sans", "bold": False, "italic": False, "font": "image text"})
    return {"boxes": out, "image": [round(v, 2) for v in bb]}


def scan_text_op(doc, page, o):
    target = _rect_of(o["bbox"])
    imgs = _scan_images(page)
    if not imgs: raise ValueError("No image under that text.")
    xref, bb = max(imgs, key=lambda t: (t[1] & target).get_area())
    im = _image_pil(doc, xref); jh = _is_jpeg_xref(doc, xref)
    kx, ky = im.width / bb.width, im.height / bb.height
    pbox = [(target.x0 - bb.x0) * kx, (target.y0 - bb.y0) * ky, (target.x1 - bb.x0) * kx, (target.y1 - bb.y0) * ky]
    if o.get("orig"):
        out, info = smart_retext(im, pbox, o["orig"], o.get("text", ""), color=o.get("color"), jpeg_hint=jh)
    else:
        out, info = retext_image(im, pbox, o.get("text", ""), color=o.get("color"), jpeg_hint=jh)
    _remove_image(doc, page, xref, bb); _insert_pil(page, bb, out, keep=False)
    return info


def magic_erase_op(doc, page, o):
    """Paint-to-remove (objects, people, stamps, watermarks, text) with LaMa."""
    import numpy as np, cv2
    pts = [(float(x), float(y)) for x, y in o.get("points", [])]
    r = float(o.get("radius") or 12)
    kind = "brush"
    if o.get("rect"):
        rr = _rect_of(o["rect"]); pts = [(rr.x0, rr.y0), (rr.x1, rr.y1)]; kind = "rect"
    if not pts: return None
    xs, ys = [p[0] for p in pts], [p[1] for p in pts]
    area = fitz.Rect(min(xs) - r, min(ys) - r, max(xs) + r, max(ys) + r) & page.rect
    imgs = [(x, b) for x, b in _scan_images(page) if (b & area).get_area() > 0.5 * area.get_area()]

    def build_mask(Wd, Ht, ox, oy, kx, ky):
        mm = np.zeros((Ht, Wd), np.uint8)
        P = [(int((px - ox) * kx), int((py - oy) * ky)) for px, py in pts]
        if kind == "rect":
            cv2.rectangle(mm, P[0], P[1], 255, -1)
        else:
            rad = max(2, int(r * kx))
            for a_, b_ in zip(P, P[1:]): cv2.line(mm, a_, b_, 255, rad * 2)
            for p_ in P: cv2.circle(mm, p_, rad, 255, -1)
        return mm

    if imgs:                                                    # erase inside the photo itself
        xref, bb = max(imgs, key=lambda t: (t[1] & area).get_area())
        im = _image_pil(doc, xref); kx, ky = im.width / bb.width, im.height / bb.height
        alpha = im.getchannel("A") if "A" in im.getbands() else None
        rgb = np.array(im.convert("RGB"))
        res = Image.fromarray(smart_fill(rgb, build_mask(rgb.shape[1], rgb.shape[0], bb.x0, bb.y0, kx, ky)))
        if alpha: res.putalpha(alpha)
        _remove_image(doc, page, xref, bb); _insert_pil(page, bb, res, keep=False)
    else:                                                       # vector page: rasterise the spot, erase, lay it back
        clip = (area + (-3 * r, -3 * r, 3 * r, 3 * r)) & page.rect
        z = 3.0; pix = page.get_pixmap(matrix=fitz.Matrix(z, z), clip=clip, alpha=False)
        rgb = np.array(Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB"))
        mm = build_mask(rgb.shape[1], rgb.shape[0], clip.x0, clip.y0, z, z)
        res = Image.fromarray(smart_fill(rgb, mm))
        for w_ in page.get_text("words"):
            wr = fitz.Rect(w_[:4]); cx, cy = int(((wr.x0 + wr.x1) / 2 - clip.x0) * z), int(((wr.y0 + wr.y1) / 2 - clip.y0) * z)
            if 0 <= cy < mm.shape[0] and 0 <= cx < mm.shape[1] and mm[cy, cx]: page.add_redact_annot(wr, fill=False)
        _apply_redactions(page, **_redact_kw())
        _insert_pil(page, clip, res, keep=False)
    return {"lama": FEATURES.get("lama")}


# ---------- chat tools ----------
@tool("image_replace_text", "Image", "Replace text in image",
      "Change or remove words inside a photo, screenshot, poster or scanned page so it looks untouched (OCR, AI background rebuild, font matching). Leave 'replace' empty to remove the text.",
      ["image", "pdf"], params=[P("find", "Find", default=""), P("replace", "Replace with (empty = remove)", default="")])
def t_image_replace_text(recs, p):
    r = recs[0]; find, rep = _s(p, "find"), _s(p, "replace")
    if not find.strip(): raise ValueError("Tell me which text to find.")
    if r["kind"] == "pdf":
        d = _open_pdf(r); n, fonts = scan_replace_doc(d, find, rep)
        if not n: d.close(); raise ValueError(f"Couldn't find '{find}' in the PDF's images.")
        return _save_pdf(d, f"{_stem(r)}_edited.pdf"), f"Replaced {n} line(s) in images · {fonts or 'font matched'}"
    im = _load_img(r); ext = _src_ext(r)
    out, n, fonts = replace_in_image(im, find, rep, jpeg_hint=ext in ("jpg", "jpeg"))
    if not n: raise ValueError(f"Couldn't find '{find}' in the image (OCR read: " + "; ".join(b["text"] for b in ocr_boxes(im)[:6]) + ").")
    q = 95 if ext in ("jpg", "jpeg", "webp") else 92
    return _save_img(out, f"{_stem(r)}_edited", ext if ext != "gif" else "png", q), f"{'Removed' if not rep else 'Replaced'} {n} line(s) · {fonts or 'font matched'}"


@tool("img_ai_upscale", "Image", "AI upscale (Real-ESRGAN)", "Sharpen and enlarge a photo 2×/3×/4× with Real-ESRGAN on the GPU — restores detail instead of stretching pixels.", ["image"],
      params=[P("scale", "Scale", "select", "2", ["2", "3", "4"])])
def t_ai_upscale(recs, p):
    im = _load_img(recs[0]); s = int(_num(p, "scale", 2))
    if im.width * im.height > 2_500_000: raise ValueError("Image is already large — try 2× on a smaller crop.")
    out = ai_upscale(im, s)
    return _save_img(out, f"{_stem(recs[0])}_x{s}_ai", _src_ext(recs[0]), 95), f"{im.width}×{im.height} → {out.width}×{out.height} · Real-ESRGAN on {FEATURES.get('esrgan', 'cpu')}"

# ==========================================================
# WEB + REDDIT SEARCH (Perplexity / ChatGPT-search style)
#   plan queries with the model → search in parallel → read the best pages and Reddit threads
#   → keep only the most relevant passages → answer with numbered citations [1] [2]
# Works free out of the box (DuckDuckGo/Bing/Brave via the `ddgs` package + Reddit's public JSON).
# Optional keys (env vars or Kaggle → Add-ons → Secrets) for better results:
#   SERPER_API_KEY (Google results, 2,500 free) · BRAVE_API_KEY · TAVILY_API_KEY
#   REDDIT_CLIENT_ID + REDDIT_CLIENT_SECRET (free "script" app at reddit.com/prefs/apps — fixes Reddit blocking cloud IPs)
# ==========================================================
import math
import concurrent.futures as _cf
from urllib.parse import urlparse, quote_plus

def _ensure_web_pkgs():
    import importlib.util, subprocess, sys
    need = {"ddgs": "ddgs", "trafilatura": "trafilatura", "youtube_transcript_api": "youtube-transcript-api"}
    miss = [pkg for mod, pkg in need.items() if importlib.util.find_spec(mod) is None]
    if miss:
        print("📦 installing", " ".join(miss), "for web search…")
        subprocess.run([sys.executable, "-m", "pip", "install", "-q", *miss], check=False)


_ensure_web_pkgs()
FEATURES["web"] = True
WEB_UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36"
REDDIT_UA = "python:meow-assistant:v1.0 (personal research bot)"
SKIP_FETCH = ("youtube.com", "youtu.be", "facebook.com", "instagram.com", "tiktok.com", "x.com", "twitter.com",
              "linkedin.com", "pinterest.", "quora.com")
_SECRETS = {}


def _secret(name):
    if name in _SECRETS: return _SECRETS[name]
    v = os.environ.get(name, "").strip() or str(globals().get(name) or "").strip()
    if not v:
        try:
            from kaggle_secrets import UserSecretsClient
            v = (UserSecretsClient().get_secret(name) or "").strip()
        except Exception:
            v = ""
    _SECRETS[name] = v
    return v


def _domain(u):
    try: return urlparse(u).netloc.lower().removeprefix("www.")
    except Exception: return ""


def _norm_url(u):
    p = urlparse(u); host = p.netloc.lower().removeprefix("www.").removeprefix("old.").removeprefix("m.")
    return host + p.path.rstrip("/")


def _http():
    import requests
    s = requests.Session(); s.headers.update({"User-Agent": WEB_UA, "Accept-Language": "en-US,en;q=0.9"})
    return s


# ---------- search engines ----------
def _ddgs_text(q, n, recency):
    from ddgs import DDGS
    tl = {"day": "d", "week": "w", "month": "m", "year": "y"}.get(recency)
    out = []
    for r in DDGS(timeout=8).text(q, max_results=n, timelimit=tl, region="us-en") or []:
        if r.get("href"): out.append({"title": r.get("title", ""), "url": r["href"], "snippet": r.get("body", "")})
    return out


def _ddgs_news(q, n, recency):
    from ddgs import DDGS
    tl = {"day": "d", "week": "w", "month": "m"}.get(recency, "m")
    out = []
    for r in DDGS(timeout=8).news(q, max_results=n, timelimit=tl) or []:
        if r.get("url"): out.append({"title": r.get("title", ""), "url": r["url"], "snippet": r.get("body", ""),
                                     "date": (r.get("date") or "")[:10], "kind": "news", "site": r.get("source", "")})
    return out


def _serper(q, n, recency):
    key = _secret("SERPER_API_KEY")
    if not key: return None
    body = {"q": q, "num": n}
    if recency in ("day", "week", "month", "year"): body["tbs"] = "qdr:" + recency[0]
    r = _http().post("https://google.serper.dev/search", json=body, headers={"X-API-KEY": key}, timeout=10)
    r.raise_for_status(); j = r.json(); out = []
    if j.get("answerBox", {}).get("answer") or j.get("answerBox", {}).get("snippet"):
        ab = j["answerBox"]; out.append({"title": ab.get("title", "Google answer"), "url": ab.get("link", "https://google.com"),
                                         "snippet": ab.get("answer") or ab.get("snippet", "")})
    for it in j.get("organic", []):
        out.append({"title": it.get("title", ""), "url": it.get("link", ""), "snippet": it.get("snippet", ""), "date": it.get("date", "")})
    return [o for o in out if o["url"]]


def _brave(q, n, recency):
    key = _secret("BRAVE_API_KEY")
    if not key: return None
    params = {"q": q, "count": min(n, 20)}
    if recency in ("day", "week", "month", "year"): params["freshness"] = {"day": "pd", "week": "pw", "month": "pm", "year": "py"}[recency]
    r = _http().get("https://api.search.brave.com/res/v1/web/search", params=params,
                    headers={"X-Subscription-Token": key, "Accept": "application/json"}, timeout=10)
    r.raise_for_status()
    return [{"title": it.get("title", ""), "url": it.get("url", ""), "snippet": re.sub(r"<[^>]+>", "", it.get("description", "")),
             "date": it.get("age", "")} for it in r.json().get("web", {}).get("results", []) if it.get("url")]


def _tavily(q, n, recency):
    key = _secret("TAVILY_API_KEY")
    if not key: return None
    body = {"api_key": key, "query": q, "max_results": min(n, 10), "search_depth": "advanced"}
    if recency in ("day", "week"): body["topic"] = "news"; body["days"] = 1 if recency == "day" else 7
    r = _http().post("https://api.tavily.com/search", json=body, timeout=15)
    r.raise_for_status()
    return [{"title": it.get("title", ""), "url": it.get("url", ""), "snippet": it.get("content", ""),
             "date": (it.get("published_date") or "")[:10]} for it in r.json().get("results", []) if it.get("url")]


def web_search(q, n=8, recency=None):
    """Best available engine first; the free ddgs metasearch is the fallback."""
    errs = []
    for eng in (_serper, _brave, _tavily, _ddgs_text):
        try:
            res = eng(q, n, recency)
            if res: return res
        except Exception as e:
            errs.append(f"{eng.__name__}: {e}")
    if errs: print("⚠️  web search:", " | ".join(errs)[:300])
    return []


# ---------- Reddit (deep) ----------
# Finding threads:  official search (if keys) + Google/Serper "site:reddit.com" + Tavily (reddit.com only) + Arctic Shift archive
# Reading threads:  official API (if keys) → public JSON → Arctic Shift archive (free, no key) → Tavily Extract (your Tavily key)
ARCTIC = "https://arctic-shift.photon-reddit.com"
_REDDIT_TOKEN = {"tok": None, "exp": 0}
_BAD_BODY = ("", "[deleted]", "[removed]")


def _reddit_get(path, params=None, timeout=10):
    """GET a Reddit JSON endpoint: OAuth app if configured (works from cloud IPs), else public JSON."""
    s = _http(); s.headers["User-Agent"] = REDDIT_UA
    params = dict(params or {}, raw_json=1)
    cid, csec = _secret("REDDIT_CLIENT_ID"), _secret("REDDIT_CLIENT_SECRET")
    if cid and csec:
        if not _REDDIT_TOKEN["tok"] or time.time() > _REDDIT_TOKEN["exp"]:
            r = s.post("https://www.reddit.com/api/v1/access_token", auth=(cid, csec),
                       data={"grant_type": "client_credentials"}, timeout=timeout)
            r.raise_for_status(); j = r.json()
            _REDDIT_TOKEN.update(tok=j["access_token"], exp=time.time() + int(j.get("expires_in", 3600)) - 60)
        r = s.get("https://oauth.reddit.com" + path, params=params, headers={"Authorization": "bearer " + _REDDIT_TOKEN["tok"]}, timeout=timeout)
        r.raise_for_status(); return r.json()
    if time.time() < _REDDIT_TOKEN.get("blocked_until", 0):
        raise RuntimeError("Reddit is blocking this server right now (using the archive instead)")
    last = None
    for host in ("https://www.reddit.com", "https://old.reddit.com"):
        try:
            r = s.get(host + path.rstrip("/") + ".json", params=params, timeout=timeout)
            if r.status_code == 200 and r.headers.get("content-type", "").startswith("application/json"): return r.json()
            last = f"HTTP {r.status_code}"
            if r.status_code in (403, 429):                    # blocked for this server — stop trying for 10 minutes
                _REDDIT_TOKEN["blocked_until"] = time.time() + 600; break
        except Exception as e:
            last = str(e)
    raise RuntimeError(f"Reddit blocked the request ({last})")


def _arctic(path, params, timeout=15):
    """Arctic Shift: a free public Reddit archive (posts + full comment trees, no key)."""
    last = None
    for attempt in range(2):
        try:
            r = _http().get(ARCTIC + path, params=params, timeout=timeout)
            if r.status_code == 200:
                j = r.json()
                if isinstance(j, dict) and j.get("error"): raise RuntimeError(str(j["error"])[:120])
                return j.get("data", j) if isinstance(j, dict) else j
            last = f"HTTP {r.status_code}"
        except Exception as e:
            last = str(e)
        time.sleep(1.5)
    raise RuntimeError(f"archive unavailable ({last})")


def _date(ts):
    try: return datetime.datetime.fromtimestamp(float(ts), datetime.timezone.utc).strftime("%Y-%m-%d")
    except Exception: return ""


def _thread_of(d):
    sub = d.get("subreddit", ""); pid = d.get("id", "")
    link = d.get("permalink") or (f"/r/{sub}/comments/{pid}/" if sub and pid else "")
    return {"title": d.get("title", ""), "url": "https://www.reddit.com" + link, "kind": "reddit",
            "sub": sub, "score": d.get("score"), "comments": d.get("num_comments"),
            "date": _date(d.get("created_utc")) if d.get("created_utc") else "",
            "snippet": (d.get("selftext") or "")[:300], "id": pid}


def _tid(url):
    m = re.search(r"/comments/([a-z0-9]+)", url or ""); return m.group(1) if m else ""


def _from_hit(r):
    m = re.search(r"reddit\.com/r/([^/]+)/comments/([a-z0-9]+)", r["url"])
    return {"title": re.sub(r"\s*[:\-|]\s*r/\w+.*$", "", r.get("title", "")).strip(), "url": r["url"].split("?")[0], "kind": "reddit",
            "sub": m.group(1) if m else "", "id": m.group(2) if m else "", "score": None, "comments": None,
            "date": r.get("date", ""), "snippet": r.get("snippet", "")}


def _tavily_reddit(q, n):
    key = _secret("TAVILY_API_KEY")
    if not key: return []
    r = _http().post("https://api.tavily.com/search", json={"api_key": key, "query": q, "max_results": min(n, 10),
                     "include_domains": ["reddit.com"]}, headers={"Authorization": "Bearer " + key}, timeout=15)
    r.raise_for_status()
    return [{"title": it.get("title", ""), "url": it.get("url", ""), "snippet": it.get("content", "")} for it in r.json().get("results", [])]


def reddit_search(q, n=6, recency=None, deep=False, subs=None):
    """Find the best threads for `q` across ALL of Reddit using every source at once; threads that several
    sources agree on rank first, then engagement (comments / upvotes)."""
    t = {"day": "day", "week": "week", "month": "month", "year": "year"}.get(recency, "all")
    kw = " ".join(_terms(q)[:5])

    def official():
        j = _reddit_get("/search", {"q": q, "sort": "relevance", "t": t, "limit": 25, "type": "link"}, timeout=8)
        return [_thread_of(c["data"]) for c in j.get("data", {}).get("children", []) if not c.get("data", {}).get("over_18")]

    def google():
        out = []
        qqs = [f"{q} site:reddit.com/r/{s}" for s in (subs or [])[:2]] or [f"{q} site:reddit.com"]
        if deep and kw and not subs: qqs.append(f"{kw} reddit")
        for qq in qqs:
            out += [_from_hit(r) for r in web_search(qq, 10 if deep else 8, "year" if recency == "year" else None)
                    if "reddit.com/r/" in r["url"] and "/comments/" in r["url"]]
        return out

    def tavily():
        return [_from_hit(r) for r in _tavily_reddit(q, 10) if "/comments/" in r["url"]]

    def archive():
        if not kw: return []
        p = {"title": kw, "limit": 25, "sort": "desc", "fields": "id,title,subreddit,score,num_comments,created_utc,selftext,permalink,over_18"}
        if subs: p["subreddit"] = subs[0]
        if recency in ("day", "week", "month", "year"):
            p["after"] = (datetime.datetime.now(datetime.timezone.utc) - datetime.timedelta(days={"day": 2, "week": 8, "month": 32, "year": 370}[recency])).strftime("%Y-%m-%d")
        return [_thread_of(d) for d in (_arctic("/api/posts/search", p) or []) if isinstance(d, dict) and not d.get("over_18")]

    found = {}
    with _cf.ThreadPoolExecutor(4) as ex:
        futs = {ex.submit(fn): name for name, fn in (("official", official), ("google", google), ("tavily", tavily), ("archive", archive))}
        try:
            for f in _cf.as_completed(futs, timeout=25):
                try:
                    res = f.result() or []
                except Exception as e:
                    print(f"⚠️  reddit {futs[f]}:", str(e)[:140]); continue
                for pos, th in enumerate(res):
                    key = th.get("id") or _norm_url(th["url"])
                    if not key: continue
                    cur = found.get(key)
                    if cur is None:
                        found[key] = {**th, "_votes": 1, "_pos": pos}
                    else:
                        cur["_votes"] += 1; cur["_pos"] = min(cur["_pos"], pos)
                        for k_, v in th.items():
                            if cur.get(k_) in (None, "") and v not in (None, ""): cur[k_] = v
        except _cf.TimeoutError:
            print("⚠️  some Reddit searches were too slow — skipped")
    out = list(found.values())
    for th in out:
        th["_rank"] = th["_pos"] - 2.0 * (th["_votes"] - 1) - 0.6 * math.log1p(th.get("comments") or 0) - 0.3 * math.log1p(max(0, th.get("score") or 0))
    return sorted(out, key=lambda th: th["_rank"])[:n]


def _walk_comments(nodes, depth=0, out=None, max_depth=3):
    """Flatten a Reddit-style comment listing (official API or archive) into dicts with replies attached."""
    out = [] if out is None else out
    if isinstance(nodes, dict): nodes = nodes.get("data", {}).get("children", nodes.get("children", []))
    for n in nodes or []:
        if not isinstance(n, dict) or n.get("kind") == "more": continue
        d = n.get("data", n)
        body = (d.get("body") or "").strip()
        if body in _BAD_BODY or d.get("stickied") or d.get("author") == "AutoModerator": continue
        c = {"score": d.get("score") or 0, "body": body, "depth": depth, "replies": [], "author": d.get("author") or ""}
        rep = d.get("replies")
        if rep and depth < max_depth: _walk_comments(rep, depth + 1, c["replies"], max_depth)
        out.append(c)
    return out


_PROMO_RE = re.compile(r"(ref(erral)?[ _-]?(link|code)|use my (code|link)|promo code|invite code|\bref=|dm me|message me|inbox me|telegram|whats ?app|t\.me/|"
                       r"link in (my )?bio|join my|my (course|group|channel|discord)|guaranteed (profit|income|returns)|100% (legit|guaranteed)|"
                       r"signals? group|limited (spots|seats)|bit\.ly/|tinyurl)", re.I)
_PROOF_RE = re.compile(r"((i|i've|i have|i'm|i am|we)\s+(made|make|earn(ed)?|been (doing|using|working)|use[d]?|using|bought|own|work(ed)?|got (paid|hired|an? offer|it)|switched|tried)|"
                       r"\b(got|bought) (it|mine|one)\b|"
                       r"(₹|rs\.?|inr|\$|usd|€|£)\s?\d|\d+\s?(k|lpa|lakh|lakhs|per month|/mo(nth)?|a month|/hr|per hour)\b|"
                       r"\b\d+\s+(months?|years?|weeks?)\b|my experience|for me it|in my case|update:)", re.I)


def comment_tags(body):
    """Cheap credibility hints for the model: first-hand evidence vs promotion/scam signals."""
    tags = []
    if _PROMO_RE.search(body): tags.append("⚠ promo/scam signs")
    elif _PROOF_RE.search(body): tags.append("first-hand")
    return (" [" + ", ".join(tags) + "]") if tags else ""


def fmt_comment(c, prefix="COMMENT", maxlen=1000):
    who = f" u/{c['author']}" if c.get("author") and c["author"] not in ("[deleted]",) else ""
    return f"{prefix}{who} (▲{c['score']}){comment_tags(c['body'])}: {c['body'][:maxlen]}"


def _pick_comments(tree, k):
    """Top comments by votes, each with its best reply (and that reply's best reply) — the useful part of a thread."""
    tops = sorted(tree, key=lambda c: -c["score"])[:k]
    res = []
    for c in tops:
        item = {"score": c["score"], "body": c["body"], "replies": [], "author": c.get("author", "")}
        for r in sorted(c["replies"], key=lambda x: -x["score"])[:2]:
            rr = {"score": r["score"], "body": r["body"], "author": r.get("author", "")}
            if r["replies"]:
                b = max(r["replies"], key=lambda x: x["score"]); rr["sub"] = {"score": b["score"], "body": b["body"], "author": b.get("author", "")}
            item["replies"].append(rr)
        res.append(item)
    return res


def _thread_official(th, k):
    j = _reddit_get(f"/comments/{th['id']}", {"sort": "top", "limit": 100, "depth": 3}, timeout=8)
    post = j[0]["data"]["children"][0]["data"]
    return post, _walk_comments(j[1])


def _thread_archive(th, k):
    post = None
    try:
        ps = _arctic("/api/posts/ids", {"ids": th["id"]})
        if ps: post = ps[0]
    except Exception:
        pass
    tree = _arctic("/api/comments/tree", {"link_id": "t3_" + th["id"], "limit": 1500}, timeout=20)
    return post, _walk_comments(tree)


def reddit_thread(th, k=12):
    """Post body + the most upvoted comments with their best replies. Tries the official API, then the archive."""
    th["id"] = th.get("id") or _tid(th["url"])
    if not th["id"]: return th
    for name, fn in (("reddit", _thread_official), ("archive", _thread_archive)):
        try:
            post, tree = fn(th, k)
        except Exception as e:
            print(f"⚠️  thread via {name}:", str(e)[:120]); continue
        if post:
            th.update({k_: v for k_, v in _thread_of(post).items() if v not in (None, "")})
            th["body"] = (post.get("selftext") or "").strip()
        if tree or post:
            th["top_comments"] = _pick_comments(tree, k); th["via"] = name
            return th
    return th


def tavily_extract(urls):
    """Read pages through Tavily's servers (works when a site blocks this machine). 1 credit per 5 pages."""
    key = _secret("TAVILY_API_KEY")
    if not key or not urls: return {}
    r = _http().post("https://api.tavily.com/extract", json={"api_key": key, "urls": urls[:20], "extract_depth": "advanced"},
                     headers={"Authorization": "Bearer " + key}, timeout=40)
    r.raise_for_status()
    return {it.get("url"): it.get("raw_content") or "" for it in r.json().get("results", [])}


def read_threads(threads, k=12):
    """Read every thread in parallel; whatever Reddit and the archive can't give, Tavily reads in one batch."""
    with _cf.ThreadPoolExecutor(8) as ex:
        list(ex.map(lambda th: _safe_thread(th, k), threads))
    miss = [th for th in threads if not th.get("top_comments") and not th.get("body")]
    if miss:
        try:
            got = tavily_extract([th["url"] for th in miss])
            for th in miss:
                txt = got.get(th["url"]) or next((v for u, v in got.items() if _tid(u) == th["id"]), "")
                if txt: th["page_text"] = _clean_reddit_page(txt); th["via"] = "tavily"
        except Exception as e:
            print("⚠️  tavily extract:", str(e)[:140])
    return threads


def _safe_thread(th, k):
    try: return reddit_thread(th, k)
    except Exception as e: print("⚠️  thread:", str(e)[:120]); return th


def _clean_reddit_page(txt):
    lines = [l.strip() for l in txt.splitlines()]
    junk = re.compile(r"^(reply|share|report|save|follow|upvote|downvote|vote|award|more replies|continue this thread|sort by:?.*|"
                      r"open comment sort options|best|top|new|controversial|old|q&a|level \d+|\d+ (points?|replies|comments?)|"
                      r"\d+[hdmy] ago|•|edited.*|log in|sign up|get the app|promoted|advertise.*|privacy policy|user agreement)$", re.I)
    out = [l for l in lines if l and not junk.match(l) and len(l) > 2]
    return "\n".join(out)


def reddit_comment_search(terms, subs, per_sub=40):
    """Deep mode: comments anywhere in the relevant subreddits that mention the topic (archive search)."""
    kw = " ".join(terms[:3])
    out = []
    for sub in subs[:2]:
        try:
            res = _arctic("/api/comments/search", {"body": kw, "subreddit": sub, "limit": per_sub, "sort": "desc",
                                                   "fields": "body,score,link_id,subreddit,created_utc,permalink,author"})
        except Exception as e:
            print("⚠️  comment search:", str(e)[:120]); continue
        for d in res or []:
            body = (d.get("body") or "").strip()
            if body in _BAD_BODY or len(body) < 40: continue
            out.append({"sub": d.get("subreddit", sub), "score": d.get("score") or 0, "body": body, "date": _date(d.get("created_utc")), "author": d.get("author") or "",
                        "url": "https://www.reddit.com" + (d.get("permalink") or f"/r/{sub}/comments/{(d.get('link_id') or '')[3:]}/")})
    return out


# ---------- page reading ----------
def fetch_page(url, max_bytes=4_000_000):
    """Download a page (or PDF) and return its main text."""
    if any(s in _domain(url) for s in SKIP_FETCH): return ""
    r = _http().get(url, timeout=(5, 10), stream=True, allow_redirects=True)
    if r.status_code != 200: return ""
    ct = r.headers.get("content-type", "").lower(); data = b""
    for chunk in r.iter_content(65536):
        data += chunk
        if len(data) > max_bytes: break
    if "pdf" in ct or url.lower().endswith(".pdf"):
        with fitz.open(stream=data, filetype="pdf") as d:
            return "\n".join(p.get_text() for p in list(d)[:25])
    if "html" not in ct and "text" not in ct: return ""
    html_ = data.decode(r.encoding or "utf-8", errors="replace")
    try:
        import trafilatura
        txt = trafilatura.extract(html_, url=url, favor_recall=True, include_comments=False, include_tables=True, deduplicate=True)
        if txt and len(txt) > 200: return txt
    except Exception:
        pass
    html_ = re.sub(r"(?is)<(script|style|noscript|nav|footer|header|svg|form)[^>]*>.*?</\1>", " ", html_)
    txt = re.sub(r"(?s)<[^>]+>", "\n", html_)
    txt = html.unescape(re.sub(r"\n\s*\n+", "\n\n", re.sub(r"[ \t]+", " ", txt)))
    return txt.strip()


_STOP = set("a an the and or of to in on for with is are was were be been it this that what which who whom how why when "
            "where do does did can could should would will i you me my your we our they their at by from as about into vs "
            "please tell give show find best latest new".split())


def _terms(s):
    return [w for w in re.findall(r"[a-z0-9][a-z0-9+#.\-]*", s.lower()) if w not in _STOP and len(w) > 1]


def _chunks(text, words=140):
    paras = [p.strip() for p in re.split(r"\n\s*\n|\n(?=[-•*#\d])", text) if len(p.strip()) > 30]
    out, cur = [], []
    for p in paras:
        w = p.split()
        if len(w) > words * 1.6:                            # long paragraph → split
            for i in range(0, len(w), words): out.append(" ".join(w[i:i + words]))
            continue
        cur += w
        if len(cur) >= words: out.append(" ".join(cur)); cur = []
    if cur: out.append(" ".join(cur))
    return out


def rank_passages(query_terms, docs, budget_tokens, per_doc=4, reddit_share=0.45):
    """BM25 over all passages; every source gets its best passage before any source gets a second."""
    items = []
    for di, d in enumerate(docs):
        for pi, c in enumerate(d.get("chunks", [])): items.append((di, pi, c, _terms(c)))
    if not items: return {}
    N = len(items); avgdl = sum(len(t) for *_, t in items) / N
    df = {}
    for *_, t in items:
        for w in set(t): df[w] = df.get(w, 0) + 1
    qset = list(dict.fromkeys(query_terms))
    scored = {}
    for di, pi, c, t in items:
        tf = {}
        for w in t: tf[w] = tf.get(w, 0) + 1
        s = 0.0
        for w in qset:
            if w in tf:
                idf = math.log(1 + (N - df[w] + 0.5) / (df[w] + 0.5))
                s += idf * tf[w] * 2.2 / (tf[w] + 1.2 * (0.25 + 0.75 * len(t) / avgdl))
        s += 0.15 * (pi == 0)                               # intros usually summarise the page
        scored.setdefault(di, []).append((s, pi, c))
    for di in scored: scored[di].sort(key=lambda x: -x[0])
    picked, used = {}, 0
    for di, d in enumerate(docs):                          # Reddit: the post + its top comments, in upvote order
        if d.get("kind") != "reddit": continue
        n_red = max(1, sum(1 for x in docs if x.get("kind") == "reddit"))
        cap, got = max(250, int(budget_tokens * reddit_share / n_red * 1.3)), []
        for pi, c in enumerate(d.get("chunks", [])):
            t = len(c) // 4 + 8
            if got and (sum(len(x) // 4 + 8 for _, x in got) + t > cap or used + t > budget_tokens * reddit_share): break
            got.append((pi, c)); used += t
        if got: picked[di] = got
        scored.pop(di, None)
    for rnd in range(per_doc):
        for di in sorted(scored, key=lambda d: docs[d].get("order", d)):
            lst = scored[di]
            if rnd >= len(lst) or (rnd > 0 and lst[rnd][0] < 0.35 * lst[0][0]) or (rnd > 0 and lst[rnd][0] <= 0): continue
            c = lst[rnd][2]; t = len(c) // 4 + 8
            if used + t > budget_tokens: continue
            picked.setdefault(di, []).append((lst[rnd][1], c)); used += t
    return {di: [c for _, c in sorted(v)] for di, v in picked.items()}


# ---------- planning ----------
AUTO_SEARCH_RE = re.compile(
    r"\b(latest|news|today|tonight|yesterday|tomorrow|this (week|month|year)|current(ly)?|right now|recent(ly)?|price|prices|cost|"
    r"rate|stock|share price|score|weather|release[sd]?|launch(ed)?|update[sd]?|version|who (is|won|are|was)|when (is|does|will|did)|"
    r"20(2[4-9]|3\d)|search|google|look ?up|find (me|out)|reddit|reviews?|vs\.?|versus|compare|best|recommend|worth it|"
    r"trending|election|match|ipl|budget|specs|benchmark|deadline|schedule|result|ranking|available|near me|"
    r"aaj|abhi|kab|kitne|kaun)\b", re.I)


def _llm_text(prompt, max_tokens=200, temperature=0.2):
    out = ""
    for ch in llm.create_completion(prompt=prompt, max_tokens=max_tokens, stream=True, temperature=temperature, top_p=0.9,
                                    stop=["<|im_end|>", "<|endoftext|>", "<|im_start|>"]):
        out += ch["choices"][0].get("text", "")
    return out


def plan_search(user, history, mode):
    """Ask the model for focused search queries (resolves 'it'/'that' from the chat). Heuristic fallback."""
    today = datetime.datetime.now().strftime("%d %B %Y")
    ctx = ""
    for m in history[-4:]:
        c = THINK_RE.sub("", str(m.get("content") or "")).strip()
        if c: ctx += f"{m.get('role')}: {c[:500]}\n"
    nq = 4 if mode == "deep" else 3
    sys_ = (f"You plan web searches for an AI assistant. Today is {today}. Reply with ONLY one JSON object, no prose:\n"
            '{"search": true|false, "queries": ["..."], "reddit": true|false, "recency": "day|week|month|year|any", "sources": [], "images": true|false}\n'
            "- search=true when a good answer needs current, factual, niche or checkable info (news, prices, people, products, "
            "releases, sports, weather, docs, statistics, reviews). false for chit-chat, jokes, creative writing, maths, or code the user pasted.\n"
            f"- 1-{nq} short keyword queries like a human would type into Google. Resolve pronouns from the chat. Add the year for time-sensitive topics. Use English unless the topic is local.\n"
            "- reddit=true for opinions, experiences, recommendations, 'is X worth it', comparisons, troubleshooting, or when the user mentions Reddit.\n"
            "- recency: how fresh results must be.\n"
            "- sources: extra specialist sources worth searching, any of: \"youtube\" (tutorials, how-to, talks, product reviews, "
            "music/video topics), \"wiki\" (definitions, concepts, history, people, places, organisations), \"papers\" (science, "
            "medicine, ML/AI research, academic questions), \"code\" (programming errors, libraries, APIs, repos — searches GitHub + "
            "Stack Overflow), \"hn\" (tech industry, startups, developer tools discussions), \"news\" (current events). Usually 0-2.\n"
            "- images=true when seeing pictures helps (what something looks like, designs, places, products, people, animals)."
            + ("\n- This is a REDDIT search: write 2-3 different queries phrased like Reddit post titles or how people ask there "
               "(e.g. 'is X worth it', 'X vs Y', 'best X for Y', 'X experience'), not SEO keywords." if mode == "reddit" else "")
            + "\n- For 'which is best / what do people actually use / how do people make money / is X legit' questions, make one of the "
              "queries hunt for warnings and proof (e.g. 'X scam', 'X legit experience', 'X income proof'), so bad options get exposed.")
    p = f"<|im_start|>system\n{sys_}<|im_end|>\n<|im_start|>user\n{('Chat so far:' + chr(10) + ctx + chr(10)) if ctx else ''}Latest message: {user[:1500]}<|im_end|>\n<|im_start|>assistant\n<think>\n\n</think>\n\n"
    plan = None
    try:
        raw = _llm_text(p, 220)
        m = re.search(r"\{.*\}", raw, re.S)
        if m: plan = json.loads(m.group(0))
    except Exception as e:
        print("⚠️  search planner:", str(e)[:120])
    if isinstance(plan, dict) and plan.get("search") is False and mode == "auto":
        return {"search": False, "queries": [], "reddit": False, "recency": None}
    if not isinstance(plan, dict) or not plan.get("queries"):
        q = re.sub(r"\s+", " ", user).strip()[:200]
        plan = {"search": True, "queries": [q], "reddit": bool(re.search(r"reddit|worth|recommend|review|vs|best|experience", user, re.I)), "recency": "any"}
    plan["queries"] = [str(q).strip()[:200] for q in plan.get("queries", []) if str(q).strip()][:nq]
    srcs = plan.get("sources") if isinstance(plan.get("sources"), list) else []
    plan["sources"] = [x for x in dict.fromkeys(str(v).lower() for v in srcs) if x in ("youtube", "wiki", "papers", "code", "hn", "news")]
    if mode == "deep" and "wiki" not in plan["sources"] and not plan.get("recency"): plan["sources"].append("wiki")
    if mode == "reddit": plan["sources"] = []
    plan["images"] = bool(plan.get("images")) and mode != "reddit"
    if mode in ("web", "deep", "reddit"): plan["search"] = True
    if mode == "reddit": plan["reddit"] = True
    if mode == "deep": plan["reddit"] = True
    if mode == "web" and re.search(r"\breddit\b", user, re.I): plan["reddit"] = True
    rc = str(plan.get("recency", "any")).lower()
    plan["recency"] = rc if rc in ("day", "week", "month", "year") else None
    return plan


WEB_RULES = """
WEB SEARCH — results for this message are in <web_results>, fetched moments ago.
- Base the answer on them. Cite sources inline right after the claim they support, using their numbers like [1] or [2][4]. Only cite numbers that exist. No "Sources" list at the end — the app shows the source cards.
- Prefer the newest and most authoritative sources for anything time-sensitive, and mention dates when they matter. If sources disagree, say so and say which is more reliable.
- Reddit threads are people's opinions and experiences, not facts: summarise them as what Redditors say, mention the consensus and notable disagreements, and weigh by upvotes (▲). When the answer is mostly from Reddit, organise it as: the consensus → the top picks/answers and how strongly they're backed (e.g. "recommended in 4 of 6 threads") → disagreements and caveats → a couple of short real quotes in quotation marks, each cited.
- Sources marked "Link you shared" are pages the user gave you: treat them as the main subject (summarise, answer about or analyse them as asked).
- YouTube sources are video transcripts with [m:ss] timestamps — mention the timestamp when pointing to a moment. Papers: mention year and how established the finding is. Stack Overflow: prefer accepted / highly voted answers and note the version they apply to.
- RECOMMENDATIONS FROM PEOPLE ("best X", "which version", "what do people actually use", "how are people genuinely making money",
  "is X legit"): don't just repeat what was said — verify it like a sceptical friend.
  • Build a ranked list of SPECIFIC options (exact names, versions, sites, programs, tools). For each, give: how many DIFFERENT users
    (u/names) and threads back it, the upvotes behind it, the concrete first-hand evidence (real numbers, how long they've done it,
    results), and its main downside.
  • Comments tagged [first-hand] carry more weight; comments tagged [⚠ promo/scam signs] (referral links, "DM me", Telegram groups,
    guaranteed income, selling a course) must NOT count as recommendations — mention them only as warnings.
  • Treat as unverified anything backed by only one person, by vague claims without numbers, or by too-good-to-be-true income.
  • Organise as: ✅ Genuinely recommended (several independent first-hand reports) → 🤔 Mixed / depends on → ⚠️ Avoid or scam
    reports. For money-making: give realistic numbers people report (ranges, time to first income) and what it takes, never hype.
- If the results don't cover the question, say that plainly, then answer from your own knowledge and mark it as such."""

# ==========================================================
# CODE INTERPRETER (like ChatGPT's): a persistent Python session per chat.
# The model writes ```meow-python blocks; the server runs them, streams the output (text, tables, charts, files)
# back into the answer, and the model continues with the real results.
# ==========================================================
import subprocess, queue, sys as _sys, shutil

PY_TIMEOUT = 90            # seconds per run
PY_MAX_ROUNDS = 4          # code runs per answer
_PY_MARK = "\x00MEOW-RESULT\x00"

_KERNEL_SRC = r'''
import sys, os, io, json, ast, traceback, contextlib, uuid, time
os.environ.setdefault("MPLBACKEND", "Agg")
MARK = "\x00MEOW-RESULT\x00"
G = {"__name__": "__main__"}
real_out = sys.stdout

def _figs(out_dir):
    paths = []
    if "matplotlib.pyplot" in sys.modules:
        plt = sys.modules["matplotlib.pyplot"]
        plt.show = lambda *a, **k: None
        for n in plt.get_fignums():
            p = os.path.join(out_dir, f"figure_{uuid.uuid4().hex[:6]}.png")
            try:
                plt.figure(n).savefig(p, dpi=130, bbox_inches="tight", facecolor="white"); paths.append(p)
            except Exception:
                pass
        plt.close("all")
    return paths

def _show(v):
    try:
        import pandas as pd
        if isinstance(v, (pd.DataFrame, pd.Series)):
            with pd.option_context("display.max_rows", 40, "display.max_columns", 20, "display.width", 160):
                return repr(v)
    except Exception:
        pass
    try:
        if getattr(v, "shape", None) == () and hasattr(v, "item"): v = v.item()     # numpy scalar → plain number
    except Exception:
        pass
    r = repr(v)
    return r if len(r) < 4000 else r[:4000] + " …"

def _snap(d):
    out = {}
    for root, _, fs in os.walk(d):
        for f in fs:
            p = os.path.join(root, f)
            try: out[p] = os.path.getmtime(p)
            except OSError: pass
    return out

for line in sys.stdin:
    try:
        req = json.loads(line)
    except Exception:
        continue
    out_dir = req.get("out_dir") or os.getcwd(); os.makedirs(out_dir, exist_ok=True)
    G["FILES"] = req.get("files", {}); G["OUT_DIR"] = out_dir
    if "matplotlib.pyplot" in sys.modules: sys.modules["matplotlib.pyplot"].show = lambda *a, **k: None
    before = _snap(out_dir)
    buf = io.StringIO(); result, err = None, None
    t0 = time.time()
    try:
        import linecache
        src = req["code"]; linecache.cache["<cell>"] = (len(src), None, src.splitlines(True), "<cell>")
        tree = ast.parse(src, filename="<cell>", mode="exec")
        last = tree.body.pop() if tree.body and isinstance(tree.body[-1], ast.Expr) else None
        with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(buf):
            exec(compile(tree, "<cell>", "exec"), G)
            if last is not None:
                v = eval(compile(ast.Expression(last.value), "<cell>", "eval"), G)
                if v is not None: result = _show(v)
    except SystemExit:
        err = "SystemExit (exit() is ignored here)"
    except BaseException:
        et, ev, tb = sys.exc_info()
        while tb is not None and tb.tb_frame.f_code.co_filename == "<string>": tb = tb.tb_next   # hide the kernel itself
        lines = "".join(traceback.format_exception(et, ev, tb)).splitlines()
        err = "\n".join(l for l in lines if 'File "<frozen' not in l)[-3000:]
    figs = _figs(out_dir)
    after = _snap(out_dir)
    new = [p for p, m in after.items() if p not in figs and (p not in before or before[p] != m)]
    real_out.write(MARK + json.dumps({"stdout": buf.getvalue()[-8000:], "result": result, "error": err, "figures": figs,
                                      "files": new, "secs": round(time.time() - t0, 2)}) + "\n")
    real_out.flush()
'''


class PyKernel:
    def __init__(self, workdir):
        self.workdir = workdir; os.makedirs(workdir, exist_ok=True)
        self.proc = None; self.q = queue.Queue(); self.used = time.time()

    def start(self):
        env = dict(os.environ, MPLBACKEND="Agg", PYTHONUNBUFFERED="1")
        self.proc = subprocess.Popen([_sys.executable, "-u", "-c", _KERNEL_SRC], stdin=subprocess.PIPE, stdout=subprocess.PIPE,
                                     stderr=subprocess.STDOUT, text=True, cwd=self.workdir, env=env, bufsize=1)
        self.q = queue.Queue()
        threading.Thread(target=self._pump, args=(self.proc, self.q), daemon=True).start()

    @staticmethod
    def _pump(proc, q):
        for line in proc.stdout:
            q.put(line)
        q.put(None)

    def alive(self):
        return self.proc is not None and self.proc.poll() is None

    def kill(self):
        if self.proc is not None:
            try: self.proc.kill()
            except Exception: pass
        self.proc = None

    def run(self, code, files, timeout=PY_TIMEOUT, stop=lambda: False):
        if not self.alive(): self.start()
        self.used = time.time()
        out_dir = os.path.join(self.workdir, "out")
        self.proc.stdin.write(json.dumps({"code": code, "files": files, "out_dir": out_dir}) + "\n"); self.proc.stdin.flush()
        stray, t0 = [], time.time()
        while True:
            left = timeout - (time.time() - t0)
            if left <= 0 or stop():
                self.kill()
                return {"stdout": "".join(stray)[-4000:], "result": None, "figures": [], "files": [], "restarted": True,
                        "error": f"Stopped: the code ran longer than {timeout}s (session restarted, variables were reset)." if left <= 0 else "Stopped by user."}
            try:
                line = self.q.get(timeout=min(0.5, left))
            except queue.Empty:
                continue
            if line is None:
                self.proc = None
                return {"stdout": "".join(stray)[-4000:], "result": None, "figures": [], "files": [], "restarted": True,
                        "error": "The Python session crashed (out of memory?). It was restarted; variables were reset."}
            if line.startswith(_PY_MARK):
                res = json.loads(line[len(_PY_MARK):])
                if stray: res["stdout"] = "".join(stray) + res.get("stdout", "")
                return res
            stray.append(line)                                # output printed by C extensions / subprocesses


_KERNELS = {}


def kernel_for(chat_id):
    chat_id = re.sub(r"[^\w-]", "", chat_id or "default")[:40] or "default"
    k = _KERNELS.get(chat_id)
    if k is None:
        if len(_KERNELS) >= 4:                                # keep memory in check: drop the least recently used
            old = min(_KERNELS, key=lambda c: _KERNELS[c].used); _KERNELS.pop(old).kill()
        k = _KERNELS[chat_id] = PyKernel(os.path.join(WORK_DIR, "py", chat_id))
    return k


def _py_files(file_ids):
    out = {}
    for fid in dict.fromkeys(file_ids or []):
        r = FILES.get(fid)
        if r and r.get("origin") != "draft" and os.path.exists(r["path"]):
            out[r["name"]] = r["path"]; out[fid] = r["path"]
    return out


PY_BLOCK_RE = re.compile(r"```meow-python[^\n]*\n(.*?)\n?```", re.S)


def run_python(code, chat_id, file_ids, stop=lambda: False):
    """Run one cell; returns (markdown to append to the answer, list of new file records)."""
    res = kernel_for(chat_id).run(code, _py_files(file_ids), stop=stop)
    text = (res.get("stdout") or "").rstrip()
    if res.get("result") is not None: text = (text + "\n" if text else "") + res["result"]
    if res.get("error"): text = (text + "\n" if text else "") + res["error"]
    if len(text) > 6000: text = text[:2500] + "\n… (output trimmed) …\n" + text[-3000:]
    md = "\n```meow-output" + ("-error" if res.get("error") else "") + "\n" + (text or "(no output)") + "\n```\n"
    recs = []
    for p in res.get("figures", []) + res.get("files", []):
        try:
            if not os.path.isfile(p) or os.path.getsize(p) == 0: continue
            fid, dst = new_path(os.path.basename(p)); shutil.copy(p, dst)
            rec = register(fid, dst, os.path.basename(p), origin="tool"); recs.append(rec)
            if p in res.get("figures", []): md += f"\n![{rec['name']}](/api/files/{fid})\n"
        except Exception as e:
            print("⚠️  python output file:", e)
    return md, recs, res


FEATURES["python"] = True

# ==========================================================
# SEARCH ANYTHING: specialist sources + link reader + deep research (2 rounds)
#   YouTube (transcripts) · Wikipedia · papers (Semantic Scholar + arXiv) · GitHub · Stack Overflow · Hacker News
#   · image results · any link the user pastes (pages, PDFs, YouTube, Reddit threads, GitHub repos)
# ==========================================================
URL_RE = re.compile(r"https?://[^\s<>\"'`)\]]+", re.I)
YT_RE = re.compile(r"(?:youtube\.com/(?:watch\?(?:.*&)?v=|shorts/|live/|embed/)|youtu\.be/)([\w-]{11})")


def _html2text(s):
    s = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", s or "")
    s = re.sub(r"(?i)<br\s*/?>|</p>|</li>|</h\d>", "\n", s)
    s = re.sub(r"(?i)<pre[^>]*>|</pre>", "\n```\n", s)
    s = re.sub(r"<[^>]+>", "", s)
    return html.unescape(re.sub(r"\n{3,}", "\n\n", s)).strip()


def _safe_fetch(url):
    try: return fetch_page(url)
    except Exception as e: print("⚠️  read:", str(e)[:100]); return ""


def _get_json(url, params=None, headers=None, timeout=12):
    r = _http().get(url, params=params, headers=headers or {}, timeout=timeout)
    r.raise_for_status(); return r.json()


# ---------- YouTube ----------
def yt_transcript(vid, max_chars=24000):
    try:
        from youtube_transcript_api import YouTubeTranscriptApi
    except ImportError:
        subprocess.run([_sys.executable, "-m", "pip", "install", "-q", "youtube-transcript-api"], check=False)
        from youtube_transcript_api import YouTubeTranscriptApi
    api = YouTubeTranscriptApi()
    tl = api.list(vid)
    tr = None
    for langs in (["en", "en-US", "en-GB", "en-IN"], ["hi"]):
        try: tr = tl.find_transcript(langs); break
        except Exception: pass
    if tr is None: tr = next(iter(tl))
    if not tr.language_code.startswith(("en", "hi")) and getattr(tr, "is_translatable", False):
        try: tr = tr.translate("en")
        except Exception: pass
    parts, t_mark = [], -60
    for sn in tr.fetch():
        start = getattr(sn, "start", 0) if not isinstance(sn, dict) else sn.get("start", 0)
        text = getattr(sn, "text", "") if not isinstance(sn, dict) else sn.get("text", "")
        if start - t_mark >= 60:                              # a timestamp every minute so the model can cite moments
            parts.append(f"\n[{int(start // 60)}:{int(start % 60):02d}] "); t_mark = start
        parts.append(text.replace("\n", " ") + " ")
    return re.sub(r"[ \t]+", " ", "".join(parts)).strip()[:max_chars]


def yt_meta(url):
    try:
        j = _get_json("https://www.youtube.com/oembed", {"url": url, "format": "json"}, timeout=8)
        return {"title": j.get("title", ""), "channel": j.get("author_name", "")}
    except Exception:
        return {}


def youtube_search(q, n=6):
    key = _secret("SERPER_API_KEY")
    if key:
        try:
            r = _http().post("https://google.serper.dev/videos", json={"q": q, "num": n}, headers={"X-API-KEY": key}, timeout=10)
            r.raise_for_status()
            out = [{"title": v.get("title", ""), "url": v.get("link", ""), "snippet": v.get("snippet", ""), "date": v.get("date", ""),
                    "channel": v.get("channel", ""), "thumb": v.get("imageUrl", ""), "duration": v.get("duration", "")}
                   for v in r.json().get("videos", [])]
            out = [v for v in out if YT_RE.search(v["url"])]
            if out: return out[:n]
        except Exception as e:
            print("⚠️  serper videos:", str(e)[:120])
    try:
        from ddgs import DDGS
        out = []
        for v in DDGS(timeout=8).videos(q, max_results=n) or []:
            u = v.get("content") or v.get("embed_url") or ""
            if YT_RE.search(u):
                out.append({"title": v.get("title", ""), "url": u, "snippet": v.get("description", ""), "date": (v.get("published") or "")[:10],
                            "channel": v.get("uploader", ""), "thumb": (v.get("images") or {}).get("medium", ""), "duration": v.get("duration", "")})
        return out
    except Exception as e:
        print("⚠️  video search:", str(e)[:120]); return []


def youtube_docs(q, n_read=2):
    vids = youtube_search(q, 6); docs = []

    def one(v):
        vid = YT_RE.search(v["url"]).group(1)
        try:
            txt = yt_transcript(vid)
        except Exception as e:
            print("⚠️  transcript:", str(e)[:100]); txt = ""
        return {**v, "kind": "youtube", "site": v.get("channel", ""),
                "text": (f"VIDEO TRANSCRIPT ({v.get('duration') or 'length ?'}):\n" + txt) if txt else (v.get("snippet") or "")}
    with _cf.ThreadPoolExecutor(3) as ex:
        docs = list(ex.map(one, vids[:n_read]))
    docs += [{**v, "kind": "youtube", "text": v.get("snippet", "")} for v in vids[n_read:4] if v.get("snippet")]
    return [d for d in docs if d.get("text")]


# ---------- Wikipedia ----------
def wiki_docs(q, n=2, lang="en"):
    api = f"https://{lang}.wikipedia.org/w/api.php"
    j = _get_json(api, {"action": "query", "list": "search", "srsearch": q, "srlimit": n, "format": "json"},
                  headers={"User-Agent": REDDIT_UA})
    titles = [h["title"] for h in j.get("query", {}).get("search", [])][:n]
    if not titles: return []
    j = _get_json(api, {"action": "query", "prop": "extracts|info", "explaintext": 1, "inprop": "url", "titles": "|".join(titles),
                        "format": "json", "redirects": 1}, headers={"User-Agent": REDDIT_UA})
    out = []
    for p in j.get("query", {}).get("pages", {}).values():
        txt = (p.get("extract") or "").strip()
        if txt:
            txt = re.sub(r"\n=+ (References|External links|See also|Further reading|Notes|Sources)[\s\S]*$", "", txt)
            out.append({"title": p.get("title", "") + " — Wikipedia", "url": p.get("fullurl") or f"https://{lang}.wikipedia.org/wiki/{quote_plus(p.get('title', ''))}",
                        "kind": "wiki", "text": txt[:30000]})
    out.sort(key=lambda d: titles.index(d["title"].replace(" — Wikipedia", "")) if d["title"].replace(" — Wikipedia", "") in titles else 9)
    return out


# ---------- research papers ----------
def paper_docs(q, n=6):
    out, seen = [], set()
    try:
        j = _get_json("https://api.semanticscholar.org/graph/v1/paper/search",
                      {"query": q, "limit": n, "fields": "title,abstract,year,citationCount,url,venue,authors,openAccessPdf,tldr"}, timeout=12)
        for p in j.get("data", []):
            if not p.get("abstract") and not (p.get("tldr") or {}).get("text"): continue
            t = p.get("title", ""); seen.add(t.lower()[:60])
            au = ", ".join(a.get("name", "") for a in (p.get("authors") or [])[:3]) + (" et al." if len(p.get("authors") or []) > 3 else "")
            out.append({"title": t, "url": (p.get("openAccessPdf") or {}).get("url") or p.get("url", ""), "kind": "paper",
                        "date": str(p.get("year") or ""), "site": p.get("venue") or "Semantic Scholar",
                        "text": f"{au} ({p.get('year')}), {p.get('venue') or ''} — cited {p.get('citationCount', 0)} times.\n"
                                + (f"TL;DR: {p['tldr']['text']}\n" if (p.get('tldr') or {}).get('text') else "") + f"Abstract: {p.get('abstract') or ''}"})
    except Exception as e:
        print("⚠️  semantic scholar:", str(e)[:120])
    try:
        r = _http().get("http://export.arxiv.org/api/query", params={"search_query": "all:" + q, "max_results": 5, "sortBy": "relevance"}, timeout=12)
        for e in re.findall(r"<entry>(.*?)</entry>", r.text, re.S):
            g_ = lambda tag: html.unescape(re.sub(r"\s+", " ", (re.search(rf"<{tag}[^>]*>(.*?)</{tag}>", e, re.S) or [None, ""])[1])).strip()
            t = g_("title")
            if not t or t.lower()[:60] in seen: continue
            au = re.findall(r"<name>(.*?)</name>", e)
            out.append({"title": t, "url": g_("id"), "kind": "paper", "date": g_("published")[:10], "site": "arXiv",
                        "text": f"{', '.join(au[:3])}{' et al.' if len(au) > 3 else ''} ({g_('published')[:10]}, arXiv preprint)\nAbstract: {g_('summary')}"})
    except Exception as e:
        print("⚠️  arxiv:", str(e)[:120])
    return out[:n + 3]


# ---------- code: GitHub + Stack Overflow ----------
def code_docs(q):
    out = []
    try:
        hdr = {"Accept": "application/vnd.github+json"}
        if _secret("GITHUB_TOKEN"): hdr["Authorization"] = "Bearer " + _secret("GITHUB_TOKEN")
        j = _get_json("https://api.github.com/search/repositories", {"q": q, "sort": "stars", "per_page": 5}, headers=hdr)
        lines = [f"- {it['full_name']} ★{it.get('stargazers_count', 0):,} — {it.get('description') or ''} (language: {it.get('language') or '?'}, "
                 f"updated {str(it.get('pushed_at', ''))[:10]}) {it.get('html_url')}" for it in j.get("items", [])[:5]]
        if lines:
            out.append({"title": f"GitHub repositories for “{q}”", "url": f"https://github.com/search?q={quote_plus(q)}&type=repositories",
                        "kind": "code", "site": "github.com", "text": "\n".join(lines)})
    except Exception as e:
        print("⚠️  github:", str(e)[:120])
    try:
        base = "https://api.stackexchange.com/2.3"
        j = _get_json(base + "/search/advanced", {"order": "desc", "sort": "relevance", "q": q, "site": "stackoverflow",
                                                   "pagesize": 4, "filter": "withbody"})
        qs_ = [it for it in j.get("items", []) if it.get("answer_count", 0) > 0][:3]
        if qs_:
            ids = ";".join(str(it["question_id"]) for it in qs_)
            ans = _get_json(f"{base}/questions/{ids}/answers", {"order": "desc", "sort": "votes", "site": "stackoverflow", "filter": "withbody", "pagesize": 12})
            by = {}
            for a in ans.get("items", []): by.setdefault(a["question_id"], []).append(a)
            for it in qs_:
                al = sorted(by.get(it["question_id"], []), key=lambda a: (not a.get("is_accepted"), -a.get("score", 0)))[:2]
                body = f"QUESTION (▲{it.get('score', 0)}): {_html2text(it.get('body', ''))[:1500]}\n" + "\n".join(
                    f"{'ACCEPTED ' if a.get('is_accepted') else ''}ANSWER (▲{a.get('score', 0)}): {_html2text(a.get('body', ''))[:2500]}" for a in al)
                out.append({"title": html.unescape(it.get("title", "")), "url": it.get("link", ""), "kind": "code", "site": "stackoverflow.com",
                            "date": _date(it.get("last_activity_date")), "text": body})
    except Exception as e:
        print("⚠️  stack overflow:", str(e)[:120])
    return out


# ---------- Hacker News ----------
def hn_docs(q, n=3):
    j = _get_json("https://hn.algolia.com/api/v1/search", {"query": q, "tags": "story", "hitsPerPage": 8})
    hits = sorted([h for h in j.get("hits", []) if h.get("num_comments")], key=lambda h: -(h.get("points") or 0))[:n]
    out = []

    def one(h):
        try:
            it = _get_json(f"https://hn.algolia.com/api/v1/items/{h['objectID']}", timeout=12)
            kids = sorted([c for c in it.get("children", []) if c.get("text")], key=lambda c: -len(c.get("children", [])))[:10]
            cm = "\n".join(f"COMMENT ({len(c.get('children', []))} replies): {_html2text(c['text'])[:700]}" for c in kids)
        except Exception:
            cm = ""
        return {"title": h.get("title", ""), "url": f"https://news.ycombinator.com/item?id={h['objectID']}", "kind": "hn", "site": "Hacker News",
                "date": (h.get("created_at") or "")[:10], "text": f"▲{h.get('points', 0)} points, {h.get('num_comments', 0)} comments. "
                + (f"Link: {h.get('url')}\n" if h.get("url") else "\n") + cm}
    with _cf.ThreadPoolExecutor(3) as ex:
        out = list(ex.map(one, hits))
    return out


# ---------- images ----------
def image_search(q, n=8):
    key = _secret("SERPER_API_KEY")
    if key:
        try:
            r = _http().post("https://google.serper.dev/images", json={"q": q, "num": n}, headers={"X-API-KEY": key}, timeout=10)
            r.raise_for_status()
            return [{"img": i.get("imageUrl"), "thumb": i.get("thumbnailUrl") or i.get("imageUrl"), "title": i.get("title", ""),
                     "link": i.get("link", ""), "site": i.get("domain") or i.get("source", "")} for i in r.json().get("images", [])][:n]
        except Exception as e:
            print("⚠️  serper images:", str(e)[:120])
    try:
        from ddgs import DDGS
        return [{"img": i.get("image"), "thumb": i.get("thumbnail") or i.get("image"), "title": i.get("title", ""), "link": i.get("url", ""),
                 "site": i.get("source", "")} for i in DDGS(timeout=8).images(q, max_results=n) or []]
    except Exception as e:
        print("⚠️  image search:", str(e)[:120]); return []


# ---------- links the user pasted ----------
def read_link(url):
    url = url.rstrip(".,;:!?")
    m = YT_RE.search(url)
    if m:
        meta = yt_meta(url)
        try:
            txt = yt_transcript(m.group(1), 60000)
        except Exception as e:
            txt = ""; print("⚠️  transcript:", str(e)[:120])
        if not txt:
            try: txt = tavily_extract([url]).get(url, "")
            except Exception: txt = ""
        return {"title": meta.get("title") or "YouTube video", "url": url, "kind": "youtube", "site": meta.get("channel", "YouTube"),
                "text": ("VIDEO TRANSCRIPT:\n" + txt) if txt else "(No transcript available for this video — YouTube may be blocking this server.)"}
    if re.search(r"reddit\.com/r/[^/]+/comments/", url):
        th = {"url": url.split("?")[0], "id": _tid(url), "title": "", "kind": "reddit", "sub": (re.search(r"/r/([^/]+)", url) or [None, ""])[1]}
        read_threads([th], 30)
        parts = ([f"POST: {th['body'][:4000]}"] if th.get("body") else [])
        for c in th.get("top_comments", []):
            s_ = fmt_comment(c, maxlen=1200)
            for r in c.get("replies", []): s_ += "\n  ↳ " + fmt_comment(r, "REPLY", 500)
            parts.append(s_)
        if not parts and th.get("page_text"): parts = [th["page_text"]]
        return {**th, "text": "\n\n".join(parts) or "(Couldn't read this thread.)"}
    gh = re.match(r"https?://github\.com/([\w.-]+)/([\w.-]+)/?$", url)
    if gh:
        o, r_ = gh.group(1), gh.group(2).removesuffix(".git")
        meta, readme = "", ""
        try:
            j = _get_json(f"https://api.github.com/repos/{o}/{r_}")
            meta = f"{j.get('full_name')} ★{j.get('stargazers_count', 0):,} · {j.get('language') or ''} · {j.get('description') or ''} · updated {str(j.get('pushed_at', ''))[:10]}\n"
        except Exception:
            pass
        for br in ("HEAD", "main", "master"):
            try:
                rr = _http().get(f"https://raw.githubusercontent.com/{o}/{r_}/{br}/README.md", timeout=10)
                if rr.status_code == 200: readme = rr.text; break
            except Exception:
                pass
        if meta or readme:
            return {"title": f"{o}/{r_} (GitHub)", "url": url, "kind": "link", "site": "github.com", "text": meta + readme[:40000]}
    txt = ""
    try:
        txt = fetch_page(url, max_bytes=12_000_000)
    except Exception as e:
        print("⚠️  link:", str(e)[:120])
    if len(txt) < 400:
        try:
            alt = tavily_extract([url]).get(url, "")
            if len(alt) > len(txt): txt = alt
        except Exception:
            pass
    title = ""
    m2 = re.match(r"\s*(.{5,140}?)\n", txt or "")
    if m2: title = m2.group(1).strip()
    return {"title": title or _domain(url), "url": url, "kind": "link", "site": _domain(url),
            "text": txt or "(This page couldn't be read — it may need a login, block bots, or be JavaScript-only.)"}


def _link_passages(text, terms, cap_tokens):
    """Whole text if it fits; otherwise the opening part plus the passages most relevant to the question."""
    if len(text) // 4 <= cap_tokens: return [text]
    ch = _chunks(text, 150)
    head, used = [], 0
    for c in ch:
        if used + len(c) // 4 > cap_tokens * 0.35: break
        head.append(c); used += len(c) // 4
    rest = ch[len(head):]
    picked = rank_passages(terms, [{"chunks": rest, "order": 0}], cap_tokens - used, per_doc=60).get(0, []) if rest else []
    return head + (["…"] if picked else []) + picked


# ---------- deep research: a second round for what's still missing ----------
def plan_followups(user, docs, n=3):
    today = datetime.datetime.now().strftime("%d %B %Y")
    found = "\n".join(f"- {d.get('title', '')[:90]}: {(d.get('chunks') or [d.get('text', '')])[0][:160]}" for d in docs[:12])
    p = (f"<|im_start|>system\nYou are a research planner. Today is {today}. Given a question and what the first search round found, "
         f"list up to {n} NEW web searches that would fill the most important gaps (missing facts, other viewpoints, primary sources, "
         'recent updates). Reply ONLY with JSON: {"queries": ["..."]} — use [] if the findings already cover it.<|im_end|>\n'
         f"<|im_start|>user\nQuestion: {user[:1200]}\n\nFound so far:\n{found}<|im_end|>\n<|im_start|>assistant\n<think>\n\n</think>\n\n")
    try:
        m = re.search(r"\{.*\}", _llm_text(p, 160), re.S)
        qs = json.loads(m.group(0)).get("queries", []) if m else []
        return [str(q).strip()[:200] for q in qs if str(q).strip()][:n]
    except Exception as e:
        print("⚠️  follow-up planner:", str(e)[:120]); return []


SPECIALISTS = {"youtube": youtube_docs, "wiki": wiki_docs, "papers": paper_docs, "code": code_docs, "hn": hn_docs}
SRC_LABEL = {"youtube": "YouTube", "wiki": "Wikipedia", "papers": "research papers", "code": "GitHub + Stack Overflow", "hn": "Hacker News", "news": "news"}


# ---------- finding THE answer: relevance-ranked Reddit comments + "is it answered yet?" loop ----------
def _bm25(texts, terms):
    toks = [_terms(t) for t in texts]
    N = max(1, len(toks)); avg = sum(len(t) for t in toks) / N or 1
    df = {}
    for t in toks:
        for w in set(t): df[w] = df.get(w, 0) + 1
    out = []
    for t in toks:
        tf = {}
        for w in t: tf[w] = tf.get(w, 0) + 1
        s = 0.0
        for w in set(terms):
            if w in tf:
                s += math.log(1 + (N - df[w] + 0.5) / (df[w] + 0.5)) * tf[w] * 2.2 / (tf[w] + 1.2 * (0.25 + 0.75 * len(t) / avg))
        out.append(s)
    return out


def rerank_threads(threads, terms, keep=25):
    """Order each thread's comments by (relevance to the question + upvotes + first-hand), promo last; order threads
    by how much they actually talk about the question — so the model sees the comments that ANSWER it."""
    allc = [(th, c) for th in threads for c in th.get("top_comments", [])]
    if not allc: return threads
    texts = [c["body"] + " " + " ".join(r["body"] for r in c.get("replies", [])) for _, c in allc]
    rel = _bm25(texts, terms); top = max(rel) or 1.0
    for (th, c), r in zip(allc, rel):
        tag = comment_tags(c["body"])
        c["_rank"] = 2.2 * r / top + 0.35 * math.log1p(max(0, c["score"])) + (0.6 if "first-hand" in tag else 0) - (3 if "promo" in tag else 0)
    for th in threads:
        cs = sorted(th.get("top_comments", []), key=lambda c: -c.get("_rank", 0))
        th["top_comments"] = cs[:keep]
        trel = _bm25([th.get("title", "") + " " + th.get("body", "")], terms)[0] if th.get("title") else 0
        th["_rel"] = sum(c.get("_rank", 0) for c in cs[:4]) + 0.5 * trel
    return threads


def judge_coverage(user, evidence, mode):
    """Ask the model whether what we found really answers the question; if not, what to search next (and where)."""
    today = datetime.datetime.now().strftime("%d %B %Y")
    where = "Reddit" if mode == "reddit" else "the web and Reddit"
    p = (f"<|im_start|>system\nYou check research for an AI assistant. Today is {today}. Given the user's question and the evidence found "
         f"so far on {where}, decide if it's enough to give a precise, well-supported answer. Reply ONLY with JSON:\n"
         '{"answered": true|false, "missing": "what is still unknown, in a few words", "queries": ["..."], "subreddits": ["..."]}\n'
         "- answered=false if the evidence is off-topic, too thin (only 1-2 people), outdated for a time-sensitive question, or misses part of the question.\n"
         "- queries: up to 3 NEW searches phrased differently from before (synonyms, how people phrase it on Reddit, specific names/versions "
         "mentioned in the evidence, 'X vs Y', 'X experience', 'X scam'). subreddits: up to 3 communities most likely to discuss it (names without r/).<|im_end|>\n"
         f"<|im_start|>user\nQuestion: {user[:1500]}\n\nEvidence so far:\n{evidence}<|im_end|>\n<|im_start|>assistant\n<think>\n\n</think>\n\n")
    try:
        m = re.search(r"\{.*\}", _llm_text(p, 220), re.S)
        j = json.loads(m.group(0)) if m else {}
    except Exception as e:
        print("⚠️  coverage judge:", str(e)[:120]); j = {}
    j["queries"] = [str(q).strip()[:200] for q in (j.get("queries") or []) if str(q).strip()][:3]
    j["subreddits"] = [re.sub(r"^/?r/", "", str(s).strip()).strip("/ ") for s in (j.get("subreddits") or []) if str(s).strip()][:3]
    j["subreddits"] = [s for s in j["subreddits"] if re.fullmatch(r"\w{2,21}", s)]
    return j


def _evidence(threads, docs, terms, max_tokens=2600):
    out, used = [], 0
    for th in sorted([t for t in threads if t.get("top_comments") or t.get("snippet")], key=lambda t: -t.get("_rel", 0)):
        head = f"• r/{th.get('sub', '?')} — {th.get('title', '')[:110]}"
        lines = [head] + [f"  - (▲{c['score']}) {c['body'][:220]}" for c in th.get("top_comments", [])[:3]]
        if not th.get("top_comments") and th.get("snippet"): lines.append("  - " + th["snippet"][:220])
        t = sum(len(l) for l in lines) // 4
        if used + t > max_tokens * 0.7: break
        out += lines; used += t
    for d in docs[:8]:
        l = f"• {d.get('title', '')[:100]} ({_domain(d['url'])}): {(d.get('chunks') or [''])[0][:260]}"
        if used + len(l) // 4 > max_tokens: break
        out.append(l); used += len(l) // 4
    return "\n".join(out) or "(nothing useful found yet)"


# ---------- the whole research run (a generator of progress events) ----------
def web_research(user, history, mode, budget_tokens, rid="", urls=None, search=True):
    """Yields {"status"}, {"search"}, {"sources"}, {"images"} events; the final event is {"_ctx": prompt_block}."""
    stopped = lambda: bool(rid and rid in STOP)
    urls = list(dict.fromkeys(u.rstrip(".,;:!?") for u in (urls or [])))[:4]
    link_docs = []
    if urls:
        yield {"search": {"stage": "links", "links": urls}}
        yield {"status": f"Reading {len(urls)} link{'s' if len(urls) > 1 else ''}…"}
        with _cf.ThreadPoolExecutor(4) as ex:
            link_docs = list(ex.map(read_link, urls))
    plan = {"search": False, "queries": [], "sources": [], "reddit": False, "recency": None, "images": False}
    if search and not stopped():
        yield {"search": {"stage": "planning", "links": urls}}
        plan = plan_search(user, history, mode)
    if not plan.get("search") and not link_docs:
        yield {"search": {"stage": "skipped"}}; yield {"_ctx": ""}; return
    qs, rec, want_reddit, extra = plan["queries"], plan["recency"], plan.get("reddit"), plan.get("sources", [])
    web_on = plan.get("search") and mode != "reddit"
    hits, threads, spec_docs, images = [], [], [], []
    if plan.get("search"):
        labels = (["the web"] if web_on else []) + (["Reddit"] if want_reddit else []) + [SRC_LABEL[x] for x in extra if x in SRC_LABEL and x != "news"]
        yield {"search": {"stage": "searching", "queries": qs, "reddit": bool(want_reddit), "sources": extra, "links": urls}}
        yield {"status": "Searching " + ", ".join(labels[:-1]) + (" and " if len(labels) > 1 else "") + labels[-1] + "…"}
        n_per = 10 if mode == "deep" else 8
        news_q = "news" in extra or rec in ("day", "week") or bool(re.search(r"\bnews\b|\btoday\b|\blatest\b", user, re.I))
        with _cf.ThreadPoolExecutor(14) as ex:
            futs = {}
            if web_on:
                for q in qs: futs[ex.submit(web_search, q, n_per, rec)] = "web"
                if news_q: futs[ex.submit(_ddgs_news, qs[0], 6, rec)] = "news"
            if want_reddit:
                deep_r = mode in ("deep", "reddit")
                for q in qs[:3 if deep_r else 2]: futs[ex.submit(reddit_search, q, 10 if deep_r else 6, rec, deep_r)] = "reddit"
            for x in extra:
                if x in SPECIALISTS: futs[ex.submit(SPECIALISTS[x], qs[0])] = "spec"
            if plan.get("images"): futs[ex.submit(image_search, qs[0], 10)] = "images"
            try:
                for f in _cf.as_completed(futs, timeout=45):
                    try:
                        res = f.result() or []
                    except Exception as e:
                        print(f"⚠️  {futs[f]} search:", str(e)[:120]); continue
                    kind = futs[f]
                    if kind == "reddit": threads.extend(res)
                    elif kind == "spec": spec_docs.extend(res)
                    elif kind == "images": images = res
                    else: hits.extend(res)
            except _cf.TimeoutError:
                print("⚠️  some searches were too slow — skipped")
        if images:
            yield {"images": [i for i in images if i.get("img")][:10]}
    if stopped(): yield {"_ctx": ""}; return
    # merge + de-duplicate web hits (results that several queries agree on rise)
    pool = {}
    n_per = 10 if mode == "deep" else 8
    for i, h in enumerate(hits):
        u = _norm_url(h["url"])
        if not u or "reddit.com" in u and "/comments/" in u:
            if "reddit.com" in u: threads.append(_from_hit(h))
            continue
        if YT_RE.search(h["url"]) and "youtube" in extra: continue
        if u in pool: pool[u]["votes"] += 1; pool[u]["pos"] = min(pool[u]["pos"], i % n_per)
        else: pool[u] = {**h, "votes": 1, "pos": i % n_per, "kind": h.get("kind", "web")}
    linked = {_norm_url(d["url"]) for d in link_docs}
    web = [h for h in sorted(pool.values(), key=lambda h: (h["pos"] - 1.5 * (h["votes"] - 1))) if _norm_url(h["url"]) not in linked]
    tseen, tl = {}, []
    for i, th in enumerate(threads):
        key = th.get("id") or _tid(th["url"]) or _norm_url(th["url"])
        if key in tseen: tseen[key]["_hits"] += 1; continue
        th = {**th, "_hits": 1, "_i": i}; tseen[key] = th; tl.append(th)
    tl.sort(key=lambda th: th["_i"] - 3 * (th["_hits"] - 1))
    n_read = 10 if mode == "deep" else 6
    n_thr = {"reddit": 10, "deep": 6}.get(mode, 3) if want_reddit else 0
    if not n_thr and tl and plan.get("search"): n_thr = 1
    k_cm = {"reddit": 25, "deep": 14}.get(mode, 10)
    web_read, thr_read = web[:n_read], tl[:n_thr]
    side, subs = [], []
    if web_read or thr_read:
        msg = ([f"{len(web_read)} page{'s' if len(web_read) != 1 else ''}"] if web_read else []) + ([f"{len(thr_read)} Reddit threads with their top comments"] if thr_read else [])
        yield {"status": "Reading " + " + ".join(msg) + "…"}
        yield {"search": {"stage": "reading", "queries": qs, "reddit": bool(want_reddit), "sources": extra, "links": urls,
                          "found": len(web) + len(tl) + len(spec_docs)}}
        cnt = {}
        for th in thr_read:
            if th.get("sub"): cnt[th["sub"]] = cnt.get(th["sub"], 0) + 1
        subs = sorted(cnt, key=lambda s_: -cnt[s_])
        with _cf.ThreadPoolExecutor(12) as ex:
            fw = {ex.submit(fetch_page, h["url"]): h for h in web_read}
            ft = ex.submit(read_threads, thr_read, max(k_cm, 40) if mode in ("reddit", "deep") else k_cm) if thr_read else None
            fc = ex.submit(reddit_comment_search, _terms(" ".join(qs)), subs) if (mode in ("reddit", "deep") and subs) else None
            try:
                for f in _cf.as_completed(list(fw) + [x for x in (ft, fc) if x], timeout=45):
                    try:
                        res = f.result()
                    except Exception as e:
                        res = None; print("⚠️  read:", str(e)[:120])
                    if f in fw: fw[f]["text"] = res or ""
                    elif f is fc: side = res or []
            except _cf.TimeoutError:
                print("⚠️  some pages were too slow — skipped")
        # pages that blocked this server → read them through Tavily in one batch
        blocked = [h for h in web_read if len(h.get("text") or "") < 300 and not any(s in _domain(h["url"]) for s in SKIP_FETCH)]
        if blocked and _secret("TAVILY_API_KEY") and not stopped():
            try:
                got = tavily_extract([h["url"] for h in blocked])
                for h in blocked:
                    if len(got.get(h["url"], "")) > len(h.get("text") or ""): h["text"] = got[h["url"]]
            except Exception as e:
                print("⚠️  tavily extract:", str(e)[:120])
    if stopped(): yield {"_ctx": ""}; return
    docs = []
    for h in web[:max(n_read, 8)]:
        body = (h.get("text") or "").strip()
        ch = _chunks(body) if body else []
        if not ch and h.get("snippet"): ch = [h["snippet"]]
        if ch: docs.append({**h, "chunks": ch})
    # keep searching until the question is actually answered (Reddit + Deep modes): up to 2 extra rounds
    terms = _terms(" ".join(qs) + " " + user)
    read = [th for th in tl if th.get("top_comments") or th.get("body") or th.get("page_text")]
    rerank_threads(read, terms, 40)
    rounds = 1
    if mode in ("reddit", "deep") and plan.get("search") and not stopped():
        seen_t = {th.get("id") or _tid(th["url"]) for th in tl}
        seen_u = {_norm_url(d["url"]) for d in docs}
        for rnd in (2, 3):
            if stopped(): break
            yield {"search": {"stage": "checking", "round": rnd - 1, "queries": qs, "reddit": bool(want_reddit), "sources": extra, "links": urls,
                              "found": len(web) + len(tl) + len(spec_docs)}}
            yield {"status": "Checking if this really answers your question…"}
            j = judge_coverage(user, _evidence(read, docs, terms), mode)
            if j.get("answered") or not (j.get("queries") or j.get("subreddits")): break
            nq = j["queries"] or [qs[0]]
            subs_r = list(dict.fromkeys(j.get("subreddits", []) + [s for s in subs if s]))[:3]
            qs = qs + [q for q in nq if q not in qs]
            yield {"search": {"stage": "deeper", "round": rnd, "queries": qs, "reddit": bool(want_reddit), "sources": extra, "links": urls,
                              "missing": j.get("missing", ""), "subs": subs_r, "found": len(web) + len(tl) + len(spec_docs)}}
            yield {"status": f"Round {rnd}: " + (j.get("missing") or "digging deeper")[:90] + (f" · searching r/{', r/'.join(subs_r)}" if subs_r and want_reddit else "")}
            new_thr, new_web = [], []
            with _cf.ThreadPoolExecutor(10) as ex:
                futs = []
                if want_reddit:
                    for q in nq: futs.append(("r", ex.submit(reddit_search, q, 8, rec, True, None)))
                    for s_ in subs_r[:2]: futs.append(("r", ex.submit(reddit_search, nq[0], 8, rec, True, [s_])))
                if mode == "deep":
                    for q in nq: futs.append(("w", ex.submit(web_search, q, 6, rec)))
                for kind, f_ in futs:
                    try:
                        res = f_.result(timeout=40) or []
                    except Exception as e:
                        print("⚠️  round search:", str(e)[:100]); continue
                    for h in res:
                        if kind == "r":
                            k_ = h.get("id") or _tid(h["url"])
                            if k_ and k_ not in seen_t: seen_t.add(k_); new_thr.append(h)
                        else:
                            u = _norm_url(h["url"])
                            if u not in seen_u and "reddit.com" not in u: seen_u.add(u); new_web.append(h)
            new_thr = sorted(new_thr, key=lambda th: th.get("_rank", 0))[:6]
            new_web = new_web[:5]
            if not new_thr and not new_web: break
            rounds = rnd
            with _cf.ThreadPoolExecutor(8) as ex:
                ft = ex.submit(read_threads, new_thr, 40) if new_thr else None
                texts = list(ex.map(_safe_fetch, [h["url"] for h in new_web]))
                if ft:
                    try: ft.result(timeout=60)
                    except Exception as e: print("⚠️  round read:", str(e)[:100])
            for h, t_ in zip(new_web, texts):
                ch = _chunks(t_) if t_ else ([h["snippet"]] if h.get("snippet") else [])
                if ch: docs.append({**h, "kind": "web", "chunks": ch})
            for th in new_thr:
                tl.append(th)
                if th.get("top_comments") or th.get("body") or th.get("page_text"): read.append(th)
            terms = _terms(" ".join(qs) + " " + user)
            rerank_threads(read, terms, 40)
    # the threads that best answer the question go first
    keep_c = {"reddit": 25, "deep": 14}.get(mode, 10)
    unread = [th for th in tl if th not in read][:2 if want_reddit else 0]
    reddit_final = sorted(read, key=lambda th: -th.get("_rel", 0))[:{"reddit": 12, "deep": 8}.get(mode, n_thr)] + unread
    for d in spec_docs:
        ch = _chunks(d["text"], 150) if d.get("kind") not in ("code", "hn") else [d["text"][i:i + 2400] for i in range(0, len(d["text"]), 2400)]
        if ch: docs.append({**d, "chunks": ch})
    for th in reddit_final:
        parts = []
        if th.get("body"): parts.append("POST: " + th["body"][:2000])
        for c in th.get("top_comments", [])[:keep_c]:
            s_ = fmt_comment(c, maxlen=1000)
            for r in c.get("replies", []):
                s_ += "\n  ↳ " + fmt_comment(r, "REPLY", 450)
                if r.get("sub"): s_ += "\n    ↳ " + fmt_comment(r["sub"], "REPLY", 300)
            parts.append(s_)
        if not parts and th.get("page_text"): parts = _chunks(th["page_text"], 110)
        if not parts and th.get("snippet"): parts = [th["snippet"]]
        if parts: docs.append({**th, "chunks": parts})
    if side:
        qt = set(_terms(" ".join(qs)))
        side.sort(key=lambda c: -(c["score"] + 25 * len(qt & set(_terms(c["body"])))))
        bysub = {}
        for c in side[:24]: bysub.setdefault(c["sub"], []).append(c)
        for sub, cs in bysub.items():
            docs.append({"title": f"More comments in r/{sub} on this topic", "url": f"https://www.reddit.com/r/{sub}/search/?q={quote_plus(' '.join(qs[:1]))}",
                         "kind": "reddit", "sub": sub, "score": None, "comments": len(cs), "date": "",
                         "chunks": [fmt_comment(c, f"COMMENT ({c['date']})", 700) for c in cs]})
    for i, d in enumerate(docs): d["order"] = i
    qterms = _terms(" ".join(qs) + " " + user)
    # links the user pasted come first and get the biggest share
    link_budget = int(budget_tokens * (0.85 if not docs else 0.5)) if link_docs else 0
    blocks, sources = [], []

    def add_source(d, passages):
        n = len(sources) + 1
        src = {"n": n, "title": (d.get("title") or _domain(d["url"]))[:160], "url": d["url"], "domain": _domain(d["url"]),
               "kind": d.get("kind", "web"), "date": d.get("date", ""), "snippet": (d.get("snippet") or passages[0])[:220]}
        if d.get("kind") == "reddit":
            src.update(sub=d.get("sub", ""), score=d.get("score"), comments=d.get("comments"))
            meta = ", ".join(x for x in [f"▲{d['score']}" if d.get("score") is not None else "", f"{d['comments']} comments" if d.get("comments") else "", d.get("date", "")] if x)
            head = f"[{n}] Reddit r/{d.get('sub', '?')} — \"{d.get('title', '')}\"" + (f" ({meta})" if meta else "")
        else:
            tag = {"youtube": "YouTube", "wiki": "Wikipedia", "paper": "Paper", "code": d.get("site") or "Code", "hn": "Hacker News",
                   "link": "Link you shared", "news": "News"}.get(d.get("kind"), "")
            if d.get("kind") == "youtube" and d.get("site"): tag += f" · {d['site']}"
            head = f"[{n}] {(tag + ': ') if tag else ''}{d.get('title', '')} — {_domain(d['url'])}" + (f" ({d['date']})" if d.get("date") else "")
        sep = "\n" if d.get("kind") in ("reddit", "code", "hn") else "\n…\n"
        blocks.append(head + f"\nURL: {d['url']}\n" + sep.join(passages)); sources.append(src)

    for d in link_docs:
        add_source(d, _link_passages(d["text"], qterms, max(400, link_budget // max(1, len(link_docs)))))
    picked = rank_passages(qterms, docs, budget_tokens - link_budget - 60 * len(docs),
                           reddit_share={"reddit": 0.92, "deep": 0.5}.get(mode, 0.4))
    for di, d in enumerate(docs):
        if di in picked: add_source(d, picked[di])
    yield {"sources": sources}
    yield {"search": {"stage": "done", "rounds": rounds, "queries": qs, "reddit": bool(want_reddit), "sources": extra, "links": urls,
                      "found": len(web) + len(tl) + len(spec_docs), "used": len(sources)}}
    if not blocks:
        yield {"_ctx": "<web_results>\n(The search returned nothing usable — the network may be off, or the engines rate-limited this server.)\n</web_results>"}
        return
    today = datetime.datetime.now().strftime("%d %B %Y")
    yield {"_ctx": f'<web_results searched="{today}" queries="{" | ".join(qs)}">\n' + "\n\n".join(blocks) + "\n</web_results>"}

# ==========================================================
# MORE FILE TOOLS: PDF → Word, searchable scans, form filling, PDF compare, photo restore
# ==========================================================
import difflib


def _ocr_rows(img):
    """OCR lines grouped into reading-order rows (same baseline → one row), blank line between blocks."""
    boxes = sorted(ocr_boxes(img), key=lambda b: (b["bbox"][1] + b["bbox"][3]) / 2)
    rows = []
    for b in boxes:
        yc, h = (b["bbox"][1] + b["bbox"][3]) / 2, b["bbox"][3] - b["bbox"][1]
        if rows and abs(yc - rows[-1]["yc"]) < 0.55 * max(h, rows[-1]["h"]): rows[-1]["items"].append(b)
        else: rows.append({"yc": yc, "h": h, "items": [b]})
    out, prev = [], None
    for r in rows:
        if prev and r["yc"] - prev["yc"] > 2.2 * max(r["h"], prev["h"]): out.append("")
        out.append("    ".join(b["text"] for b in sorted(r["items"], key=lambda b: b["bbox"][0]))); prev = r
    return "\n".join(out)


@tool("pdf_to_word", "PDF", "PDF → Word", "Convert a PDF into an editable Word (.docx) file, keeping layout, tables and images. Scanned PDFs are OCR'd.",
      PDF, params=[P("pages", "Pages", default="all")])
def t_pdf_to_word(recs, p):
    rec = recs[0]; stem = _stem(rec)
    fid, path = new_path(f"{stem}.docx")
    with fitz.open(rec["path"]) as d:
        n = d.page_count; pages = parse_pages(p.get("pages"), n)
        scanned = sum(len(d[i].get_text("text").strip()) < 25 for i in pages) > len(pages) / 2
    if not scanned:
        try:
            try:
                from pdf2docx import Converter
            except ImportError:
                subprocess.run([_sys.executable, "-m", "pip", "install", "-q", "pdf2docx"], check=False)
                from pdf2docx import Converter
            cv = Converter(rec["path"])
            cv.convert(path, pages=pages); cv.close()
            return register(fid, path, f"{stem}.docx", "tool"), f"Converted {len(pages)} page(s) to Word with layout"
        except Exception as e:
            print("⚠️  pdf2docx:", e)
    import docx                                               # scanned (or fallback): OCR text → clean Word document
    doc = docx.Document()
    with fitz.open(rec["path"]) as d:
        for k, i in enumerate(pages):
            page = d[i]; t = page.get_text("text").strip()
            if len(t) < 25:
                pix = page.get_pixmap(dpi=200)
                t = _ocr_rows(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
            for para in re.split(r"\n\s*\n", t):
                if para.strip(): doc.add_paragraph(para.strip())
            if k < len(pages) - 1: doc.add_page_break()
    doc.save(path)
    return register(fid, path, f"{stem}.docx", "tool"), f"Converted {len(pages)} page(s) to Word" + (" (text read with OCR)" if scanned else "")


@tool("pdf_make_searchable", "PDF", "Make scan searchable", "OCR a scanned PDF and add an invisible text layer so you can search, select and copy text.",
      PDF, params=[P("pages", "Pages", default="all")])
def t_pdf_searchable(recs, p):
    d = _open_pdf(recs[0]); pages = parse_pages(p.get("pages"), d.page_count); done = words = 0
    for i in pages:
        page = d[i]
        if len(page.get_text("text").strip()) > 50: continue  # already has real text
        z = 200 / 72; pix = page.get_pixmap(matrix=fitz.Matrix(z, z))
        boxes = ocr_boxes(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
        for b in boxes:
            x0, y0, x1, y1 = [v / z for v in b["bbox"]]
            h = max(4.0, y1 - y0); txt = b["text"]
            fs = h * 0.82; w = fitz.get_text_length(txt, "helv", fs)
            if w > 0 and w > (x1 - x0) * 1.02: fs *= (x1 - x0) / w
            try:
                page.insert_text((x0, y1 - h * 0.2), txt, fontsize=max(3, fs), fontname="helv", render_mode=3)   # 3 = invisible
                words += len(txt.split())
            except Exception:
                pass
        done += 1
    rec = _save_pdf(d, f"{_stem(recs[0])}_searchable.pdf", garbage=3, deflate=True)
    return rec, (f"Added a searchable text layer to {done} page(s) (~{words} words)" if done else "This PDF already has selectable text — nothing to do")


@tool("pdf_fill_form", "PDF", "Fill PDF form", "Fill a fillable PDF form. values = JSON object {field name: value}; checkboxes take true/false. Field names are listed in the file's content.",
      PDF, params=[P("values", "Values (JSON)", "textarea", "{}"), P("flatten", "Flatten (lock) after filling", "select", "no", ["no", "yes"])])
def t_pdf_fill(recs, p):
    vals = p.get("values") or {}
    if isinstance(vals, str):
        try: vals = json.loads(vals or "{}")
        except Exception: raise ValueError("values must be a JSON object like {\"Name\": \"Tarun\"}")
    norm = lambda s: re.sub(r"[\s_:.-]+", "", str(s).lower())
    want = {norm(k): v for k, v in vals.items()}
    d = _open_pdf(recs[0]); filled, missing = [], set(want)
    for page in d:
        for w in page.widgets() or []:
            key = norm(w.field_name)
            if key not in want: continue
            v = want[key]
            if w.field_type in (fitz.PDF_WIDGET_TYPE_CHECKBOX, fitz.PDF_WIDGET_TYPE_RADIOBUTTON):
                on = v if isinstance(v, bool) else str(v).strip().lower() in ("1", "true", "yes", "on", "x", "✓", "checked")
                w.field_value = w.on_state() if on else "Off"
            else:
                w.field_value = str(v)
            w.update(); filled.append(w.field_name); missing.discard(key)
    if not filled:
        names = [w.field_name for pg in d for w in (pg.widgets() or [])]
        raise ValueError("No matching fields. This form's fields are: " + (", ".join(names[:40]) or "none (not a fillable form — use pdf_add_text instead)"))
    if _s(p, "flatten", "no") == "yes":
        try: d.bake()
        except Exception: pass
    rec = _save_pdf(d, f"{_stem(recs[0])}_filled.pdf")
    return rec, f"Filled {len(filled)} field(s)" + (f"; not found: {', '.join(sorted(missing))}" if missing else "")


def _words(doc):
    out = []
    for i, page in enumerate(doc):
        for w in page.get_text("words", sort=True):
            out.append((w[4], i, fitz.Rect(w[:4])))
    return out


@tool("pdf_compare", "PDF", "Compare two PDFs", "Find every difference between two versions of a document: side-by-side PDF with removed text in red and added text in green, plus a summary.",
      PDF, multi=True)
def t_pdf_compare(recs, p):
    if len(recs) < 2: raise ValueError("Pick two PDFs: the old version first, then the new one.")
    A, B = fitz.open(recs[0]["path"]), fitz.open(recs[1]["path"])
    wa, wb = _words(A), _words(B)
    sm = difflib.SequenceMatcher(None, [w[0].lower() for w in wa], [w[0].lower() for w in wb], autojunk=False)
    red, green, changes = [], [], []
    for op, a0, a1, b0, b1 in sm.get_opcodes():
        if op == "equal": continue
        if op in ("replace", "delete"): red += wa[a0:a1]
        if op in ("replace", "insert"): green += wb[b0:b1]
        old = " ".join(w[0] for w in wa[a0:a1]); new = " ".join(w[0] for w in wb[b0:b1])
        pg = (wb[b0][1] if b0 < len(wb) else wa[a0][1] if a0 < len(wa) else 0) + 1
        changes.append({"replace": f"p{pg}: “{old[:80]}” → “{new[:80]}”", "delete": f"p{pg}: removed “{old[:100]}”",
                        "insert": f"p{pg}: added “{new[:100]}”"}[op])
    out = fitz.open()
    gap = 24
    for i in range(max(A.page_count, B.page_count)):
        ra = A[i].rect if i < A.page_count else fitz.Rect(0, 0, 0, 0)
        rb = B[i].rect if i < B.page_count else fitz.Rect(0, 0, 0, 0)
        W, H = ra.width + gap + rb.width, max(ra.height, rb.height) + 30
        pg = out.new_page(width=W, height=H)
        pg.insert_text((10, 18), f"OLD · {recs[0]['name']} · page {i+1}", fontsize=9, color=(0.75, 0.1, 0.1))
        pg.insert_text((ra.width + gap + 10, 18), f"NEW · {recs[1]['name']} · page {i+1}", fontsize=9, color=(0.05, 0.5, 0.2))
        if i < A.page_count: pg.show_pdf_page(fitz.Rect(0, 30, ra.width, 30 + ra.height), A, i)
        if i < B.page_count: pg.show_pdf_page(fitz.Rect(ra.width + gap, 30, ra.width + gap + rb.width, 30 + rb.height), B, i)
        pg.draw_line((ra.width + gap / 2, 30), (ra.width + gap / 2, H), color=(0.6, 0.6, 0.6), width=0.6)
        for _, pi, r in red:
            if pi == i: pg.draw_rect(r + (0, 30, 0, 30), color=None, fill=(1, 0.25, 0.25), fill_opacity=0.3, overlay=True)
        for _, pi, r in green:
            if pi == i: pg.draw_rect(r + (ra.width + gap, 30, ra.width + gap, 30), color=None, fill=(0.1, 0.8, 0.3), fill_opacity=0.3, overlay=True)
    fid, path = new_path(f"compare_{_stem(recs[0])}_vs_{_stem(recs[1])}.pdf")
    out.save(path, garbage=3, deflate=True); A.close(); B.close(); out.close()
    rec = register(fid, path, os.path.basename(path).split("_", 1)[1], "tool")
    if not changes: return rec, "The two PDFs have identical text"
    sim = sm.ratio() * 100
    return rec, f"{len(changes)} change(s) · {sim:.1f}% text similarity. " + " | ".join(changes[:12]) + (" …" if len(changes) > 12 else "")


@tool("image_restore", "Image", "Restore photo", "One-click photo fix: removes noise and JPEG blockiness, fixes colour cast, recovers shadows/contrast, sharpens, and optionally AI-upscales small photos.",
      IMG, params=[P("strength", "Strength", "select", "medium", ["light", "medium", "strong"]), P("upscale", "AI upscale if small", "select", "yes", ["yes", "no"])])
def t_image_restore(recs, p):
    import numpy as np, cv2
    im = _load_img(recs[0]); alpha = im.getchannel("A") if "A" in im.getbands() else None
    rgb = np.array(im.convert("RGB"))
    st = {"light": 0.6, "medium": 1.0, "strong": 1.6}.get(_s(p, "strength", "medium"), 1.0)
    sig = _noise_sigma(rgb, np.ones(rgb.shape[:2], bool))
    h = float(np.clip(sig * 1.3 * st, 2, 14))
    bgr = cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR)
    bgr = cv2.fastNlMeansDenoisingColored(bgr, None, h, h, 7, 21)
    lab = cv2.cvtColor(bgr, cv2.COLOR_BGR2LAB).astype(np.float32)
    for c in (1, 2):                                          # grey-world white balance (gentle)
        lab[..., c] = lab[..., c] - (lab[..., c].mean() - 128) * min(1.0, 0.7 * st)
    lo, hi = np.percentile(lab[..., 0], (0.5, 99.5))           # auto levels: real blacks and whites
    if hi - lo > 40: lab[..., 0] = lab[..., 0] * 0.4 + np.clip((lab[..., 0] - lo) * 255.0 / (hi - lo), 0, 255) * 0.6
    L = np.clip(lab[..., 0], 0, 255).astype(np.uint8)
    L = cv2.createCLAHE(clipLimit=1.4 + 0.8 * st, tileGridSize=(8, 8)).apply(L)
    lab[..., 0] = L
    out = cv2.cvtColor(np.clip(lab, 0, 255).astype(np.uint8), cv2.COLOR_LAB2RGB)
    blur = cv2.GaussianBlur(out, (0, 0), 1.2)
    out = cv2.addWeighted(out, 1 + 0.45 * st, blur, -0.45 * st, 0)
    res = Image.fromarray(out); note = f"denoised (σ≈{sig:.1f}), colour-balanced, contrast recovered, sharpened"
    if _s(p, "upscale", "yes") == "yes" and max(res.size) < 1400:
        try:
            res = ai_upscale(res, 2); note += ", AI-upscaled 2×"
            if alpha: alpha = alpha.resize(res.size, Image.LANCZOS)
        except Exception as e:
            print("⚠️  upscale:", e)
    if alpha: res.putalpha(alpha)
    return _save_img(res, f"{_stem(recs[0])}_restored", _src_ext(recs[0])), note.capitalize()

# ==========================================================
# FRONTEND (single page)
# ==========================================================
MEOW_HTML = r"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<meta name="theme-color" content="#05060f">
<title>Meow 🐈</title>
<link rel="icon" id="favicon" href="data:,">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&family=Space+Grotesk:wght@500;600;700&family=JetBrains+Mono:wght@400;500&display=swap" rel="stylesheet">
<link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/highlight.js/11.9.0/styles/atom-one-dark.min.css">
<link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/KaTeX/0.16.9/katex.min.css">
<script src="https://cdn.jsdelivr.net/npm/marked@12.0.2/marked.min.js"></script>
<script src="https://cdn.jsdelivr.net/npm/dompurify@3.1.6/dist/purify.min.js"></script>
<script src="https://cdnjs.cloudflare.com/ajax/libs/highlight.js/11.9.0/highlight.min.js"></script>
<script src="https://cdnjs.cloudflare.com/ajax/libs/KaTeX/0.16.9/katex.min.js"></script>
<style>
:root{--bg:#04050d;--panel:rgba(14,16,34,.6);--line:rgba(255,255,255,.08);--line2:rgba(255,255,255,.14);--text:#e8eaf6;--muted:#9097b5;--dim:#646b8a;--violet:#a78bfa;--blue:#60a5fa;--pink:#f0abfc;--good:#34d399;--bad:#f87171}
*{box-sizing:border-box}
html,body{height:100%;margin:0}
body{background:radial-gradient(1200px 800px at 75% 115%,#121a4d 0%,transparent 60%),radial-gradient(900px 700px at -10% -10%,#2b0f4e 0%,transparent 55%),var(--bg);color:var(--text);font-family:Inter,system-ui,sans-serif;overflow:hidden;-webkit-font-smoothing:antialiased}
button{font-family:inherit}
#sky{position:fixed;inset:0;width:100%;height:100%;z-index:0;pointer-events:none}
.nebula{position:fixed;border-radius:50%;filter:blur(90px);z-index:0;pointer-events:none;opacity:.3}
.n1{width:60vmax;height:60vmax;left:-20vmax;top:-25vmax;background:radial-gradient(circle,#7c3aed,transparent 62%);animation:drift 46s ease-in-out infinite alternate}
.n2{width:55vmax;height:55vmax;right:-22vmax;bottom:-28vmax;background:radial-gradient(circle,#2563eb,transparent 62%);animation:drift 52s ease-in-out infinite alternate-reverse}
.n3{width:32vmax;height:32vmax;left:42%;top:30%;background:radial-gradient(circle,#db2777,transparent 65%);opacity:.14;animation:drift 64s ease-in-out infinite alternate}
@keyframes drift{to{transform:translate(6vmax,4vmax) scale(1.18)}}
.app{position:relative;z-index:1;display:flex;height:100vh;height:100dvh}
.glass{background:var(--panel);backdrop-filter:blur(20px) saturate(150%);-webkit-backdrop-filter:blur(20px) saturate(150%);border:1px solid var(--line)}
.brand{font-family:'Space Grotesk',Inter,sans-serif;font-weight:700;letter-spacing:-.02em;background:linear-gradient(90deg,#f5f3ff,#c4b5fd,#93c5fd,#f5d0fe,#f5f3ff);background-size:250% 100%;-webkit-background-clip:text;background-clip:text;color:transparent;animation:shine 10s linear infinite}
@keyframes shine{to{background-position:250% 0}}
.spin{display:inline-block;width:18px;height:18px;border:2px solid rgba(255,255,255,.15);border-top-color:var(--violet);border-radius:50%;animation:spin .8s linear infinite}
@keyframes spin{to{transform:rotate(360deg)}}
/* sidebar */
.sidebar{width:280px;flex-shrink:0;display:flex;flex-direction:column;background:rgba(9,10,24,.72);backdrop-filter:blur(22px);-webkit-backdrop-filter:blur(22px);border-right:1px solid var(--line);transition:transform .3s ease,margin .3s ease}
.app.sb-collapsed .sidebar{margin-left:-280px}
.sb-top{padding:16px 14px 10px}
.logo-row{display:flex;align-items:center;gap:10px;padding:2px 4px 14px}
.logo-row .lg{width:34px;height:34px;filter:drop-shadow(0 0 12px rgba(167,139,250,.4))}
.logo-row .brand{font-size:1.4rem}
.logo-row .sp{flex:1}
.new-chat{width:100%;display:flex;align-items:center;gap:10px;padding:11px 14px;border-radius:14px;border:1px solid var(--line2);background:linear-gradient(135deg,rgba(167,139,250,.18),rgba(96,165,250,.1));color:var(--text);font-weight:500;font-size:.9rem;cursor:pointer;transition:.2s}
.new-chat:hover{border-color:rgba(167,139,250,.5);box-shadow:0 0 26px rgba(167,139,250,.22)}
.new-chat kbd{margin-left:auto;font-size:.66rem;color:var(--muted);font-family:inherit;border:1px solid var(--line2);border-radius:6px;padding:1px 6px}
.search{margin-top:10px;display:flex;align-items:center;gap:8px;padding:8px 12px;border-radius:12px;background:rgba(255,255,255,.04);border:1px solid var(--line);color:var(--muted)}
.search input{flex:1;min-width:0;background:none;border:0;outline:0;color:var(--text);font-size:.85rem;font-family:inherit}
.chat-list{flex:1;overflow-y:auto;padding:4px 8px 8px}
.cl-group{font-size:.68rem;text-transform:uppercase;letter-spacing:.08em;color:var(--dim);padding:12px 10px 6px}
.cl-item{display:flex;align-items:center;gap:9px;padding:9px 10px;border-radius:10px;color:var(--muted);font-size:.86rem;cursor:pointer;transition:.15s}
.cl-item:hover{background:rgba(255,255,255,.05);color:var(--text)}
.cl-item.active{background:rgba(167,139,250,.15);color:#fff}
.cl-item .t{flex:1;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
.cl-item .del{opacity:0;display:flex;padding:3px;border-radius:6px;color:var(--muted);background:none;border:0;cursor:pointer}
.cl-item:hover .del{opacity:1}.cl-item .del:hover{color:var(--bad)}
.sb-bottom{padding:10px 14px 14px;border-top:1px solid var(--line)}
.sb-link .ver{margin-left:auto;font-size:.66rem;color:var(--dim);font-family:'JetBrains Mono',monospace}
.sb-link.power:hover{color:#fca5a5;background:rgba(248,113,113,.08)}
.news-ov{position:fixed;inset:0;z-index:75;background:rgba(3,3,10,.7);backdrop-filter:blur(6px);display:flex;align-items:center;justify-content:center;padding:18px}
.news-card{width:min(560px,100%);max-height:80vh;overflow:auto;border-radius:20px;padding:22px 24px}
.news-card h3{margin:0 0 14px;font-family:'Space Grotesk',Inter,sans-serif;display:flex;align-items:center;gap:10px}
.news-card h4{margin:16px 0 6px;font-size:.82rem;color:#c4b5fd;font-family:'JetBrains Mono',monospace}.news-card h4.cur::after{content:' · current';color:#86efac}
.news-card ul{margin:0;padding-left:20px;color:#d6d8ee;font-size:.9rem;line-height:1.6}
.news-card .x{margin-left:auto;background:none;border:0;color:var(--muted);cursor:pointer;display:flex}
.sb-link{width:100%;display:flex;align-items:center;gap:10px;padding:9px 10px;border-radius:10px;background:none;border:0;color:var(--muted);cursor:pointer;font-size:.86rem}
.sb-link:hover{background:rgba(255,255,255,.05);color:var(--text)}
.health{display:flex;align-items:center;gap:8px;font-size:.74rem;color:var(--dim);padding:8px 10px 0}
.dot{width:8px;height:8px;border-radius:50%;background:#f59e0b;flex-shrink:0}
.dot.ok{background:var(--good);box-shadow:0 0 10px var(--good)}.dot.bad{background:var(--bad);box-shadow:0 0 10px var(--bad)}
.scrim{display:none}
/* main */
.main{flex:1;display:flex;flex-direction:column;min-width:0;position:relative}
.topbar{height:60px;display:flex;align-items:center;gap:10px;padding:0 18px;flex-shrink:0}
.title-wrap{display:flex;align-items:center;gap:10px;flex:1;min-width:0}
.title-wrap .brand{font-size:1.15rem;white-space:nowrap}
.model-badge{font-size:.72rem;color:var(--muted);padding:4px 10px;border-radius:999px;border:1px solid var(--line2);background:rgba(255,255,255,.03);white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
.icon-btn{display:inline-flex;align-items:center;justify-content:center;width:38px;height:38px;border-radius:12px;border:0;background:transparent;color:var(--muted);cursor:pointer;transition:.15s;flex-shrink:0}
.icon-btn:hover{background:rgba(255,255,255,.07);color:var(--text)}
.pill-btn{display:inline-flex;align-items:center;gap:7px;padding:8px 14px;border-radius:999px;border:1px solid var(--line2);background:rgba(255,255,255,.04);color:var(--text);font-size:.82rem;cursor:pointer;transition:.15s}
.pill-btn:hover{border-color:rgba(167,139,250,.5);background:rgba(167,139,250,.1)}
.scroller{flex:1;overflow-y:auto}
.messages,.welcome{max-width:820px;margin:0 auto;padding:0 20px}
.messages{padding-bottom:30px}
/* welcome */
.welcome{padding-top:8vh;text-align:center}
.hero-cat{width:112px;height:112px;margin:0 auto 12px;filter:drop-shadow(0 0 34px rgba(167,139,250,.5));animation:float 6s ease-in-out infinite}
@keyframes float{50%{transform:translateY(-9px)}}
.hero{font-size:clamp(2.6rem,7vw,4.3rem);margin:0;line-height:1.05}
.sub{color:var(--muted);font-size:1.05rem;margin:12px 0 34px}
.suggest-grid{display:grid;grid-template-columns:repeat(2,1fr);gap:12px;text-align:left}
.suggest{padding:16px;border-radius:18px;cursor:pointer;transition:.2s;border:1px solid var(--line);background:rgba(14,16,34,.5);backdrop-filter:blur(14px);-webkit-backdrop-filter:blur(14px)}
.suggest:hover{transform:translateY(-2px);border-color:rgba(167,139,250,.45);box-shadow:0 12px 40px -12px rgba(124,58,237,.5)}
.suggest .e{font-size:1.3rem}.suggest .h{font-weight:600;font-size:.92rem;margin-top:6px}.suggest .d{color:var(--muted);font-size:.8rem;margin-top:3px;line-height:1.45}
/* messages */
.msg{display:flex;gap:14px;margin-top:26px;animation:rise .35s ease both}
@keyframes rise{from{opacity:0;transform:translateY(8px)}}
.msg.user{justify-content:flex-end}
.bubble{max-width:82%;padding:12px 16px;border-radius:20px 20px 6px 20px;background:linear-gradient(135deg,rgba(124,58,237,.34),rgba(37,99,235,.26));border:1px solid rgba(167,139,250,.25);white-space:pre-wrap;word-wrap:break-word;line-height:1.6;font-size:.95rem;backdrop-filter:blur(10px)}
.bubble .att{display:flex;flex-wrap:wrap;gap:8px;margin-bottom:8px;justify-content:flex-end;white-space:normal}
.avatar{width:34px;height:34px;flex-shrink:0;margin-top:2px}
.avatar.busy svg{animation:bob 1.2s ease-in-out infinite}
@keyframes bob{50%{transform:translateY(-3px) rotate(-5deg)}}
.a-body{flex:1;min-width:0}
.a-name{font-size:.78rem;font-weight:600;color:var(--muted);margin-bottom:6px;display:flex;align-items:center;gap:8px}
.mode-tag{font-size:.64rem;padding:2px 7px;border-radius:999px;border:1px solid var(--line2);font-weight:500}
.mode-tag.m-think{color:#d8b4fe;border-color:rgba(216,180,254,.35)}.mode-tag.m-fast{color:#fcd34d;border-color:rgba(252,211,77,.3)}
.status-line{font-weight:400;color:var(--dim)}
/* think */
.think{margin:2px 0 12px;border-radius:16px;border:1px solid rgba(167,139,250,.25);background:rgba(76,29,149,.14);overflow:hidden}
.think summary{list-style:none;cursor:pointer;display:flex;align-items:center;gap:9px;padding:10px 14px;font-size:.8rem;color:#d8b4fe;user-select:none}
.think summary::-webkit-details-marker{display:none}
.think .chev{margin-left:auto;display:flex;transition:transform .2s}.think[open] .chev{transform:rotate(180deg)}
.think-body{padding:10px 16px 14px;font-size:.82rem;line-height:1.65;color:#aaa4cc;white-space:pre-wrap;max-height:380px;overflow-y:auto;border-top:1px solid rgba(167,139,250,.14)}
.orb{width:14px;height:14px;border-radius:50%;background:conic-gradient(#a78bfa,#60a5fa,#f0abfc,#a78bfa);animation:spin 1.1s linear infinite;flex-shrink:0;box-shadow:0 0 12px rgba(167,139,250,.6)}
.orb.done{animation:none;background:#a78bfa;opacity:.6;box-shadow:none}
.fups{display:flex;flex-direction:column;gap:6px;margin-top:14px}.fups:empty{display:none}
.fup{display:flex;align-items:center;gap:8px;text-align:left;padding:9px 13px;border-radius:12px;border:1px solid var(--line);background:rgba(255,255,255,.03);color:#d6d8ee;font:inherit;font-size:.86rem;cursor:pointer;transition:.15s}
.fup svg{transform:rotate(-90deg);color:#a78bfa;flex-shrink:0}.fup:hover{border-color:rgba(167,139,250,.5);background:rgba(167,139,250,.08);color:#fff}
.edit-bar{display:flex;align-items:center;gap:8px;margin:-2px 0 8px;padding:7px 10px;border-radius:12px;background:rgba(167,139,250,.1);border:1px solid rgba(167,139,250,.28);font-size:.78rem;color:#ddd6fe}
.edit-bar span{flex:1}.edit-bar button{background:none;border:1px solid rgba(255,255,255,.18);color:var(--text);border-radius:8px;padding:3px 10px;cursor:pointer;font:inherit;font-size:.76rem}
.u-wrap{display:flex;flex-direction:column;align-items:flex-end;max-width:82%;min-width:0}.u-wrap .bubble{max-width:100%}
.u-acts{display:flex;gap:2px;margin-top:4px;opacity:0;transition:.15s}.msg.user:hover .u-acts,.u-acts:focus-within{opacity:1}
.u-acts button{background:none;border:0;color:var(--dim);cursor:pointer;padding:5px;border-radius:8px;display:flex}.u-acts button:hover{color:var(--text);background:rgba(255,255,255,.07)}
.touch-only{display:none}
@media (pointer:coarse){.touch-only{display:inline-flex}.u-acts{opacity:.7}}
.voice-ov{position:fixed;inset:0;z-index:70;background:radial-gradient(ellipse at center,rgba(40,20,90,.92),rgba(3,3,12,.97));display:flex;flex-direction:column;align-items:center;justify-content:center;gap:22px;padding:24px;text-align:center}
.vo-orb{position:relative;width:170px;height:170px}
.vo-orb span{position:absolute;inset:0;border-radius:50%;background:conic-gradient(from 0deg,#a78bfa,#60a5fa,#f0abfc,#a78bfa);filter:blur(2px);opacity:.85;animation:spin 6s linear infinite}
.vo-orb span:nth-child(2){inset:18px;filter:blur(10px);opacity:.6;animation-duration:4s;animation-direction:reverse}
.vo-orb span:nth-child(3){inset:46px;background:#0b0b1e;filter:none;opacity:1;animation:none;box-shadow:inset 0 0 30px rgba(167,139,250,.5)}
.voice-ov.listening .vo-orb{animation:voPulse 1.6s ease-in-out infinite}.voice-ov.speaking .vo-orb{animation:voPulse .7s ease-in-out infinite}
.voice-ov.thinking .vo-orb span{animation-duration:1.2s}.voice-ov.paused .vo-orb{filter:grayscale(.8) brightness(.6)}
@keyframes voPulse{0%,100%{transform:scale(1)}50%{transform:scale(1.07)}}
.vo-state{font-family:'Space Grotesk',Inter,sans-serif;font-size:1.3rem;color:#fff}
.vo-text{max-width:640px;min-height:3em;color:#c7c9e8;font-size:.98rem;line-height:1.55;max-height:30vh;overflow:auto}
.vo-btns{display:flex;gap:16px}.vo-btns button{width:58px;height:58px;border-radius:50%;border:1px solid rgba(255,255,255,.2);background:rgba(255,255,255,.08);color:#fff;cursor:pointer;display:flex;align-items:center;justify-content:center}
.vo-btns button.end{background:#dc2626;border-color:#dc2626}.vo-btns button.off{background:rgba(255,255,255,.25)}
.vo-hint{font-size:.76rem;color:var(--dim)}
/* web search */
.web-pick{position:relative}
.web-btn{display:inline-flex;align-items:center;gap:6px;height:34px;padding:0 12px;border-radius:999px;border:1px solid var(--line);background:rgba(255,255,255,.05);color:var(--muted);font:inherit;font-size:.8rem;cursor:pointer;transition:.15s;white-space:nowrap}
.web-btn:hover{color:var(--text);background:rgba(255,255,255,.08)}
.web-btn.on{color:#7dd3fc;border-color:rgba(125,211,252,.4);background:rgba(14,165,233,.12)}
.web-btn.rd{color:#fdba74;border-color:rgba(253,186,116,.42);background:rgba(249,115,22,.12)}
.web-btn.off{opacity:.7}
.web-menu{position:absolute;bottom:44px;left:0;width:300px;padding:6px;border-radius:16px;z-index:30;display:flex;flex-direction:column;gap:2px;background:rgba(13,13,27,.97);border:1px solid var(--line2);box-shadow:0 20px 50px -10px rgba(0,0,0,.7)}
.web-menu button{display:grid;grid-template-columns:24px 1fr;align-items:center;column-gap:8px;text-align:left;padding:9px 10px;border-radius:10px;border:0;background:none;color:var(--text);font:inherit;font-size:.86rem;cursor:pointer}
.web-menu button svg{grid-row:span 2;color:var(--muted)}
.web-menu button small{grid-column:2;color:var(--dim);font-size:.72rem;line-height:1.35}
.web-menu button:hover{background:rgba(255,255,255,.06)}
.web-menu button.sel{background:rgba(125,211,252,.1)}.web-menu button.sel svg{color:#7dd3fc}
.web-panel{margin:2px 0 12px}
.wp-head{display:flex;align-items:center;gap:8px;font-size:.8rem;color:#7dd3fc}
.wp-head .orb{background:conic-gradient(#38bdf8,#a78bfa,#f97316,#38bdf8)}
.wp-q{display:flex;flex-wrap:wrap;gap:6px;margin-top:8px}
.wp-q span{font-size:.72rem;padding:3px 10px;border-radius:999px;background:rgba(125,211,252,.07);border:1px solid rgba(125,211,252,.2);color:#bae6fd;display:inline-flex;align-items:center;gap:5px}
.wp-q span.r{background:rgba(249,115,22,.08);border-color:rgba(249,115,22,.28);color:#fed7aa}
.src-row{display:flex;gap:8px;overflow-x:auto;padding:9px 0 4px;scrollbar-width:thin}
.src{flex:0 0 200px;display:flex;flex-direction:column;gap:6px;padding:9px 11px;border-radius:12px;background:rgba(255,255,255,.035);border:1px solid var(--line);text-decoration:none;color:var(--text);font-size:.76rem;line-height:1.35;transition:.15s}
.src:hover{border-color:rgba(125,211,252,.5);background:rgba(125,211,252,.06)}
.src .st{display:-webkit-box;-webkit-line-clamp:2;-webkit-box-orient:vertical;overflow:hidden;font-weight:500}
.src .sd{display:flex;align-items:center;gap:6px;color:var(--dim);font-size:.7rem;min-width:0}
.src .sd img{width:14px;height:14px;border-radius:3px;flex-shrink:0}
.src .sd .dm{overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
.src .sd .n{margin-left:auto;flex-shrink:0;font-weight:600;color:#7dd3fc}
.wp-q span.x{background:rgba(167,139,250,.08);border-color:rgba(167,139,250,.28);color:#ddd6fe}
.src .kb{align-self:flex-start;font-size:.6rem;font-weight:600;letter-spacing:.02em;padding:1px 7px;border-radius:999px;background:rgba(255,255,255,.07);color:var(--muted)}
.src.k-yt .kb{background:rgba(239,68,68,.15);color:#fca5a5}.src.k-wk .kb{background:rgba(226,232,240,.12);color:#e2e8f0}
.src.k-pp .kb{background:rgba(52,211,153,.14);color:#6ee7b7}.src.k-cd .kb{background:rgba(96,165,250,.15);color:#93c5fd}
.src.k-hn .kb{background:rgba(251,146,60,.15);color:#fdba74}.src.k-lk .kb{background:rgba(167,139,250,.18);color:#ddd6fe}.src.k-lk{border-color:rgba(167,139,250,.4)}
.img-row{display:flex;gap:8px;overflow-x:auto;padding:10px 0 2px;scrollbar-width:thin}
.img-row a{flex:0 0 auto;height:118px;border-radius:12px;overflow:hidden;border:1px solid var(--line);background:rgba(255,255,255,.04)}
.img-row img{height:100%;width:auto;max-width:240px;object-fit:cover;display:block;transition:.2s}.img-row a:hover img{transform:scale(1.04)}
.src.reddit{border-color:rgba(249,115,22,.28)}.src.reddit:hover{border-color:rgba(249,115,22,.6);background:rgba(249,115,22,.06)}.src.reddit .n{color:#fdba74}
.cite{display:inline-flex;align-items:center;justify-content:center;min-width:17px;height:17px;padding:0 4px;margin:0 1px;border-radius:6px;font-size:.66rem;font-weight:600;line-height:1;vertical-align:2px;background:rgba(125,211,252,.14);color:#7dd3fc!important;text-decoration:none!important;transition:.12s}
.cite.r{background:rgba(249,115,22,.16);color:#fdba74!important}
.cite:hover{background:#7dd3fc;color:#0b1020!important}.cite.r:hover{background:#fb923c}
/* markdown */
.md{font-size:.96rem;line-height:1.72;color:#e3e5f3;word-wrap:break-word}
.md>*:first-child{margin-top:0}.md>*:last-child{margin-bottom:0}
.md p{margin:.65em 0}
.md h1,.md h2,.md h3,.md h4{font-family:'Space Grotesk',Inter,sans-serif;line-height:1.3;margin:1.3em 0 .5em;color:#fff;letter-spacing:-.01em}
.md h1{font-size:1.5rem}.md h2{font-size:1.25rem;padding-bottom:.3em;border-bottom:1px solid var(--line)}.md h3{font-size:1.08rem}.md h4{font-size:.98rem}
.md ul,.md ol{padding-left:1.4em;margin:.6em 0}.md li{margin:.28em 0}.md li::marker{color:var(--violet)}
.md a{color:#93c5fd;text-decoration:underline;text-underline-offset:3px}
.md strong{color:#fff;font-weight:600}
.md blockquote{margin:.9em 0;padding:.4em 1em;border-left:3px solid var(--violet);background:rgba(167,139,250,.07);border-radius:0 10px 10px 0;color:#c7c9de}
.md hr{border:0;height:1px;background:linear-gradient(90deg,transparent,var(--line2),transparent);margin:1.4em 0}
.md :not(pre)>code{font-family:'JetBrains Mono',monospace;font-size:.84em;padding:.15em .45em;border-radius:6px;background:rgba(167,139,250,.14);color:#e9d5ff}
.md table{border-collapse:separate;border-spacing:0;margin:1em 0;font-size:.88rem;border:1px solid var(--line2);border-radius:12px;display:block;overflow-x:auto;max-width:100%}
.md th{background:rgba(167,139,250,.12);font-weight:600;text-align:left}
.md th,.md td{padding:9px 13px;border-bottom:1px solid var(--line)}
.md tr:last-child td{border-bottom:0}
.md pre{margin:1em 0;border-radius:14px;overflow:hidden;background:rgba(6,8,20,.9);border:1px solid var(--line2)}
.md pre code{display:block;padding:14px 16px;overflow-x:auto;font-family:'JetBrains Mono',monospace;font-size:.82rem;line-height:1.6;background:transparent!important}
.code-bar{display:flex;align-items:center;justify-content:space-between;padding:5px 8px 5px 14px;background:rgba(255,255,255,.035);border-bottom:1px solid var(--line);font-size:.72rem;color:var(--muted);font-family:'JetBrains Mono',monospace}
.code-bar button{display:inline-flex;align-items:center;gap:5px;background:none;border:0;color:var(--muted);cursor:pointer;padding:4px 8px;border-radius:7px;font-size:.72rem}
.code-bar button:hover{background:rgba(255,255,255,.07);color:var(--text)}
.code-bar span:first-child{display:inline-flex;align-items:center;gap:6px}
.md pre.py-code{border-color:rgba(96,165,250,.35)}.md pre.py-code .code-bar span:first-child{color:#93c5fd}
.md pre.py-out{background:rgba(2,6,16,.75);border-color:rgba(74,222,128,.25);margin-top:-.4em}
.md pre.py-out .code-bar{color:#86efac;background:rgba(74,222,128,.05)}
.md pre.py-out code{color:#d1fae5;max-height:420px;overflow:auto;white-space:pre}
.md pre.py-out.err{border-color:rgba(248,113,113,.35)}.md pre.py-out.err .code-bar{color:#fca5a5;background:rgba(248,113,113,.06)}.md pre.py-out.err code{color:#fecaca}
.md img{max-width:100%;max-height:480px;width:auto;border-radius:12px;border:1px solid var(--line);background:#fff;cursor:zoom-in;margin:.4em 0}
.mermaid-box{margin:1em 0;border-radius:14px;border:1px solid var(--line2);background:rgba(6,8,20,.6);overflow:hidden}
.mm-bar{display:flex;align-items:center;justify-content:space-between;padding:5px 8px 5px 14px;border-bottom:1px solid var(--line);font-size:.72rem;color:#c4b5fd}
.mm-bar span:first-child{display:inline-flex;align-items:center;gap:6px}
.mm-bar button{display:inline-flex;align-items:center;gap:5px;background:none;border:0;color:var(--muted);cursor:pointer;padding:4px 8px;border-radius:7px;font-size:.72rem}
.mm-bar button:hover{background:rgba(255,255,255,.07);color:var(--text)}
.mm-svg{padding:16px;overflow-x:auto;display:flex;justify-content:center}.mm-svg svg{max-width:100%;height:auto}
.html-prev{position:fixed;inset:0;z-index:80;background:rgba(3,3,10,.72);backdrop-filter:blur(6px);display:flex;align-items:center;justify-content:center;padding:20px}
.hp-card{width:min(1100px,100%);height:min(780px,100%);border-radius:18px;display:flex;flex-direction:column;overflow:hidden}
.hp-bar{display:flex;align-items:center;justify-content:space-between;padding:10px 14px;border-bottom:1px solid var(--line);font-size:.85rem}
.hp-bar b{display:inline-flex;align-items:center;gap:8px}.hp-bar span{display:flex;gap:6px}
.hp-bar .btn{display:inline-flex;align-items:center;gap:6px;padding:7px 11px;border-radius:10px;border:1px solid var(--line2);background:rgba(255,255,255,.05);color:var(--text);cursor:pointer;font:inherit;font-size:.78rem}
.hp-bar .btn:hover{background:rgba(255,255,255,.1)}
.html-prev iframe{flex:1;border:0;background:#fff;width:100%}
.katex-display{overflow-x:auto;overflow-y:hidden;padding:.3em 0}
.tool-chip{display:inline-flex;align-items:center;gap:7px;margin:.4em 0;padding:6px 12px;border-radius:999px;font-size:.78rem;background:rgba(96,165,250,.12);border:1px solid rgba(96,165,250,.3);color:#bfdbfe}
.caret{display:inline-block;width:7px;height:1em;vertical-align:-2px;background:linear-gradient(var(--violet),var(--blue));border-radius:2px;animation:blink 1s steps(2) infinite;margin-left:3px}
@keyframes blink{50%{opacity:0}}
.dots{display:inline-flex;gap:5px;padding:8px 0}.dots span{width:7px;height:7px;border-radius:50%;background:var(--violet);animation:dot 1.2s infinite ease-in-out}.dots span:nth-child(2){animation-delay:.15s;background:var(--blue)}.dots span:nth-child(3){animation-delay:.3s;background:var(--pink)}
@keyframes dot{0%,80%,100%{transform:scale(.6);opacity:.4}40%{transform:scale(1);opacity:1}}
/* file cards */
.fcard{display:flex;align-items:center;gap:12px;padding:10px 12px;border-radius:14px;background:rgba(255,255,255,.04);border:1px solid var(--line2);margin-top:10px;max-width:540px}
.th{width:46px;height:46px;border-radius:10px;object-fit:cover;flex-shrink:0;background:rgba(167,139,250,.1);display:flex;align-items:center;justify-content:center;color:var(--violet)}
img.th{cursor:zoom-in}
.fcard .info{flex:1;min-width:0}.fcard .nm{font-size:.86rem;font-weight:500;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}.fcard .mt{font-size:.74rem;color:var(--muted);margin-top:2px}
.fcard .acts{display:flex;gap:2px}
.fcard .acts a,.fcard .acts button{display:inline-flex;align-items:center;justify-content:center;width:32px;height:32px;border-radius:9px;color:var(--muted);background:none;border:0;cursor:pointer}
.fcard .acts a:hover,.fcard .acts button:hover{background:rgba(255,255,255,.08);color:var(--text)}
.fcard.err{border-color:rgba(248,113,113,.35)}
.actions{display:flex;align-items:center;gap:2px;margin-top:10px;opacity:.7;transition:.2s;flex-wrap:wrap}
.msg:hover .actions{opacity:1}
.actions button{display:inline-flex;align-items:center;justify-content:center;width:30px;height:30px;border-radius:8px;border:0;background:none;color:var(--muted);cursor:pointer}
.actions button:hover{background:rgba(255,255,255,.07);color:var(--text)}
.stats{font-size:.72rem;color:var(--dim);margin-left:8px}
/* composer */
.composer-wrap{padding:8px 20px 14px;flex-shrink:0}
.composer{max-width:820px;margin:0 auto;border-radius:26px;padding:12px 12px 10px 16px;transition:.25s;box-shadow:0 20px 60px -20px rgba(0,0,0,.6)}
.composer:focus-within{border-color:rgba(167,139,250,.45);box-shadow:0 0 0 1px rgba(167,139,250,.18),0 20px 60px -15px rgba(124,58,237,.4)}
.composer textarea{width:100%;background:none;border:0;outline:0;resize:none;color:var(--text);font:inherit;font-size:.98rem;line-height:1.55;max-height:220px;padding:4px 2px}
.composer textarea::placeholder{color:var(--dim)}
.composer-row{display:flex;align-items:center;justify-content:space-between;margin-top:6px;gap:8px}
.composer-row .left,.composer-row .right{display:flex;align-items:center;gap:6px}
.mode-toggle{position:relative;display:flex;padding:3px;border-radius:999px;background:rgba(255,255,255,.05);border:1px solid var(--line)}
.mode-toggle button{position:relative;z-index:1;display:inline-flex;align-items:center;gap:6px;padding:6px 13px;border:0;background:none;color:var(--muted);font-size:.8rem;font-weight:500;cursor:pointer;border-radius:999px;transition:color .2s}
.mode-toggle button.on{color:#fff}
.mode-pill{position:absolute;top:3px;bottom:3px;left:3px;width:0;border-radius:999px;transition:all .35s cubic-bezier(.4,1.3,.5,1);background:linear-gradient(135deg,rgba(251,191,36,.28),rgba(245,158,11,.16));border:1px solid rgba(251,191,36,.35)}
.mode-toggle.think .mode-pill{background:linear-gradient(135deg,rgba(167,139,250,.42),rgba(124,58,237,.24));border-color:rgba(167,139,250,.5)}
.send-btn{width:40px;height:40px;border-radius:50%;border:0;display:inline-flex;align-items:center;justify-content:center;cursor:pointer;color:#fff;background:linear-gradient(135deg,#8b5cf6,#3b82f6);box-shadow:0 6px 20px -4px rgba(139,92,246,.6);transition:.2s}
.send-btn:hover{transform:scale(1.06)}
.send-btn.stop{background:linear-gradient(135deg,#ef4444,#db2777);box-shadow:0 6px 20px -4px rgba(239,68,68,.6)}
.mic-on{color:#fff!important;background:#ef4444!important;animation:pulse 1.4s infinite}
@keyframes pulse{0%{box-shadow:0 0 0 0 rgba(239,68,68,.6)}70%{box-shadow:0 0 0 10px rgba(239,68,68,0)}100%{box-shadow:0 0 0 0 rgba(239,68,68,0)}}
.hint{text-align:center;font-size:.7rem;color:var(--dim);margin:8px 0 0}
.pending{display:flex;flex-wrap:wrap;gap:8px;padding-bottom:10px;margin-bottom:6px;border-bottom:1px solid var(--line)}
.chip{position:relative;display:flex;align-items:center;gap:9px;padding:6px 32px 6px 6px;border-radius:12px;background:rgba(255,255,255,.05);border:1px solid var(--line2);max-width:250px}
.chip .th{width:36px;height:36px;border-radius:8px}
.chip .nm{font-size:.78rem;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}.chip .mt{font-size:.68rem;color:var(--muted)}
.chip .x{position:absolute;right:6px;top:50%;transform:translateY(-50%);width:20px;height:20px;border-radius:50%;border:0;background:rgba(0,0,0,.45);color:#fff;display:flex;align-items:center;justify-content:center;cursor:pointer}
.chip.err{border-color:rgba(248,113,113,.5)}
.att-chip{display:inline-flex;align-items:center;gap:6px;padding:5px 10px;border-radius:10px;background:rgba(0,0,0,.25);font-size:.78rem;color:#dbeafe}
.att-img{height:110px;max-width:220px;border-radius:12px;object-fit:cover;cursor:zoom-in}
.drop-overlay{position:absolute;inset:12px;z-index:30;border-radius:28px;border:2px dashed rgba(167,139,250,.6);background:rgba(10,8,30,.75);backdrop-filter:blur(6px);display:flex;align-items:center;justify-content:center;font-size:1.4rem;font-family:'Space Grotesk',sans-serif;pointer-events:none}
.hidden{display:none!important}
/* modal */
.modal{position:fixed;inset:0;z-index:50;background:rgba(3,4,12,.6);backdrop-filter:blur(6px);display:flex;align-items:center;justify-content:center;padding:20px}
.modal-card{width:min(1000px,100%);height:min(700px,100%);border-radius:24px;display:flex;flex-direction:column;overflow:hidden;background:rgba(12,13,30,.9);border:1px solid var(--line2);backdrop-filter:blur(24px)}
.modal-head{display:flex;align-items:center;gap:10px;padding:12px 14px 12px 20px;border-bottom:1px solid var(--line);font-weight:600}
.modal-head .sp{flex:1}
.modal-body{flex:1;display:flex;min-height:0}
.tool-list{width:240px;flex-shrink:0;overflow-y:auto;border-right:1px solid var(--line);padding:8px}
.tl-group{font-size:.68rem;letter-spacing:.08em;text-transform:uppercase;color:var(--dim);padding:12px 10px 6px}
.tl-item{padding:8px 10px;border-radius:10px;font-size:.85rem;color:var(--muted);cursor:pointer}
.tl-item:hover{background:rgba(255,255,255,.05);color:var(--text)}.tl-item.on{background:rgba(167,139,250,.16);color:#fff}
.tool-pane{flex:1;overflow-y:auto;padding:20px 24px 28px}
.tool-pane h3{margin:0;font-family:'Space Grotesk',sans-serif;font-size:1.3rem}
.tool-pane .desc{color:var(--muted);font-size:.86rem;margin:6px 0 10px}
.flabel{display:flex;align-items:center;justify-content:space-between;font-size:.78rem;color:var(--muted);margin:16px 0 7px;font-weight:500}
.tool-files{display:grid;grid-template-columns:repeat(auto-fill,minmax(210px,1fr));gap:8px;max-height:230px;overflow-y:auto}
.tf{display:flex;align-items:center;gap:9px;padding:7px 30px 7px 7px;border-radius:12px;border:1px solid var(--line);background:rgba(255,255,255,.03);cursor:pointer;font-size:.8rem;position:relative}
.tf.on{border-color:rgba(167,139,250,.65);background:rgba(167,139,250,.13)}
.tf .th{width:34px;height:34px;border-radius:8px;cursor:pointer}
.tf .nm{white-space:nowrap;overflow:hidden;text-overflow:ellipsis}.tf .mt{font-size:.68rem;color:var(--muted)}
.tf .ord{position:absolute;right:8px;top:8px;width:18px;height:18px;border-radius:50%;background:var(--violet);color:#fff;font-size:.65rem;display:flex;align-items:center;justify-content:center}
.inp{width:100%;padding:9px 12px;border-radius:11px;border:1px solid var(--line2);background:rgba(255,255,255,.04);color:var(--text);font:inherit;font-size:.88rem;outline:0}
.inp:focus{border-color:rgba(167,139,250,.6)}
textarea.inp{min-height:120px;resize:vertical}
select.inp option{background:#151733}
.run-btn{margin-top:22px;display:inline-flex;align-items:center;gap:8px;padding:11px 22px;border-radius:999px;border:0;cursor:pointer;color:#fff;font-weight:600;background:linear-gradient(135deg,#8b5cf6,#3b82f6);box-shadow:0 8px 24px -6px rgba(139,92,246,.6)}
.run-btn:disabled{opacity:.55;cursor:default}
.small-btn{display:inline-flex;align-items:center;gap:5px;padding:4px 10px;border-radius:8px;border:1px solid var(--line2);background:none;color:var(--muted);font-size:.74rem;cursor:pointer}
.small-btn:hover{color:var(--text);border-color:rgba(167,139,250,.5)}
.empty{color:var(--dim);font-size:.82rem;padding:8px 2px}
#toasts{position:fixed;bottom:110px;left:50%;transform:translateX(-50%);z-index:80;display:flex;flex-direction:column;gap:8px;align-items:center;pointer-events:none}
.toast{padding:10px 16px;border-radius:12px;background:rgba(20,22,46,.96);border:1px solid var(--line2);font-size:.85rem;animation:rise .25s ease;box-shadow:0 10px 30px rgba(0,0,0,.5)}
.toast.err{border-color:rgba(248,113,113,.55)}
.lightbox{position:fixed;inset:0;z-index:90;background:rgba(0,0,0,.86);display:flex;align-items:center;justify-content:center;cursor:zoom-out;padding:20px}
.lightbox img{max-width:100%;max-height:100%;border-radius:12px;box-shadow:0 20px 80px rgba(0,0,0,.6)}
::-webkit-scrollbar{width:8px;height:8px}::-webkit-scrollbar-thumb{background:rgba(255,255,255,.1);border-radius:99px}::-webkit-scrollbar-thumb:hover{background:rgba(255,255,255,.18)}::-webkit-scrollbar-track{background:transparent}
.cat-eyes{transform-box:fill-box;transform-origin:center;animation:catblink 6s infinite}
@keyframes catblink{0%,94%,100%{transform:scaleY(1)}96%{transform:scaleY(.08)}}
@media (max-width:820px){
  .sidebar{position:fixed;z-index:40;top:0;bottom:0;left:0;transform:translateX(-100%)}
  .app.sb-open .sidebar{transform:none}
  .app.sb-collapsed .sidebar{margin-left:0}
  .app.sb-open .scrim{display:block;position:fixed;inset:0;z-index:35;background:rgba(0,0,0,.5)}
  .suggest-grid{grid-template-columns:1fr}
  .welcome{padding-top:4vh}
  .bubble{max-width:90%}.u-wrap{max-width:90%}
  .messages,.welcome{padding:0 14px}
  .composer-wrap{padding:6px 10px 10px}
  .topbar{padding:0 10px}
  .pill-btn .lbl{display:none}
  .web-btn{padding:0 9px}#webLbl{display:none}#micBtn{display:none}.composer-row .left,.composer-row .right{gap:2px}.composer{padding:10px 8px 8px 12px}.web-menu{position:fixed;left:12px;right:12px;bottom:118px;width:auto}
  .modal{padding:0}.modal-card{height:100%;border-radius:0}
  .modal-body{flex-direction:column}.tool-list{width:100%;max-height:32%;border-right:0;border-bottom:1px solid var(--line)}
  .msg{gap:10px}.avatar{width:28px;height:28px}
}
@media (min-width:821px){#sbClose{display:none}}
@media (prefers-reduced-motion:reduce){*{animation:none!important;transition:none!important}}
</style>
</head>
<body>
<svg width="0" height="0" style="position:absolute" aria-hidden="true"><defs>
  <radialGradient id="catFur" cx="50%" cy="36%" r="68%"><stop offset="0" stop-color="#d4d8de"/><stop offset=".55" stop-color="#9ea4ad"/><stop offset="1" stop-color="#6b717a"/></radialGradient>
  <radialGradient id="catEye" cx="40%" cy="35%" r="70%"><stop offset="0" stop-color="#e2fbb0"/><stop offset="1" stop-color="#5fa317"/></radialGradient>
</defs></svg>
<canvas id="sky"></canvas>
<div class="nebula n1"></div><div class="nebula n2"></div><div class="nebula n3"></div>

<div class="app" id="app">
  <aside class="sidebar" id="sidebar">
    <div class="sb-top">
      <div class="logo-row"><span class="lg cat-slot"></span><span class="brand">Meow</span><span>🐈</span><span class="sp"></span><button class="icon-btn" id="sbClose" title="Close"><i data-i="x"></i></button></div>
      <button class="new-chat" id="newChatBtn"><i data-i="plus"></i>New chat<kbd>Ctrl K</kbd></button>
      <div class="search"><i data-i="search" data-s="15"></i><input id="chatSearch" placeholder="Search chats" autocomplete="off"></div>
    </div>
    <div class="chat-list" id="chatList"></div>
    <div class="sb-bottom">
      <button class="sb-link" id="sbTools"><i data-i="wrench" data-s="16"></i>File tools &amp; library</button>
      <button class="sb-link" id="sbNews"><i data-i="sparkle" data-s="16"></i>What's new<span class="ver" id="verTag"></span></button>
      <button class="sb-link power" id="sbPower"><i data-i="power" data-s="16"></i>Power off Meow</button>
      <div class="health"><span class="dot" id="hdot"></span><span id="healthText">Connecting to Kaggle…</span></div>
    </div>
  </aside>
  <div class="scrim" id="scrim"></div>

  <main class="main" id="main">
    <header class="topbar">
      <button class="icon-btn" id="menuBtn" title="Menu"><i data-i="menu"></i></button>
      <div class="title-wrap"><span class="brand">Meow 🐈</span><span class="model-badge" id="modelBadge">Qwen · Kaggle GPU</span></div>
      <button class="pill-btn" id="toolsBtn" title="File tools"><i data-i="wrench" data-s="15"></i><span class="lbl">Tools</span></button>
      <button class="icon-btn" id="exportBtn" title="Export chat as Markdown"><i data-i="share"></i></button>
    </header>

    <div class="scroller" id="scroller">
      <div class="welcome" id="welcome">
        <div class="hero-cat cat-slot"></div>
        <h1 class="brand hero">Meow</h1>
        <p class="sub" id="greet">What are we working on today?</p>
        <div class="suggest-grid">
          <div class="suggest" data-prompt="Summarize this PDF into the key points, then list any action items or deadlines." data-attach="1"><div class="e">📄</div><div class="h">Summarize a PDF</div><div class="d">Key points, action items, deadlines</div></div>
          <div class="suggest" data-prompt="Describe this image in detail and pull out any text in it." data-attach="1"><div class="e">🖼️</div><div class="h">Analyze an image</div><div class="d">Scene, text (OCR), colors, metadata</div></div>
          <div class="suggest" data-mode="think" data-prompt="Explain how a hash map handles collisions — chaining vs open addressing — with time complexity and a clean Python implementation."><div class="e">🧠</div><div class="h">Think it through</div><div class="d">Deep reasoning for hard problems</div></div>
          <div class="suggest" data-web="reddit" data-prompt="What does Reddit recommend as the best budget laptop for coding and college right now?"><div class="e">🟠</div><div class="h">Ask Reddit</div><div class="d">Real opinions from across all subreddits</div></div>
          <div class="suggest" data-prompt="Analyze this spreadsheet with Python: key statistics, trends and two clear charts, then summarize the insights." data-attach="1"><div class="e">📊</div><div class="h">Analyze data</div><div class="d">Runs real Python — stats &amp; charts</div></div>
          <div class="suggest" data-prompt="Create a one-page PDF study guide on the OSI model: each layer, what it does, protocols, and a memory trick."><div class="e">✨</div><div class="h">Create a PDF</div><div class="d">Reports, notes and guides as real files</div></div>
        </div>
      </div>
      <div class="messages" id="messages"></div>
    </div>

    <div class="composer-wrap">
      <div class="composer glass" id="composer">
        <div class="edit-bar hidden" id="editBar"><i data-i="pen" data-s="14"></i><span>Editing your message — sending replaces everything after it</span><button id="editCancel">Cancel</button></div>
        <div class="pending hidden" id="pending"></div>
        <textarea id="input" rows="1" placeholder="Message Meow…"></textarea>
        <div class="composer-row">
          <div class="left">
            <input type="file" id="fileInput" multiple hidden>
            <button class="icon-btn" id="attachBtn" title="Attach files (PDF, images, docs, sheets, code)"><i data-i="clip" data-s="19"></i></button>
            <input type="file" id="camInput" accept="image/*" capture="environment" hidden>
            <button class="icon-btn touch-only" id="camBtn" title="Take a photo"><i data-i="camera" data-s="19"></i></button>
            <div class="mode-toggle" id="modeToggle"><span class="mode-pill" id="modePill"></span>
              <button data-mode="fast"><i data-i="zap" data-s="14"></i><span>Fast</span></button>
              <button data-mode="think"><i data-i="brain" data-s="14"></i><span>Think</span></button>
            </div>
            <div class="web-pick" id="webPick">
              <button class="web-btn" id="webBtn" title="Web search"><i data-i="globe" data-s="15"></i><span id="webLbl">Auto</span></button>
              <div class="web-menu hidden" id="webMenu">
                <button data-web="auto"><i data-i="sparkle" data-s="17"></i><b>Auto</b><small>Searches only when the question needs fresh info</small></button>
                <button data-web="web"><i data-i="globe" data-s="17"></i><b>Web search</b><small>Always search the web and cite sources</small></button>
                <button data-web="reddit"><i data-i="reddit" data-s="17"></i><b>Reddit</b><small>What real people say, across all subreddits</small></button>
                <button data-web="deep"><i data-i="layers" data-s="17"></i><b>Deep research</b><small>More queries, more pages, plus Reddit (slower)</small></button>
                <button data-web="off"><i data-i="globeoff" data-s="17"></i><b>Off</b><small>Answer from the model's own knowledge</small></button>
              </div>
            </div>
          </div>
          <div class="right">
            <button class="icon-btn" id="voiceBtn" title="Voice mode — talk to Meow hands-free"><i data-i="waves" data-s="19"></i></button>
            <button class="icon-btn" id="micBtn" title="Voice input"><i data-i="mic" data-s="19"></i></button>
            <button class="send-btn" id="sendBtn" title="Send"><i data-i="send" data-s="18"></i></button>
          </div>
        </div>
      </div>
      <p class="hint">Meow can make mistakes — double-check anything important.</p>
    </div>
    <div class="voice-ov hidden" id="voiceOv">
      <div class="vo-orb" id="voOrb"><span></span><span></span><span></span></div>
      <div class="vo-state" id="voState">Listening…</div>
      <div class="vo-text" id="voText"></div>
      <div class="vo-btns"><button id="voMute" title="Pause / resume listening"><i data-i="mic" data-s="18"></i></button><button id="voEnd" class="end" title="End voice mode"><i data-i="x" data-s="20"></i></button></div>
      <div class="vo-hint">Speak naturally — Meow answers out loud, then listens again. Works best in Chrome / Edge.</div>
    </div>
    <div class="drop-overlay hidden" id="dropOverlay">🐈 Drop files for Meow</div>
  </main>
</div>

<div class="modal hidden" id="toolsModal">
  <div class="modal-card">
    <div class="modal-head"><i data-i="wrench" data-s="17"></i><span>File tools</span><span class="sp"></span><button class="icon-btn" id="toolsClose" title="Close"><i data-i="x"></i></button></div>
    <div class="modal-body">
      <div class="tool-list" id="toolList"></div>
      <div class="tool-pane">
        <h3 id="toolTitle">Tools</h3>
        <p class="desc" id="toolDesc"></p>
        <div id="toolFilesWrap">
          <div class="flabel"><span id="toolFilesHint">Pick a file</span><button class="small-btn" id="toolUpload"><i data-i="plus" data-s="13"></i>Upload</button></div>
          <input type="file" id="toolFileInput" multiple hidden>
          <div class="tool-files" id="toolFiles"></div>
        </div>
        <div id="toolParams"></div>
        <button class="run-btn" id="toolRun"><i data-i="sparkles" data-s="16"></i>Run</button>
        <div id="toolResult"></div>
      </div>
    </div>
  </div>
</div>
<style>
.pe{position:fixed;inset:0;z-index:65;display:flex;flex-direction:column;background:rgba(6,7,18,.97);backdrop-filter:blur(10px)}
.pe-top{display:flex;align-items:center;gap:10px;padding:8px 12px;border-bottom:1px solid var(--line);flex-wrap:wrap;background:rgba(12,13,30,.9)}
.pe-title{display:flex;align-items:center;gap:8px;min-width:0;max-width:240px}
.pe-title .cat-slot{width:26px;height:26px;flex-shrink:0}
.pe-title #peName{font-weight:600;font-size:.88rem;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
.pe-status{font-size:.72rem;color:var(--dim);white-space:nowrap}
.pe-group{display:flex;align-items:center;gap:2px;padding:3px;border-radius:12px;background:rgba(255,255,255,.04);border:1px solid var(--line)}
.pe-group button{display:inline-flex;align-items:center;justify-content:center;gap:6px;height:34px;min-width:34px;padding:0 9px;border-radius:9px;border:0;background:none;color:var(--muted);cursor:pointer;font-size:.78rem;white-space:nowrap}
.pe-group button:hover{background:rgba(255,255,255,.07);color:var(--text)}
.pe-group button.on{background:linear-gradient(135deg,rgba(139,92,246,.45),rgba(59,130,246,.3));color:#fff}
.pe-group button:disabled{opacity:.35;cursor:default}
.pe-group .lbl{display:none}
@media (min-width:1600px){.pe-tools .lbl{display:inline}}
.pe-props input[type=color]{width:30px;height:30px;border:0;background:none;padding:0;cursor:pointer}
.pe-props select,.pe-props input[type=number]{height:30px;border-radius:8px;border:1px solid var(--line2);background:rgba(255,255,255,.05);color:var(--text);font:inherit;font-size:.78rem;padding:0 6px}
.pe-props input[type=number]{width:54px}
.pe-props button{font-weight:700;font-family:Georgia,serif}
.pe-sp{flex:1}
.pe-save{height:36px;display:inline-flex;align-items:center;gap:7px;padding:0 16px;border-radius:999px;border:0;cursor:pointer;color:#fff;font-weight:600;font-size:.84rem;background:linear-gradient(135deg,#8b5cf6,#3b82f6);box-shadow:0 6px 18px -6px rgba(139,92,246,.7)}
.pe-body{flex:1;display:flex;min-height:0}
.pe-thumbs{width:150px;flex-shrink:0;overflow-y:auto;padding:12px 10px;border-right:1px solid var(--line)}
.pe-th{position:relative;margin-bottom:14px;cursor:pointer}
.pe-th img{width:100%;display:block;border-radius:6px;background:#fff;border:2px solid transparent;box-shadow:0 4px 14px rgba(0,0,0,.4)}
.pe-th.on img{border-color:#8b5cf6}
.pe-th .no{text-align:center;font-size:.7rem;color:var(--muted);margin-top:4px}
.pe-th .ops{position:absolute;top:4px;right:4px;display:none;flex-direction:column;gap:3px}
.pe-th:hover .ops{display:flex}
.pe-th .ops button{width:24px;height:24px;border-radius:7px;border:0;background:rgba(10,10,25,.85);color:#fff;display:flex;align-items:center;justify-content:center;cursor:pointer}
.pe-th .ops button:hover{background:#8b5cf6}
.pe-addpage{width:100%;padding:9px;border-radius:10px;border:1px dashed var(--line2);background:none;color:var(--muted);font-size:.76rem;cursor:pointer}
.pe-addpage:hover{color:var(--text);border-color:#8b5cf6}
.pe-stage{flex:1;overflow:auto;display:flex;justify-content:center;align-items:flex-start;padding:28px}
.pe-page{position:relative;flex-shrink:0;background:#fff;box-shadow:0 20px 60px rgba(0,0,0,.6);user-select:none}
.pe-page img{position:absolute;inset:0;width:100%;height:100%;pointer-events:none}
.pe-page canvas{position:absolute;inset:0;width:100%;height:100%;pointer-events:none}
.pe-ov{position:absolute;inset:0;touch-action:none}
.pe-ov[data-tool=addtext]{cursor:text}.pe-ov[data-tool=draw]{cursor:crosshair}.pe-ov[data-tool=highlight],.pe-ov[data-tool=shape],.pe-ov[data-tool=erase],.pe-ov[data-tool=magic]{cursor:crosshair}.pe-ov[data-tool=place]{cursor:copy}
.pe-run{position:absolute;border-radius:2px;cursor:text}
.pe-ov[data-tool=select] .pe-run:hover,.pe-ov[data-tool=text] .pe-run:hover{outline:1.5px dashed rgba(124,58,237,.8);background:rgba(167,139,250,.12)}
.pe-ov[data-tool=text] .pe-run{outline:1px dashed rgba(124,58,237,.28)}
.pe-run.scan{outline:1px dashed rgba(245,158,11,.55)}
.pe-ov[data-tool=select] .pe-run.scan:hover,.pe-ov[data-tool=text] .pe-run.scan:hover{outline:1.5px dashed #f59e0b;background:rgba(245,158,11,.12)}
.pe-run.sel{outline:2px solid #8b5cf6!important;background:rgba(167,139,250,.15)}
.pe-img{position:absolute;cursor:move}
.pe-ov[data-tool=select] .pe-img:hover{outline:1.5px dashed rgba(59,130,246,.9)}
.pe-img.sel{outline:2px solid #3b82f6}
.pe-img .h{position:absolute;right:-7px;bottom:-7px;width:14px;height:14px;background:#3b82f6;border:2px solid #fff;border-radius:4px;cursor:nwse-resize;display:none}
.pe-img.sel .h{display:block}
.pe-ov:not([data-tool=select]):not([data-tool=text]) .pe-run,.pe-ov:not([data-tool=select]) .pe-img{pointer-events:none}
.pe-ov[data-tool=text] .pe-img{pointer-events:none}
.pe-edit{position:absolute;outline:2px solid #8b5cf6;white-space:pre;min-width:24px;z-index:5;padding:0 1px;margin:0;border-radius:2px;cursor:text;box-shadow:0 0 0 4px rgba(139,92,246,.2)}
.pe-rubber{position:absolute;border:1.5px dashed #8b5cf6;background:rgba(139,92,246,.1);pointer-events:none}
.pe-ghost{position:absolute;opacity:.6;pointer-events:none;outline:1.5px dashed #8b5cf6}
.pe-side{width:270px;flex-shrink:0;overflow-y:auto;padding:16px;border-left:1px solid var(--line);font-size:.84rem}
.pe-side h4{margin:0 0 4px;font-family:'Space Grotesk',sans-serif;font-size:1rem}
.pe-side .muted{color:var(--muted);font-size:.78rem;line-height:1.5}
.pe-side .grid{display:grid;grid-template-columns:1fr 1fr;gap:6px;margin-top:10px}
.pe-side .btn{display:flex;align-items:center;justify-content:center;gap:6px;padding:8px;border-radius:10px;border:1px solid var(--line2);background:rgba(255,255,255,.04);color:var(--text);cursor:pointer;font-size:.76rem}
.pe-side .btn:hover{border-color:#8b5cf6;background:rgba(139,92,246,.14)}
.pe-side .btn.danger:hover{border-color:var(--bad);background:rgba(248,113,113,.12)}
.pe-side .btn.wide{grid-column:1/-1}
.pe-side .prev{width:100%;max-height:140px;object-fit:contain;border-radius:10px;background:repeating-conic-gradient(#2a2c40 0 25%,#1d1f30 0 50%) 0/16px 16px;margin-top:10px}
.pe-side .sec{margin-top:16px;font-size:.68rem;letter-spacing:.08em;text-transform:uppercase;color:var(--dim)}
.pe-side kbd{font-family:inherit;font-size:.7rem;border:1px solid var(--line2);border-radius:5px;padding:0 5px;color:var(--muted)}
.pe-side .row{display:flex;align-items:center;gap:8px;margin-top:8px}
.pe-side .row label{font-size:.76rem;color:var(--muted);min-width:64px}
.pe-busy{position:absolute;right:18px;bottom:18px;display:flex;align-items:center;gap:8px;padding:8px 14px;border-radius:999px;background:rgba(20,22,46,.95);border:1px solid var(--line2);font-size:.78rem;z-index:6}
.sig-pad{width:100%;height:200px;border-radius:14px;background:#fff;touch-action:none;cursor:crosshair}
@media (max-width:820px){.pe-thumbs{display:none}.pe-side{position:absolute;left:0;right:0;bottom:0;width:auto;max-height:42%;background:rgba(12,13,30,.97);border-left:0;border-top:1px solid var(--line);z-index:7}.pe-side.empty{display:none}.pe-stage{padding:12px}.pe-title{max-width:120px}}
</style>
<div class="pe hidden" id="pdfEditor">
  <div class="pe-top">
    <div class="pe-title"><span class="cat-slot"></span><span id="peName">document.pdf</span><span class="pe-status" id="peStatus"></span></div>
    <div class="pe-group pe-tools" id="peTools">
      <button data-tool="select" title="Select & move (V)"><i data-i="cursor" data-s="16"></i><span class="lbl">Select</span></button>
      <button data-tool="text" title="Edit text (E)"><i data-i="type" data-s="16"></i><span class="lbl">Edit text</span></button>
      <button data-tool="addtext" title="Add text (T)"><i data-i="typeplus" data-s="16"></i><span class="lbl">Add text</span></button>
      <button data-tool="image" title="Add image (I)"><i data-i="image" data-s="16"></i><span class="lbl">Image</span></button>
      <button data-tool="draw" title="Draw (D)"><i data-i="pen" data-s="16"></i><span class="lbl">Draw</span></button>
      <button data-tool="highlight" title="Highlight (H)"><i data-i="marker" data-s="16"></i><span class="lbl">Highlight</span></button>
      <button data-tool="shape" title="Shapes (R)"><i data-i="square" data-s="16"></i><span class="lbl">Shape</span></button>
      <button data-tool="erase" title="Erase area (X)"><i data-i="eraser" data-s="16"></i><span class="lbl">Erase</span></button>
      <button data-tool="magic" title="Magic eraser (M)"><i data-i="wand" data-s="16"></i><span class="lbl">Magic</span></button>
      <button data-tool="sign" title="Signature (S)"><i data-i="sign" data-s="16"></i><span class="lbl">Sign</span></button>
    </div>
    <div class="pe-group pe-props" title="Style for new text, drawing and edits">
      <input type="color" id="peColor" value="#111111" title="Colour">
      <select id="peFamily" title="Font"><option value="sans">Sans</option><option value="serif">Serif</option><option value="mono">Mono</option></select>
      <input type="number" id="peSize" value="12" min="4" max="200" step="0.5" title="Size (pt)">
      <button id="peBold" title="Bold">B</button><button id="peItalic" title="Italic"><i style="font-family:Georgia,serif">I</i></button>
    </div>
    <div class="pe-group">
      <button id="peUndo" title="Undo (Ctrl+Z)"><i data-i="undo" data-s="16"></i></button>
      <button id="peRedo" title="Redo (Ctrl+Y)"><i data-i="redo" data-s="16"></i></button>
    </div>
    <div class="pe-group">
      <button id="peZoomOut" title="Zoom out"><i data-i="zout" data-s="16"></i></button>
      <button id="peZoom" title="Fit width" style="min-width:52px">100%</button>
      <button id="peZoomIn" title="Zoom in"><i data-i="zin" data-s="16"></i></button>
    </div>
    <span class="pe-sp"></span>
    <div class="pe-group"><button id="peDownload" title="Download PDF"><i data-i="download" data-s="16"></i></button></div>
    <button class="pe-save" id="peSave"><i data-i="check" data-s="16"></i>Save</button>
    <button class="icon-btn" id="peClose" title="Close (Esc)"><i data-i="x"></i></button>
  </div>
  <div class="pe-body">
    <div class="pe-thumbs" id="peThumbs"></div>
    <div class="pe-stage" id="peStage">
      <div class="pe-page" id="pePage"><img id="peImg" alt=""><canvas id="peInk"></canvas><div class="pe-ov" id="peOv" data-tool="select"></div></div>
    </div>
    <div class="pe-side" id="peSide"></div>
  </div>
  <div class="pe-busy hidden" id="peBusy"><span class="spin" style="width:16px;height:16px"></span><span id="peBusyText">Applying…</span></div>
</div>
<input type="file" id="peFile" accept="image/*,.heic,.heif" hidden>
<div class="modal hidden" id="signModal" style="z-index:70">
  <div class="modal-card" style="height:auto;width:min(560px,100%)">
    <div class="modal-head"><i data-i="sign" data-s="17"></i><span>Draw your signature</span><span class="sp"></span><button class="icon-btn" id="signClose"><i data-i="x"></i></button></div>
    <div style="padding:18px">
      <canvas class="sig-pad" id="sigPad"></canvas>
      <div style="display:flex;gap:10px;margin-top:14px;align-items:center">
        <input type="color" id="sigColor" value="#0b1f6b" style="width:34px;height:34px;border:0;background:none">
        <button class="small-btn" id="sigClear">Clear</button>
        <button class="small-btn" id="sigUpload">Upload image instead</button>
        <span class="pe-sp"></span>
        <button class="pe-save" id="sigUse">Place signature</button>
      </div>
    </div>
  </div>
</div>
<div id="toasts"></div>
<div class="lightbox hidden" id="lightbox"><img alt=""></div>
<script>
(() => {
'use strict';
const USER_NAME = '__USER_NAME__';
const HDR = {'ngrok-skip-browser-warning': '1'};
const $ = s => document.querySelector(s);
const esc = s => String(s ?? '').replace(/[&<>"']/g, m => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[m]));
const uid = () => Math.random().toString(36).slice(2, 10);
const isTouch = matchMedia('(pointer: coarse)').matches;

/* ---------------- icons ---------------- */
const ICONS = {
  plus:'<path d="M12 5v14M5 12h14"/>', x:'<path d="M18 6 6 18M6 6l12 12"/>',
  send:'<path d="M22 2 11 13"/><path d="M22 2 15 22l-4-9-9-4 20-7z"/>', stop:'<rect x="6" y="6" width="12" height="12" rx="2" fill="currentColor"/>',
  mic:'<rect x="9" y="2" width="6" height="12" rx="3"/><path d="M19 10a7 7 0 0 1-14 0M12 17v5M8 22h8"/>',
  clip:'<path d="m21.44 11.05-9.19 9.19a6 6 0 0 1-8.49-8.49l8.57-8.57A4 4 0 1 1 18 8.84l-8.59 8.57a2 2 0 0 1-2.83-2.83l8.49-8.48"/>',
  zap:'<path d="M13 2 3 14h9l-1 8 10-12h-9l1-8z"/>',
  brain:'<path d="M9.5 2A2.5 2.5 0 0 0 7 4.5v.5a3 3 0 0 0-3 3v1a3 3 0 0 0 0 6v1a3 3 0 0 0 3 3 2.5 2.5 0 0 0 5 0V4.5A2.5 2.5 0 0 0 9.5 2z"/><path d="M14.5 2A2.5 2.5 0 0 1 17 4.5v.5a3 3 0 0 1 3 3v1a3 3 0 0 1 0 6v1a3 3 0 0 1-3 3 2.5 2.5 0 0 1-5 0"/>',
  copy:'<rect x="9" y="9" width="13" height="13" rx="2"/><path d="M5 15H4a2 2 0 0 1-2-2V4a2 2 0 0 1 2-2h9a2 2 0 0 1 2 2v1"/>',
  download:'<path d="M21 15v4a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2v-4M7 10l5 5 5-5M12 15V3"/>',
  refresh:'<path d="M21 12a9 9 0 1 1-3-6.7L21 8"/><path d="M21 3v5h-5"/>',
  volume:'<path d="M11 5 6 9H2v6h4l5 4V5z"/><path d="M15.5 8.5a5 5 0 0 1 0 7M19 5a10 10 0 0 1 0 14"/>',
  trash:'<path d="M3 6h18M8 6V4h8v2M19 6l-1 14H6L5 6"/>', menu:'<path d="M4 6h16M4 12h16M4 18h16"/>',
  search:'<circle cx="11" cy="11" r="7"/><path d="m21 21-4.3-4.3"/>',
  globe:'<circle cx="12" cy="12" r="10"/><path d="M2 12h20M12 2a15.3 15.3 0 0 1 4 10 15.3 15.3 0 0 1-4 10 15.3 15.3 0 0 1-4-10 15.3 15.3 0 0 1 4-10z"/>',
  globeoff:'<circle cx="12" cy="12" r="10"/><path d="M2 12h20M12 2a15.3 15.3 0 0 1 4 10 15.3 15.3 0 0 1-4 10 15.3 15.3 0 0 1-4-10 15.3 15.3 0 0 1 4-10z"/><path d="m4 4 16 16"/>',
  sparkle:'<path d="M12 3l1.8 5.2L19 10l-5.2 1.8L12 17l-1.8-5.2L5 10l5.2-1.8z"/><path d="M19 17l.7 1.8 1.8.7-1.8.7L19 22l-.7-1.8-1.8-.7 1.8-.7z"/>',
  reddit:'<circle cx="12" cy="14" r="7"/><circle cx="9.5" cy="13.5" r="1" fill="currentColor"/><circle cx="14.5" cy="13.5" r="1" fill="currentColor"/><path d="M9.5 16.5c1.4 1 3.6 1 5 0M12 7l1.2-4 3.8 1"/><circle cx="18" cy="4" r="1.3"/><circle cx="19.5" cy="10" r="1.6"/><circle cx="4.5" cy="10" r="1.6"/>',
  layers:'<path d="m12 2 10 5-10 5L2 7z"/><path d="m2 17 10 5 10-5M2 12l10 5 10-5"/>',
  wrench:'<path d="M14.7 6.3a1 1 0 0 0 0 1.4l1.6 1.6a1 1 0 0 0 1.4 0l3.77-3.77a6 6 0 0 1-7.94 7.94l-6.91 6.91a2.12 2.12 0 0 1-3-3l6.91-6.91a6 6 0 0 1 7.94-7.94l-3.76 3.76z"/>',
  file:'<path d="M14 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V8z"/><path d="M14 2v6h6M8 13h8M8 17h5"/>',
  image:'<rect x="3" y="3" width="18" height="18" rx="2"/><circle cx="9" cy="9" r="2"/><path d="m21 15-5-5L5 21"/>',
  table:'<rect x="3" y="3" width="18" height="18" rx="2"/><path d="M3 9h18M3 15h18M9 3v18"/>',
  code:'<path d="m16 18 6-6-6-6M8 6l-6 6 6 6"/>', folder:'<path d="M4 20h16a2 2 0 0 0 2-2V8a2 2 0 0 0-2-2h-7.9a2 2 0 0 1-1.7-.9L9.6 3.9A2 2 0 0 0 7.9 3H4a2 2 0 0 0-2 2v13a2 2 0 0 0 2 2z"/>',
  check:'<path d="M20 6 9 17l-5-5"/>', power:'<path d="M12 2v10"/><path d="M18.4 6.6a9 9 0 1 1-12.8 0"/>', camera:'<path d="M23 19a2 2 0 0 1-2 2H3a2 2 0 0 1-2-2V8a2 2 0 0 1 2-2h4l2-3h6l2 3h4a2 2 0 0 1 2 2z"/><circle cx="12" cy="13" r="4"/>', waves:'<path d="M2 10v3M6 6v11M10 3v18M14 8v7M18 5v13M22 10v3"/>', terminal:'<path d="m4 17 6-6-6-6M12 19h8"/>', code:'<path d="m16 18 6-6-6-6M8 6l-6 6 6 6"/>', chev:'<path d="m6 9 6 6 6-6"/>',
  sparkles:'<path d="M12 3l1.9 5.8L20 10l-6.1 1.2L12 17l-1.9-5.8L4 10l6.1-1.2z"/><path d="M19 17v4M17 19h4"/>',
  eye:'<path d="M2 12s3.5-7 10-7 10 7 10 7-3.5 7-10 7S2 12 2 12z"/><circle cx="12" cy="12" r="3"/>',
  share:'<path d="M4 12v8a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2v-8M16 6l-4-4-4 4M12 2v13"/>',
  chat:'<path d="M21 15a2 2 0 0 1-2 2H7l-4 4V5a2 2 0 0 1 2-2h14a2 2 0 0 1 2 2z"/>',
  slides:'<rect x="2" y="3" width="20" height="14" rx="2"/><path d="M8 21h8M12 17v4"/>'
};
const ic = (n, s = 16) => `<svg width="${s}" height="${s}" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round">${ICONS[n] || ''}</svg>`;
function hydrateIcons(root = document) { root.querySelectorAll('[data-i]').forEach(el => { el.outerHTML = ic(el.dataset.i, +el.dataset.s || 18); }); }

/* ---------------- the grey cat (original mascot) ---------------- */
const CAT_DEFS = '<defs><radialGradient id="catFur" cx="50%" cy="36%" r="68%"><stop offset="0" stop-color="#d4d8de"/><stop offset=".55" stop-color="#9ea4ad"/><stop offset="1" stop-color="#6b717a"/></radialGradient><radialGradient id="catEye" cx="40%" cy="35%" r="70%"><stop offset="0" stop-color="#e2fbb0"/><stop offset="1" stop-color="#5fa317"/></radialGradient></defs>';
const CAT = `<svg viewBox="0 0 64 64" width="100%" height="100%">
<path d="M9 30 L11 5 L28 17 Z" fill="url(#catFur)"/><path d="M55 30 L53 5 L36 17 Z" fill="url(#catFur)"/>
<path d="M13.5 23 L14.5 11 L23 17.5 Z" fill="#f2b8c6"/><path d="M50.5 23 L49.5 11 L41 17.5 Z" fill="#f2b8c6"/>
<ellipse cx="32" cy="37" rx="24" ry="21" fill="url(#catFur)"/>
<path d="M32 17v7M26.5 18.5l1.8 5.2M37.5 18.5l-1.8 5.2" stroke="#5f666e" stroke-width="2.2" stroke-linecap="round"/>
<g class="cat-eyes"><ellipse cx="22.5" cy="35" rx="5.2" ry="5.8" fill="url(#catEye)"/><ellipse cx="41.5" cy="35" rx="5.2" ry="5.8" fill="url(#catEye)"/>
<ellipse cx="22.5" cy="35.4" rx="1.7" ry="4.6" fill="#111"/><ellipse cx="41.5" cy="35.4" rx="1.7" ry="4.6" fill="#111"/>
<circle cx="24.2" cy="32.6" r="1.3" fill="#fff"/><circle cx="43.2" cy="32.6" r="1.3" fill="#fff"/></g>
<ellipse cx="16.5" cy="44" rx="3.5" ry="2" fill="#f2b8c6" opacity=".35"/><ellipse cx="47.5" cy="44" rx="3.5" ry="2" fill="#f2b8c6" opacity=".35"/>
<path d="M29.4 43.2h5.2L32 46z" fill="#f2a1b5"/>
<path d="M32 46q-1.8 3.2-5 2.2M32 46q1.8 3.2 5 2.2" stroke="#3f444a" stroke-width="1.5" fill="none" stroke-linecap="round"/>
<path d="M19 44.5 4 42.5M19 47.5 5 50M45 44.5l15-2M45 47.5l14 2.5" stroke="#eef0f3" stroke-width="1" stroke-linecap="round" opacity=".75"/>
</svg>`;
document.querySelectorAll('.cat-slot').forEach(el => el.innerHTML = CAT);
$('#favicon').href = 'data:image/svg+xml,' + encodeURIComponent(CAT.replace(/<svg[^>]*>/, m => m.replace('<svg', '<svg xmlns="http://www.w3.org/2000/svg"') + CAT_DEFS));

/* ---------------- state ---------------- */
const LS = 'meow_chats_v1';
let chats = []; try { chats = JSON.parse(localStorage.getItem(LS) || '[]'); } catch {}
let activeId = null;
let mode = 'fast'; try { mode = localStorage.getItem('meow_mode') || 'fast'; } catch {}
let pending = [];
let busy = false, abortCtl = null, currentRid = null;
let liveEl = null, liveMsg = null, rTimer = null;
let lib = [];

function save() {
  const clean = chats.slice().sort((a, b) => b.updated - a.updated).map(c => ({...c,
    messages: c.messages.map(m => Object.fromEntries(Object.entries(m).filter(([k]) => !k.startsWith('_'))))}));
  for (let i = 0; i < 8; i++) {
    try { localStorage.setItem(LS, JSON.stringify(clean)); return; }
    catch { if (clean.length > 1) clean.pop(); else return; }
  }
}
const activeChat = () => chats.find(c => c.id === activeId);
function newChat(title) { const c = {id: uid(), title: title || 'New chat', created: Date.now(), updated: Date.now(), messages: [], fileIds: []}; chats.unshift(c); return c; }

/* ---------------- helpers ---------------- */
const KIND_ICON = {pdf: 'file', image: 'image', docx: 'file', pptx: 'slides', sheet: 'table', text: 'code', zip: 'folder'};
const EXT = {python:'py',py:'py',javascript:'js',js:'js',typescript:'ts',ts:'ts',tsx:'tsx',jsx:'jsx',html:'html',css:'css',json:'json',bash:'sh',sh:'sh',shell:'sh',zsh:'sh',powershell:'ps1',cpp:'cpp','c++':'cpp',c:'c',java:'java',kotlin:'kt',swift:'swift',go:'go',rust:'rs',ruby:'rb',php:'php',sql:'sql',yaml:'yml',yml:'yml',markdown:'md',md:'md',csv:'csv',xml:'xml',dart:'dart',r:'r',text:'txt',plaintext:'txt'};
function guessKind(name) { const e = (name.split('.').pop() || '').toLowerCase();
  if (e === 'pdf') return 'pdf'; if (['png','jpg','jpeg','webp','gif','bmp','tif','tiff','heic','heif'].includes(e)) return 'image';
  if (e === 'docx') return 'docx'; if (e === 'pptx') return 'pptx'; if (['xlsx','xlsm','xls','csv','tsv'].includes(e)) return 'sheet'; if (e === 'zip') return 'zip'; return 'text'; }
const humanSize = b => !b ? '' : b < 1024 ? b + ' B' : b < 1048576 ? (b / 1024).toFixed(0) + ' KB' : (b / 1048576).toFixed(1) + ' MB';
function metaStr(f) { const m = f.meta || {}; const p = [(f.kind || '').toUpperCase()]; if (m.pages) p.push(m.pages + (m.pages === 1 ? ' page' : ' pages')); if (m.dims) p.push(m.dims); if (f.size) p.push(humanSize(f.size)); return p.filter(Boolean).join(' · '); }
function thumb(f) {
  if (f.id && (f.kind === 'image' || f.kind === 'pdf')) return `<img class="th" src="/api/thumb/${f.id}?size=200" data-zoom="/api/thumb/${f.id}?size=1600" alt="" onerror="this.style.opacity=0">`;
  if (f.previewUrl) return `<img class="th" src="${f.previewUrl}" alt="">`;
  return `<div class="th">${ic(KIND_ICON[f.kind] || 'file', 18)}</div>`;
}
function toast(msg, type = '') { const t = document.createElement('div'); t.className = 'toast ' + type; t.textContent = msg; $('#toasts').appendChild(t); setTimeout(() => t.remove(), 3200); }
async function copyText(t) { try { await navigator.clipboard.writeText(t); } catch { const a = document.createElement('textarea'); a.value = t; document.body.appendChild(a); a.select(); document.execCommand('copy'); a.remove(); } }
function download(blob, name) { const a = document.createElement('a'); a.href = URL.createObjectURL(blob); a.download = name; document.body.appendChild(a); a.click(); a.remove(); setTimeout(() => URL.revokeObjectURL(a.href), 4000); }
async function api(path, body) {
  const r = await fetch(path, {method: 'POST', headers: {...HDR, 'Content-Type': 'application/json'}, body: JSON.stringify(body)});
  if (r.status === 401) { location.href = '/login'; throw new Error('Login required'); }
  if (!r.ok) { let t = ''; try { t = (await r.json()).detail; } catch {} throw new Error(t || 'HTTP ' + r.status); }
  return r.json();
}
const scroller = $('#scroller');
const nearBottom = () => scroller.scrollHeight - scroller.scrollTop - scroller.clientHeight < 140;
const scrollBottom = () => { scroller.scrollTop = scroller.scrollHeight; };

/* ---------------- markdown + math ---------------- */
/* graceful fallbacks if a CDN is slow/blocked — the app still works */
const HAS = {md: !!window.marked, pur: !!window.DOMPurify, hl: !!window.hljs, kx: !!window.katex};
if (HAS.pur) DOMPurify.addHook('afterSanitizeAttributes', n => { if (n.tagName === 'A') { n.setAttribute('target', '_blank'); n.setAttribute('rel', 'noopener noreferrer'); } });
if (HAS.hl) hljs.configure({ignoreUnescapedHTML: true});
const sanitize = h => HAS.pur ? DOMPurify.sanitize(h) : h;
const mdParse = s => HAS.md ? marked.parse(s, {gfm: true, breaks: false})
  : esc(s).replace(/```(\w*)\n([\s\S]*?)```/g, (_, l, c) => `<pre><code class="language-${l || 'text'}">${c}</code></pre>`).replace(/\n/g, '<br>');
function mathHTML(src) {
  let display = false, body;
  if (src.startsWith('$$')) { display = true; body = src.slice(2, -2); }
  else if (src.startsWith('\\[')) { display = true; body = src.slice(2, -2); }
  else if (src.startsWith('\\(')) body = src.slice(2, -2);
  else body = src.slice(1, -1);
  if (!HAS.kx) return `<code>${esc(src)}</code>`;
  try { return katex.renderToString(body, {displayMode: display, throwOnError: false}); } catch { return esc(src); }
}
function splitThink(raw) {
  const r = splitThink0(raw);
  const m = r.answer.match(/<<\s*follow-?ups?\s*:([\s\S]*?)>>/i);
  r.followups = m ? m[1].split('|').map(s => s.trim().replace(/^["'“]|["'”]$/g, '')).filter(s => s && s.length < 160).slice(0, 4) : [];
  r.answer = r.answer.replace(/<<\s*f[\s\S]*$/i, '').replace(/\n?<<?\s*$/, '').trimEnd();
  return r;
}
function splitThink0(raw) {
  raw = raw || '';
  const s = raw.indexOf('<think>');
  if (s === -1) return {thought: '', answer: raw, thinking: false, has: false};
  const e = raw.indexOf('</think>');
  if (e === -1) return {thought: raw.slice(s + 7).trim(), answer: raw.slice(0, s).trim(), thinking: true, has: true};
  return {thought: raw.slice(s + 7, e).trim(), answer: (raw.slice(0, s) + raw.slice(e + 8)).trim(), thinking: false, has: true};
}
function renderMD(src) {
  const tools = [];
  src = src.replace(/```meow-tool[^\n]*\n?([\s\S]*?)(?:```|$)/g, (m, body) => {
    tools.push((body.match(/"op"\s*:\s*"([\w-]+)"/) || [])[1] || 'tool');
    return `\n\n⟦T${tools.length - 1}⟧\n\n`;
  });
  const codes = [];
  src = src.replace(/```[\s\S]*?(?:```|$)|`[^`\n]+`/g, m => { codes.push(m); return `⟦C${codes.length - 1}⟧`; });
  const maths = [];
  src = src.replace(/\$\$[\s\S]+?\$\$|\\\[[\s\S]+?\\\]|\\\([\s\S]+?\\\)|\$(?=[^\s$])[^$\n]+?(?<=[^\s\\])\$(?!\d)/g, m => { maths.push(m); return `⟦M${maths.length - 1}⟧`; });
  src = src.replace(/⟦C(\d+)⟧/g, (_, i) => codes[+i]);
  let out = sanitize(mdParse(src));
  out = out.replace(/⟦M(\d+)⟧/g, (_, i) => mathHTML(maths[+i]));
  out = out.replace(/<p>\s*⟦T(\d+)⟧\s*<\/p>|⟦T(\d+)⟧/g, (_, a, b) => `<div class="tool-chip">${ic('wrench', 13)} Using tool <b>${esc(tools[+(a ?? b)])}</b></div>`);
  return out;
}
let mermaidP = null;
function loadMermaid() {
  if (!mermaidP) mermaidP = new Promise((res, rej) => {
    const s = document.createElement('script'); s.src = 'https://cdn.jsdelivr.net/npm/mermaid@10.9.1/dist/mermaid.min.js';
    s.onload = () => { try { mermaid.initialize({startOnLoad: false, theme: 'dark', securityLevel: 'strict', fontFamily: 'Arial, Helvetica, sans-serif'}); res(mermaid); } catch (e) { rej(e); } };
    s.onerror = () => { mermaidP = null; rej(new Error('mermaid failed to load')); }; document.head.appendChild(s);
  });
  return mermaidP;
}
async function renderMermaid(pre, src) {
  try {
    const m = await loadMermaid(); const id = 'mm' + Math.random().toString(36).slice(2);
    const {svg} = await m.render(id, src);
    if (!pre.isConnected) return;
    const box = document.createElement('div'); box.className = 'mermaid-box';
    box.innerHTML = `<div class="mm-bar"><span>${ic('layers', 13)} Diagram</span><span><button data-act="mm-code">${ic('code', 13)}Code</button><button data-act="mm-svg">${ic('download', 13)}SVG</button></span></div><div class="mm-svg">${svg}</div>`;
    box._src = src; box._svg = svg; pre.replaceWith(box); box._pre = pre;
  } catch (e) { pre.querySelector('.code-bar span').textContent = 'mermaid (couldn\'t render)'; }
}
function enhance(root, live) {
  root.querySelectorAll('img').forEach(img => { if (!img.closest('a')) { img.dataset.zoom = img.getAttribute('src'); img.loading = 'lazy'; } });
  root.querySelectorAll('pre').forEach(pre => {
    const code = pre.querySelector('code'); if (!code) return;
    let lang = ((code.className.match(/language-([\w+#.-]+)/) || [])[1] || 'text').toLowerCase();
    const bar = document.createElement('div'); bar.className = 'code-bar';
    if (lang.startsWith('meow-output')) {
      const err = lang.includes('error'); pre.classList.add('py-out'); if (err) pre.classList.add('err');
      bar.innerHTML = `<span>${ic(err ? 'x' : 'terminal', 13)} ${err ? 'Error' : 'Output'}</span><span><button data-act="copy-code">${ic('copy', 13)}Copy</button></span>`;
      pre.prepend(bar); return;
    }
    let label = lang, extra = '';
    if (lang === 'meow-python') { pre.classList.add('py-code'); label = 'python · ran by Meow'; lang = 'python'; }
    if (lang === 'html' || lang === 'svg') extra = `<button data-act="preview-html">${ic('eye', 13)}Preview</button>`;
    bar.innerHTML = `<span>${esc(label)}</span><span>${extra}<button data-act="copy-code">${ic('copy', 13)}Copy</button><button data-act="save-code" data-lang="${esc(lang)}">${ic('download', 13)}Save</button></span>`;
    pre.prepend(bar);
    if (lang === 'mermaid' && !live) { renderMermaid(pre, code.textContent); return; }
    if (!HAS.hl) return;
    if (lang !== 'text' && hljs.getLanguage(lang)) { try { code.className = 'language-' + lang; hljs.highlightElement(code); } catch {} }
    else if (lang === 'text' && code.textContent.length < 20000) { try { const r = hljs.highlightAuto(code.textContent); code.innerHTML = r.value; code.classList.add('hljs'); } catch {} }
  });
}
function previewHTML(src) {
  let ov = $('#htmlPrev');
  if (!ov) {
    ov = document.createElement('div'); ov.id = 'htmlPrev'; ov.className = 'html-prev hidden';
    ov.innerHTML = `<div class="hp-card glass"><div class="hp-bar"><b>${ic('eye', 15)} Live preview</b><span><button class="btn" data-hp="tab">${ic('share', 14)}New tab</button><button class="btn" data-hp="close">${ic('x', 14)}</button></span></div><iframe sandbox="allow-scripts allow-forms allow-modals allow-popups" referrerpolicy="no-referrer"></iframe></div>`;
    document.body.appendChild(ov);
    ov.addEventListener('click', e => {
      const b = e.target.closest('[data-hp]');
      if (e.target === ov || b?.dataset.hp === 'close') ov.classList.add('hidden');
      else if (b?.dataset.hp === 'tab') window.open(URL.createObjectURL(new Blob([ov._src], {type: 'text/html'})), '_blank');
    });
  }
  ov._src = /<html|<body|<!doctype/i.test(src) ? src : `<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"></head><body>${src}</body></html>`;
  ov.querySelector('iframe').srcdoc = ov._src; ov.classList.remove('hidden');
}
function addCaret(root) {
  let n = root.lastElementChild;
  while (n && ['UL', 'OL', 'LI', 'BLOCKQUOTE'].includes(n.tagName) && n.lastElementChild) n = n.lastElementChild;
  if (n && !['PRE', 'TABLE', 'DIV', 'HR'].includes(n.tagName)) n.insertAdjacentHTML('beforeend', '<span class="caret"></span>');
  else root.insertAdjacentHTML('beforeend', '<span class="caret"></span>');
}

/* ---------------- chat list ---------------- */
function renderChatList() {
  const q = $('#chatSearch').value.trim().toLowerCase(), list = $('#chatList');
  const items = chats.slice().sort((a, b) => b.updated - a.updated)
    .filter(c => !q || c.title.toLowerCase().includes(q) || c.messages.some(m => (m.content || '').toLowerCase().includes(q)));
  if (!items.length) { list.innerHTML = `<div class="empty" style="padding:14px">${q ? 'No matches' : 'No chats yet — say meow 🐾'}</div>`; return; }
  const today = new Date(new Date().toDateString()).getTime();
  const grp = t => { const d = (today - new Date(new Date(t).toDateString()).getTime()) / 864e5; return d < 1 ? 'Today' : d < 2 ? 'Yesterday' : d < 7 ? 'This week' : 'Earlier'; };
  let h = '', last = '';
  for (const c of items) {
    const g = grp(c.updated); if (g !== last) { h += `<div class="cl-group">${g}</div>`; last = g; }
    h += `<div class="cl-item ${c.id === activeId ? 'active' : ''}" data-id="${c.id}">${ic('chat', 15)}<span class="t">${esc(c.title)}</span><button class="del" data-del="${c.id}" title="Delete">${ic('trash', 14)}</button></div>`;
  }
  list.innerHTML = h;
}
$('#chatList').addEventListener('click', e => {
  const d = e.target.closest('[data-del]');
  if (d) { e.stopPropagation(); if (busy && d.dataset.del === activeId) return toast('Stop the reply first');
    if (confirm('Delete this chat?')) { chats = chats.filter(c => c.id !== d.dataset.del); if (activeId === d.dataset.del) activeId = null; save(); renderChatList(); renderChat(); } return; }
  const it = e.target.closest('.cl-item'); if (!it) return;
  if (busy) return toast('Wait for the reply to finish (or stop it)');
  activeId = it.dataset.id; renderChatList(); renderChat(); closeSidebarMobile();
});
$('#chatSearch').addEventListener('input', renderChatList);

/* ---------------- rendering messages ---------------- */
function userEl(m) {
  const d = document.createElement('div'); d.className = 'msg user';
  const att = (m.files || []).map(f => f.kind === 'image'
    ? `<img class="att-img" src="/api/thumb/${f.id}?size=600" data-zoom="/api/thumb/${f.id}?size=1600" alt="${esc(f.name)}" onerror="this.style.display='none'">`
    : f.kind === 'pdf' ? `<button class="att-chip" style="border:0;cursor:pointer" data-act="edit-pdf" data-id="${f.id}" title="Open in PDF editor">${ic('file', 14)}${esc(f.name)} ${ic('pen', 12)}</button>`
    : `<span class="att-chip">${ic(KIND_ICON[f.kind] || 'file', 14)}${esc(f.name)}</span>`).join('');
  d._msg = m;
  d.innerHTML = `<div class="u-wrap"><div class="bubble">${att ? `<div class="att">${att}</div>` : ''}${esc(m.content)}</div>
    <div class="u-acts"><button data-act="copy-user" title="Copy">${ic('copy', 14)}</button><button data-act="edit-user" title="Edit and resend">${ic('pen', 14)}</button></div></div>`;
  return d;
}
function assistantEl(m) {
  const d = document.createElement('div'); d.className = 'msg assistant'; d._msg = m;
  d.innerHTML = `<div class="avatar">${CAT}</div><div class="a-body">
    <div class="a-name">Meow <span class="mode-tag ${m.mode === 'think' ? 'm-think' : 'm-fast'}">${m.mode === 'think' ? '🧠 Think' : '⚡ Fast'}</span><span class="status-line"></span></div>
    <div class="web-panel hidden"></div>
    <details class="think hidden"><summary><span class="orb"></span><span class="tl">Thinking…</span><span class="chev">${ic('chev', 14)}</span></summary><div class="think-body"></div></details>
    <div class="md answer"></div><div class="tools"></div><div class="fups"></div><div class="actions"></div></div>`;
  const det = d.querySelector('details.think');
  det.addEventListener('toggle', () => { if (det.open !== det._want) det._user = true; });
  updateAssistant(d, m, false); renderTools(d, m); renderActions(d, m);
  return d;
}
function updateAssistant(d, m, live) {
  const {thought, answer, thinking, has, followups} = splitThink(m.content);
  const fu = d.querySelector('.fups');
  if (fu) fu.innerHTML = !live && followups.length ? followups.map(q => `<button class="fup" data-act="followup" data-q="${esc(q)}">${ic('chev', 13)}<span>${esc(q)}</span></button>`).join('') : '';
  const det = d.querySelector('details.think');
  if (has) {
    det.classList.remove('hidden');
    const body = det.querySelector('.think-body');
    const stick = body.scrollHeight - body.scrollTop - body.clientHeight < 40;
    body.textContent = thought || '…';
    if (stick) body.scrollTop = body.scrollHeight;
    const secs = thinking && live && m._t0 ? Math.round((Date.now() - m._t0) / 1000) : m.thinkSecs;
    det.querySelector('.tl').textContent = thinking && live ? `Thinking deeply… ${secs || 0}s` : (secs != null ? `Thought for ${secs}s` : 'Thought process');
    det.querySelector('.orb').classList.toggle('done', !(thinking && live));
    if (!det._user) { const want = !!(thinking && live); if (det.open !== want) { det._want = want; det.open = want; } }
  } else det.classList.add('hidden');
  renderWeb(d, m, live);
  const ans = d.querySelector('.answer');
  if (answer) { ans.innerHTML = renderMD(answer); linkCites(ans, m.sources); enhance(ans, live); if (live) addCaret(ans); }
  else if (live && !thinking) ans.innerHTML = '<div class="dots"><span></span><span></span><span></span></div>';
  else ans.innerHTML = (!live && m.stopped && !answer) ? '<em style="color:var(--dim)">Stopped.</em>' : '';
  d.querySelector('.status-line').textContent = live && m._status ? '· ' + m._status : '';
  d.querySelector('.avatar').classList.toggle('busy', !!live);
}
const fav = dm => `https://www.google.com/s2/favicons?domain=${encodeURIComponent(dm)}&sz=32`;
const SRC_KIND = {youtube: ['YouTube', 'yt'], wiki: ['Wikipedia', 'wk'], paper: ['Paper', 'pp'], code: ['Code', 'cd'], hn: ['Hacker News', 'hn'], link: ['Your link', 'lk'], news: ['News', 'nw']};
function renderWeb(d, m, live) {
  const p = d.querySelector('.web-panel'), s = m.search, src = m.sources || [], imgs = m.images || [];
  if (!s || (s.stage === 'skipped' && !src.length)) { p.classList.add('hidden'); return; }
  const busyW = live && !['done', 'error', 'skipped'].includes(s.stage);
  const key = [s.stage, s.round, src.length, busyW, (s.queries || []).length, imgs.length].join('|'); if (p._key === key) return; p._key = key;
  p.classList.remove('hidden');
  const nq = (s.queries || []).length, nr = src.filter(x => x.kind === 'reddit').length, nl = (s.links || []).length;
  const xs = (s.sources || []).filter(x => x !== 'news').map(x => ({youtube: 'YouTube', wiki: 'Wikipedia', papers: 'papers', code: 'GitHub + Stack Overflow', hn: 'Hacker News'}[x])).filter(Boolean);
  const where = ['the web'].concat(s.reddit ? ['Reddit'] : [], xs).join(', ').replace(/, ([^,]*)$/, ' and $1');
  const label = s.stage === 'links' ? `Reading ${nl} link${nl > 1 ? 's' : ''} you shared…`
    : s.stage === 'planning' ? 'Figuring out what to search…'
    : s.stage === 'searching' ? `Searching ${where}…`
    : s.stage === 'reading' ? `Reading the best of ${s.found || 0} results…`
    : s.stage === 'checking' ? 'Checking if this really answers your question…'
    : s.stage === 'deeper' ? `Round ${s.round || 2}: ${esc(s.missing || 'digging deeper')}${(s.subs || []).length ? ' · searching ' + s.subs.map(x => 'r/' + esc(x)).join(', ') : ''}`
    : s.stage === 'error' ? `Search failed — answering from memory (${esc(s.error || '')})`
    : (s.rounds > 1 ? `${s.rounds} search rounds · ` : '') + (nq ? `Searched ${nq} quer${nq === 1 ? 'y' : 'ies'} · ` : '') + `used ${src.length} source${src.length === 1 ? '' : 's'}${nr ? ` (${nr} from Reddit)` : ''}${nl ? ` · read ${nl} link${nl > 1 ? 's' : ''}` : ''}`;
  const q = (s.queries || []).map(x => `<span>${ic('search', 11)}${esc(x)}</span>`).join('') + (s.reddit ? `<span class="r">${ic('reddit', 11)}Reddit</span>` : '')
    + xs.map(x => `<span class="x">${ic('layers', 11)}${esc(x)}</span>`).join('');
  const cards = src.map(x => { const k = SRC_KIND[x.kind];
    const sub = x.kind === 'reddit' ? `r/${esc(x.sub || '?')}${x.score != null ? ` · ▲${x.score}` : ''}` : esc(x.domain);
    return `<a class="src ${x.kind === 'reddit' ? 'reddit' : ''} ${k ? 'k-' + k[1] : ''}" href="${esc(x.url)}" target="_blank" rel="noopener" title="${esc(x.snippet || x.title)}">
    ${k ? `<span class="kb">${k[0]}</span>` : ''}<span class="st">${esc(x.title)}</span>
    <span class="sd"><img src="${fav(x.domain)}" alt="" loading="lazy" onerror="this.style.display='none'"><span class="dm">${sub}${x.date ? ' · ' + esc(String(x.date).slice(0, 10)) : ''}</span><span class="n">${x.n}</span></span></a>`; }).join('');
  const strip = imgs.length ? `<div class="img-row">${imgs.map(i => `<a href="${esc(i.link || i.img)}" target="_blank" rel="noopener" title="${esc(i.title)} — ${esc(i.site || '')}"><img src="${esc(i.thumb || i.img)}" alt="${esc(i.title)}" loading="lazy" referrerpolicy="no-referrer" onerror="this.parentElement.remove()"></a>`).join('')}</div>` : '';
  p.innerHTML = `<div class="wp-head">${busyW ? '<span class="orb"></span>' : ic('globe', 15)}<span>${label}</span></div>${q ? `<div class="wp-q">${q}</div>` : ''}${strip}${cards ? `<div class="src-row">${cards}</div>` : ''}`;
}
function linkCites(root, src) {
  if (!src || !src.length) return;
  const by = {}; src.forEach(s => by[s.n] = s);
  const w = document.createTreeWalker(root, NodeFilter.SHOW_TEXT, {acceptNode: n => n.parentElement.closest('pre,code,a,.katex') ? NodeFilter.FILTER_REJECT : (/\[\d{1,2}\]/.test(n.nodeValue) ? NodeFilter.FILTER_ACCEPT : NodeFilter.FILTER_SKIP)});
  const nodes = []; while (w.nextNode()) nodes.push(w.currentNode);
  nodes.forEach(n => {
    const h = esc(n.nodeValue).replace(/\[(\d{1,2})\]/g, (x, k) => { const s = by[+k]; return s ? `<a class="cite${s.kind === 'reddit' ? ' r' : ''}" href="${esc(s.url)}" target="_blank" rel="noopener" title="${esc(s.title)} — ${esc(s.kind === 'reddit' ? 'r/' + (s.sub || '') : s.domain)}">${k}</a>` : x; });
    const sp = document.createElement('span'); sp.innerHTML = h; n.replaceWith(...sp.childNodes);
  });
}
function fileCard(f, msg) {
  return `<div class="fcard">${thumb(f)}<div class="info"><div class="nm">${esc(f.name)}</div><div class="mt">${esc(msg || metaStr(f))}</div></div>
  <div class="acts"><a href="/api/files/${f.id}" target="_blank" rel="noopener" title="Open">${ic('eye', 16)}</a><a href="/api/files/${f.id}?dl=1" title="Download">${ic('download', 16)}</a>${f.kind === 'pdf' || f.kind === 'image' ? `<button data-act="edit-pdf" data-id="${f.id}" title="Open in editor">${ic('pen', 16)}</button>` : ''}<button data-act="attach" data-id="${f.id}" title="Attach to next message">${ic('clip', 16)}</button></div></div>`;
}
function renderTools(d, m) {
  d.querySelector('.tools').innerHTML = (m.tools || []).map(t => {
    if (t.status === 'running') return `<div class="fcard"><div class="th"><span class="spin"></span></div><div class="info"><div class="nm">Running ${esc(t.op)}…</div><div class="mt">Working on your file</div></div></div>`;
    if (t.status === 'error') return `<div class="fcard err"><div class="th" style="color:var(--bad)">${ic('x', 18)}</div><div class="info"><div class="nm">${esc(t.op)} failed</div><div class="mt">${esc(t.error || 'Unknown error')}</div></div></div>`;
    return fileCard(t.file, t.message);
  }).join('');
}
function renderActions(d, m) {
  const chat = activeChat(); const isLast = chat && chat.messages[chat.messages.length - 1] === m;
  const st = m.stats && m.stats.tokens ? `<span class="stats">${m.stats.tps} tok/s · ${m.stats.tokens} tokens · ${m.stats.secs}s</span>` : '';
  d.querySelector('.actions').innerHTML = `<button data-act="copy-msg" title="Copy">${ic('copy', 15)}</button>${isLast ? `<button data-act="regen" title="Regenerate">${ic('refresh', 15)}</button>` : ''}<button data-act="speak" title="Read aloud">${ic('volume', 15)}</button>${st}`;
}
function renderChat() {
  const chat = activeChat(), box = $('#messages');
  box.innerHTML = '';
  $('#welcome').classList.toggle('hidden', !!(chat && chat.messages.length));
  if (!chat) return;
  chat.messages.forEach(m => box.appendChild(m.role === 'user' ? userEl(m) : assistantEl(m)));
  scrollBottom();
}

$('#messages').addEventListener('click', e => {
  const z = e.target.closest('[data-zoom]'); if (z) return openLightbox(z.dataset.zoom);
  const b = e.target.closest('[data-act]'); if (!b) return;
  const act = b.dataset.act, m = b.closest('.msg')?._msg;
  if (act === 'copy-code') { copyText(b.closest('pre').querySelector('code').innerText); b.innerHTML = ic('check', 13) + 'Copied'; setTimeout(() => b.innerHTML = ic('copy', 13) + 'Copy', 1500); }
  else if (act === 'save-code') { const l = b.dataset.lang; download(new Blob([b.closest('pre').querySelector('code').innerText], {type: 'text/plain'}), 'meow-snippet.' + (EXT[l] || (/^[a-z0-9]{1,6}$/.test(l) ? l : 'txt'))); }
  else if (act === 'copy-msg' && m) { copyText(splitThink(m.content).answer.replace(/```meow-tool[\s\S]*?```/g, '').trim()); toast('Copied 🐾'); }
  else if (act === 'regen') regenerate();
  else if (act === 'speak' && m) speak(splitThink(m.content).answer);
  else if (act === 'preview-html') previewHTML(b.closest('pre').querySelector('code').innerText);
  else if (act === 'mm-code') { const box = b.closest('.mermaid-box'); box.replaceWith(box._pre); }
  else if (act === 'mm-svg') { const box = b.closest('.mermaid-box'); download(new Blob([box._svg], {type: 'image/svg+xml'}), 'meow-diagram.svg'); }
  else if (act === 'copy-user' && m) { copyText(m.content); toast('Copied 🐾'); }
  else if (act === 'edit-user' && m) startEditMsg(m);
  else if (act === 'followup') { if (busy) return; input.value = b.dataset.q; autosize(); sendMessage(); }
  else if (act === 'attach') attachExisting(b.dataset.id);
  else if (act === 'edit-pdf') openEditor(b.dataset.id);
});

/* ---------------- uploads ---------------- */
async function uploadFiles(fileList, toComposer = true) {
  const arr = [...fileList]; if (!arr.length) return [];
  const items = arr.map(f => ({tmp: uid(), name: f.name, kind: guessKind(f.name), size: f.size, status: 'uploading', meta: {},
    previewUrl: f.type.startsWith('image/') && !/hei[cf]$/i.test(f.name) ? URL.createObjectURL(f) : null}));
  if (toComposer) { pending.push(...items); renderPending(); }
  const fd = new FormData(); arr.forEach(f => fd.append('files', f, f.name));
  try {
    const r = await fetch('/api/upload', {method: 'POST', body: fd, headers: HDR});
    if (!r.ok) throw new Error('HTTP ' + r.status);
    (await r.json()).files.forEach((rec, i) => Object.assign(items[i], rec, {status: 'ready'}));
  } catch (e) { items.forEach(it => it.status = 'error'); toast('Upload failed: ' + e.message, 'err'); }
  if (toComposer) renderPending();
  return items;
}
function renderPending() {
  const box = $('#pending'); box.classList.toggle('hidden', !pending.length);
  box.innerHTML = pending.map((p, i) => `<div class="chip ${p.status === 'error' ? 'err' : ''}">${p.status === 'uploading' ? '<div class="th"><span class="spin"></span></div>' : thumb(p)}
    <div style="min-width:0"><div class="nm">${esc(p.name)}</div><div class="mt">${p.status === 'uploading' ? 'Reading…' : p.status === 'error' ? 'Upload failed' : esc(metaStr(p))}</div></div>
    <button class="x" data-rm="${i}" title="Remove">${ic('x', 12)}</button></div>`).join('');
}
$('#pending').addEventListener('click', e => {
  const r = e.target.closest('[data-rm]'); if (r) { pending.splice(+r.dataset.rm, 1); renderPending(); return; }
  const z = e.target.closest('[data-zoom]'); if (z) openLightbox(z.dataset.zoom);
});
$('#attachBtn').onclick = () => $('#fileInput').click();
$('#fileInput').onchange = e => { uploadFiles(e.target.files); e.target.value = ''; };
async function loadLib() { try { const r = await fetch('/api/files', {headers: HDR}); lib = (await r.json()).files || []; } catch { lib = []; } }
async function attachExisting(id) {
  let f = lib.find(x => x.id === id); if (!f) { await loadLib(); f = lib.find(x => x.id === id); }
  if (!f) return toast('File not found on server', 'err');
  if (!pending.some(p => p.id === id)) pending.push({...f, status: 'ready'});
  renderPending(); closeTools(); toast('Attached ' + f.name); $('#input').focus();
}

/* ---------------- generation ---------------- */
function historyFor(chat, upto) {
  return chat.messages.slice(0, upto).map(m => {
    let c = m.content || '';
    if (m.role === 'assistant') {
      c = splitThink(c).answer;
      if (m.tools && m.tools.length) c += '\n\n[Tool results: ' + m.tools.map(t => t.status === 'ok' ? `${t.op} → ${t.file.name} (id=${t.file.id})` : `${t.op} failed: ${t.error}`).join('; ') + ']';
    } else if (m.files && m.files.length) c = `[Attached: ${m.files.map(f => `${f.name} (id=${f.id})`).join(', ')}]\n${c}`;
    return {role: m.role, content: c};
  }).filter(x => x.content.trim());
}
function setBusy(b) {
  busy = b; const s = $('#sendBtn');
  s.classList.toggle('stop', b); s.innerHTML = ic(b ? 'stop' : 'send', 18); s.title = b ? 'Stop (Esc)' : 'Send';
  warpTarget = b ? (mode === 'think' ? 0.5 : 0.16) : 0;
}
function schedule() {
  if (rTimer) return;
  rTimer = setTimeout(() => { rTimer = null; if (!liveEl) return; const stick = nearBottom(); updateAssistant(liveEl, liveMsg, true); if (stick) scrollBottom(); }, 45);
}
let editing = null;
function startEditMsg(m) {
  if (busy) return toast('Stop the current reply first');
  const chat = activeChat(); const idx = chat ? chat.messages.indexOf(m) : -1; if (idx < 0) return;
  editing = {idx};
  input.value = m.content; autosize(); input.focus();
  pending = (m.files || []).map(f => ({...f, status: 'ready'})); renderPending();
  $('#editBar').classList.remove('hidden');
}
function cancelEdit() { editing = null; $('#editBar').classList.add('hidden'); input.value = ''; pending = []; renderPending(); autosize(); }
$('#editCancel').onclick = cancelEdit;
async function sendMessage() {
  if (busy) return stopGen();
  if (pending.some(p => p.status === 'uploading')) return toast('Hold on — files are still uploading 🐾');
  const input = $('#input'), text = input.value.trim(), ready = pending.filter(p => p.status === 'ready');
  if (!text && !ready.length) return;
  let chat = activeChat(); if (!chat) { chat = newChat(); activeId = chat.id; }
  if (editing) { chat.messages = chat.messages.slice(0, editing.idx); editing = null; $('#editBar').classList.add('hidden'); }
  ready.forEach(f => { if (!chat.fileIds.includes(f.id)) chat.fileIds.push(f.id); });
  chat.messages.push({role: 'user', content: text, files: ready.map(({id, name, kind, size, meta}) => ({id, name, kind, size, meta}))});
  if (chat.messages.length === 1) chat.title = (text || ready.map(f => f.name).join(', ')).slice(0, 48);
  pending = []; renderPending(); input.value = ''; autosize();
  await generate(chat);
}
async function generate(chat) {
  const m = {role: 'assistant', content: '', mode, web: webMode, tools: []};
  chat.messages.push(m); chat.updated = Date.now();
  renderChatList(); renderChat();
  const el = $('#messages').lastElementChild;
  liveEl = el; liveMsg = m; setBusy(true); updateAssistant(el, m, true); scrollBottom();
  const rid = uid(); currentRid = rid; abortCtl = new AbortController();
  const idx = chat.messages.length - 2;
  const body = {prompt: chat.messages[idx].content || '', mode, web: webMode, history: historyFor(chat, idx), file_ids: chat.fileIds, rid, chat_id: chat.id};
  try {
    const res = await fetch('/api/chat', {method: 'POST', headers: {...HDR, 'Content-Type': 'application/json'}, body: JSON.stringify(body), signal: abortCtl.signal});
    if (!res.ok) throw new Error('server returned ' + res.status);
    const reader = res.body.getReader(), dec = new TextDecoder(); let buf = '';
    while (true) {
      const {done, value} = await reader.read(); if (done) break;
      buf += dec.decode(value, {stream: true});
      let i;
      while ((i = buf.indexOf('\n\n')) !== -1) {
        const evt = buf.slice(0, i); buf = buf.slice(i + 2);
        for (const line of evt.split('\n')) {
          if (!line.startsWith('data:')) continue;
          const data = line.slice(5).trim(); if (!data || data === '[DONE]') continue;
          let j; try { j = JSON.parse(data); } catch { continue; }
          if (j.t != null) {
            m.content += j.t;
            if (!m._t0 && m.content.includes('<think>')) m._t0 = Date.now();
            if (m._t0 && m.thinkSecs == null && m.content.includes('</think>')) m.thinkSecs = Math.round((Date.now() - m._t0) / 1000);
            schedule();
          } else if (j.status != null) { m._status = j.status; schedule(); }
          else if (j.search) { m.search = Object.assign(m.search || {}, j.search); schedule(); }
          else if (j.sources) { m.sources = j.sources; schedule(); }
          else if (j.images) { m.images = j.images; schedule(); }
          else if (j.rollback) { m.content = m.content.slice(0, Math.max(0, m.content.length - j.rollback)); schedule(); }
          else if (j.files) { j.files.forEach(f => { if (!chat.fileIds.includes(f.id)) chat.fileIds.push(f.id); if (!/^figure_/.test(f.name)) m.tools.push({op: 'python', status: 'ok', file: f, message: 'Created with Python'}); }); renderTools(el, m); }
          else if (j.stats) m.stats = j.stats;
        }
      }
    }
  } catch (err) {
    if (err.name === 'AbortError') m.stopped = true;
    else m.content += (m.content ? '\n\n' : '') + `⚠️ **Couldn't reach Meow's server** (${err.message}). Check that the Kaggle notebook cell is still running.`;
  } finally {
    if (m._t0 && m.thinkSecs == null) m.thinkSecs = Math.round((Date.now() - m._t0) / 1000);
    clearTimeout(rTimer); rTimer = null; liveEl = null; m._status = '';
    updateAssistant(el, m, false);
    if (!m.stopped) await runToolCalls(chat, m, el);
    setBusy(false); renderActions(el, m);
    if (voiceOn) voiceAnswer(m);
    chat.updated = Date.now(); save(); renderChatList();
  }
}
function extractToolCalls(text) { const out = []; const re = /```meow-tool[^\n]*\n([\s\S]*?)```/g; let x; while ((x = re.exec(text))) out.push(x[1].trim()); return out; }
async function runToolCalls(chat, m, el) {
  const calls = extractToolCalls(splitThink(m.content).answer); if (!calls.length) return;
  let last = '';
  for (const raw of calls) {
    const item = {op: (raw.match(/"op"\s*:\s*"([\w-]+)"/) || [])[1] || 'tool', status: 'running'};
    m.tools.push(item); renderTools(el, m); if (nearBottom()) scrollBottom();
    try {
      const r = await api('/api/tool', {raw, last});
      if (r.ok) { Object.assign(item, {status: 'ok', file: r.file, message: r.message}); last = r.file.id; if (!chat.fileIds.includes(r.file.id)) chat.fileIds.push(r.file.id); }
      else Object.assign(item, {status: 'error', error: r.error});
    } catch (e) { Object.assign(item, {status: 'error', error: e.message}); }
    renderTools(el, m);
  }
  if (nearBottom()) scrollBottom();
}
function stopGen() {
  if (!busy) return;
  if (currentRid) api('/api/stop', {rid: currentRid}).catch(() => {});
  abortCtl?.abort();
}
function regenerate() {
  if (busy) return; const chat = activeChat(); if (!chat) return;
  if (chat.messages[chat.messages.length - 1]?.role === 'assistant') chat.messages.pop();
  if (chat.messages.length) generate(chat);
}
function speak(text) {
  if (!('speechSynthesis' in window)) return toast('Speech not supported here', 'err');
  if (speechSynthesis.speaking) { speechSynthesis.cancel(); return; }
  const plain = text.replace(/```[\s\S]*?```/g, ' (code block) ').replace(/[#*_`>|~]/g, '').replace(/\[(.*?)\]\(.*?\)/g, '$1').replace(/\$+/g, '');
  speechSynthesis.speak(new SpeechSynthesisUtterance(plain));
}

/* ---------------- composer ---------------- */
const input = $('#input');
function autosize() { input.style.height = 'auto'; input.style.height = Math.min(input.scrollHeight, 220) + 'px'; }
input.addEventListener('input', autosize);
input.addEventListener('keydown', e => { if (e.key === 'Enter' && !e.shiftKey && !e.isComposing && !isTouch) { e.preventDefault(); sendMessage(); } });
input.addEventListener('paste', e => { const fs = [...(e.clipboardData?.files || [])]; if (fs.length) { e.preventDefault(); uploadFiles(fs); } });
$('#sendBtn').onclick = sendMessage;
function setMode(m) {
  mode = m; try { localStorage.setItem('meow_mode', m); } catch {}
  const t = $('#modeToggle'); t.classList.toggle('think', m === 'think');
  t.querySelectorAll('button').forEach(b => b.classList.toggle('on', b.dataset.mode === m));
  placePill();
  input.placeholder = m === 'think' ? 'Ask something hard — Meow will think it through…' : 'Message Meow…';
}
function placePill() { const b = $(`#modeToggle button[data-mode="${mode}"]`); if (!b) return; const p = $('#modePill'); p.style.left = b.offsetLeft + 'px'; p.style.width = b.offsetWidth + 'px'; }
$('#modeToggle').addEventListener('click', e => { const b = e.target.closest('button'); if (b) setMode(b.dataset.mode); });

/* voice */
const SR = window.SpeechRecognition || window.webkitSpeechRecognition; let rec = null, listening = false, baseText = '';
if (SR) {
  rec = new SR(); rec.interimResults = true; rec.lang = navigator.language || 'en-IN';
  rec.onstart = () => { listening = true; baseText = input.value; $('#micBtn').classList.add('mic-on'); };
  rec.onresult = e => { let t = ''; for (let i = 0; i < e.results.length; i++) t += e.results[i][0].transcript; input.value = (baseText ? baseText + ' ' : '') + t; autosize(); };
  rec.onend = () => { listening = false; $('#micBtn').classList.remove('mic-on'); };
  rec.onerror = e => { if (e.error !== 'no-speech' && e.error !== 'aborted') toast('Mic: ' + e.error, 'err'); };
}
$('#camBtn').onclick = () => $('#camInput').click();
$('#camInput').onchange = e => { const fs = [...e.target.files]; e.target.value = ''; if (fs.length) { uploadFiles(fs); if (!input.value.trim()) { input.value = 'What is in this photo? Read any text and explain it.'; autosize(); } } };

/* ---------------- voice mode (hands-free, Astra-style) ---------------- */
let voiceOn = false, vRec = null, vState = 'idle', vPaused = false, vHeard = '';
function vSet(state, text) {
  vState = state; const ov = $('#voiceOv');
  ov.classList.remove('listening', 'thinking', 'speaking', 'paused'); ov.classList.add(vPaused ? 'paused' : state);
  $('#voState').textContent = vPaused ? 'Paused — tap the mic to resume' : {listening: 'Listening…', thinking: 'Thinking…', speaking: 'Speaking…'}[state] || '';
  if (text != null) $('#voText').textContent = text;
}
function voiceStart() {
  if (!SR) return toast('Voice mode needs Chrome or Edge (speech recognition)', 'err');
  if (busy) return toast('Wait for the current reply to finish');
  voiceOn = true; vPaused = false; $('#voiceOv').classList.remove('hidden'); vListen();
}
function voiceStop() {
  voiceOn = false; try { vRec && vRec.abort(); } catch {} try { speechSynthesis.cancel(); } catch {}
  $('#voiceOv').classList.add('hidden'); if (busy) stopGen();
}
function vListen() {
  if (!voiceOn || vPaused) return;
  vHeard = ''; vSet('listening', '');
  vRec = new SR(); vRec.lang = navigator.language || 'en-IN'; vRec.interimResults = true; vRec.continuous = false;
  vRec.onresult = e => { let t = ''; for (let i = 0; i < e.results.length; i++) t += e.results[i][0].transcript; vHeard = t; $('#voText').textContent = t; };
  vRec.onerror = e => { if (e.error === 'not-allowed') { toast('Microphone permission denied', 'err'); voiceStop(); } };
  vRec.onend = () => {
    if (!voiceOn || vPaused || vState !== 'listening') return;
    const t = vHeard.trim();
    if (!t) return setTimeout(vListen, 250);
    vSet('thinking', t); input.value = t; autosize(); sendMessage();
  };
  try { vRec.start(); } catch { setTimeout(vListen, 400); }
}
function speakable(md) {
  return md.replace(/```[\s\S]*?```/g, ' (I put the code on screen) ').replace(/!\[[^\]]*\]\([^)]*\)/g, '')
    .replace(/\[(\d{1,2})\]/g, '').replace(/\[([^\]]+)\]\([^)]*\)/g, '$1').replace(/<<[\s\S]*$/, '')
    .replace(/^\s*\|.*\|\s*$/gm, '').replace(/[#*_`>~|]/g, '').replace(/\$\$?([^$]+)\$\$?/g, '$1').replace(/\n{2,}/g, '. ').replace(/\s+/g, ' ').trim();
}
function voiceAnswer(m) {
  if (!voiceOn) return;
  const text = speakable(splitThink(m.content).answer).slice(0, 2400);
  if (!text) return vListen();
  vSet('speaking', text);
  const parts = text.match(/[^.!?।]+[.!?।]+["')]*|\S[^.!?।]*$/g) || [text];
  const hindi = /[\u0900-\u097F]/.test(text);
  const voices = speechSynthesis.getVoices();
  const voice = voices.find(v => hindi ? /^hi/i.test(v.lang) : /en-IN/i.test(v.lang)) || voices.find(v => /^en/i.test(v.lang) && /Google|Natural|Neural/i.test(v.name)) || null;
  speechSynthesis.cancel();
  parts.forEach((p, i) => {
    const u = new SpeechSynthesisUtterance(p.trim()); if (voice) u.voice = voice; u.lang = voice?.lang || (hindi ? 'hi-IN' : 'en-IN'); u.rate = 1.04;
    if (i === parts.length - 1) u.onend = () => { if (voiceOn) setTimeout(vListen, 200); };
    speechSynthesis.speak(u);
  });
}
$('#voiceBtn').onclick = voiceStart;
$('#voEnd').onclick = voiceStop;
$('#voMute').onclick = () => {
  vPaused = !vPaused; $('#voMute').classList.toggle('off', vPaused);
  if (vPaused) { try { vRec && vRec.abort(); } catch {} try { speechSynthesis.cancel(); } catch {} vSet(vState); }
  else vListen();
};
try { speechSynthesis.getVoices(); speechSynthesis.onvoiceschanged = () => {}; } catch {}

$('#micBtn').onclick = () => { if (!rec) return toast('Voice input needs Chrome or Edge', 'err'); listening ? rec.stop() : rec.start(); };

/* drag & drop */
let dragDepth = 0; const main = $('#main');
window.addEventListener('dragover', e => e.preventDefault()); window.addEventListener('drop', e => e.preventDefault());
main.addEventListener('dragenter', e => { if (![...(e.dataTransfer?.types || [])].includes('Files')) return; dragDepth++; $('#dropOverlay').classList.remove('hidden'); });
main.addEventListener('dragleave', () => { if (--dragDepth <= 0) { dragDepth = 0; $('#dropOverlay').classList.add('hidden'); } });
main.addEventListener('drop', e => { e.preventDefault(); dragDepth = 0; $('#dropOverlay').classList.add('hidden'); if (e.dataTransfer.files.length) uploadFiles(e.dataTransfer.files); });

/* welcome suggestions */
document.querySelectorAll('.suggest').forEach(s => s.addEventListener('click', () => {
  if (s.dataset.mode) setMode(s.dataset.mode);
  if (s.dataset.web) setWeb(s.dataset.web);
  input.value = s.dataset.prompt; autosize(); input.focus();
  if (s.dataset.attach) $('#fileInput').click();
}));
$('#greet').textContent = USER_NAME ? `Hey ${USER_NAME} 🐾 What are we working on today?` : 'What are we working on today?';

/* ---------------- sidebar, export, lightbox, keys ---------------- */
const appEl = $('#app');
const isMobile = () => innerWidth <= 820;
$('#menuBtn').onclick = () => isMobile() ? appEl.classList.toggle('sb-open') : appEl.classList.toggle('sb-collapsed');
function closeSidebarMobile() { appEl.classList.remove('sb-open'); }
$('#sbClose').onclick = closeSidebarMobile; $('#scrim').onclick = closeSidebarMobile;
function newChatUI() { if (busy) return toast('Stop the current reply first'); activeId = null; pending = []; renderPending(); renderChatList(); renderChat(); closeSidebarMobile(); input.focus(); }
$('#newChatBtn').onclick = newChatUI;
$('#exportBtn').onclick = () => {
  const c = activeChat(); if (!c || !c.messages.length) return toast('Nothing to export yet');
  const md = `# ${c.title}\n\n` + c.messages.map(m => `### ${m.role === 'user' ? 'You' : 'Meow 🐈'}\n\n${m.role === 'user' ? m.content : splitThink(m.content).answer}`).join('\n\n---\n\n');
  download(new Blob([md], {type: 'text/markdown'}), (c.title.replace(/[^\w-]+/g, '_').slice(0, 40) || 'chat') + '.md');
};
let webMode = 'auto'; try { webMode = localStorage.getItem('meow_web') || 'auto'; } catch {}
const WEB_LBL = {auto: 'Auto', web: 'Web', reddit: 'Reddit', deep: 'Deep', off: 'Off'};
function setWeb(w) {
  if (!WEB_LBL[w]) w = 'auto';
  webMode = w; try { localStorage.setItem('meow_web', w); } catch {}
  const b = $('#webBtn'); $('#webLbl').textContent = WEB_LBL[w];
  b.classList.toggle('on', w === 'web' || w === 'deep'); b.classList.toggle('rd', w === 'reddit'); b.classList.toggle('off', w === 'off');
  b.title = {auto: 'Web search: Auto — searches when needed', web: 'Web search: always on', reddit: 'Searching Reddit', deep: 'Deep research', off: 'Web search off'}[w];
  document.querySelectorAll('#webMenu button').forEach(x => x.classList.toggle('sel', x.dataset.web === w));
}
$('#webBtn').addEventListener('click', e => { e.stopPropagation(); $('#webMenu').classList.toggle('hidden'); });
$('#webMenu').addEventListener('click', e => { const b = e.target.closest('button[data-web]'); if (!b) return; setWeb(b.dataset.web); $('#webMenu').classList.add('hidden'); input.focus(); });
document.addEventListener('click', e => { if (!e.target.closest('#webPick')) $('#webMenu').classList.add('hidden'); });

function openLightbox(src) { const lb = $('#lightbox'); lb.querySelector('img').src = src; lb.classList.remove('hidden'); }
$('#lightbox').onclick = () => $('#lightbox').classList.add('hidden');
document.addEventListener('keydown', e => {
  if (!$('#pdfEditor').classList.contains('hidden')) return;
  if (e.key === 'Escape') {
    if (editing) cancelEdit();
    else if (!$('#voiceOv').classList.contains('hidden')) voiceStop();
    else if (!$('#webMenu').classList.contains('hidden')) $('#webMenu').classList.add('hidden');
    else if (!$('#lightbox').classList.contains('hidden')) $('#lightbox').classList.add('hidden');
    else if (!$('#toolsModal').classList.contains('hidden')) closeTools();
    else if (busy) stopGen();
  }
  if ((e.ctrlKey || e.metaKey) && e.key.toLowerCase() === 'k') { e.preventDefault(); newChatUI(); }
});

/* ---------------- tools modal ---------------- */
let toolSpecs = null, curTool = null, toolSel = [];
async function openTools(op) {
  $('#toolsModal').classList.remove('hidden'); closeSidebarMobile();
  if (!toolSpecs) {
    try { toolSpecs = (await (await fetch('/api/tools', {headers: HDR})).json()).tools;
      toolSpecs.unshift({op: '__editor', group: 'PDF', label: '✏️ Visual editor (PDF & images)', desc: 'Click any text to edit it in its original font — even text inside photos, screenshots and scanned pages. Move, resize, replace or adjust images; draw, highlight, add shapes, erase, sign and manage pages.', accepts: ['pdf', 'image'], multi: false, params: []}); }
    catch { toast('Could not load tools — is the server running?', 'err'); return; }
    const groups = [...new Set(toolSpecs.map(t => t.group))];
    $('#toolList').innerHTML = groups.map(g => `<div class="tl-group">${esc(g)}</div>` + toolSpecs.filter(t => t.group === g).map(t => `<div class="tl-item" data-op="${t.op}">${esc(t.label)}</div>`).join('')).join('');
  }
  await loadLib();
  selectTool(op || curTool?.op || toolSpecs[0].op);
}
function closeTools() { $('#toolsModal').classList.add('hidden'); }
$('#toolsBtn').onclick = () => openTools(); $('#sbTools').onclick = () => openTools(); $('#toolsClose').onclick = closeTools;
$('#toolsModal').addEventListener('click', e => { if (e.target.id === 'toolsModal') closeTools(); });
$('#toolList').addEventListener('click', e => { const it = e.target.closest('.tl-item'); if (it) selectTool(it.dataset.op); });
function selectTool(op) {
  curTool = toolSpecs.find(t => t.op === op); if (!curTool) return;
  document.querySelectorAll('.tl-item').forEach(x => x.classList.toggle('on', x.dataset.op === op));
  $('#toolTitle').textContent = curTool.label; $('#toolDesc').textContent = curTool.desc;
  toolSel = toolSel.filter(id => { const f = lib.find(x => x.id === id); return f && curTool.accepts.includes(f.kind); });
  if (!curTool.multi) toolSel = toolSel.slice(0, 1);
  renderToolFiles();
  $('#toolParams').innerHTML = curTool.params.map(paramHTML).join('');
  $('#toolResult').innerHTML = '';
  $('#toolRun').innerHTML = curTool.op === '__editor' ? ic('pen', 16) + 'Open editor' : ic('sparkles', 16) + 'Run';
}
function renderToolFiles() {
  const wrap = $('#toolFilesWrap');
  if (!curTool.accepts.length) { wrap.classList.add('hidden'); return; }
  wrap.classList.remove('hidden');
  $('#toolFilesHint').textContent = curTool.multi ? 'Pick files (click in the order you want)' : 'Pick a file';
  const fs = lib.filter(f => curTool.accepts.includes(f.kind));
  $('#toolFiles').innerHTML = fs.length ? fs.map(f => { const o = toolSel.indexOf(f.id);
    return `<div class="tf ${o > -1 ? 'on' : ''}" data-fid="${f.id}">${thumb(f)}<div style="min-width:0"><div class="nm">${esc(f.name)}</div><div class="mt">${esc(metaStr(f))}</div></div>${o > -1 ? `<span class="ord">${curTool.multi ? o + 1 : '✓'}</span>` : ''}</div>`; }).join('')
    : `<div class="empty">No ${curTool.accepts.join(' / ')} files yet — upload one.</div>`;
}
$('#toolFiles').addEventListener('click', e => {
  const t = e.target.closest('.tf'); if (!t) return; const id = t.dataset.fid, i = toolSel.indexOf(id);
  if (curTool.multi) { i > -1 ? toolSel.splice(i, 1) : toolSel.push(id); } else toolSel = i > -1 ? [] : [id];
  renderToolFiles();
});
function paramHTML(p) {
  const id = 'tp_' + p.name; let el;
  if (p.type === 'select') el = `<select class="inp" id="${id}">${p.options.map(o => `<option ${String(o) === String(p.default) ? 'selected' : ''}>${esc(o)}</option>`).join('')}</select>`;
  else if (p.type === 'textarea') el = `<textarea class="inp" id="${id}" placeholder="${esc(p.hint)}">${esc(p.default)}</textarea>`;
  else el = `<input class="inp" id="${id}" type="${p.type === 'number' ? 'number' : p.type === 'password' ? 'password' : 'text'}" step="any" value="${esc(p.default)}" placeholder="${esc(p.hint)}">`;
  return `<div class="flabel"><span>${esc(p.label)}</span></div>${el}`;
}
$('#toolUpload').onclick = () => $('#toolFileInput').click();
$('#toolFileInput').onchange = async e => {
  const items = await uploadFiles(e.target.files, false); e.target.value = '';
  await loadLib();
  items.filter(x => x.status === 'ready' && curTool.accepts.includes(x.kind)).forEach(x => { if (curTool.multi) toolSel.push(x.id); else toolSel = [x.id]; });
  renderToolFiles();
};
$('#toolRun').onclick = async () => {
  if (!curTool) return;
  if (curTool.accepts.length && !toolSel.length) return toast('Pick a file first', 'err');
  if (curTool.op === '__editor') return openEditor(toolSel[0]);
  const params = {}; curTool.params.forEach(p => params[p.name] = document.getElementById('tp_' + p.name).value);
  const btn = $('#toolRun'); btn.disabled = true; btn.innerHTML = '<span class="spin" style="width:16px;height:16px"></span>Working…';
  try {
    const r = await api('/api/tool', {op: curTool.op, files: toolSel, params});
    $('#toolResult').innerHTML = r.ok ? fileCard(r.file, r.message)
      : `<div class="fcard err"><div class="th" style="color:var(--bad)">${ic('x', 18)}</div><div class="info"><div class="nm">That didn't work</div><div class="mt">${esc(r.error)}</div></div></div>`;
    if (r.ok) { await loadLib(); renderToolFiles(); }
  } catch (e) { toast(e.message, 'err'); }
  finally { btn.disabled = false; btn.innerHTML = ic('sparkles', 16) + 'Run'; }
};
$('#toolResult').addEventListener('click', e => {
  const b = e.target.closest('[data-act="attach"]'); if (b) attachExisting(b.dataset.id);
  const ep = e.target.closest('[data-act="edit-pdf"]'); if (ep) openEditor(ep.dataset.id);
  const z = e.target.closest('[data-zoom]'); if (z) openLightbox(z.dataset.zoom);
});

/* ---------------- what's new + power ---------------- */
let META = {};
function showNews() {
  const log = META.changelog || []; if (!log.length) return toast('No release notes yet');
  let ov = $('#newsOv');
  if (!ov) { ov = document.createElement('div'); ov.id = 'newsOv'; ov.className = 'news-ov'; document.body.appendChild(ov);
    ov.addEventListener('click', e => { if (e.target === ov || e.target.closest('.x')) ov.remove(); }); }
  ov.innerHTML = `<div class="news-card glass"><h3>🐈 What's new in Meow<button class="x" title="Close">${ic('x', 18)}</button></h3>${log.map((r, i) =>
    `<h4 class="${i === 0 ? 'cur' : ''}">v${esc(r.v)}</h4><ul>${r.items.map(x => `<li>${esc(x)}</li>`).join('')}</ul>`).join('')}</div>`;
}
$('#sbNews').onclick = showNews;
$('#sbPower').onclick = async () => {
  if (!confirm(META.batch ? 'Power off Meow and release the Kaggle GPU? You can start it again anytime.' : 'Stop the Meow server? You can restart it by running the Kaggle cell again.')) return;
  try { await api('/api/shutdown', {}); } catch {}
  document.body.insertAdjacentHTML('beforeend', '<div class="news-ov"><div class="news-card glass" style="text-align:center"><h3 style="justify-content:center">⏻ Meow is powering off</h3><p style="color:#9a9cc0">GPU released. Start it again with your Start Meow button or the Kaggle cell.</p></div></div>');
};

/* ---------------- health ---------------- */
async function health() {
  try {
    const j = await (await fetch('/api/health', {headers: HDR})).json();
    $('#hdot').className = 'dot ok';
    $('#healthText').textContent = `Online · ${Math.round(j.ctx / 1024)}k context`;
    $('#modelBadge').textContent = `${j.model_short} · Kaggle GPU`;
    META = j; $('#verTag').textContent = 'v' + (j.version || '');
    try { const seen = localStorage.getItem('meow_seen_ver'); if (seen && seen !== j.version) { showNews(); } localStorage.setItem('meow_seen_ver', j.version); } catch {}
  } catch { $('#hdot').className = 'dot bad'; $('#healthText').textContent = 'Offline — is the Kaggle cell running?'; }
}

/* ---------------- galaxy: 3D starfield + shooting stars ---------------- */
const sky = $('#sky'), sx = sky.getContext('2d');
let W = 0, H = 0, stars = [], shooting = [], warp = 0, warpTarget = 0, nextShoot = 2500;
const reduceMotion = matchMedia('(prefers-reduced-motion: reduce)').matches;
const STAR_COLORS = ['255,255,255', '205,218,255', '232,214,255', '255,236,214', '196,232,255'];
function mkStar(fresh) { return {x: (Math.random() * 2 - 1) * W, y: (Math.random() * 2 - 1) * H, z: fresh ? W : Math.random() * W + 1, tw: Math.random() * 6.28, c: STAR_COLORS[Math.random() * STAR_COLORS.length | 0], px: null, py: null}; }
function resizeSky() {
  const dpr = Math.min(window.devicePixelRatio || 1, 2); W = innerWidth; H = innerHeight;
  sky.width = W * dpr; sky.height = H * dpr; sx.setTransform(dpr, 0, 0, dpr, 0, 0);
  stars = Array.from({length: Math.min(1100, Math.round(W * H / 1500))}, () => mkStar(false));
}
function tick(t) {
  requestAnimationFrame(tick);
  if (document.hidden) return;
  warp += (warpTarget - warp) * 0.035;
  sx.clearRect(0, 0, W, H);
  const cx = W / 2, cy = H / 2, f = W * 0.5, speed = reduceMotion ? 0 : 0.35 + warp * 9;
  for (const s of stars) {
    s.z -= speed;
    if (s.z < 1) { Object.assign(s, mkStar(true)); continue; }
    const k = f / s.z, x = cx + s.x * k, y = cy + s.y * k;
    if (x < -40 || x > W + 40 || y < -40 || y > H + 40) { Object.assign(s, mkStar(true)); continue; }
    const depth = 1 - s.z / W, r = Math.max(0.25, depth * 1.9);
    const a = (0.3 + 0.7 * depth) * (0.65 + 0.35 * Math.sin(t * 0.003 + s.tw));
    if (warp > 0.08 && s.px != null) {
      sx.strokeStyle = `rgba(${s.c},${a})`; sx.lineWidth = r; sx.beginPath(); sx.moveTo(s.px, s.py); sx.lineTo(x, y); sx.stroke();
    } else {
      sx.fillStyle = `rgba(${s.c},${a})`; sx.beginPath(); sx.arc(x, y, r, 0, 6.283); sx.fill();
      if (depth > 0.86) { sx.fillStyle = `rgba(${s.c},${a * 0.12})`; sx.beginPath(); sx.arc(x, y, r * 4, 0, 6.283); sx.fill(); }
    }
    s.px = x; s.py = y;
  }
  if (!reduceMotion && t > nextShoot && shooting.length < 2) {
    nextShoot = t + 3000 + Math.random() * 6000;
    const ang = Math.PI * (0.12 + Math.random() * 0.22), v = 9 + Math.random() * 7;
    shooting.push({x: Math.random() * W * 0.8, y: Math.random() * H * 0.35, vx: Math.cos(ang) * v, vy: Math.sin(ang) * v, life: 1});
  }
  shooting = shooting.filter(s => s.life > 0);
  for (const s of shooting) {
    s.x += s.vx; s.y += s.vy; s.life -= 0.018;
    const tx = s.x - s.vx * 10, ty = s.y - s.vy * 10, g = sx.createLinearGradient(s.x, s.y, tx, ty);
    g.addColorStop(0, `rgba(255,255,255,${s.life})`); g.addColorStop(1, 'rgba(167,139,250,0)');
    sx.strokeStyle = g; sx.lineWidth = 2; sx.beginPath(); sx.moveTo(s.x, s.y); sx.lineTo(tx, ty); sx.stroke();
  }
}

/* ================= PRO PDF EDITOR ================= */
Object.assign(ICONS, {
  cursor: '<path d="m4 4 7 17 2.5-7.5L21 11z"/>',
  type: '<path d="M4 7V4h16v3M9 20h6M12 4v16"/>',
  typeplus: '<path d="M3 7V4h12v3M9 4v14M6 18h6M19 13v8M15 17h8"/>',
  pen: '<path d="M12 20h9"/><path d="M16.5 3.5a2.1 2.1 0 0 1 3 3L7 19l-4 1 1-4z"/>',
  marker: '<path d="m9 11-6 6v3h9l3-3"/><path d="m22 12-4.6 4.6a2 2 0 0 1-2.8 0l-5.2-5.2a2 2 0 0 1 0-2.8L14 4"/>',
  square: '<rect x="3" y="3" width="18" height="18" rx="2"/>',
  eraser: '<path d="m7 21-4.3-4.3a1 1 0 0 1 0-1.4l10-10a1 1 0 0 1 1.4 0l5.6 5.6a1 1 0 0 1 0 1.4L13 19"/><path d="M22 21H7M5 11l9 9"/>',
  wand: '<path d="m15 4 5 5M3 21l11-11"/><path d="M14 4V2M18 8h2M17 3l1-1M20 6l1-1M9 3l.6 1.6L11 5l-1.4.6L9 7l-.6-1.4L7 5l1.4-.4z"/>',
  sign: '<path d="M3 17c3-5 5-8 7-8s-1 8 1 8 3-4 4-4 1 3 3 3 2-1 3-2"/><path d="M3 21h18"/>',
  undo: '<path d="M9 14 4 9l5-5"/><path d="M4 9h10.5a5.5 5.5 0 0 1 0 11H11"/>',
  redo: '<path d="m15 14 5-5-5-5"/><path d="M20 9H9.5a5.5 5.5 0 0 0 0 11H13"/>',
  zin: '<circle cx="11" cy="11" r="7"/><path d="m21 21-4.3-4.3M11 8v6M8 11h6"/>',
  zout: '<circle cx="11" cy="11" r="7"/><path d="m21 21-4.3-4.3M8 11h6"/>',
  up: '<path d="m18 15-6-6-6 6"/>'
});
const ed = {src: null, draft: null, pages: [], page: 0, zoom: 1, base: 1, tool: 'select', layout: null, sel: null,
  undo: [], redo: [], editing: null, busy: false, dirty: false, placeFile: null, placeW: 150, placeAR: 1, fileMode: 'add',
  hl: '#facc15', magicR: 10, shape: {kind: 'rect', width: 1.5, fill: false, fillColor: '#e9d5ff'}, bgc: null, bgk: 1, scan: {}};
const peOv = $('#peOv'), pePage = $('#pePage'), peImg = $('#peImg'), peInk = $('#peInk'), peSide = $('#peSide');
const FAMCSS = {serif: "'Times New Roman', Times, serif", sans: 'Helvetica, Arial, sans-serif', mono: "'Courier New', monospace"};
const IMG_FX = [['auto_enhance', 'Auto-enhance'], ['grayscale', 'Black & white'], ['brighter', 'Brighter'], ['darker', 'Darker'],
  ['contrast', 'More contrast'], ['sharpen', 'Sharpen'], ['blur', 'Blur'], ['sepia', 'Sepia'], ['invert', 'Invert'],
  ['rotate', 'Rotate 90°'], ['flip_h', 'Flip ↔'], ['flip_v', 'Flip ↕'], ['remove_bg', 'Remove BG']];
const S = () => ed.base * ed.zoom;
const edOpen = () => !$('#pdfEditor').classList.contains('hidden');
const place = (el, b, s) => { el.style.left = b[0] * s + 'px'; el.style.top = b[1] * s + 'px'; el.style.width = (b[2] - b[0]) * s + 'px'; el.style.height = (b[3] - b[1]) * s + 'px'; };
const norm = b => [Math.min(b[0], b[2]), Math.min(b[1], b[3]), Math.max(b[0], b[2]), Math.max(b[1], b[3])];
const toPdf = e => { const r = pePage.getBoundingClientRect(), s = S(); return [(e.clientX - r.left) / s, (e.clientY - r.top) / s]; };
function showBusy(t) { $('#peBusyText').textContent = t || 'Applying…'; $('#peBusy').classList.remove('hidden'); }
function hideBusy() { $('#peBusy').classList.add('hidden'); }

async function openEditor(fileId) {
  if (!fileId) return;
  closeTools();
  $('#pdfEditor').classList.remove('hidden'); showBusy('Opening…');
  try {
    const j = await api('/api/pdf/open', {file_id: fileId});
    Object.assign(ed, {src: j.source, draft: j.draft, pages: j.pages, page: 0, zoom: 1, undo: [], redo: [], dirty: false, sel: null, editing: null, placeFile: null, scan: {}});
    $('#peName').textContent = j.source.name;
    computeBase(); setTool('select'); renderThumbs(); await loadPage(); updateUndo();
  } catch (e) { toast('Could not open PDF: ' + e.message, 'err'); $('#pdfEditor').classList.add('hidden'); }
  finally { hideBusy(); }
}
function closeEditor(force) {
  if (!force && ed.dirty && !confirm('Close without saving your edits?')) return;
  if (ed.editing) cancelEdit();
  $('#pdfEditor').classList.add('hidden'); ed.dirty = false;
}
function computeBase() {
  const pw = ed.pages[ed.page]?.w || 595, avail = Math.max(260, $('#peStage').clientWidth - 56);
  ed.base = Math.min(avail / pw, 1.7);
}
function loadImg(url) { return new Promise((res, rej) => { peImg.onload = () => res(); peImg.onerror = () => rej(new Error('Page render failed')); peImg.src = url; }); }
async function loadPage(reselect) {
  const pg = ed.pages[ed.page]; if (!pg) return;
  const s = S(), dpr = Math.min(devicePixelRatio || 1, 2), rs = Math.min(4.5, +(s * dpr).toFixed(2));
  pePage.style.width = pg.w * s + 'px'; pePage.style.height = pg.h * s + 'px';
  const [lay] = await Promise.all([
    fetch(`/api/pdf/${ed.draft.id}/layout/${ed.page}`, {headers: HDR}).then(r => r.json()),
    loadImg(`/api/pdf/${ed.draft.id}/render/${ed.page}?scale=${rs}`)]);
  ed.layout = lay;
  peInk.width = pg.w * s * dpr; peInk.height = pg.h * s * dpr;
  buildBg(); drawOverlays(); markThumb(); updateStatus();
  const pageArea = pg.w * pg.h, big = lay.images.find(m => (m.bbox[2] - m.bbox[0]) * (m.bbox[3] - m.bbox[1]) > pageArea * 0.4);
  if (!lay.runs.length && big && !ed.scan[ed.page]) runOcr(big.bbox, true);
  $('#peZoom').textContent = Math.round(ed.zoom * 100) + '%';
  if (reselect) reselectAt(reselect); else select(null);
}
function buildBg() {
  try { const c = document.createElement('canvas'); c.width = peImg.naturalWidth; c.height = peImg.naturalHeight;
    ed.bgc = c.getContext('2d', {willReadFrequently: true}); ed.bgc.drawImage(peImg, 0, 0); ed.bgk = peImg.naturalWidth / ed.pages[ed.page].w; } catch { ed.bgc = null; }
}
function sampleBg(b) {
  if (!ed.bgc) return '#ffffff';
  const k = ed.bgk, pts = [[b[0] - 2, (b[1] + b[3]) / 2], [b[2] + 2, (b[1] + b[3]) / 2], [b[0], b[1] - 1.5]];
  for (const [x, y] of pts) {
    const px = Math.round(x * k), py = Math.round(y * k);
    if (px < 0 || py < 0 || px >= ed.bgc.canvas.width || py >= ed.bgc.canvas.height) continue;
    const d = ed.bgc.getImageData(px, py, 1, 1).data; return `rgb(${d[0]},${d[1]},${d[2]})`;
  }
  return '#ffffff';
}
function drawOverlays() {
  const s = S(); peOv.innerHTML = '';
  for (const m of ed.layout.images) { const d = document.createElement('div'); d.className = 'pe-img'; d.dataset.img = m.index; place(d, m.bbox, s); d.innerHTML = '<span class="h"></span>'; peOv.appendChild(d); }
  for (const r of ed.layout.runs) { const d = document.createElement('div'); d.className = 'pe-run'; d.dataset.run = r.id; place(d, r.bbox, s); peOv.appendChild(d); }
  for (const r of ed.scan[ed.page] || []) { const d = document.createElement('div'); d.className = 'pe-run scan'; d.dataset.run = r.id; place(d, r.bbox, s); peOv.appendChild(d); }
}
async function runOcr(imgBbox, auto) {
  const pageAt = ed.page, draftAt = ed.draft.id;
  showBusy('Reading text in the image…');
  try {
    const q = imgBbox ? `?x0=${imgBbox[0]}&y0=${imgBbox[1]}&x1=${imgBbox[2]}&y1=${imgBbox[3]}` : '';
    const j = await (await fetch(`/api/pdf/${draftAt}/ocr/${pageAt}${q}`, {headers: HDR})).json();
    if (ed.draft.id !== draftAt || ed.page !== pageAt) return;
    const old = (ed.scan[pageAt] || []).filter(b => !j.boxes.some(n => overlap(n.bbox, b.bbox) > 0.3));
    ed.scan[pageAt] = old.concat(j.boxes.map((b, i) => ({...b, id: `s${Date.now() % 1e6}_${i}`})));
    drawOverlays(); if (ed.sel) select(ed.sel);
    if (j.boxes.length) { if (!auto || ed.tool === 'select') setTool('text'); toast(`Found ${j.boxes.length} line${j.boxes.length > 1 ? 's' : ''} of text — click one to edit it`); }
    else if (!auto) toast(j.error ? 'Text reading failed: ' + j.error : 'No readable text found in that image', 'err');
  } catch (e) { if (!auto) toast('Text reading failed: ' + e.message, 'err'); }
  finally { hideBusy(); }
}
function overlap(a, b) {
  const ix = Math.max(0, Math.min(a[2], b[2]) - Math.max(a[0], b[0])), iy = Math.max(0, Math.min(a[3], b[3]) - Math.max(a[1], b[1]));
  return ix * iy / Math.max(1e-6, Math.min((a[2] - a[0]) * (a[3] - a[1]), (b[2] - b[0]) * (b[3] - b[1])));
}
const runOf = el => el.classList.contains('scan') ? (ed.scan[ed.page] || []).find(r => r.id === el.dataset.run) : ed.layout.runs.find(r => String(r.id) === el.dataset.run);
const imgOf = el => ed.layout.images.find(m => String(m.index) === el.dataset.img);
function select(sel) {
  ed.sel = sel;
  peOv.querySelectorAll('.sel').forEach(x => x.classList.remove('sel'));
  if (sel?.type === 'img') peOv.querySelector(`.pe-img[data-img="${sel.m.index}"]`)?.classList.add('sel');
  if (sel?.type === 'run') peOv.querySelector(`.pe-run[data-run="${CSS.escape(String(sel.r.id))}"]`)?.classList.add('sel');
  renderSide();
}
function reselectAt(rect) {
  let best = null, sc = 0; const r = norm(rect);
  for (const m of ed.layout.images) {
    const b = m.bbox, ix = Math.max(0, Math.min(b[2], r[2]) - Math.max(b[0], r[0])), iy = Math.max(0, Math.min(b[3], r[3]) - Math.max(b[1], r[1]));
    const inter = ix * iy, uni = (b[2] - b[0]) * (b[3] - b[1]) + (r[2] - r[0]) * (r[3] - r[1]) - inter, v = inter / Math.max(1e-6, uni);
    if (v > sc) { sc = v; best = m; }
  }
  select(best && sc > 0.2 ? {type: 'img', m: best} : null);
}
function updateStatus() { $('#peStatus').textContent = `Page ${ed.page + 1}/${ed.pages.length}${ed.dirty ? ' · unsaved' : ''}`; }
function updateUndo() { $('#peUndo').disabled = !ed.undo.length; $('#peRedo').disabled = !ed.redo.length; updateStatus(); }

/* ---------- apply edits on the server (every edit is a new draft → exact preview + undo) ---------- */
async function applyOps(ops, opt = {}) {
  if (ed.busy) return false;
  ed.busy = true; showBusy(opt.label || 'Applying…');
  try {
    const r = await api(`/api/pdf/${ed.draft.id}/edit`, {ops});
    if (!r.ok) { toast(r.error, 'err'); return false; }
    ed.undo.push({draft: ed.draft, pages: ed.pages}); if (ed.undo.length > 80) ed.undo.shift(); ed.redo = [];
    ed.draft = r.draft; ed.pages = r.pages; ed.dirty = true;
    if (ops.some(o => o.op.startsWith('page_'))) ed.scan = {};
    else if (!ops.every(o => o.op === 'scan_text')) ops.forEach(o => { if (['move_image', 'image_filter', 'replace_image', 'delete_image', 'erase', 'magic_erase'].includes(o.op)) delete ed.scan[o.page]; });
    if (opt.onOk) opt.onOk();
    if (opt.page != null) ed.page = opt.page;
    ed.page = Math.max(0, Math.min(ed.page, ed.pages.length - 1));
    await loadPage(opt.reselect);
    if (ops.some(o => o.op.startsWith('page_'))) renderThumbs(); else refreshThumb(ed.page);
    return true;
  } catch (e) { toast(e.message, 'err'); return false; }
  finally { ed.busy = false; hideBusy(); updateUndo(); }
}
async function stepHistory(from, to) {
  if (ed.busy || !from.length) return;
  if (ed.editing) cancelEdit();
  to.push({draft: ed.draft, pages: ed.pages});
  const st = from.pop(); ed.draft = st.draft; ed.pages = st.pages; ed.dirty = true; ed.scan = {};
  ed.page = Math.min(ed.page, ed.pages.length - 1);
  renderThumbs(); await loadPage(); updateUndo();
}

/* ---------- thumbnails & page operations ---------- */
function renderThumbs() {
  $('#peThumbs').innerHTML = ed.pages.map((p, i) => `<div class="pe-th ${i === ed.page ? 'on' : ''}" data-p="${i}">
    <img loading="lazy" src="/api/pdf/${ed.draft.id}/render/${i}?scale=0.26" style="aspect-ratio:${p.w}/${p.h}" alt="">
    <div class="no">${i + 1}</div><div class="ops">
    <button data-pop="rotate" title="Rotate page">${ic('refresh', 13)}</button>${i > 0 ? `<button data-pop="up" title="Move up">${ic('up', 13)}</button>` : ''}
    ${i < ed.pages.length - 1 ? `<button data-pop="down" title="Move down">${ic('chev', 13)}</button>` : ''}<button data-pop="dup" title="Duplicate page">${ic('copy', 13)}</button>
    ${ed.pages.length > 1 ? `<button data-pop="del" title="Delete page">${ic('trash', 13)}</button>` : ''}</div></div>`).join('')
    + '<button class="pe-addpage" id="peAddPage">+ Blank page</button>';
}
function refreshThumb(i) { const img = $(`#peThumbs .pe-th[data-p="${i}"] img`); if (img) img.src = `/api/pdf/${ed.draft.id}/render/${i}?scale=0.26`; }
function markThumb() { document.querySelectorAll('#peThumbs .pe-th').forEach(t => t.classList.toggle('on', +t.dataset.p === ed.page)); }
async function gotoPage(i) {
  if (ed.busy || i === ed.page || i < 0 || i >= ed.pages.length) return;
  if (ed.editing) await commitEdit();
  ed.page = i; computeBase(); await loadPage();
  $(`#peThumbs .pe-th[data-p="${i}"]`)?.scrollIntoView({block: 'nearest'});
}
$('#peThumbs').addEventListener('click', e => {
  if (e.target.closest('#peAddPage')) return applyOps([{op: 'page_insert', after: ed.page}], {page: ed.page + 1, label: 'Adding page…'});
  const th = e.target.closest('.pe-th'); if (!th) return; const i = +th.dataset.p;
  const b = e.target.closest('[data-pop]');
  if (!b) return gotoPage(i);
  const a = b.dataset.pop;
  if (a === 'rotate') applyOps([{op: 'page_rotate', page: i, angle: 90}], {page: i});
  else if (a === 'up') applyOps([{op: 'page_move', page: i, to: i - 1}], {page: i - 1});
  else if (a === 'down') applyOps([{op: 'page_move', page: i, to: i + 1}], {page: i + 1});
  else if (a === 'dup') applyOps([{op: 'page_duplicate', page: i}], {page: i + 1});
  else if (a === 'del' && confirm(`Delete page ${i + 1}?`)) applyOps([{op: 'page_delete', page: i}], {page: Math.max(0, i - 1)});
});

/* ---------- tools & style props ---------- */
function setTool(t) {
  if (ed.editing) commitEdit();
  if (t !== 'place') { ed.placeFile = null; $('.pe-ghost')?.remove(); }
  if (t === 'image') { ed.fileMode = 'add'; $('#peFile').click(); return; }
  if (t === 'sign') { openSign(); return; }
  ed.tool = t; peOv.dataset.tool = t;
  document.querySelectorAll('#peTools button').forEach(b => b.classList.toggle('on', b.dataset.tool === t));
  if (t !== 'select') select(null); else renderSide();
}
$('#peTools').addEventListener('click', e => { const b = e.target.closest('[data-tool]'); if (b) setTool(b.dataset.tool); });
const props = () => ({color: $('#peColor').value, family: $('#peFamily').value, size: parseFloat($('#peSize').value) || 12,
  bold: $('#peBold').classList.contains('on'), italic: $('#peItalic').classList.contains('on')});
function setProps(p) {
  if (p.color) $('#peColor').value = /^#[0-9a-f]{6}$/i.test(p.color) ? p.color : '#111111';
  if (p.family) $('#peFamily').value = p.family;
  if (p.size) $('#peSize').value = +(+p.size).toFixed(1);
  $('#peBold').classList.toggle('on', !!p.bold); $('#peItalic').classList.toggle('on', !!p.italic);
}
function liveStyle() {
  const E = ed.editing; if (!E) return; const p = props(), s = S();
  Object.assign(E.el.style, {color: p.color, fontFamily: FAMCSS[p.family], fontSize: p.size * s + 'px', fontWeight: p.bold ? '700' : '400', fontStyle: p.italic ? 'italic' : 'normal'});
  if (!E.run) E.el.style.lineHeight = p.size * 1.2 * s + 'px';
}
['peColor', 'peFamily', 'peSize'].forEach(id => $('#' + id).addEventListener('input', liveStyle));
['peBold', 'peItalic'].forEach(id => $('#' + id).addEventListener('mousedown', e => { e.preventDefault(); $('#' + id).classList.toggle('on'); liveStyle(); }));

/* ---------- text editing ---------- */
function selectAll(el) { const r = document.createRange(); r.selectNodeContents(el); const s = getSelection(); s.removeAllRanges(); s.addRange(r); }
function bindEditKeys(el) {
  el.addEventListener('keydown', e => {
    e.stopPropagation();
    if (e.key === 'Enter' && !e.shiftKey) { e.preventDefault(); commitEdit(); }
    else if (e.key === 'Escape') { e.preventDefault(); cancelEdit(); }
  });
  el.addEventListener('paste', e => { e.preventDefault(); document.execCommand('insertText', false, (e.clipboardData || window.clipboardData).getData('text/plain')); });
}
function startEdit(r) {
  if (ed.editing) commitEdit();
  select({type: 'run', r});
  const s = S(), el = document.createElement('div'), h = r.bbox[3] - r.bbox[1];
  el.className = 'pe-edit'; el.contentEditable = 'true'; el.spellcheck = false;
  Object.assign(el.style, {left: r.bbox[0] * s + 'px', top: r.bbox[1] * s + 'px', minWidth: (r.bbox[2] - r.bbox[0]) * s + 'px',
    lineHeight: h * s + 'px', background: r.bg || sampleBg(r.bbox)});
  el.textContent = r.text; peOv.appendChild(el);
  setProps(r); ed.editing = {el, run: r, start: props()}; liveStyle();
  bindEditKeys(el); el.focus(); selectAll(el);
}
function startAdd(p) {
  const s = S(), el = document.createElement('div');
  el.className = 'pe-edit'; el.contentEditable = 'true'; el.spellcheck = false;
  Object.assign(el.style, {left: p[0] * s + 'px', top: p[1] * s + 'px', background: 'rgba(255,255,255,.55)'});
  peOv.appendChild(el); ed.editing = {el, run: null, at: p, start: props()}; liveStyle();
  bindEditKeys(el); setTimeout(() => el.focus(), 0);
}
function cancelEdit() { if (ed.editing) { ed.editing.el.remove(); ed.editing = null; } }
async function commitEdit() {
  const E = ed.editing; if (!E) return; ed.editing = null;
  const text = E.el.innerText.replace(/ /g, ' ').replace(/\n+$/, ''); E.el.remove();
  const p = props();
  if (E.run?.scan) {
    if (text === E.run.text && p.color === E.start.color) return;
    const r = E.run, op = {op: 'scan_text', page: ed.page, bbox: r.bbox, text, orig: r.text};
    if (p.color !== E.start.color) op.color = p.color;
    await applyOps([op], {label: 'Rewriting text in image…', onOk: () => { r.text = text; if (op.color) r.color = op.color; }});
  } else if (E.run) {
    const diff = ['color', 'family', 'size', 'bold', 'italic'].filter(k => String(p[k]) !== String(E.start[k]));
    if (text === E.run.text && !diff.length) return;
    const op = {op: 'edit_text', page: ed.page, bbox: E.run.bbox, text}; diff.forEach(k => op[k] = p[k]);
    await applyOps([op], {label: 'Updating text…'});
  } else if (text.trim()) {
    await applyOps([{op: 'add_text', page: ed.page, x: E.at[0], y: E.at[1] + p.size * 0.88, text, size: p.size, color: p.color,
      family: p.family, bold: p.bold, italic: p.italic}], {label: 'Adding text…'});
  }
}

/* ---------- pointer interactions ---------- */
peOv.addEventListener('pointerdown', e => {
  if (ed.busy || e.button > 0 || e.target.closest('.pe-edit')) return;
  if (ed.editing) { e.preventDefault(); commitEdit(); return; }
  const p = toPdf(e), t = ed.tool, runEl = e.target.closest('.pe-run'), imgEl = e.target.closest('.pe-img');
  if (t === 'text') { if (runEl) { e.preventDefault(); startEdit(runOf(runEl)); } else { e.preventDefault(); startAdd(p); } return; }
  if (t === 'select') {
    if (imgEl) { e.preventDefault(); select({type: 'img', m: imgOf(imgEl)}); startDrag(e, imgEl, e.target.classList.contains('h')); return; }
    if (runEl) { select({type: 'run', r: runOf(runEl)}); return; }
    return select(null);
  }
  e.preventDefault();
  if (t === 'addtext') startAdd(p);
  else if (t === 'draw') startInk(p);
  else if (t === 'magic') startMagic(p);
  else if (['highlight', 'shape', 'erase'].includes(t)) startRubber(p, t);
  else if (t === 'place') placeImage(p);
});
peOv.addEventListener('dblclick', e => { const r = e.target.closest('.pe-run'); if (r && ed.tool === 'select') startEdit(runOf(r)); });
peOv.addEventListener('pointermove', e => {
  if (ed.tool !== 'place' || !ed.placeFile) return;
  let g = $('.pe-ghost');
  if (!g) { g = document.createElement('img'); g.className = 'pe-ghost'; g.src = `/api/thumb/${ed.placeFile.id}?size=400`; peOv.appendChild(g); }
  const p = toPdf(e), w = ed.placeW, h = w * ed.placeAR; place(g, [p[0] - w / 2, p[1] - h / 2, p[0] + w / 2, p[1] + h / 2], S());
});
function track(move, up) {
  const mv = ev => move(ev), u = ev => { removeEventListener('pointermove', mv); removeEventListener('pointerup', u); up(ev); };
  addEventListener('pointermove', mv); addEventListener('pointerup', u);
}
function startDrag(e, el, resize) {
  const s = S(), m = imgOf(el), b0 = [...m.bbox], sx = e.clientX, sy = e.clientY, ar = (b0[3] - b0[1]) / Math.max(1, b0[2] - b0[0]);
  let moved = false, nb = null;
  track(ev => {
    const dx = (ev.clientX - sx) / s, dy = (ev.clientY - sy) / s;
    if (Math.abs(dx) + Math.abs(dy) > 1.5) moved = true;
    if (resize) { const w = Math.max(8, b0[2] - b0[0] + dx); nb = [b0[0], b0[1], b0[0] + w, b0[1] + (ev.shiftKey ? Math.max(8, b0[3] - b0[1] + dy) : w * ar)]; }
    else nb = [b0[0] + dx, b0[1] + dy, b0[2] + dx, b0[3] + dy];
    place(el, nb, s);
  }, () => {
    if (moved && nb) { nb = nb.map(v => +v.toFixed(2)); applyOps([{op: 'move_image', page: ed.page, xref: m.xref, bbox: m.bbox, rect: nb}], {reselect: nb, label: resize ? 'Resizing…' : 'Moving…'}); }
  });
}
function startInk(p) {
  const s = S(), dpr = Math.min(devicePixelRatio || 1, 2), ctx = peInk.getContext('2d'), pr = props(), w = Math.max(0.6, pr.size / 6), pts = [p];
  ctx.setTransform(dpr * s, 0, 0, dpr * s, 0, 0); Object.assign(ctx, {strokeStyle: pr.color, lineWidth: w, lineCap: 'round', lineJoin: 'round'});
  ctx.beginPath(); ctx.moveTo(p[0], p[1]);
  track(ev => { const q = toPdf(ev), l = pts[pts.length - 1]; if (Math.hypot(q[0] - l[0], q[1] - l[1]) < 0.7) return; pts.push(q); ctx.lineTo(q[0], q[1]); ctx.stroke(); },
    async () => {
      if (pts.length > 1) await applyOps([{op: 'ink', page: ed.page, points: pts.map(q => [+q[0].toFixed(2), +q[1].toFixed(2)]), color: pr.color, width: w}], {label: 'Saving drawing…'});
      ctx.setTransform(1, 0, 0, 1, 0, 0); ctx.clearRect(0, 0, peInk.width, peInk.height);
    });
}
function startMagic(p) {
  const s = S(), dpr = Math.min(devicePixelRatio || 1, 2), ctx = peInk.getContext('2d'), r = ed.magicR, pts = [p];
  ctx.setTransform(dpr * s, 0, 0, dpr * s, 0, 0);
  Object.assign(ctx, {strokeStyle: 'rgba(168,85,247,.55)', fillStyle: 'rgba(168,85,247,.55)', lineWidth: r * 2, lineCap: 'round', lineJoin: 'round'});
  ctx.globalCompositeOperation = 'copy'; ctx.beginPath(); ctx.arc(p[0], p[1], r, 0, 7); ctx.fill(); ctx.globalCompositeOperation = 'source-over';
  const redraw = () => { ctx.save(); ctx.setTransform(1, 0, 0, 1, 0, 0); ctx.clearRect(0, 0, peInk.width, peInk.height); ctx.restore();
    ctx.beginPath(); ctx.moveTo(pts[0][0], pts[0][1]); pts.forEach(q => ctx.lineTo(q[0], q[1])); if (pts.length === 1) ctx.lineTo(p[0] + .01, p[1]); ctx.stroke(); };
  redraw();
  track(ev => { const q = toPdf(ev), l = pts[pts.length - 1]; if (Math.hypot(q[0] - l[0], q[1] - l[1]) < Math.max(1, r / 4)) return; pts.push(q); redraw(); },
    async () => {
      await applyOps([{op: 'magic_erase', page: ed.page, points: pts.map(q => [+q[0].toFixed(2), +q[1].toFixed(2)]), radius: +r.toFixed(2)}], {label: 'Magic erasing… rebuilding the background'});
      ctx.setTransform(1, 0, 0, 1, 0, 0); ctx.clearRect(0, 0, peInk.width, peInk.height);
    });
}
function startRubber(p, t) {
  const s = S(), el = document.createElement('div'); el.className = 'pe-rubber';
  if (t === 'highlight') el.style.background = 'rgba(250,204,21,.35)';
  if (t === 'erase') el.style.background = 'rgba(255,255,255,.75)';
  peOv.appendChild(el); let raw = [p[0], p[1], p[0], p[1]];
  track(ev => { const q = toPdf(ev); raw = [p[0], p[1], q[0], q[1]]; place(el, norm(raw), s); }, () => {
    el.remove(); const b = norm(raw).map(v => +v.toFixed(2));
    if (b[2] - b[0] < 3 && b[3] - b[1] < 3) return;
    if (t === 'highlight') applyOps([{op: 'highlight', page: ed.page, rect: b, color: ed.hl}], {label: 'Highlighting…'});
    else if (t === 'erase') applyOps([{op: 'erase', page: ed.page, rect: b}], {label: 'Erasing…'});
    else applyOps([{op: 'shape', page: ed.page, kind: ed.shape.kind, rect: ed.shape.kind === 'line' ? raw.map(v => +v.toFixed(2)) : b,
      color: props().color, width: ed.shape.width, fill: ed.shape.fill ? ed.shape.fillColor : null}], {label: 'Drawing…'});
  });
}

/* ---------- images: add / place / replace ---------- */
const arOf = f => { const d = (f.meta?.dims || '').split('×').map(Number); return d.length === 2 && d[0] ? d[1] / d[0] : 1; };
$('#peFile').onchange = async e => {
  const files = e.target.files; if (!files.length) return; e.target.value = '';
  showBusy('Uploading…'); const [f] = await uploadFiles(files, false); hideBusy();
  if (!f || f.status !== 'ready' || f.kind !== 'image') return toast('That file isn\'t a usable image', 'err');
  if (ed.fileMode === 'replace' && ed.sel?.type === 'img') {
    const m = ed.sel.m; return applyOps([{op: 'replace_image', page: ed.page, xref: m.xref, bbox: m.bbox, file_id: f.id}], {reselect: m.bbox, label: 'Replacing image…'});
  }
  startPlace(f, ed.fileMode === 'sign' ? 150 : Math.min(ed.pages[ed.page].w * 0.4, 240));
};
function startPlace(f, w) {
  ed.placeFile = f; ed.placeW = w; ed.placeAR = arOf(f);
  ed.tool = 'place'; peOv.dataset.tool = 'place';
  document.querySelectorAll('#peTools button').forEach(b => b.classList.remove('on'));
  toast('Click on the page where it should go'); renderSide();
}
function placeImage(p) {
  const f = ed.placeFile; if (!f) return;
  const w = ed.placeW, h = w * ed.placeAR, r = [p[0] - w / 2, p[1] - h / 2, p[0] + w / 2, p[1] + h / 2].map(v => +v.toFixed(2));
  ed.placeFile = null; $('.pe-ghost')?.remove(); setTool('select');
  applyOps([{op: 'add_image', page: ed.page, file_id: f.id, rect: r}], {reselect: r, label: 'Placing image…'});
}

/* ---------- signature pad ---------- */
const sig = $('#sigPad'); let sigCtx = null, sigBox = null, sigOn = false;
function openSign() {
  $('#signModal').classList.remove('hidden');
  const r = sig.getBoundingClientRect(), dpr = Math.min(devicePixelRatio || 1, 2);
  sig.width = r.width * dpr; sig.height = r.height * dpr; sigCtx = sig.getContext('2d'); sigCtx.scale(dpr, dpr);
  Object.assign(sigCtx, {lineCap: 'round', lineJoin: 'round', lineWidth: 2.8}); sigBox = null;
}
const sigPos = e => { const r = sig.getBoundingClientRect(); return [e.clientX - r.left, e.clientY - r.top]; };
const grow = (p) => { sigBox = sigBox ? [Math.min(sigBox[0], p[0]), Math.min(sigBox[1], p[1]), Math.max(sigBox[2], p[0]), Math.max(sigBox[3], p[1])] : [p[0], p[1], p[0], p[1]]; };
sig.addEventListener('pointerdown', e => { sigOn = true; const p = sigPos(e); sigCtx.strokeStyle = $('#sigColor').value; sigCtx.beginPath(); sigCtx.moveTo(...p); grow(p); sig.setPointerCapture(e.pointerId); });
sig.addEventListener('pointermove', e => { if (!sigOn) return; const p = sigPos(e); sigCtx.lineTo(...p); sigCtx.stroke(); grow(p); });
sig.addEventListener('pointerup', () => sigOn = false);
$('#sigClear').onclick = () => { sigCtx.clearRect(0, 0, sig.width, sig.height); sigBox = null; };
$('#signClose').onclick = () => $('#signModal').classList.add('hidden');
$('#sigUpload').onclick = () => { $('#signModal').classList.add('hidden'); ed.fileMode = 'sign'; $('#peFile').click(); };
$('#sigUse').onclick = async () => {
  if (!sigBox) return toast('Draw your signature first ✍️');
  const dpr = sig.width / sig.getBoundingClientRect().width, pad = 6;
  const x = Math.max(0, (sigBox[0] - pad) * dpr), y = Math.max(0, (sigBox[1] - pad) * dpr);
  const w = Math.min(sig.width - x, (sigBox[2] - sigBox[0] + pad * 2) * dpr), h = Math.min(sig.height - y, (sigBox[3] - sigBox[1] + pad * 2) * dpr);
  const c = document.createElement('canvas'); c.width = w; c.height = h; c.getContext('2d').drawImage(sig, x, y, w, h, 0, 0, w, h);
  const blob = await new Promise(r => c.toBlob(r, 'image/png'));
  $('#signModal').classList.add('hidden'); showBusy('Preparing signature…');
  const [f] = await uploadFiles([new File([blob], 'signature.png', {type: 'image/png'})], false); hideBusy();
  if (f?.status === 'ready') startPlace(f, 150); else toast('Signature upload failed', 'err');
};

/* ---------- inspector panel ---------- */
function renderSide() {
  const sel = ed.sel, t = ed.tool; let h = '';
  if (sel?.type === 'img') {
    const m = sel.m, b = m.bbox;
    h = `<h4>Image ${m.index}</h4><div class="muted">${Math.round(b[2] - b[0])} × ${Math.round(b[3] - b[1])} pt on page${m.w ? ` · ${m.w}×${m.h}px` : ''}</div>
      ${m.xref ? `<img class="prev" src="/api/pdf/${ed.draft.id}/image/${m.xref}" alt="">` : ''}
      <div class="grid"><button class="btn wide" data-se="ocr">${ic('type', 14)}Edit text in this image</button><button class="btn" data-se="replace">${ic('refresh', 14)}Replace</button>
      ${m.xref ? `<a class="btn" href="/api/pdf/${ed.draft.id}/image/${m.xref}" download="image.png">${ic('download', 14)}Download</a>` : '<span></span>'}
      <button class="btn" data-se="center">Center</button><button class="btn danger" data-se="delimg">${ic('trash', 14)}Delete</button></div>
      <div class="sec">Adjust image</div><div class="grid">${IMG_FX.map(([k, l]) => `<button class="btn" data-fx="${k}">${l}</button>`).join('')}</div>
      <div class="muted" style="margin-top:12px">Drag to move · corner handle resizes (hold Shift to stretch) · <kbd>Del</kbd> deletes</div>`;
  } else if (sel?.type === 'run' && sel.r.scan) {
    const r = sel.r;
    h = `<h4>Text in image</h4><div class="muted" style="margin:6px 0;color:var(--text)">“${esc(r.text)}”</div>
      <div class="muted">Read with OCR. When you change it, Meow wipes the old letters, rebuilds the background and redraws your text in the closest matching font, size and colour.</div>
      <div class="grid"><button class="btn wide" data-se="edit">${ic('type', 14)}Edit text</button><button class="btn danger wide" data-se="scanerase">${ic('eraser', 14)}Remove text</button></div>
      <div class="muted" style="margin-top:12px">Tip: if OCR misread a letter, just type the whole line correctly. Colour can be changed in the toolbar; the font is matched automatically.</div>`;
  } else if (sel?.type === 'run') {
    const r = sel.r;
    h = `<h4>Text</h4><div class="muted" style="margin:6px 0;color:var(--text)">“${esc(r.text)}”</div>
      <div class="muted">${esc((r.font || '').split('+').pop())} · ${+r.size.toFixed(1)}pt · ${r.family}${r.bold ? ' · bold' : ''}${r.italic ? ' · italic' : ''}</div>
      <div class="grid"><button class="btn wide" data-se="edit">${ic('type', 14)}Edit text</button><button class="btn danger wide" data-se="deltext">${ic('trash', 14)}Delete text</button></div>
      <div class="muted" style="margin-top:12px">While editing you can change colour, font, size, bold and italic in the toolbar. New text keeps the original font when possible. <kbd>Enter</kbd> saves · <kbd>Shift+Enter</kbd> new line · <kbd>Esc</kbd> cancels.</div>`;
  } else if (t === 'shape') {
    const sh = ed.shape;
    h = `<h4>Shape</h4><div class="muted">Drag on the page to draw. Stroke colour comes from the toolbar.</div>
      <div class="row"><label>Type</label><select class="inp" id="shKind">${['rect', 'ellipse', 'line'].map(k => `<option ${k === sh.kind ? 'selected' : ''}>${k}</option>`).join('')}</select></div>
      <div class="row"><label>Line width</label><input class="inp" id="shW" type="number" min="0.5" step="0.5" value="${sh.width}"></div>
      <div class="row"><label>Fill</label><input type="checkbox" id="shFill" ${sh.fill ? 'checked' : ''}><input type="color" id="shFillC" value="${sh.fillColor}"></div>`;
  } else if (t === 'magic') {
    h = `<h4>Magic eraser</h4><div class="muted">Paint over anything — a person, stamp, logo, watermark, scribble or words — and Meow rebuilds what was behind it with AI (GPU when free). Works on photos, screenshots and scans.</div>
      <div class="row"><label>Brush</label><input type="range" id="mgR" min="3" max="60" step="1" value="${ed.magicR}" style="flex:1"><span id="mgRv">${ed.magicR}</span></div>
      <div class="muted">Tip: cover the whole object plus a little margin. Two short strokes beat one huge one.</div>`;
  } else if (t === 'highlight') {
    h = `<h4>Highlight</h4><div class="muted">Drag across text to highlight it.</div><div class="row" style="flex-wrap:wrap">${['#facc15', '#4ade80', '#60a5fa', '#f472b6', '#fb923c'].map(c => `<button class="btn" data-hl="${c}" style="width:36px;height:36px;padding:0;background:${c};${c === ed.hl ? 'outline:2px solid #fff' : ''}"></button>`).join('')}</div>`;
  } else {
    const tips = {select: 'Click an image to move, resize, replace or adjust it. Double-click text to edit it.',
      text: 'Click any text to edit it in place, in its original font. Click empty space to add new text.',
      addtext: 'Click anywhere to type new text. Pick font, size and colour in the toolbar first.',
      draw: 'Draw freehand with the toolbar colour. Size controls pen thickness.',
      erase: 'Drag a box to permanently remove everything inside it (text, images, lines).',
      place: 'Click where the image should go. You can move and resize it afterwards.'};
    h = `<h4>${{select: 'Select', text: 'Edit text', addtext: 'Add text', draw: 'Draw', erase: 'Erase', place: 'Place image'}[t] || 'Editor'}</h4>
      <div class="muted">${tips[t] || ''}</div><div class="sec">Shortcuts</div>
      <div class="muted" style="line-height:1.9"><kbd>V</kbd> select · <kbd>E</kbd> edit text · <kbd>T</kbd> add text · <kbd>I</kbd> image<br><kbd>D</kbd> draw · <kbd>H</kbd> highlight · <kbd>R</kbd> shape · <kbd>X</kbd> erase · <kbd>M</kbd> magic · <kbd>S</kbd> sign<br><kbd>Ctrl Z</kbd> undo · <kbd>Ctrl Y</kbd> redo · <kbd>Del</kbd> delete · <kbd>Esc</kbd> close</div>
      <div class="sec">Ask Meow</div><div class="muted">You can also just tell Meow in chat — “replace the logo with this image”, “change the date to 16 Oct 2026”, “make page 2 photo black & white”.</div>`;
  }
  peSide.innerHTML = h;
  peSide.classList.toggle('empty', !sel && ['select', 'text'].includes(t));
}
peSide.addEventListener('click', e => {
  const hl = e.target.closest('[data-hl]'); if (hl) { ed.hl = hl.dataset.hl; return renderSide(); }
  const fx = e.target.closest('[data-fx]');
  if (fx && ed.sel?.type === 'img') { const m = ed.sel.m; return applyOps([{op: 'image_filter', page: ed.page, xref: m.xref, bbox: m.bbox, action: fx.dataset.fx}], {reselect: m.bbox, label: fx.textContent + '…'}); }
  const b = e.target.closest('[data-se]'); if (!b) return;
  const a = b.dataset.se, sel = ed.sel;
  if (a === 'replace') { ed.fileMode = 'replace'; $('#peFile').click(); }
  else if (a === 'delimg') delSelection();
  else if (a === 'center' && sel?.type === 'img') { const m = sel.m, pg = ed.pages[ed.page], w = m.bbox[2] - m.bbox[0], h = m.bbox[3] - m.bbox[1];
    const nb = [(pg.w - w) / 2, (pg.h - h) / 2, (pg.w + w) / 2, (pg.h + h) / 2].map(v => +v.toFixed(2));
    applyOps([{op: 'move_image', page: ed.page, xref: m.xref, bbox: m.bbox, rect: nb}], {reselect: nb, label: 'Centering…'}); }
  else if (a === 'edit' && sel?.type === 'run') startEdit(sel.r);
  else if (a === 'ocr' && sel?.type === 'img') runOcr(sel.m.bbox, false);
  else if (a === 'scanerase' && sel?.type === 'run') { const r = sel.r; applyOps([{op: 'scan_text', page: ed.page, bbox: r.bbox, text: '', orig: r.text}], {label: 'Removing text…', onOk: () => { ed.scan[ed.page] = (ed.scan[ed.page] || []).filter(b => b !== r); }}); }
  else if (a === 'deltext') delSelection();
});
peSide.addEventListener('input', e => {
  if (e.target.id === 'shKind') ed.shape.kind = e.target.value;
  if (e.target.id === 'mgR') { ed.magicR = parseFloat(e.target.value) || 10; const v = $('#mgRv'); if (v) v.textContent = ed.magicR; return; }
  if (e.target.id === 'shW') ed.shape.width = parseFloat(e.target.value) || 1.5;
  if (e.target.id === 'shFill') ed.shape.fill = e.target.checked;
  if (e.target.id === 'shFillC') ed.shape.fillColor = e.target.value;
});
function delSelection() {
  const s = ed.sel; if (!s) return;
  if (s.type === 'img') applyOps([{op: 'delete_image', page: ed.page, xref: s.m.xref, bbox: s.m.bbox}], {label: 'Deleting image…'});
  if (s.type === 'run' && s.r.scan) { const r = s.r; applyOps([{op: 'scan_text', page: ed.page, bbox: r.bbox, text: '', orig: r.text}], {label: 'Removing text…', onOk: () => { ed.scan[ed.page] = (ed.scan[ed.page] || []).filter(b => b !== r); }}); }
  else if (s.type === 'run') applyOps([{op: 'delete_text', page: ed.page, bbox: s.r.bbox}], {label: 'Deleting text…'});
}

/* ---------- top bar ---------- */
$('#peUndo').onclick = () => stepHistory(ed.undo, ed.redo);
$('#peRedo').onclick = () => stepHistory(ed.redo, ed.undo);
const zoomTo = z => { ed.zoom = Math.max(0.3, Math.min(4, z)); loadPage(); };
$('#peZoomIn').onclick = () => zoomTo(ed.zoom * 1.2);
$('#peZoomOut').onclick = () => zoomTo(ed.zoom / 1.2);
$('#peZoom').onclick = () => zoomTo(1);
$('#peClose').onclick = () => closeEditor(false);
$('#peDownload').onclick = async () => { if (ed.editing) await commitEdit(); if (ed.src.kind === 'image') return $('#peSave').click(); location.href = `/api/files/${ed.draft.id}?dl=1`; };
$('#peSave').onclick = async () => {
  if (ed.editing) await commitEdit();
  if (ed.busy) return;
  showBusy('Saving…');
  try {
    const isImg = ed.src.kind === 'image', stem = ed.src.name.replace(/\.[^.]+$/, '').replace(/_edited$/, '');
    const ext = isImg ? ((ed.src.name.split('.').pop() || 'png').toLowerCase().replace(/^(heic|heif|gif|bmp|tiff?)$/, 'jpg')) : 'pdf';
    const r = await api(`/api/pdf/${ed.draft.id}/${isImg ? 'save_image' : 'save'}`, {name: `${stem}_edited.${ext}`});
    const f = r.file;
    let chat = activeChat(); if (!chat) { chat = newChat('Edited ' + ed.src.name); activeId = chat.id; }
    if (!chat.fileIds.includes(f.id)) chat.fileIds.push(f.id);
    chat.messages.push({role: 'assistant', mode: 'fast', content: `Here's your edited ${isImg ? 'image' : 'PDF'}, **${f.name}** ✨ Open it again with ✏️ to keep editing, or tell me what else to change.`,
      tools: [{op: 'editor', status: 'ok', file: f, message: 'Edited in the editor · ' + metaStr(f)}]});
    chat.updated = Date.now(); save(); renderChatList(); renderChat();
    closeEditor(true); toast('Saved ' + f.name);
  } catch (e) { toast('Save failed: ' + e.message, 'err'); }
  finally { hideBusy(); }
};
let resizeT = null;
addEventListener('resize', () => { if (!edOpen()) return; clearTimeout(resizeT); resizeT = setTimeout(() => { computeBase(); loadPage(); }, 200); });
document.addEventListener('keydown', e => {
  if (!edOpen()) return;
  if (!$('#signModal').classList.contains('hidden')) { if (e.key === 'Escape') $('#signModal').classList.add('hidden'); e.stopImmediatePropagation(); return; }
  if (ed.editing) return;                                   // the text box handles its own keys
  if (/INPUT|SELECT|TEXTAREA/.test(document.activeElement?.tagName || '')) { if (e.key === 'Escape') document.activeElement.blur(); return; }
  e.stopImmediatePropagation();
  const k = e.key.toLowerCase(), mod = e.ctrlKey || e.metaKey;
  if (mod && k === 'z' && !e.shiftKey) { e.preventDefault(); stepHistory(ed.undo, ed.redo); }
  else if (mod && (k === 'y' || (k === 'z' && e.shiftKey))) { e.preventDefault(); stepHistory(ed.redo, ed.undo); }
  else if (mod) return;
  else if (k === 'delete' || k === 'backspace') { if (ed.sel) { e.preventDefault(); delSelection(); } }
  else if (k === 'escape') { if (ed.sel) select(null); else if (ed.tool === 'place') setTool('select'); else closeEditor(false); }
  else if (k === 'arrowright' || k === 'pagedown') gotoPage(ed.page + 1);
  else if (k === 'arrowleft' || k === 'pageup') gotoPage(ed.page - 1);
  else { const map = {v: 'select', e: 'text', t: 'addtext', i: 'image', d: 'draw', h: 'highlight', r: 'shape', x: 'erase', m: 'magic', s: 'sign'}; if (map[k]) setTool(map[k]); }
}, true);

/* ---------------- boot ---------------- */
hydrateIcons();
setMode(mode); setWeb(webMode);
document.fonts?.ready.then(placePill);
addEventListener('resize', () => { resizeSky(); placePill(); });
resizeSky(); requestAnimationFrame(tick);
renderChatList(); renderChat();
health(); setInterval(health, 25000);
input.focus();
})();
</script>
</body>
</html>
"""

# ==========================================================
# PROMPTING
# ==========================================================
def _tool_catalog():
    lines = []
    for t in TOOLS.values():
        need = "no file" if not t["accepts"] else (f"files: {'/'.join(t['accepts'])}" + (" (several, in order)" if t["multi"] else " (one)"))
        ps = ", ".join(p["name"] for p in t["params"]) or "—"
        lines.append(f"- {t['op']} [{need}; params: {ps}] — {t['desc']}")
    return "\n".join(lines)


def system_prompt(think):
    today = datetime.datetime.now().strftime("%A, %d %B %Y")
    mode = (
        "THINK MODE is on: before answering, reason carefully in your thinking — work out what's really being asked, break the "
        "problem down, consider alternatives and edge cases, check every calculation and line of code, and catch your own "
        "mistakes. Then write a polished final answer that stands on its own (never mention your reasoning process)."
        if think else
        "FAST MODE is on: answer directly and efficiently. Keep it tight, but never trade away correctness."
    )
    py = ("" if not FEATURES.get("python") else """
PYTHON — you can run real Python code (a persistent Jupyter-like session: variables survive between runs). Use it whenever exact
computation beats mental maths: arithmetic beyond trivial, statistics, unit conversions, dates, data analysis on attached
CSV/Excel/JSON files, plots and charts, simulations, checking code you wrote. Write a block like this and STOP — the app runs it
and shows you the output, then you continue:
```meow-python
import pandas as pd
df = pd.read_csv(FILES["sales.csv"])      # FILES maps each attached file name → its path
df.describe()
```
- The last expression is displayed (like Jupyter). print() works. matplotlib figures are shown to the user automatically
  (plt.show() is optional); use clear titles, labels and a readable style. Save any file you want to give the user into OUT_DIR
  (e.g. df.to_excel(f"{OUT_DIR}/report.xlsx")) and it appears as a download card.
- Available: numpy, pandas, matplotlib, scipy, sympy, sklearn (if installed), openpyxl, PIL. No internet from Python.
- Never invent output. Read the real output before concluding; if there's an error, fix it and run again (max 3 tries).
- Don't use Python for things that need no computation (explanations, writing, simple facts).
""")
    return f"""You are Meow 🐈 — a brilliant, warm and genuinely helpful AI assistant with the calm focus of a grey cat. You're talking with {USER_NAME or 'the user'}. Today is {today}.

{mode}

HOW TO ANSWER (this matters most)
- First understand the real need behind the message: what would an expert friend say? Answer that, not just the literal words.
- Lead with the answer or the key insight in the first sentence or two. Then give the supporting detail, reasoning or steps.
- Match length and shape to the question. Casual message → short, natural, friendly reply in plain sentences. Quick factual
  question → a direct answer plus only the context that adds value. Explanations, comparisons, plans, research → well-organised
  answers with clear structure. Never pad. Never repeat the question back. No "Great question!", no "I hope this helps".
- Formatting: use Markdown when it truly aids reading — headings only for long multi-part answers, numbered steps for procedures,
  bullets for genuinely parallel items, tables for comparisons across several attributes, **bold** for the few things that
  matter most. Explanations read best as clear prose paragraphs; don't chop reasoning into fragmented bullets.
- Teach like the best teachers: build intuition first, then precision. Use a concrete example, a small worked case or a vivid
  analogy when an idea is abstract. Define jargon the first time if the user may not know it.
- For decisions and advice: give a clear recommendation and why, the key trade-offs, and when the other option wins.
  Be specific (names, numbers, versions, commands, prices with dates) rather than generic.
- For code: complete, runnable, idiomatic, with a language tag on every block. Explain the approach briefly and point out the
  non-obvious parts. When debugging, find the root cause, explain it in one or two lines, then give the fix. When editing a
  user's file, return the whole updated file in one block.
- Maths/science: show the key steps, LaTeX with $inline$ and $$display$$, units throughout, and a sanity check of the result.
- Be honest and calibrated. If you're unsure, say how sure you are. If something can't be known, say so. If the user's premise is
  wrong, say so kindly and explain. Never invent facts, quotes, sources, numbers, file contents or tool results.
- If a request is genuinely ambiguous and a wrong guess would waste effort, ask one short clarifying question; otherwise make the
  most reasonable assumption, state it in a few words, and proceed.
- End cleanly. Instead of generic offers, you may close with one specific, useful next step.
- Reply in the user's language and register — English, Hindi or Hinglish — mirroring how they write.
- FOLLOW-UPS: after research, learning, planning or open-ended answers, end with one final line exactly in this form:
  <<followups: short question 1 | short question 2 | short question 3>>
  (questions the user would naturally ask next, written in their voice). Skip it for chit-chat, quick facts, or code-only replies.

WEB
- The app can search the web, Reddit, YouTube, Wikipedia, research papers, GitHub and Stack Overflow for you (🌐 button: Auto,
  Web, Reddit, Deep research), and it reads any link the user pastes. When <web_results> or <link> content is attached, use it
  and cite it.
- When nothing is attached and the question needs fresh information (news, prices, recent releases), answer with what you know,
  say it may be outdated, and suggest turning on 🌐 Web.

FILES
- Attached file contents arrive inside <file> tags, with ids listed above them.
- Images arrive as an automated analysis (caption, OCR text, colours, EXIF) — you don't see pixels. Ground what you say about an
  image in that analysis, and say when a detail isn't covered by it.

DIAGRAMS — for flows, architectures, timelines, mind-maps or sequences, a ```mermaid block renders as a real diagram. For a web
page, UI or visual demo, a single self-contained ```html block gets a live Preview button.
{py}
TOOLS — you can edit, convert and create real files. To run one, put a block like this in your reply:
```meow-tool
{{"op": "pdf_watermark", "files": ["<file id>"], "params": {{"text": "CONFIDENTIAL"}}}}
```
The app runs it after your reply and shows a download card. Use "$last" as a file id to chain onto the previous tool's output. Several blocks run in order. Only use tools when the user wants a file changed, converted or created; say in one line what you're doing, and never claim a result before it runs. Use make_pdf for documents/reports the user wants as a PDF. To change words in a PDF use pdf_replace_text (it keeps the original font, size and colour); for images inside a PDF use the pdf_*_image tools with the image numbers listed in the file. For words inside photos, screenshots, posters or scanned pages use image_replace_text (it reads the image, cleans the background and redraws the text in a matching font); pdf_replace_text also falls back to this automatically for scanned PDFs. For hands-on edits, the user can open any PDF or image in the visual editor (✏️ on its file card).
Available tools:
{_tool_catalog()}"""


def file_context(file_ids, budget):
    recs = [FILES[i] for i in dict.fromkeys(file_ids) if i in FILES]
    if not recs:
        return ""

    def meta(r):
        m = r["meta"]; bits = [r["kind"]]
        if m.get("pages"): bits.append(f"{m['pages']} pages")
        if m.get("dims"): bits.append(m["dims"])
        if r["origin"] == "tool": bits.append("created by a tool")
        return ", ".join(bits)

    header = "Files in this chat (ids for tools):\n" + "\n".join(f"- {r['name']} — id={r['id']} ({meta(r)})" for r in recs)
    remaining = budget - ntok(header)
    blocks = {}
    for r in reversed([r for r in recs if r["origin"] != "tool"]):   # newest files get priority
        if remaining < 200:
            break
        if r.get("text") is None:
            extract(r)
        body = r["text"] or "(no extractable content)"
        t = ntok(body)
        if t > remaining - 40:
            body = trim_tokens(body, remaining - 60); t = remaining
        blocks[r["id"]] = f'<file name="{r["name"]}" id="{r["id"]}" type="{r["kind"]}">\n{body}\n</file>'
        remaining -= t + 30
    return header + "\n\n" + "\n\n".join(blocks[r["id"]] for r in recs if r["id"] in blocks)


def build_prompt(req, web_ctx=""):
    think = req.mode == "think"
    max_new = min(MAX_NEW_THINK if think else MAX_NEW_FAST, CTX // 2)
    budget = CTX - max_new - 64
    system = system_prompt(think) + (WEB_RULES if web_ctx else "")
    user = (req.prompt or "").strip() or (
        "Analyze the attached file(s) and give me a clear, well-organized overview of what's in them." if req.file_ids else "Hi!")
    sys_t = ntok(system)
    if ntok(user) > budget * 0.6:
        user = trim_tokens(user, int(budget * 0.6))
    avail = budget - sys_t - ntok(user)
    if web_ctx:
        if ntok(web_ctx) > int(avail * 0.7): web_ctx = trim_tokens(web_ctx, int(avail * 0.7)) + "\n</web_results>"
        avail -= ntok(web_ctx) + 8
    fctx = file_context(req.file_ids, int(avail * 0.75)) if req.file_ids else ""
    avail -= ntok(fctx)
    turns = []
    for m in reversed(req.history[-24:]):
        role, c = m.get("role"), str(m.get("content") or "")
        if role not in ("user", "assistant") or not c.strip():
            continue
        if role == "assistant":
            c = THINK_RE.sub("", c).strip()
        t = ntok(c) + 10
        if t > avail:
            break
        turns.insert(0, (role, c)); avail -= t
    ctx = "\n\n".join(x for x in (fctx, web_ctx) if x)
    user_full = f"{ctx}\n\n---\n\n{user}" if ctx else user
    return system, turns, user_full, max_new, think


def chatml(system, turns, user, think):
    s = f"<|im_start|>system\n{system}<|im_end|>\n"
    for role, c in turns:
        s += f"<|im_start|>{role}\n{c}<|im_end|>\n"
    s += f"<|im_start|>user\n{user}<|im_end|>\n<|im_start|>assistant\n"
    # Native Qwen thinking switch: open a think block, or pre-close it for fast mode
    return s + ("<think>\n" if think else "<think>\n\n</think>\n\n")


def sse(obj):
    return f"data: {json.dumps(obj, ensure_ascii=False)}\n\n"

# ==========================================================
# API
# ==========================================================
app = FastAPI(title="Meow")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])

# ---------- password gate, activity tracking, power-off ----------
import hashlib, hmac
from fastapi import Request
from fastapi.responses import JSONResponse, RedirectResponse
MEOW_PASSWORD = os.environ.get("MEOW_PASSWORD", "")
STATE = {"last": time.time(), "started": time.time(), "shutdown": False}
BATCH = os.environ.get("MEOW_BATCH") == "1"


def _auth_token():
    return hashlib.sha256(("meow-auth:" + MEOW_PASSWORD).encode()).hexdigest()[:40]


LOGIN_HTML = """<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Meow · sign in</title><style>
body{margin:0;min-height:100vh;display:flex;align-items:center;justify-content:center;background:radial-gradient(ellipse at top,#1e1240,#05050d 70%);
font-family:Inter,system-ui,sans-serif;color:#e8e8f5}form{width:min(340px,88vw);padding:30px 26px;border-radius:22px;background:rgba(255,255,255,.05);
border:1px solid rgba(167,139,250,.25);text-align:center;box-shadow:0 30px 80px -20px rgba(124,58,237,.45)}h1{margin:6px 0 4px;font-size:1.6rem}
p{margin:0 0 20px;color:#9a9cc0;font-size:.9rem}input{width:100%;box-sizing:border-box;padding:12px 14px;border-radius:12px;border:1px solid rgba(255,255,255,.15);
background:rgba(0,0,0,.35);color:#fff;font-size:1rem;outline:none}input:focus{border-color:#a78bfa}button{margin-top:14px;width:100%;padding:12px;border:0;
border-radius:12px;background:linear-gradient(135deg,#7c3aed,#2563eb);color:#fff;font-size:1rem;font-weight:600;cursor:pointer}.err{color:#fca5a5;font-size:.85rem;margin-top:10px}
</style></head><body><form method="post" action="/login"><div style="font-size:2.6rem">🐈</div><h1>Meow</h1><p>Enter your password</p>
<input type="password" name="password" autofocus autocomplete="current-password"><button>Open Meow</button>__ERR__</form></body></html>"""


@app.middleware("http")
async def _gate(request: Request, call_next):
    path = request.url.path
    if path != "/api/health": STATE["last"] = time.time()
    if MEOW_PASSWORD and path not in ("/login", "/api/health", "/favicon.ico"):
        if not hmac.compare_digest(request.cookies.get("meow_auth", ""), _auth_token()):
            if path.startswith("/api/"): return JSONResponse({"detail": "Login required — reload the page"}, status_code=401)
            return RedirectResponse("/login", status_code=303)
    return await call_next(request)


@app.get("/login", response_class=HTMLResponse)
def login_page():
    return HTMLResponse(LOGIN_HTML.replace("__ERR__", ""))


@app.post("/login")
async def login(request: Request):
    form = await request.form()
    if MEOW_PASSWORD and hmac.compare_digest(str(form.get("password", "")), MEOW_PASSWORD):
        r = RedirectResponse("/", status_code=303)
        https = request.headers.get("x-forwarded-proto", "") == "https" or request.url.scheme == "https"
        r.set_cookie("meow_auth", _auth_token(), max_age=30 * 86400, httponly=True, samesite="lax", secure=https)
        return r
    time.sleep(1.2)                                        # slow down guessing
    return HTMLResponse(LOGIN_HTML.replace("__ERR__", '<div class="err">Wrong password</div>'), status_code=401)


@app.post("/api/shutdown")
def shutdown():
    STATE["shutdown"] = True
    if not BATCH and globals().get("MEOW_SERVER") is not None:
        threading.Timer(1.0, lambda: setattr(MEOW_SERVER, "should_exit", True)).start()
    return {"ok": True, "batch": BATCH}


def announce(msg, url=None):
    """Push a phone/desktop notification through ntfy.sh (topic = MEOW_NTFY_TOPIC). Tap it to open Meow."""
    topic = os.environ.get("MEOW_NTFY_TOPIC", "").strip()
    if not topic: return
    try:
        import requests
        h = {"Title": "Meow", "Tags": "cat", "Priority": "high" if url else "default"}
        if url: h["Click"] = url; h["Actions"] = f"view, Open Meow, {url}"
        requests.post(f"https://ntfy.sh/{topic}", data=msg.encode("utf-8"), headers=h, timeout=10)
    except Exception as e:
        print("⚠️  notification failed:", e)


class ChatReq(BaseModel):
    prompt: str = ""
    mode: str = "fast"
    history: list = []
    file_ids: List[str] = []
    rid: str = ""
    web: str = "off"          # off | auto | web | reddit | deep
    chat_id: str = ""


class ToolReq(BaseModel):
    op: str = ""
    files: List[str] = []
    params: dict = {}
    raw: str = ""
    last: str = ""


@app.get("/", response_class=HTMLResponse)
def index():
    return HTMLResponse(MEOW_HTML.replace("__USER_NAME__", html.escape(USER_NAME or "")))


@app.get("/api/health")
def health():
    return {"status": "online", "version": MEOW_VERSION, "batch": BATCH, "auth": bool(MEOW_PASSWORD),
            "changelog": MEOW_CHANGELOG[:8], "model": GGUF_FILENAME, "model_short": MODEL_LABEL, "ctx": CTX,
            "chatml": USE_CHATML, "files": len(FILES), **FEATURES}


@app.get("/api/tools")
def list_tools():
    return {"tools": [{k: v for k, v in t.items() if k != "fn"} for t in TOOLS.values()]}


@app.get("/api/files")
def list_files():
    return {"files": [public(r) for r in sorted(FILES.values(), key=lambda r: -r["created"]) if r["origin"] != "draft"]}


@app.get("/api/files/{fid}")
def get_file(fid: str, dl: int = 0):
    rec = FILES.get(fid)
    if not rec or not os.path.exists(rec["path"]):
        raise HTTPException(404, "File not found (server may have restarted)")
    mt = mimetypes.guess_type(rec["name"])[0] or "application/octet-stream"
    return FileResponse(rec["path"], media_type=mt, filename=rec["name"],
                        content_disposition_type="attachment" if dl else "inline")


@app.get("/api/thumb/{fid}")
def thumb(fid: str, size: int = 400):
    rec = FILES.get(fid)
    if not rec or rec["kind"] not in ("image", "pdf"):
        raise HTTPException(404)
    size = max(64, min(int(size), 2000))
    tp = os.path.join(WORK_DIR, f".thumb_{fid}_{size}.jpg")
    if not os.path.exists(tp):
        try:
            if rec["kind"] == "pdf":
                with fitz.open(rec["path"]) as d:
                    pix = d[0].get_pixmap(dpi=int(72 * min(3, max(0.5, size / 600))))
                    im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            else:
                im = ImageOps.exif_transpose(Image.open(rec["path"]))
            im = _flatten(im); im.thumbnail((size, size), Image.LANCZOS); im.save(tp, "JPEG", quality=85)
        except Exception:
            raise HTTPException(404)
    return FileResponse(tp, media_type="image/jpeg")


@app.post("/api/upload")
async def upload(files: List[UploadFile] = File(...)):
    out = []
    for f in files:
        data = await f.read()
        rec = await run_in_threadpool(ingest, f.filename or "file", data)
        out.append(public(rec))
    return {"files": out}


@app.post("/api/tool")
async def api_tool(req: ToolReq):
    try:
        op, fids, params = req.op, req.files, req.params
        if req.raw:
            d = json.loads(req.raw.strip(), strict=False)
            op, fids, params = d.get("op", ""), d.get("files") or [], d.get("params") or {}
            if isinstance(fids, str): fids = [fids]
        fids = [req.last if f == "$last" else f for f in fids]
        rec, msg = await run_in_threadpool(run_tool, op, fids, params)
        return {"ok": True, "op": op, "file": public(rec), "message": msg}
    except json.JSONDecodeError as e:
        return {"ok": False, "error": f"Tool call wasn't valid JSON ({e.msg})."}
    except Exception as e:
        return {"ok": False, "error": str(e)}


@app.post("/api/stop")
def stop(d: dict):
    if d.get("rid"): STOP.add(d["rid"])
    return {"ok": True}


@app.post("/api/chat")
def chat(req: ChatReq):
    def gen():
        if not LLM_LOCK.acquire(blocking=False):
            yield sse({"status": "Waiting for the GPU…"})
            LLM_LOCK.acquire()
        try:
            web_ctx, wm, q = "", (req.web or "off").lower(), (req.prompt or "").strip()
            urls = URL_RE.findall(q)
            do_search = bool(q) and wm in ("auto", "web", "reddit", "deep") and (wm != "auto" or bool(AUTO_SEARCH_RE.search(re.sub(URL_RE, " ", q))))
            if urls or do_search:
                mx = min(MAX_NEW_THINK if req.mode == "think" else MAX_NEW_FAST, CTX // 2)
                big = wm in ("deep", "reddit")
                wb = max(700, int(min(8500 if big else 5000, (CTX - mx - 2600) * (0.8 if big else 0.65))))
                try:
                    for ev in web_research(q, req.history, wm, wb, req.rid, urls=urls, search=do_search):
                        if "_ctx" in ev: web_ctx = ev["_ctx"]
                        else: yield sse(ev)
                except Exception as e:
                    print("⚠️  web research failed:", e)
                    yield sse({"search": {"stage": "error", "error": str(e)[:200]}})
            if req.rid and req.rid in STOP:
                return
            if req.file_ids:
                yield sse({"status": "Reading your files…"})
            system, turns, user, max_new, think = build_prompt(req, web_ctx)
            yield sse({"status": ""})
            stopped = lambda: bool(req.rid and req.rid in STOP)
            n, t_first = 0, None
            if USE_CHATML:
                if think:
                    yield sse({"t": "<think>\n"})
                base = chatml(system, turns, user, think)
                acc, scan_from, rounds = "", 0, 0
                py_on = bool(FEATURES.get("python"))
                while True:
                    budget_left = max_new - n
                    room = CTX - ntok(base + acc) - 16
                    if budget_left < 32 or room < 64: break
                    stream = llm.create_completion(
                        prompt=base + acc, max_tokens=min(budget_left, room), stream=True,
                        temperature=0.6 if think else 0.7, top_p=0.95 if think else 0.8, top_k=20, min_p=0.0,
                        presence_penalty=0.5 if think else 0.3, repeat_penalty=1.0,
                        stop=["<|im_end|>", "<|endoftext|>", "<|im_start|>", "```meow-output"])
                    code = None
                    for ch in stream:
                        if stopped(): break
                        tok = ch["choices"][0].get("text", "")
                        if not tok: continue
                        if t_first is None: t_first = time.time()
                        n += 1; acc += tok
                        yield sse({"t": tok})
                        if py_on and "`" in tok and "```meow-python" in acc[scan_from:]:
                            te = acc.find("</think>")
                            if think and te < 0: continue                         # still thinking — never run code from thoughts
                            m = PY_BLOCK_RE.search(acc, max(scan_from, te + 8 if think else 0))
                            if m:
                                code = m.group(1); extra = len(acc) - m.end()
                                if extra > 0: acc = acc[:m.end()]; yield sse({"rollback": extra})   # drop text after the fence
                                scan_from = m.end(); break
                    if code is None or stopped() or rounds >= PY_MAX_ROUNDS: break
                    rounds += 1
                    if not acc.endswith("\n"): acc += "\n"; yield sse({"t": "\n"})
                    yield sse({"status": "Running Python…"})
                    try:
                        md, recs, _res = run_python(code, req.chat_id or req.rid, req.file_ids, stop=stopped)
                    except Exception as e:
                        md, recs = f"\n```meow-output error\nCouldn't run the code: {e}\n```\n", []
                    acc += md; scan_from = len(acc)
                    yield sse({"t": md})
                    if recs: yield sse({"files": [public(r) for r in recs]})
                    yield sse({"status": ""})
            else:
                msgs = [{"role": "system", "content": system}] + [{"role": r, "content": c} for r, c in turns] + \
                       [{"role": "user", "content": user + (" /think" if think else " /no_think")}]
                stream = llm.create_chat_completion(messages=msgs, max_tokens=max_new, stream=True,
                                                    temperature=0.6 if think else 0.7, top_p=0.95 if think else 0.8)
                for ch in stream:
                    if stopped(): break
                    tok = ch["choices"][0]["delta"].get("content", "") or ""
                    if tok:
                        if t_first is None: t_first = time.time()
                        n += 1
                        yield sse({"t": tok})
            dt = (time.time() - t_first) if t_first else 0
            yield sse({"stats": {"tokens": n, "secs": round(dt, 1), "tps": round(n / dt, 1) if dt > 0 else 0}})
        except Exception as e:
            yield sse({"t": f"\n\n⚠️ **Server error:** `{e}`"})
        finally:
            STOP.discard(req.rid)
            LLM_LOCK.release()
        yield "data: [DONE]\n\n"

    return StreamingResponse(gen(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


# ==========================================================
# VISUAL PDF EDITOR API
# ==========================================================
class EditReq(BaseModel):
    ops: list = []


class SaveReq(BaseModel):
    name: str = ""


def _need_pdf(fid, allow_image=False):
    rec = FILES.get(fid)
    if not rec or rec["kind"] not in (("pdf", "image") if allow_image else ("pdf",)) or not os.path.exists(rec["path"]):
        raise HTTPException(404, "PDF not found (server may have restarted)")
    return rec


@app.post("/api/pdf/open")
async def pdf_open(d: dict):
    rec = _need_pdf(d.get("file_id", ""), allow_image=True)
    try:
        dr = await run_in_threadpool(open_draft, rec)
    except Exception as e:
        raise HTTPException(400, str(e))
    return {"draft": public(dr), "pages": await run_in_threadpool(pages_info, dr), "source": public(rec)}


@app.get("/api/pdf/{fid}/render/{n}")
def pdf_render(fid: str, n: int, scale: float = 1.5):
    rec = _need_pdf(fid)
    scale = max(0.1, min(float(scale), 5.0))
    tp = os.path.join(WORK_DIR, f".r_{fid}_{n}_{scale:.2f}.png")
    if not os.path.exists(tp):
        with fitz.open(rec["path"]) as d:
            if not 0 <= n < d.page_count: raise HTTPException(404)
            d[n].get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False, annots=True).save(tp)
    return FileResponse(tp, media_type="image/png", headers={"Cache-Control": "public, max-age=86400"})


@app.get("/api/pdf/{fid}/layout/{n}")
def pdf_layout(fid: str, n: int):
    rec = _need_pdf(fid)
    with fitz.open(rec["path"]) as d:
        if not 0 <= n < d.page_count: raise HTTPException(404)
        return page_layout(d[n])


@app.get("/api/pdf/{fid}/image/{xref}")
def pdf_image(fid: str, xref: int):
    rec = _need_pdf(fid)
    with fitz.open(rec["path"]) as d:
        try:
            im = _image_pil(d, xref)
        except Exception:
            raise HTTPException(404)
    buf = io.BytesIO(); im.save(buf, "PNG")
    tp = os.path.join(WORK_DIR, f".img_{fid}_{xref}.png"); open(tp, "wb").write(buf.getvalue())
    return FileResponse(tp, media_type="image/png", filename=f"image_{xref}.png")


@app.post("/api/pdf/{fid}/edit")
async def pdf_edit(fid: str, req: EditReq):
    rec = _need_pdf(fid)
    try:
        dr = await run_in_threadpool(edit_draft, rec, req.ops)
    except Exception as e:
        return {"ok": False, "error": str(e)}
    return {"ok": True, "draft": public(dr), "pages": await run_in_threadpool(pages_info, dr)}


@app.post("/api/pdf/{fid}/save")
def pdf_save(fid: str, req: SaveReq):
    rec = _need_pdf(fid)
    name = safe_name(req.name or rec["name"])
    if not name.lower().endswith(".pdf"): name += ".pdf"
    nid, path = new_path(name)
    with fitz.open(rec["path"]) as d:
        d.save(path, garbage=4, deflate=True, clean=True)
    out = register(nid, path, name, origin="tool")
    extract(out)
    return {"ok": True, "file": public(out)}


@app.get("/api/pdf/{fid}/ocr/{n}")
async def pdf_ocr(fid: str, n: int, x0: float = None, y0: float = None, x1: float = None, y1: float = None):
    rec = _need_pdf(fid)
    def run():
        with fitz.open(rec["path"]) as d:
            if not 0 <= n < d.page_count: raise HTTPException(404)
            return pdf_ocr_layout(d, n, [x0, y0, x1, y1] if x0 is not None else None)
    try:
        return await run_in_threadpool(run)
    except HTTPException:
        raise
    except Exception as e:
        return {"boxes": [], "image": None, "error": str(e)}


@app.post("/api/pdf/{fid}/save_image")
def pdf_save_image(fid: str, req: SaveReq):
    """Export the edited page back to an image at the original resolution."""
    rec = _need_pdf(fid)
    name = safe_name(req.name or rec["name"])
    ext = os.path.splitext(name)[1].lower().lstrip(".") or "png"
    if ext not in ("png", "jpg", "jpeg", "webp"): ext = "png"; name = os.path.splitext(name)[0] + ".png"
    with fitz.open(rec["path"]) as d:
        pg = d[0]; px = rec["meta"].get("image_px") or [int(pg.rect.width / 0.75), int(pg.rect.height / 0.75)]
        z = px[0] / pg.rect.width
        pix = pg.get_pixmap(matrix=fitz.Matrix(z, z), alpha=False)
        im = Image.open(io.BytesIO(pix.tobytes("png"))); im.load()
    out = _save_img(im, os.path.splitext(name)[0], ext, 95)
    return {"ok": True, "file": public(out)}

# ==========================================================
# LAUNCH (background server + public ngrok link)
# ==========================================================
if __name__ == "__main__":
    from pyngrok import ngrok
    _old = globals().get("MEOW_SERVER")
    if _old is not None:                         # re-running the cell: stop the old server cleanly
        _old.should_exit = True
        if globals().get("MEOW_THREAD"): MEOW_THREAD.join(timeout=10)
    try: ngrok.kill()
    except Exception: pass

    warm_up_media()                              # fonts + LaMa download in the background
    MEOW_SERVER = uvicorn.Server(uvicorn.Config(app, host="0.0.0.0", port=PORT, log_level="warning"))
    MEOW_THREAD = threading.Thread(target=MEOW_SERVER.run, daemon=True)
    MEOW_THREAD.start()
    for _ in range(60):
        if MEOW_SERVER.started: break
        time.sleep(0.25)

    _tok = NGROK_AUTH_TOKEN if NGROK_AUTH_TOKEN and "PASTE" not in NGROK_AUTH_TOKEN else None
    if not _tok:
        try:
            from kaggle_secrets import UserSecretsClient
            _tok = UserSecretsClient().get_secret("NGROK_AUTH_TOKEN")
        except Exception:
            pass
    if _tok: ngrok.set_auth_token(_tok)

    try:
        public_url = ngrok.connect(PORT).public_url
    except Exception as _e:                      # ngrok failed → free Cloudflare tunnel, no token needed
        print(f"⚠️  ngrok failed ({str(_e)[:90]}…) — switching to Cloudflare tunnel")
        import subprocess
        if not os.path.exists("./cloudflared"):
            os.system("wget -q https://github.com/cloudflare/cloudflared/releases/latest/download/cloudflared-linux-amd64 -O cloudflared && chmod +x cloudflared")
        MEOW_CF = subprocess.Popen(["./cloudflared", "tunnel", "--url", f"http://localhost:{PORT}", "--no-autoupdate"],
                                   stderr=subprocess.PIPE, text=True)
        public_url = "(waiting for Cloudflare…)"
        for _line in MEOW_CF.stderr:
            _m = re.search(r"https://[\w-]+\.trycloudflare\.com", _line)
            if _m: public_url = _m.group(); break
    print("=" * 66)
    print(f"🐈  MEOW IS LIVE  (v{MEOW_VERSION})")
    print(f"🌌  Open: {public_url}")
    print(f"🧠  Model: {MODEL_LABEL} · context {CTX:,} tokens · {'ChatML' if USE_CHATML else 'chat template'}")
    print("=" * 66)
    print(f"MEOW_URL {public_url}", flush=True)                 # machine-readable line for start_meow.py
    STATE["last"] = time.time()
    announce(f"Meow v{MEOW_VERSION} is live 🐈 — tap to open" + (" (password protected)" if MEOW_PASSWORD else ""), public_url)

    if BATCH:                                                # started remotely (start_meow.py / GitHub button)
        _idle = float(os.environ.get("MEOW_IDLE_MIN", "40")) * 60
        _maxh = float(os.environ.get("MEOW_MAX_HOURS", "8")) * 3600
        print(f"🕒  Auto power-off after {_idle / 60:.0f} min without use (max {_maxh / 3600:.0f} h) to save GPU hours.", flush=True)
        while True:
            time.sleep(15)
            _now = time.time()
            _why = ("you pressed Power off" if STATE["shutdown"] else
                    f"no use for {_idle / 60:.0f} min" if _now - STATE["last"] > _idle and not LLM_LOCK.locked() else
                    "time limit reached" if _now - STATE["started"] > _maxh else None)
            if _why:
                print(f"⏻  Shutting down: {_why}", flush=True)
                announce(f"Meow is off ({_why}). GPU released — start it again anytime.")
                MEOW_SERVER.should_exit = True; time.sleep(3)
                try: ngrok.kill()
                except Exception: pass
                break
